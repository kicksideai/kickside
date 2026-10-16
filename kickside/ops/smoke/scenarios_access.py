"""Component access, observable domain actions, and QA regression contracts."""
from dataclasses import dataclass, field
import base64
import json
import hashlib
import hmac
import re
import shutil
from pathlib import Path
import sqlite3
import subprocess
import time
import uuid
from urllib.parse import quote
from urllib import request as urlrequest

from fixture_member_permissions import register_member_policy_grants
from harness import Cleanup, Failure, Scenario, require
from browser import Browser
from cleanup_evidence import await_automation_deletion, workflow_publication_ready
from access_ownership import contracts as declared_contracts, entry_owners, fixture_owner, output_events
from mocks.jira_regressions import DiscordGatewayFixture, DiscordIngress


def http(component, *routes):
    return tuple(component + '|http|' + route for route in routes)


SHARE_FEATURES = http('kickside/component', 'POST /api/v1/component-access/share/{id}',
                      'GET /api/v1/component-access/grants/{id}')
WORKSPACE_FEATURES = http('chestor/workspace', 'POST /api/v1/workspace/entries',
                          'GET /api/v1/workspace/entries/{id}',
                          'PUT /api/v1/workspace/entries/{id}',
                          'POST /api/v1/workspace/entries/{id}/move',
                          'DELETE /api/v1/workspace/entries/{id}')
SHARING_RULE = {'share': 8, 'delete': 4, 'read': 1, 'write': 3}


@dataclass(frozen=True)
class Kind:
    name: str
    package: str
    binding: str
    read_route: str = ''
    features: tuple = ()


KINDS = (
    Kind('folder', 'chestor/workspace', 'chestor.workspace:folder'),
    Kind('file', 'chestor/workspace', 'chestor.workspace:file'),
    Kind('link', 'chestor/workspace', 'chestor.workspace:link'),
    Kind('graph', 'chestor/graph', 'chestor.graph:graph_kind', '/workspace/graphs/{id}'),
    Kind('journal', 'chestor/journal', 'chestor.journal:journal_kind', '/workspace/journals/{id}/context_pack'),
    Kind('roadmap', 'chestor/roadmap', 'chestor.roadmap:roadmap_kind', '/workspace/roadmaps/{id}'),
    Kind('database', 'chestor/db-schema', 'chestor.db_schema:database', '/workspace/db-schema/{id}'),
    Kind('agent', 'kickside/agents', 'kickside.agents:user_agent', '/user-agents/{id}'),
    Kind('kb_document', 'kickside/knowledge', 'kickside.kb10:embeddable', '/knowledge/{id}'),
    Kind('kb_wiki', 'kickside/knowledge', 'kickside.kb10:embeddable', '/knowledge/{id}'),
    Kind('workflow', 'kickside/workflows', 'kickside.workflows:workflow_kind', '/workflows/{id}/state'),
    Kind('recurring', 'kickside/automation', 'kickside.automation.recurring:task', '/automations/{id}/config'),
    Kind('data_sync', 'kickside/sync', 'kickside.sync:sync', '/automations/{id}/config'),
    Kind('connection', 'kickside/atlassian', 'kickside.atlassian.connection:atlassian_connection'),
    Kind('mcp_credential', 'kickside/mcp', ''),
    Kind('uploads', 'kickside/uploads', 'kickside.uploads:uploads_component'),
    Kind('crm', 'kickside/crm', 'kickside.crm.binding:crm', '/workspace/crm/{id}/schema'),
    Kind('research', 'kickside/research', 'kickside.research:research_kind', '/research/{id}/events'),
    Kind('skill', 'kickside/skills', 'kickside.skills:skill', '/skills/{id}'),
    Kind('webhook', 'kickside/webhooks', 'kickside.webhooks:webhook_kind'),
    Kind('inbox', 'kickside/inbox', 'kickside.inbox:inbox_binding'),
    Kind('lifecycle', 'kickside/core', 'kickside.core.lifecycle:system_lifecycle'),
    Kind('dashboard', 'kickside/dashboard', ''),
    Kind('remote_mcp', 'kickside/mcp-client', 'kickside.mcp_client:remote_connection'),
    Kind('binding_kickside_dm_dm_kind', 'kickside/dm', 'kickside.dm:dm_kind', features=
         http('kickside/identity', 'POST /api/v1/identity/link-codes', 'POST /api/v1/identity/unlink')),
)
_inventory = json.loads((Path(__file__).parent / 'coverage' / 'inventory.json').read_text())
INVENTORY_IDS = {f['id'] for component in _inventory['components'].values() for f in component['features']}
_adapted = {k.binding for k in KINDS}
DISCOVERED_KINDS = tuple(Kind('binding_' + f['entry'].replace('.', '_').replace(':', '_'), package, f['entry'])
    for package, component in _inventory['components'].items() for f in component['features']
    if f['category'] == 'workspace_kind' and f['entry'] not in _adapted)
KINDS += DISCOVERED_KINDS




STORE_FIXTURES = {
    'kickside.crm.sink:crm_records': {'kind': 'record', 'object_type': 'person',
                                    'values': {'name': 'Northwind sink access'}},
}


SYSTEM_STREAM_FIXTURES = {
    'kickside.channel.events:message.inbound': {'provider': 'discord', 'identity_field': 'external_message_id',
        'event': lambda marker, text, sender, bot: DiscordIngress.mention(marker, text, bot),
        'responder': True, 'linked': False},
    'kickside.dm.events:message.inbound': {'provider': 'discord', 'identity_field': 'external_message_id',
        'event': lambda marker, text, sender, bot: DiscordIngress.direct_message(marker, text, sender),
        'responder': False, 'linked': True},
}




@dataclass
class Fixture:
    kind: Kind
    id: str
    title: str
    data: dict = field(default_factory=dict)


class Checks:
    """Retain every independent failure and fail the scenario after recording cells."""
    def __init__(self, ctx, kind, role):
        self.ctx, self.kind, self.role, self.rows = ctx, kind, role, []

    def check(self, action, operation):
        try:
            value = operation()
            self.rows.append({'kind': self.kind, 'role': self.role, 'action': action,
                              'status': 'pass', 'error': ''})
            return value
        except Exception as error:
            self.rows.append({'kind': self.kind, 'role': self.role, 'action': action,
                              'status': 'fail', 'error': self.ctx.inst.redactor.text(
                                  type(error).__name__ + ': ' + str(error))})
            return None

    def finish(self):
        self.ctx.inst.evidence.save('access-matrix-' + self.kind.replace(':', '_').replace('/', '_') + '-' + self.role, self.rows)
        errors = [r['action'] + ': ' + r['error'] for r in self.rows if r['status'] == 'fail']
        require(not errors, '\n'.join(errors))


def identity(ctx, role):
    return ctx.inst.token if role == 'owner' else ctx.state['access_users'][role]['bearer']


def request(ctx, role, method, path, body=None, *, denied=False, hidden=False, expect=None, label=None, **kwargs):
    expected = expect or ((403, 404) if hidden or denied else (200,))
    response = ctx.inst.request(method, '/api/v1' + path, body, auth=identity(ctx, role),
                                label=label or role + '-' + method.lower() + '-' + uuid.uuid4().hex[:8],
                                expect=expected, **kwargs)
    if denied or hidden:
        require(response.get('error') and response.get('success') is False,
                'Denied operation returns no observable error: ' + str(response))
    return response


def renew(ctx):
    for user in ctx.state['access_users'].values():
        token = ctx.inst.request('POST', '/api/public/user/token',
                {'email': user['email'], 'password': ctx.inst.password})['token']
        ctx.inst.redactor.add(token)
        user['bearer'] = token


def registry(ctx):
    entries = []
    offset = 0
    while True:
        page = request(ctx, 'owner', 'GET', '/keeper/registry/entries?kind=contract.binding&limit=200&offset=' + str(offset))
        entries.extend(page['entries'])
        if not page['has_more']:
            break
        require(page['count'] > 0, 'Live registry pagination makes no progress')
        offset += page['count']
    declarations = {e['id']: request(ctx, 'owner', 'GET', '/keeper/registry/entry?id=' + quote(e['id'], safe=''))['entry']
                    for e in entries if e['kind'] == 'contract.binding' and e.get('meta', {}).get('class')}
    ctx.state['access_declarations'] = declarations
    return entries


def component_rows(ctx, binding):
    """Read persistence only to count scoped instances; permission checks use APIs."""
    sql = 'SELECT component_id, impl_id, parent_id FROM kickside_components WHERE impl_id = '
    if ctx.inst.engine == 'sqlite':
        with sqlite3.connect('file:' + str(ctx.inst.root / 'data' / 'app.db') + '?mode=ro', uri=True) as db:
            db.row_factory = sqlite3.Row
            rows = [dict(r) for r in db.execute(sql + '?', (binding,))]
    else:
        pg = ctx.inst.config['pg']
        query = "SELECT coalesce(json_agg(t),'[]'::json) FROM (" + sql + "'" + binding.replace("'", "''") + "') t"
        result = subprocess.run(['psql', '-h', pg['host'], '-p', pg['port'], '-U', pg['user'],
                    '-d', ctx.inst.db, '-v', 'ON_ERROR_STOP=1', '-Atqc', query],
                    env=ctx.inst.env, capture_output=True, text=True, timeout=35)
        require(result.returncode == 0, 'Instance inventory query fails: ' + ctx.inst.redactor.text(result.stderr))
        rows = json.loads(result.stdout)
    ctx.inst.evidence.save('instance-count-' + uuid.uuid4().hex[:8], {'binding': binding, 'instances': rows})
    return rows


def instancing(declaration):
    provision = declaration.get('meta', {}).get('provision')
    if provision is None:
        return 'demand'
    scope = provision.get('scope', 'user')
    require(scope in ('system', 'user'), 'Invalid declared provision scope: ' + str(scope))
    return scope


def provision_id(scope_key, class_name):
    value = hmac.new(b'kickside.autoinit', (scope_key + ':' + class_name).encode(), hashlib.sha256).hexdigest()
    return value[:8] + '-' + value[8:12] + '-5' + value[13:16] + '-' + value[16:20] + '-' + value[20:32]


def actor_id(ctx, role):
    if role == 'owner':
        return request(ctx, role, 'GET', '/user/me')['user']['id']
    return ctx.state['access_users'][role]['id']


def provisioned(ctx, kind, role='owner'):
    declaration = ctx.state['access_declarations'][kind.binding]
    scope = instancing(declaration)
    require(scope != 'demand', 'On-demand kind has no provision declaration')
    classes = declaration['meta']['class']
    class_name = classes[0] if isinstance(classes, list) else classes
    rows = component_rows(ctx, kind.binding)
    expected = provision_id('system' if scope == 'system' else actor_id(ctx, role), class_name)
    if scope == 'user' and kind.name in ('inbox', 'uploads'):
        def present():
            current = component_rows(ctx, kind.binding)
            return current if any(row['component_id'] == expected for row in current) else None
        rows = ctx.inst.wait(present, 'Login provisions the declared ' + kind.name + ' USER instance for ' + role, seconds=35)
    if scope == 'system':
        require(len(rows) == 1, 'SYSTEM kind must have exactly one instance: ' + str(rows))
        require(rows[0]['component_id'] == expected, 'SYSTEM singleton ID differs from declared deterministic ID: ' + str(rows))
    else:
        require(sum(r['component_id'] == expected for r in rows) == 1,
                'USER kind must have exactly one deterministic instance for ' + role + ': ' + str(rows))
    f = Fixture(kind, expected, declaration['meta'].get('title') or kind.name,
                {'scope': scope, 'declaration': declaration})
    ctx.inst.evidence.save('instancing-' + role, {'kind': kind.name, 'scope': scope, 'expected_id': expected, 'instances': rows})
    return f


def install_package(ctx, package):
    specifications = json.loads((Path(__file__).parent / 'mocks' / 'access_install_parameters.json').read_text())
    specification = specifications.get(package, {})
    if 'parameters' in specification:
        return ctx.inst.install(package, parameters=specification['parameters'])
    profile_id = specification.get('profile')
    if not profile_id:
        return ctx.inst.install(package)
    from urllib.parse import quote
    profile = request(ctx, 'owner', 'GET', '/keeper/registry/entry?id=' + quote(profile_id, safe=''))['entry']
    packages = profile['data']['packages']
    declared = [row for row in packages if row['component'] == package]
    require(len(declared) == 1 and declared[0].get('parameters'),
            'Onboarding profile supplies no unique declared install parameters for ' + package)
    ctx.inst.evidence.save('package-install-design', {'design_reference': profile_id,
        'package': declared[0]})
    return ctx.inst.install(package, parameters=declared[0]['parameters'])


def access_catalog(ctx):
    checks = Checks(ctx, 'inventory', 'owner')
    packages = json.loads((Path(__file__).parent / 'coverage' / 'modules.json').read_text())['modules'] + ['kickside/skills']
    for package in packages:
        checks.check('install_' + package, lambda p=package: install_package(ctx, p))
    renew(ctx)
    entries = registry(ctx)
    ctx.inst.evidence.save('live-kind-inventory', {'workspace': request(ctx, 'owner', 'GET', '/workspace/entries/types'),
                    'components': request(ctx, 'owner', 'GET', '/mcp/components'),
                    'bindings': list(ctx.state['access_declarations'].values()), 'entries': entries})
    checks.finish()


def access_users(ctx):
    inst = ctx.inst
    install_package(ctx, 'chestor/workspace')
    install_package(ctx, 'kickside/mcp')
    ctx.state['access_users'] = {}
    for role in ('write', 'read', 'stranger'):
        email = 'northwind-' + role + '-' + uuid.uuid4().hex[:8] + '@smoke.example'
        user = inst.request('POST', '/api/v1/users', {'email': email, 'password': inst.password,
                            'full_name': 'Northwind ' + role, 'groups': ['app.security:members']},
                            expect=(201,), label='create-' + role)['user']
        ctx.state['access_users'][role] = {'email': email, 'id': user['user_id']}
    renew(ctx)
    for role in ('write', 'read', 'stranger'):
        profile = request(ctx, role, 'GET', '/user/me')['user']
        ctx.state['access_users'][role]['id'] = profile['id']
        actor_browser(ctx, role)
    actor_browser(ctx, 'owner')
    ctx.state['access_owner_id'] = actor_id(ctx, 'owner')
    scopes = request(ctx, 'owner', 'GET', '/users/assignable-scopes')
    groups = [s for s in scopes['scopes'] if s['id'] == 'app.security:user']
    require(groups, 'The member group is absent from declared assignable groups')
    vocabulary = request(ctx, 'owner', 'GET', '/components/access-model')['capabilities']
    bits = {row['key']: row['bit'] for row in vocabulary}
    require(bits == {'read': 1, 'write': 2, 'delete': 4, 'admin': 8}, 'Access vocabulary differs: ' + str(bits))
    ctx.state['access_rule'] = {'share': 'ADMIN (8)', 'delete': 'DELETE (4)',
                              'write_grant': 3, 'read_grant': 1,
                              'group': 'app.security:user'}
    inst.evidence.save('declared-access-rule', ctx.state['access_rule'])


def component_id(fixture):
    return fixture.data.get('component_id', fixture.id)


def share(ctx, fixture, subject, mask=1, *, group=False, revoke=False, role='owner', denied=False):
    sid = subject if group else ctx.state['access_users'][subject]['id']
    body = {'subject_id': sid, 'subject_type': 'group' if group else 'user'}
    body.update({'action': 'revoke'} if revoke else {'access_mask': mask})
    result = request(ctx, role, 'POST', '/component-access/share/' + component_id(fixture), body, denied=denied)
    grants = request(ctx, 'owner', 'GET', '/component-access/grants/' + component_id(fixture))
    matching = [g for g in grants['access'] if g['subject_id'] == sid]
    if denied or revoke:
        require(not matching, 'Denied/revoked sharing leaves a grant: ' + str(matching))
    else:
        require(len(matching) == 1 and matching[0]['access_mask'] == mask,
                'Sharing does not persist the exact mask: ' + str(grants))
    return result


def workspace(ctx, fixture, role='owner', hidden=False):
    return request(ctx, role, 'GET', '/workspace/entries/' + component_id(fixture), hidden=hidden)


def upload(ctx, role, filename, content, *, target_parent_id=None):
    boundary = 'access_' + uuid.uuid4().hex
    raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
           'Content-Type: text/plain\r\n\r\n').encode() + content.encode() + f'\r\n--{boundary}--\r\n'.encode()
    if target_parent_id is not None:
        metadata = json.dumps({'target_parent_id': target_parent_id})
        raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="metadata"\r\n\r\n'
               f'{metadata}\r\n').encode() + raw
    result = request(ctx, role, 'POST', '/uploads', raw=raw,
                     headers={'Content-Type': 'multipart/form-data; boundary=' + boundary}, timeout=120)
    uid = result['uuid']
    def ready():
        row = request(ctx, role, 'GET', '/uploads/' + uid)['upload']
        require(row['status'] not in ('failed', 'error'), 'Upload processing fails: ' + str(row))
        return row if row['status'] == 'completed' else None
    ctx.inst.wait(ready, 'Access fixture upload conversion', seconds=120)
    return result


ACCESS_PROVIDER_CREDENTIALS = {
    'kickside.' + provider + '.connection:' + provider + '_connection': credentials
    for provider, credentials in {
        'github': {'token': 'fictional-github-token'},
        'slack': {'bot_token': 'fictional-slack-token', 'app_token': 'fictional-slack-app'},
        'discord': {'bot_token': 'fictional-discord-token'},
        'hubspot': {'api_key': 'fictional-hubspot-token'},
        'telegram': {'bot_token': 'fictional-telegram-token', 'default_chat_id': 'NW'},
        'fireflies': {'api_key': 'fictional-fireflies-key'},
        'phantombuster': {'api_key': 'fictional-phantombuster-key'},
    }.items()
}


def setup_access_connection_fixtures(ctx):
    if ctx.inst.fixture_state.get('access_connection_bases'):
        ctx.state['access_connection_bases'] = ctx.inst.fixture_state['access_connection_bases']
        return
    import yaml
    from mocks import Fixtures
    from scenarios_providers import configure_provider_endpoints, ORIGIN_ROUTES
    names = ('jira', 'confluence', 'github', 'slack', 'discord', 'hubspot', 'telegram', 'fireflies', 'phantombuster', 'simplefin', 'google')
    for name in names:
        if name not in ('jira', 'confluence'):
            install_package(ctx, 'kickside/' + name)
    fixtures = Fixtures()
    setup_cleanup = Cleanup()
    setup_cleanup.add('access connection fixtures', fixtures.stop)
    try:
        bases = {name: fixtures.start(name) for name in names}
        fixture = ctx.inst.app / 'src' / 'access-connections'
        fixture.mkdir()
        setup_cleanup.add('access fixture source', lambda: shutil.rmtree(fixture))
        entries = [
            {'name': 'egress', 'kind': 'security.policy', 'groups': ['app.security:user'],
             'policy': {'resources': [base + '/*' for base in bases.values()],
                        'actions': ['http_client.request'], 'effect': 'allow'}},
            {'name': 'loopback', 'kind': 'security.policy', 'groups': ['app.security:user'],
             'policy': {'resources': ['127.0.0.1'], 'actions': ['http_client.private_ip'], 'effect': 'allow'}},
        ]
        (fixture / 'ingress.lua').write_text((Path(__file__).parent / 'mocks/jira_ingress.lua').read_text())
        entries.append({'name': 'ingress', 'kind': 'function.lua', 'source': 'file://ingress.lua', 'method': 'handler',
                        'modules': ['funcs'], 'meta': {'type': 'tool', 'public': True, 'llm_alias': 'NorthwindAccessIngress',
                        'title': 'Northwind local provider ingress', 'input_schema': {'type': 'object',
                        'properties': {'connection_id': {'type': 'string'}, 'event': {'type': 'object'}},
                        'required': ['connection_id', 'event']}}})
        entries.append({'name': 'ingress_access', 'kind': 'security.policy', 'groups': ['app.security:user'],
                        'policy': {'resources': ['kickside.channel:receive', 'kickside.dm:receive'],
                                   'actions': 'channel.receive.inject', 'effect': 'allow'}})
        (fixture / 'probe.lua').write_text((Path(__file__).parent / 'mocks/providers_probe.lua').read_text())
        entries.append({'name': 'probe', 'kind': 'function.lua', 'source': 'file://probe.lua', 'method': 'handler',
                        'modules': ['funcs', 'contract', 'registry', 'json'],
                        'imports': {'component': 'kickside.component:component'},
                        'meta': {'type': 'tool', 'public': True, 'llm_alias': 'NorthwindAccessConnectionProbe',
                                 'title': 'Northwind access connection probe', 'input_schema': {'type': 'object',
                                 'properties': {'connection_id': {'type': 'string'}, 'tool': {'type': 'string'},
                                                'port': {'type': 'string'}, 'write': {'type': 'boolean'},
                                                'arguments': {'type': 'object'}}, 'required': ['connection_id', 'arguments']}}})
        entries.append({'name': 'tools_access', 'kind': 'security.policy', 'groups': ['app.security:user'],
                        'policy': {'resources': ['app.smoke.access_connections:' + entry['name'] for entry in entries
                                                 if entry['kind'] == 'function.lua' and entry.get('meta', {}).get('type') == 'tool'],
                                   'actions': ['registry.get', 'registry.find', 'funcs.call', 'funcs.security'],
                                   'effect': 'allow'}})
        entries = register_member_policy_grants(ctx.inst.app, entries)
        (fixture / '_index.yaml').write_text(yaml.safe_dump({'version': '1.0', 'namespace': 'app.smoke.access_connections',
                                                           'entries': entries}, sort_keys=False))
        endpoints = {requirement: bases[name] + path for requirement, (name, path) in ORIGIN_ROUTES.items()
                     if name in bases and not requirement.startswith('kickside.sso.') and
                     not requirement.startswith('kickside.webscout:')}
        endpoints['kickside.discord:endpoint_gateway'] = DiscordGatewayFixture(bases['discord']).request('gateway')['url']
        configure_provider_endpoints(ctx, endpoints)
        ctx.restart()
        ctx.inst.manifest()
    except BaseException as failure:
        errors = setup_cleanup.run()
        if errors:
            raise Failure(str(failure) + '\n' + '\n'.join(errors)) from failure
        raise
    ctx.inst.cleanup.add('access connection fixtures', fixtures.stop)
    ctx.inst.fixture_state['access_connection_bases'] = bases
    ctx.state['access_connection_bases'] = bases


def create_oauth_access_fixture(ctx, kind, title, role):
    setup_access_connection_fixtures(ctx)
    for key, value in (('client_id', 'northwind-client'), ('client_secret', 'fictional-oidc-secret')):
        saved = request(ctx, 'owner', 'PUT', '/settings', {'namespace': 'kickside.google', 'key': key, 'value': value})
        require(saved.get('success'), 'OAuth fixture setting does not persist')
    before = {row['component_id'] for row in component_rows(ctx, kind.binding)}
    initialized = request(ctx, role, 'POST', '/oauth/providers/google/init', {'name': title, 'scopes': []})
    require(initialized['authorization_url'].startswith(ctx.state['access_connection_bases']['google']),
            'OAuth fixture authorization does not use its declared endpoint')
    browser = actor_browser(ctx, role)
    browser.settle()
    browser.page.goto(initialized['authorization_url'], wait_until='domcontentloaded')
    def created():
        ids = {row['component_id'] for row in component_rows(ctx, kind.binding)} - before
        require(len(ids) <= 1, 'One OAuth callback creates multiple connection components')
        return next(iter(ids)) if ids else None
    return ctx.inst.wait(created, 'OAuth owning flow persists its connection', pump=browser.page.wait_for_timeout)


def workflow_target_schema(schema):
    fields = schema.get('properties', {})
    reference = fields.get('flow_ref', {}).get('properties', {})
    return (schema.get('type') == 'object' and 'workflow' in fields.get('kind', {}).get('enum', [])
            and reference.get('kind', {}).get('const') == 'flow'
            and reference.get('id', {}).get('x-reference', {}).get('kind') == 'flow'
            and reference.get('version', {}).get('type') == 'integer'
            and fields.get('input', {}).get('type') == 'object')


def supported_installable(declaration):
    if 'kickside.contract:installable' not in declared_contracts(declaration):
        return False
    schema = declaration.get('meta', {}).get('inputs', {})
    fields = schema.get('properties', {})
    for name in schema.get('required', []):
        field = fields.get(name, {})
        if ('default' in field or name in ('title', 'channel_id')
                or field.get('x-reference', {}).get('kind') in ('connection', 'agent')
                or workflow_target_schema(field)):
            continue
        return False
    return bool(fields)


def fixture_kind_for_binding(ctx, binding):
    native = next((kind for kind in KINDS if kind.binding == binding), None)
    if native:
        return native
    dependencies = request(ctx, 'owner', 'GET', '/keeper/hub/dependencies?entries=1&migrations=0')['dependencies']
    package = entry_owners(dependencies).get(binding)
    require(package, 'Declared installable type has no resolved module provenance: ' + binding)
    return Kind('binding_' + binding.replace('.', '_').replace(':', '_'), package, binding)


def automation_input(ctx, schema, title, role, *, connection=None):
    fields = schema['properties']
    required = schema.get('required', [])
    needs_channel = 'channel_id' in required
    needs_connection = needs_channel or any(fields[name].get('x-reference', {}).get('kind') == 'connection'
                                           for name in required)
    values = {'agent': ctx.state['agent']}
    supplied = {'title': title}
    if needs_connection:
        setup_access_connection_fixtures(ctx)
        registry(ctx)
        reply_bindings = {entry['id'] for entry in ctx.state['access_declarations'].values()
                          if 'kickside.connection:reply_provider' in declared_contracts(entry)}
        candidates = [kind for kind in KINDS if kind.binding in reply_bindings and
                      kind.binding in ACCESS_PROVIDER_CREDENTIALS]
        require(candidates, 'Automation fixture has no validated local reply-provider connection')
        if connection is None:
            connection = create(ctx, candidates[0], role)
            ctx.inst.cleanup.add('automation fixture connection',
                lambda: request(ctx, role, 'DELETE', '/connections/' + connection.id), scope='scenario')
        values['connection'] = connection.id
        if needs_channel:
            pending, seen, channels = [{}], set(), []
            while pending:
                arguments = pending.pop(0)
                resources = request(ctx, role, 'POST', '/connections/' + connection.id + '/discover_resources', arguments)['resources']
                for resource in resources:
                    if resource.get('selectable') is True:
                        channels.append(resource)
                    elif resource.get('drillable') is True and resource['id'] not in seen:
                        seen.add(resource['id'])
                        pending.append({'parent_id': resource['id']})
            require(channels, 'Automation fixture exposes no declared selectable channel')
            supplied['channel_id'] = channels[0]['id']
    result = {'title': title} if 'title' in fields else {}
    for name in required:
        field = fields[name]
        reference = field.get('x-reference', {}).get('kind')
        if reference:
            require(reference in values, 'Automation fixture has no owning flow for reference: ' + reference)
            result[name] = values[reference]
        elif 'default' in field:
            result[name] = field['default']
        elif workflow_target_schema(field):
            from scenarios_automation import dispatch_workflow
            def owned_request(method, path, body=None, **kwargs):
                return request(ctx, role, method, path.removeprefix('/api/v1'), body, **kwargs)
            wid, version = dispatch_workflow(ctx, title=title + ' target', request=owned_request)
            ctx.inst.cleanup.add('automation fixture workflow',
                lambda: owned_request('DELETE', '/api/v1/workflows/' + wid), scope='scenario')
            result[name] = {'kind': 'workflow', 'flow_ref': {'kind': 'flow', 'id': wid, 'version': version},
                            'input': {'marker': 'NW-ACCESS-731'}}
        else:
            require(name in supplied, 'Automation fixture has no declared input data for: ' + name)
            result[name] = supplied[name]
    return result

def stream_fixture(ctx, kind, declaration, title, role):
    setup_access_connection_fixtures(ctx)
    registry(ctx)
    specs = [(event, spec) for event, spec in SYSTEM_STREAM_FIXTURES.items() if event in output_events(declaration)]
    require(len(specs) == 1, 'System stream has no unique declared inbound fixture flow: ' + kind.binding)
    event_type, spec = specs[0]
    owning = [candidate for candidate in KINDS if ctx.state['access_declarations'].get(candidate.binding, {}).get('meta', {}).get('provider') == spec['provider']
              and 'kickside.connection:reply_provider' in declared_contracts(ctx.state['access_declarations'][candidate.binding])]
    require(len(owning) == 1, 'Inbound fixture has no unique reply-provider connection')
    gateway = DiscordGatewayFixture(ctx.state['access_connection_bases'][spec['provider']])
    gateway.new_identity()
    ctx.inst.redactor.add(gateway.token)
    connection = create(ctx, owning[0], role, connection_context={'bot_token': gateway.token})
    ctx.inst.cleanup.add('inbound fixture connection', lambda: delete(ctx, connection, role), scope='scenario')
    ctx.inst.wait(lambda: gateway.request('gateway')['connections'] == 1, 'Inbound fixture gateway ready')
    if spec['responder']:
        types = request(ctx, role, 'GET', '/automations/types')['types']
        responders = [row for row in types if {'connection_id', 'agent_id', 'channel_id'} <= set(row['inputs']['required'])]
        require(len(responders) == 1, 'Inbound fixture has no unique declared responder creation flow')
        payload = automation_input(ctx, responders[0]['inputs'], title, role, connection=connection)
        rid = request(ctx, role, 'POST', '/automations', {'type': responders[0]['id'], 'input': payload})['id']
        ctx.inst.cleanup.add('inbound fixture responder', lambda: request(ctx, role, 'DELETE', '/automations/' + rid), scope='scenario')
    before = {row['component_id'] for row in component_rows(ctx, kind.binding)}
    marker = str(uuid.uuid4().int)
    sender = str(uuid.uuid4().int)
    event = spec['event'](marker, 'Northwind inbound fixture', sender, gateway.bot_id)
    if spec['linked']:
        code = request(ctx, role, 'POST', '/identity/link-codes', {})['code']
        ctx.inst.redactor.add(code)
        gateway.send(spec['event'](str(uuid.uuid4().int), 'link ' + code, sender, gateway.bot_id))
        ctx.inst.wait(lambda: any(link['provider'] == spec['provider'] and link['external_id'] == sender
                      for link in request(ctx, role, 'GET', '/identity/links')['links']),
                      'Inbound fixture links its external identity')
        ctx.inst.cleanup.add('inbound fixture identity', lambda: request(ctx, role, 'POST', '/identity/unlink',
                            {'provider': spec['provider'], 'external_id': sender}), scope='scenario')
    gateway.send(event)
    def created():
        added = {row['component_id'] for row in component_rows(ctx, kind.binding)} - before
        require(len(added) <= 1, 'One inbound message creates multiple conversation components')
        return next(iter(added)) if added else None
    fid = ctx.inst.wait(created, 'Inbound message creates its declared conversation resource')
    # The component and its thread row commit together; the fixture is ready once
    # the creator reads the inbound message in the listed conversation.
    def delivered():
        events = request(ctx, role, 'GET', '/threads/' + fid + '/events?ascending=true&limit=200')['events']
        return events if marker in json.dumps(events) else None
    ctx.inst.wait(delivered, 'Inbound message lands in its declared conversation')
    return Fixture(kind, fid, title, {'generic': True, 'declaration': declaration, 'marker': marker,
                                     'event_type': event_type, 'identity_field': spec['identity_field']})


def create(ctx, kind, role='owner', *, connection_context=None):
    inst = ctx.inst
    install_package(ctx, kind.package)
    renew(ctx)
    registry(ctx)
    declaration = ctx.state['access_declarations'].get(kind.binding)
    if kind.binding:
        require(declaration, 'Installed kind is absent from the live class-bearing registry: ' + kind.binding)
    if declaration and kind.name.startswith('binding_') and ('kickside.contract:component' not in declared_contracts(declaration) or
            (not declaration['meta'].get('component') and output_events(declaration) and
             not all(isinstance(event, dict) and event.get('system_only') for event in
                     declaration['meta']['events']['output']))):
        install_package(ctx, 'kickside/automation')
        ports = request(ctx, role, 'GET', '/automations/sources')['sources'] + request(ctx, role, 'GET', '/automations/sinks')['sinks']
        setup_access_connection_fixtures(ctx)
        oauth_providers = []
        if declaration.get('meta', {}).get('provider') and not any(
                entry.get('meta', {}).get('provider') == declaration['meta']['provider'] and
                'kickside.contract:component' in declared_contracts(entry)
                for entry in ctx.state['access_declarations'].values()):
            providers = request(ctx, role, 'GET', '/oauth/providers/discover')['providers']
            for provider in providers:
                if provider['oauth_provider'] == declaration['meta']['provider']:
                    oauth_providers.append(request(ctx, role, 'GET', '/keeper/registry/entry?id=' + quote(provider['id'], safe=''))['entry'])
        owner, port = fixture_owner(declaration, ctx.state['access_declarations'].values(), ports, oauth_providers)
        owning = next((candidate for candidate in KINDS if candidate.binding == owner), None)
        require(owning, 'Declared resource owner has no native fixture flow: ' + owner)
        context = None
        if owning.name == 'connection':
            from scenarios_providers import PULL_CONNECTION_FIXTURES
            service = PULL_CONNECTION_FIXTURES[port['id']]
            context = {'site_url': ctx.state['access_connection_bases'][service], 'deployment': 'cloud',
                       'email': 'northwind@example.invalid', 'api_token': 'fictional-jira-token'}
        fixture = create(ctx, owning, role, connection_context=context)
        fixture.data['port'] = port
        if port['surface'] == 'store':
            writer = create(ctx, kind_named('data_sync'), role)
            fixture.data['writer_origin'] = writer.id
            inst.cleanup.add('store fixture writer', lambda: delete(ctx, writer, role), scope='scenario')
        inst.evidence.save('capability-owner', {'binding': kind.binding, 'owner': owner, 'port': port, 'component_id': fixture.id})
        return fixture
    if declaration and instancing(declaration) != 'demand':
        if instancing(declaration) == 'user':
            actor_browser(ctx, role)
        fixture = provisioned(ctx, kind, role)
        if kind.name == 'uploads':
            result = upload(ctx, role, 'northwind-access.txt', 'Northwind access NW-ACCESS-731')
            require(result['component_id'] == fixture.id, 'Upload lands in a different USER instance')
            fixture.data['upload'] = result['uuid']
        if kind.name == 'inbox':
            marker = 'NW-INBOX-READ-' + uuid.uuid4().hex
            start_seq = inbox_cursor(ctx, fixture, role)
            notice = request(ctx, role, 'POST', '/inbox', {'item_type': 'notice', 'title': marker, 'content': marker})['item']
            fixture.data.update(inbox_item=notice['id'], notice_marker=marker, inbox_from_seq=start_seq)
            events = inbox_event_window(ctx, fixture, role, start_seq + 1, inbox_cursor(ctx, fixture, role))
            require(notice['id'] in json.dumps(events) and marker in json.dumps(events),
                    'Inbox fixture notice does not persist in its declared USER instance')
        return fixture
    if kind.name == 'dashboard':
        value = request(ctx, role, 'GET', '/dashboard/overview')
        require(value.get('count') == len(value.get('surfaces', [])), 'Dashboard overview content differs from its declared count')
        raise Failure('Dashboard declares no shareable component kind or provision block; /dashboard/overview is a per-actor surface')
    title = 'Northwind access ' + kind.name + ' ' + uuid.uuid4().hex[:8]
    data = {'declaration': declaration}
    if connection_context is not None:
        result = request(ctx, role, 'POST', '/connections', {'name': title, 'impl_id': kind.binding,
                         'private_context': connection_context, 'validate': True})
        return Fixture(kind, result['component_id'], title, {'connection': True, 'declaration': declaration})
    if kind.name.startswith('binding_'):
        if (not declaration['meta'].get('component') and output_events(declaration) and
                all(isinstance(event, dict) and event.get('system_only') for event in declaration['meta']['events']['output'])):
            return stream_fixture(ctx, kind, declaration, title, role)
        if 'kickside.contract:installable' in declared_contracts(declaration):
            payload = automation_input(ctx, declaration['meta']['inputs'], title, role)
            fid = request(ctx, role, 'POST', '/automations', {'type': kind.binding, 'input': payload})['id']
            fixture_kind = Kind('automation', kind.package, kind.binding, '/automations/{id}/config')
            data = {'declaration': declaration, 'input': payload}
            targets = [payload[name] for name in declaration['meta']['inputs'].get('required', [])
                       if workflow_target_schema(declaration['meta']['inputs']['properties'][name])]
            require(len(targets) <= 1, 'Automation fixture declares multiple workflow task targets')
            if targets:
                data['workflow_target'] = targets[0]
            return Fixture(fixture_kind, fid, title, data)
        classes = declaration['meta']['class']
        if 'connection' in classes:
            if kind.binding == 'kickside.oauth:oauth_connection':
                fid = create_oauth_access_fixture(ctx, kind, title, role)
                return Fixture(kind, fid, title, {'oauth': True, 'declaration': declaration})
            if kind.binding in ACCESS_PROVIDER_CREDENTIALS or kind.package == 'kickside/simplefin':
                setup_access_connection_fixtures(ctx)
                context = dict(ACCESS_PROVIDER_CREDENTIALS.get(kind.binding, {}))
                if kind.package == 'kickside/simplefin':
                    context['setup_token'] = base64.b64encode((ctx.state['access_connection_bases']['simplefin'] +
                        '/simplefin/claim/northwind').encode()).decode()
                result = request(ctx, role, 'POST', '/connections', {'name': title, 'impl_id': kind.binding,
                                 'private_context': context, 'validate': True})
                f = Fixture(kind, result['component_id'], title, {'connection': True, 'declaration': declaration})
                if context.get('default_chat_id'):
                    f.data['default_chat_id'] = context['default_chat_id']
                inst.evidence.save('created-' + kind.name, {'id': f.id, 'binding': kind.binding, 'scope': 'demand'})
                return f
            fields = declaration['meta'].get('credential_schema', {}).get('fields', [])
            require(fields, 'Connection kind declares no credential schema fixture: ' + kind.binding)
            context = {}
            for field in fields:
                if not field.get('required'):
                    continue
                key, typ = field['key'], field['type']
                if typ == 'select':
                    first = field['options'][0]
                    context[key] = first['value'] if isinstance(first, dict) else first
                elif typ == 'url': context[key] = 'https://northwind.invalid'
                elif typ == 'email': context[key] = 'northwind@smoke.example'
                elif typ == 'number': context[key] = 1
                elif typ == 'boolean': context[key] = True
                else: context[key] = 'northwind-fictional-credential'
            result = request(ctx, role, 'POST', '/connections', {'name': title, 'impl_id': kind.binding,
                             'private_context': context, 'validate': False})
            f = Fixture(kind, result['component_id'], title, {'connection': True, 'declaration': declaration})
            inst.evidence.save('created-' + kind.name, {'id': f.id, 'binding': kind.binding, 'scope': 'demand'})
            return f
        raise Failure('Discovered class binding declares no user creation/primary-action adapter: ' + kind.binding +
                      '; class=' + str(classes) + '; provision=' + str(declaration['meta'].get('provision')))
    if kind.name == 'file':
        require('create' not in declaration['meta'].get('component', {}), 'File unexpectedly declares user creation')
        install_package(ctx, 'kickside/uploads')
        renew(ctx)
        parent = request(ctx, role, 'POST', '/workspace/entries', {'type': 'folder', 'name': title + ' uploads'})['entry']
        data['upload'] = upload(ctx, role, title + '.txt', title + ' NW-ACCESS-731', target_parent_id=parent['id'])['uuid']
        def filed():
            rows = request(ctx, role, 'GET', '/workspace/entries?parent_id=' + parent['id'])['entries']
            matches = [row for row in rows if row.get('type') == 'file' and row.get('meta', {}).get('upload_id') == data['upload']]
            require(len(matches) <= 1, 'Upload pipeline creates duplicate workspace files: ' + str(matches))
            return matches[0] if matches else None
        entry = inst.wait(filed, 'Upload pipeline materializes its workspace file', seconds=60)
        fid = entry['id']
        data['upload_parent'] = parent['id']
    elif kind.name in ('folder', 'graph', 'journal', 'roadmap', 'database'):
        body = {'type': kind.name, 'name': title}
        result = request(ctx, role, 'POST', '/workspace/entries', body)
        fid = result['entry']['id']
    elif kind.name == 'agent':
        result = request(ctx, role, 'POST', '/user-agents', {'title': title, 'model': ctx.state['model'],
                         'prompt': 'Use fictional Northwind facts. Follow the requested marker.', 'traits': []})
        fid = result['component_id']
    elif kind.name.startswith('kb_'):
        result = request(ctx, role, 'POST', '/knowledge/create', {'name': title,
                          'kb_type': 'wiki' if kind.name == 'kb_wiki' else 'document',
                          'embedding_model': ctx.state['embed']})
        fid = result['id']
        request(ctx, role, 'PUT', '/knowledge/' + fid, {'config': {'analysis_model': ctx.state['model']}})
        request(ctx, role, 'POST', '/knowledge/' + fid + '/ingest', {'title': title,
                 'content': 'Fictional Northwind access marker NW-ACCESS-731. The return window is 17 days.',
                 'external_id': title, 'external_version': '1'})
        inst.wait(lambda: 'NW-ACCESS-731' in json.dumps(request(ctx, role, 'POST', '/knowledge/' + fid + '/search',
                  {'query': 'NW-ACCESS-731', 'limit': 5})), 'KB access fixture index', seconds=120)
    elif kind.name == 'workflow':
        install_package(ctx, 'kickside/code-execution')
        renew(ctx)
        fid = request(ctx, role, 'POST', '/workflows', {'title': title})['flow_id']
        schema = {'type': 'object', 'additionalProperties': True}
        commands = [
            {'type': 'node.add', 'body': {'id': 'start', 'kind': 'start', 'config': {'input_schema': schema}}},
            {'type': 'node.add', 'body': {'id': 'echo', 'kind': 'block', 'ref': 'kickside.code_execution:block.run_code',
                                       'config': {'code': 'return __cx_args', 'tool_ids': []}}},
            {'type': 'node.add', 'body': {'id': 'finish', 'kind': 'finish', 'config': {'output_schema': schema}}},
            {'type': 'edge.add', 'body': {'id': 'start-echo', 'source_id': 'start', 'target_id': 'echo',
                                       'source_port': 'out', 'target_port': 'in'}},
            {'type': 'edge.add', 'body': {'id': 'echo-finish', 'source_id': 'echo', 'target_id': 'finish',
                                       'source_port': 'result', 'target_port': 'in', 'transform': 'output.result'}}]
        request(ctx, role, 'POST', '/workflows/' + fid + '/apply', {'commands': commands})
        data['version'] = request(ctx, role, 'POST', '/workflows/' + fid + '/publish', {})['version']
    elif kind.name in ('recurring', 'data_sync'):
        data['notice'] = title + ' delivered'
        if kind.name == 'recurring':
            payload = {'title': title, 'schedule': {'type': 'interval', 'expression': '2s'}, 'enabled': False,
                       'target': {'kind': 'action', 'port': 'kickside.inbox.sink:inbox_item', 'config': {},
                                  'input': {'title': data['notice'], 'content': 'NW-ACCESS-731'}}}
        else:
            install_package(ctx, 'kickside/webhooks')
            renew(ctx)
            wh = request(ctx, role, 'POST', '/webhooks', {'label': title + ' source', 'schema': {'type': 'object'}})['webhook']
            data['source'] = wh
            inst.redactor.add(wh['secret'])
            payload = {'title': title, 'enabled': True,
                       'trigger': {'v': 2, 'source': 'kickside.webhooks:webhook_received',
                                   'config': {'thread_id': wh['webhook_id']}, 'backfill': 'none'},
                       'sink_port': 'kickside.inbox.sink:inbox_item', 'sink_config': {},
                       'transform': {'kind': 'expr', 'expr': json.dumps({'title': data['notice'], 'content': 'NW-ACCESS-731'})}}
        data['input'] = payload
        fid = request(ctx, role, 'POST', '/automations', {'type': kind.binding, 'input': payload})['id']
    elif kind.name == 'connection':
        if not ctx.state.get('access_jira'):
            ctx.state['access_jira'] = ctx.fixtures.start('jira')
            inst.cleanup.add('access Jira fixture', lambda: ctx.fixtures.stop('jira'))
        result = request(ctx, role, 'POST', '/connections', {'name': title, 'impl_id': kind.binding,
                          'private_context': {'site_url': ctx.state['access_jira'], 'deployment': 'cloud',
                                              'email': 'northwind@example.invalid', 'api_token': 'fictional-jira-token'},
                          'validate': True})
        fid = result['component_id']
    elif kind.name == 'remote_mcp':
        kb = create(ctx, KINDS[8], role)
        credential = request(ctx, role, 'POST', '/mcp/credentials', {'label': title,
                'whitelist_component_ids': [kb.id], 'allowed_tool_ids': ['kickside.knowledge.traits:kb_search_tool']})['credential']
        inst.redactor.add(credential['token'])
        data['credential'], data['kb'] = credential, kb.id
        result = request(ctx, role, 'POST', '/connections', {'name': title, 'impl_id': kind.binding,
                'private_context': {'endpoint_url': inst.base + '/api/public/mcp', 'auth_type': 'bearer',
                                    'bearer_token': credential['token']}, 'validate': True})
        fid = result['component_id']
    elif kind.name == 'mcp_credential':
        kb = create(ctx, KINDS[8], role)
        data['kb'] = kb.id
        credential = request(ctx, role, 'POST', '/mcp/credentials', {'label': title,
                'whitelist_component_ids': [kb.id], 'allowed_tool_ids': ['kickside.knowledge.traits:kb_search_tool']})['credential']
        data['credential'] = credential
        inst.redactor.add(credential['token'])
        fid = credential['credential_id']
    elif kind.name == 'uploads':
        result = upload(ctx, role, title + '.txt', title + ' NW-ACCESS-731')
        fid, data['upload'] = result['component_id'], result['uuid']
    elif kind.name == 'crm':
        fid = request(ctx, role, 'POST', '/workspace/crm', {'title': title, 'idempotency_key': uuid.uuid4().hex}, expect=(201,))['component_id']
        revision = request(ctx, role, 'GET', '/workspace/crm/' + fid + '/schema')['revision']
        declaration = {'version': '5', 'objects': [{'type': 'person', 'label': 'People'}],
                       'attributes': [{'object_type': 'person', 'attr': 'name', 'label': 'Name', 'data_type': 'text'}],
                       'pipelines': [], 'views': [], 'merge_config': {'objects': {}}}
        request(ctx, role, 'POST', '/workspace/crm/' + fid + '/schema', {'declaration': declaration, 'expected_revision': revision})
    elif kind.name == 'research':
        fid = request(ctx, role, 'POST', '/research', {'prompt': 'Add fictional Northwind 47 crates and 12 reserve crates. '
                      'Use only these facts. End with NW-ACCESS-731.', 'agent_id': 'kickside.research.agents:researcher'})['research_id']
        data['component_id'] = request(ctx, role, 'GET', '/research/' + fid + '/events')['run']['thread_id']
    elif kind.name == 'skill':
        fid = request(ctx, role, 'POST', '/skills', {'name': title.lower().replace(' ', '-'), 'description': title,
                       'body': 'Fictional Northwind skill NW-ACCESS-731', 'status': 'active'})['skill_id']
    elif kind.name == 'webhook':
        wh = request(ctx, role, 'POST', '/webhooks', {'label': title, 'schema': {'type': 'object',
                     'required': ['marker'], 'properties': {'marker': {'type': 'string'}}}})['webhook']
        fid, data['secret'] = wh['webhook_id'], wh['secret']
        inst.redactor.add(wh['secret'])
    else:
        raise Failure('No domain creation adapter for discovered kind: ' + kind.name)
    require(fid, 'Creation returns no resource ID')
    fixture = Fixture(kind, fid, title, data)
    inst.evidence.save('created-' + kind.name, {'id': fid, 'title': title, 'binding': kind.binding})
    return fixture


def read(ctx, f, role='owner', *, hidden=False):
    name = f.kind.name
    if f.data.get('connection'): name = 'connection'
    if name == 'automation' and f.data['declaration']['meta'].get('reconfigurable') is not True:
        rows = request(ctx, role, 'GET', '/automations')['automations']
        matches = [row for row in rows if row['id'] == f.id]
        require(not matches if hidden else len(matches) == 1, 'Automation catalog violates granted visibility')
        result = matches[0] if matches else {'success': False, 'error': 'resource absent from list'}
    elif f.data.get('oauth'):
        result = request(ctx, role, 'GET', '/oauth/connections/' + f.id, denied=hidden)
    elif f.kind.read_route:
        domain_denial = hidden and name in ('graph', 'database', 'roadmap')
        result = request(ctx, role, 'GET', f.kind.read_route.format(id=f.id),
                         hidden=hidden and not domain_denial, denied=domain_denial,
                         expect=(403, 404) if hidden and name in ('crm', 'workflow') else None)
        if domain_denial:
            require(not any(key in result for key in ('graph', 'schema', 'roadmap')),
                    'Denied Chestor read exposes domain content: ' + str(result))
    elif name in ('connection', 'remote_mcp', 'webhook', 'mcp_credential'):
        if name == 'remote_mcp': name = 'connection'
        key, route, id_key = {'connection': ('connections', '/connections', 'component_id'),
                'webhook': ('webhooks', '/webhooks', 'webhook_id'),
                'mcp_credential': ('credentials', '/mcp/credentials', 'credential_id')}[name]
        rows = request(ctx, role, 'GET', route)[key]
        matches = [r for r in rows if r[id_key] == f.id]
        require(not matches if hidden else len(matches) == 1,
                'Hidden resource leaks in list' if hidden else 'Granted resource is absent from list: ' + str(rows))
        result = matches[0] if matches else {'success': False, 'error': 'resource absent from list'}
    elif name == 'uploads' and 'upload' in f.data:
        result = request(ctx, role, 'GET', '/uploads/' + f.data['upload'], hidden=hidden)
    elif name in ('inbox', 'lifecycle', 'uploads') or f.data.get('generic'):
        result = request(ctx, role, 'GET', '/threads/' + f.id, hidden=hidden)
    else:
        result = workspace(ctx, f, role, hidden)
    if not hidden:
        require(f.id in json.dumps(result) or name in ('graph', 'journal', 'roadmap', 'database', 'crm'),
                'Read returns no observable resource identity: ' + str(result))
        require(bool(result), 'Read returns no content')
    return result


def mutation_snapshot(ctx, f, baseline=None):
    from data_contracts import event_window
    value = read(ctx, f)
    resource = value.get('entry', value)
    authored = {key: resource[key] for key in ('id', 'component_id', 'name', 'title', 'description', 'config',
                'content', 'body', 'label') if key in resource}
    declaration = f.data.get('declaration')
    if declaration and any(action['name'] == 'read_config' for action in declaration['meta'].get('actions', [])):
        configuration = request(ctx, 'owner', 'POST', '/automations/' + f.id + '/actions/read_config', {})['result']
        return {'authored': authored, 'parent_id': value.get('parent_id'), 'config': configuration['state']}
    if not f.kind.binding:
        return value
    if declaration is not None and not (declaration.get('meta', {}).get('thread_class') or
            any(row.get('methods', {}).get('get_status') == 'kickside.core.threads:thread_status'
                for row in declaration.get('data', {}).get('contracts', []))):
        return {'authored': authored, 'parent_id': resource.get('parent_id')}
    base = '/threads/' + component_id(f)
    source = request(ctx, 'owner', 'GET', base, expect=(200, 404))
    if 'thread' not in source:
        require(source.get('success') is False and source.get('error') == 'thread not found',
                'Canonical snapshot returns no thread or declared absence')
        return {'authored': authored, 'thread': None, 'events': []}
    thread = source['thread']
    if baseline is None:
        tail = request(ctx, 'owner', 'GET', base + '/events?ascending=false&limit=1')['events']
        to_seq = max((row['seq'] for row in tail), default=0)
    else:
        to_seq = max((row['seq'] for row in baseline['events']), default=0)
    events = event_window(lambda *args, **kwargs: request(ctx, 'owner', *args, **kwargs),
                          base + '/events', 1, to_seq)
    return {'authored': authored, 'thread': {key: thread.get(key) for key in ('id', 'kind', 'attrs', 'parent_id', 'state', 'closed_at')},
            'events': events}


def update(ctx, f, role, *, denied=False):
    marker = f.title + ' edited by ' + role
    name = f.kind.name
    if f.data.get('connection'): name = 'connection'
    if name in ('connection', 'remote_mcp') or f.data.get('oauth'):
        marker += ' ' + uuid.uuid4().hex[:8]
    if f.data.get('oauth'): route, method, body = '/oauth/connections/' + f.id, 'PUT', {'connection_name': marker}
    elif name == 'agent': route, method, body = '/user-agents/' + f.id, 'PUT', {'title': marker}
    elif name.startswith('kb_'): route, method, body = '/knowledge/' + f.id, 'PUT', {'name': marker}
    elif name in ('connection', 'remote_mcp'): route, method, body = '/connections/' + f.id, 'PUT', {'name': marker}
    elif name in ('recurring', 'data_sync', 'automation'): route, method, body = '/automations/' + f.id, 'PATCH', {'title': marker}
    elif name == 'skill': route, method, body = '/skills/' + f.id, 'PATCH', {'body': marker + ' NW-ACCESS-731'}
    elif name == 'webhook': route, method, body = '/webhooks/' + f.id, 'PATCH', {'label': marker}
    elif name == 'mcp_credential': route, method, body = '/mcp/credentials/' + f.id, 'PATCH', {'label': marker}
    elif name in ('inbox', 'lifecycle', 'uploads', 'research'):
        route, method, body = '/threads/' + component_id(f), 'PATCH', {'attrs': {'title': marker}}
    else: route, method, body = '/workspace/entries/' + f.id, 'PUT', {'name': marker}
    before = mutation_snapshot(ctx, f) if denied else None
    result = request(ctx, role, method, route, body, denied=denied)
    if denied:
        require(mutation_snapshot(ctx, f, before) == before,
                'Denied write changes observable owner content')
    else:
        ctx.inst.wait(lambda: marker in json.dumps(request(ctx, role, 'GET', '/threads/' + component_id(f))
                         if name in ('inbox', 'lifecycle', 'uploads', 'research') else read(ctx, f)) or marker in json.dumps(workspace(ctx, f)),
                      'Authorized update persists', seconds=35)
    return result


def rpc(ctx, credential, method, params):
    result = ctx.inst.request('POST', '/api/public/mcp', {'jsonrpc': '2.0', 'id': uuid.uuid4().hex,
                'method': method, 'params': params}, auth=credential['token'], timeout=120,
                headers={'Accept': 'application/json, text/event-stream'}, label='mcp-' + method.replace('/', '-'))
    require('error' not in result and not result.get('result', {}).get('isError'), 'MCP action fails: ' + str(result))
    return result['result']


def inbox_cursor(ctx, fixture, role='owner'):
    value = request(ctx, role, 'GET', '/threads/' + component_id(fixture) + '/events?ascending=false&limit=1')
    return max((event['seq'] for event in value['events']), default=0)


def inbox_event_window(ctx, fixture, role, from_seq, to_seq):
    from data_contracts import event_window
    return event_window(lambda *args, **kwargs: request(ctx, role, *args, **kwargs),
                        '/threads/' + component_id(fixture) + '/events', from_seq, to_seq)


def actor_browser(ctx, role):
    Browser.release(ctx)
    browser = ctx.browser_check()
    ctx.state['access_browser_error_start'] = len(browser.errors)
    browser.settle()
    if role == 'owner':
        browser.login()
    else:
        browser.login(ctx.state['access_users'][role]['email'], ctx.inst.password)
    browser.frame.locator('.sidebar').wait_for(timeout=30000)
    browser.settle()
    splash = browser.frame.locator('kickside-agent-onboarding')
    if splash.count() and splash.is_visible():
        splash.get_by_label('Agent name', exact=True).fill('Northwind private ' + role + ' assistant')
        with browser.page.expect_response(lambda r: '/api/v1/user-agents' in r.url and r.request.method == 'POST') as response:
            splash.get_by_role('button', name='Create agent', exact=True).click()
        require(response.value.status == 200 and response.value.json().get('component_id'), 'Member first-agent creation fails')
        browser.frame.locator('wippy-chat').wait_for(timeout=30000)
        browser.settle()
    return browser


def chat(ctx, agent, role, prompt='Give twenty numbered fictional Northwind office tasks. Write two full sentences for each task. End with NW-ACCESS-731.', marker='NW-ACCESS-731'):
    """Start the granted agent through its card and assert persisted assistant content."""
    detail = request(ctx, role, 'GET', '/user-agents/' + agent)['agent']
    browser = actor_browser(ctx, role)
    browser.frame.locator('.sidebar').get_by_text('Agents', exact=True).click()
    panel = browser.frame.locator('kickside-agents')
    panel.wait_for(timeout=30000)
    before = request(ctx, role, 'GET', '/sessions?limit=100')['sessions']
    panel.locator('.kx-card').filter(has_text=detail['title']).get_by_role('button', name='Chat', exact=True).click()
    browser.frame.locator('wippy-chat').get_by_placeholder('Type a message', exact=True).wait_for(timeout=30000)
    browser.stream(prompt, marker, 'granted-agent-chat-' + role)
    after = request(ctx, role, 'GET', '/sessions?limit=100')['sessions']
    fresh = [row for row in after if row['session_id'] not in {r['session_id'] for r in before}]
    require(len(fresh) == 1, 'Agent card chat creates no unique durable session: ' + str(after))
    sid = fresh[0]['session_id']
    messages = request(ctx, role, 'GET', '/sessions/messages?session_id=' + sid)['messages']
    require(any(m['type'] == 'assistant' and marker in str(m.get('data')) for m in messages),
            'Granted agent chat lacks persisted assistant marker: ' + str(messages))
    require(sum(m['type'] == 'user' for m in messages) == 1, 'Agent chat duplicates/lacks its user event')
    browser.clean(ctx.state['access_browser_error_start'])
    ctx.state['access_last_chat'] = {'browser': browser, 'session': sid, 'messages': messages}
    return messages


def completed_workflow(ctx, wid, run_id, role):
    def completed():
        value = request(ctx, role, 'GET', '/workflows/' + wid + '/runs/' + run_id)
        require(value['run']['status'] != 'failed', 'Granted workflow run fails: ' + str(value))
        return value if value['run']['status'] == 'completed' and workflow_publication_ready(ctx.inst, wid) else None
    value = ctx.inst.wait(completed, 'Granted workflow run completes', seconds=60)
    require('NW-ACCESS-731' in json.dumps(value['run']['output']), 'Granted workflow loses input/output marker')
    return value


def workflow_task_action(ctx, fixture, role, *, denied=False, readonly=False):
    if readonly and not denied:
        return read(ctx, fixture, role)
    wid = fixture.data['workflow_target']['flow_ref']['id']
    before = {row['run_id'] for row in request(ctx, 'owner', 'GET', '/workflows/' + wid + '/runs')['runs']}
    snapshot = mutation_snapshot(ctx, fixture) if denied else None
    result = request(ctx, role, 'POST', '/automations/' + fixture.id + '/control',
        {'control': 'run_now', 'request_id': uuid.uuid4().hex}, denied=denied)
    def fresh_runs():
        return [row for row in request(ctx, 'owner', 'GET', '/workflows/' + wid + '/runs')['runs']
                if row['run_id'] not in before]
    if denied:
        require(mutation_snapshot(ctx, fixture) == snapshot and not fresh_runs(),
                'Denied workflow task control changes its configuration or creates a workflow run')
        return result
    def launched():
        rows = fresh_runs()
        require(len(rows) <= 1, 'One admitted workflow task control creates multiple runs')
        return rows[0] if rows else None
    run = ctx.inst.wait(launched, 'Granted workflow task starts its published target', seconds=60)
    return completed_workflow(ctx, wid, run['run_id'], 'owner')


def primary(ctx, f, role, *, denied=False, readonly=False):
    port = f.data.get('port')
    if port and port['surface'] in ('collection', 'store'):
        from scenarios_providers import PULL_FIXTURES
        writable = port['surface'] == 'store'
        if writable:
            marker = uuid.uuid4().hex
            config = {key: f.id for key, field in port['config_schema'].items() if field.get('class')}
            require(config, 'Store fixture declares no owning resource reference')
            arguments = {'origin': f.data['writer_origin'], 'sink_op': 'upsert', 'identity_key': marker, 'idempotency_key': marker,
                         'config': config, 'input': {**STORE_FIXTURES[port['id']], 'record_id': marker}}
            denied = denied or readonly
            before = mutation_snapshot(ctx, f) if denied else None
        else:
            config, marker = PULL_FIXTURES[port['id']]
            arguments = {'limit': 100, 'config': config}
        credential = request(ctx, role, 'POST', '/mcp/credentials', {
            'label': 'Northwind capability access', **({'whitelist_component_ids': [f.id]} if not denied else {}),
            'allowed_tool_ids': ['app.smoke.access_connections:probe']})['credential']
        ctx.inst.redactor.add(credential['token'])
        cleanup = ctx.inst.cleanup.add('capability credential', lambda: request(ctx, role, 'DELETE',
                    '/mcp/credentials/' + credential['credential_id']), scope='scenario')
        from mcp_client import MCPClient
        client = MCPClient(ctx.inst, credential['token'], '/api/public/mcp', stateful=False)
        result = client.request('tools/call', {'name': 'NorthwindAccessConnectionProbe', 'arguments': {
                    'connection_id': f.id, 'port': port['id'], 'write': writable, 'arguments': arguments}},
                    'capability-port-' + role)
        failed = 'error' in result or result.get('result', {}).get('isError')
        if denied:
            require(failed and 'access' in json.dumps(result).lower(), 'Denied capability action admits the foreign resource')
            if writable:
                require(mutation_snapshot(ctx, f) == before, 'Denied store write changes the destination')
        else:
            require(not failed, 'Granted capability read fails: ' + str(result))
            envelope = client.decode_text(client.text(result['result']))
            value = envelope['value']
            if writable:
                require(value.get('success') is True and value.get('dest_ref', {}).get('record_id') == marker,
                        'Granted store write loses its declared destination reference: ' + str(value))
                ctx.inst.wait(lambda: 'Northwind sink access' in json.dumps(request(ctx, role, 'POST',
                    '/workspace/crm/' + f.id + '/query', {'kind': 'record', 'record_id': marker})),
                    'Granted store write appears in its owning resource', seconds=60)
            else:
                require(value.get('items') and marker in json.dumps(value['items']),
                        'Granted capability read loses its declared fixture rows: ' + str(value))
        cleanup()
        ctx.inst.cleanup.remove(cleanup)
        return result
    name = f.kind.name
    if name == 'file':
        path = '/api/v1/uploads/' + f.data['upload'] + '/raw'
        if denied:
            return request(ctx, role, 'GET', path.removeprefix('/api/v1'), denied=True)
        fetch = urlrequest.Request(ctx.inst.base + path, headers={'Authorization': 'Bearer ' + identity(ctx, role)})
        with urlrequest.urlopen(fetch, timeout=35) as response:
            require(response.status == 200, 'Workspace file viewer returns no successful raw response')
            content = response.read().decode('utf-8')
        ctx.inst.evidence.save('workspace-file-use-' + role, {'path': path, 'content': content})
        require('NW-ACCESS-731' in content, 'Workspace file viewer omits the original uploaded content')
        return content
    if f.data.get('connection'): name = 'connection'
    if f.data.get('oauth'):
        if denied:
            return read(ctx, f, role, hidden=True)
        credential = request(ctx, role, 'POST', '/mcp/credentials', {
            'label': 'Northwind OAuth access probe', 'whitelist_component_ids': [f.id],
            'allowed_tool_ids': ['app.smoke.access_connections:probe']})['credential']
        ctx.inst.redactor.add(credential['token'])
        ctx.inst.cleanup.add('OAuth access probe credential', lambda: request(
            ctx, role, 'DELETE', '/mcp/credentials/' + credential['credential_id']))
        result = rpc(ctx, credential, 'tools/call', {'name': 'NorthwindAccessConnectionProbe',
                     'arguments': {'connection_id': f.id, 'tool': 'kickside.google.traits:gmail_read_tool',
                                   'arguments': {'action': 'get_message', 'message_id': '731'}}})
        require('Northwind' in json.dumps(result), 'Granted OAuth reader cannot use its fixture-backed provider')
        return result
    if name == 'automation':
        if f.data.get('workflow_target'):
            return workflow_task_action(ctx, f, role, denied=denied, readonly=readonly)
        if denied:
            before = mutation_snapshot(ctx, f)
            result = request(ctx, role, 'POST', '/automations/' + f.id + '/control', {'control': 'pause'}, denied=True)
            require(mutation_snapshot(ctx, f) == before, 'Denied automation control changes its durable configuration')
            return result
        value = read(ctx, f, role)
        require(f.data['input']['connection_id'] in json.dumps(value) and f.data['input']['channel_id'] in json.dumps(value),
                'Granted automation status loses its owning connection and channel')
        return value
    if name == 'agent':
        if denied:
            return request(ctx, role, 'PUT', '/user-agents/' + f.id, {'prompt': 'unauthorized'}, denied=True)
        return chat(ctx, f.id, role)
    if name.startswith('kb_'):
        route, method, body = '/knowledge/' + f.id + ('/ingest' if denied else '/ask'), 'POST', (
            {'title': 'Denied access marker', 'content': 'NW-DENIED-731', 'external_id': 'denied', 'external_version': '1'}
            if denied else {'query': 'What is the Northwind return window and access marker? Cite the source.'})
        result = request(ctx, role, method, route, body, denied=denied, timeout=120)
        if not denied:
            require('17' in str(result.get('answer')) and 'NW-ACCESS-731' in str(result.get('answer')),
                    'KB use does not return grounded facts: ' + str(result))
            sources = result.get('results', [])
            references = [int(index) for index in re.findall(r'\[(\d+)\]', result['answer'])]
            require(sources and references and all(1 <= index <= len(sources) for index in references),
                    'KB read-only use returns no valid numbered source citations: ' + str(result))
            require(any('17' in sources[index - 1]['content'] and 'NW-ACCESS-731' in sources[index - 1]['content']
                        for index in references), 'KB citation does not identify the returned grounded facts: ' + str(result))
        return result
    if name == 'workflow':
        if readonly: return read(ctx, f, role)
        run = request(ctx, role, 'POST', '/workflows/' + f.id + '/run',
                      {'version': f.data['version'], 'input': {'marker': 'NW-ACCESS-731'}}, denied=denied)
        if denied: return run
        return completed_workflow(ctx, f.id, run['run_id'], role)
    if name in ('recurring', 'data_sync'):
        if readonly: return read(ctx, f, role)
        previous = {i['id'] for i in request(ctx, 'owner', 'GET', '/inbox?status=all&limit=200')['items']}
        result = request(ctx, role, 'POST', '/automations/' + f.id + '/control',
                         {'control': 'resume', 'request_id': uuid.uuid4().hex}, denied=denied)
        if denied: return result
        if name == 'data_sync':
            source = f.data['source']
            ctx.inst.request('POST', '/api/public/hooks/' + source['webhook_id'], {'marker': 'NW-ACCESS-731'},
                             auth=source['secret'], expect=(202,), label='sync-trigger')
        def notices():
            items = request(ctx, 'owner', 'GET', '/inbox?status=all&limit=200')['items']
            fresh = [i for i in items if i['id'] not in previous and i.get('title') == f.data['notice'] and 'NW-ACCESS-731' in json.dumps(i)]
            return fresh or None
        delivered = ctx.inst.wait(notices, 'Granted automation produces an observable inbox event', seconds=60)
        request(ctx, role, 'POST', '/automations/' + f.id + '/control', {'control': 'pause', 'request_id': uuid.uuid4().hex})
        return delivered
    if name == 'research':
        if denied:
            return update(ctx, f, role, denied=True)
        def completed():
            value = read(ctx, f, role)
            require(value['run'].get('status') != 'failed', 'Research run fails: ' + str(value))
            return value if value['run'].get('status') == 'complete' else None
        value = ctx.inst.wait(completed, 'Research primary action completes', seconds=180)
        results = [e['body'].get('markdown') for e in value['events']
                   if e['type'] == 'kickside.research.events:result']
        report = value['run'].get('result_md')
        require(isinstance(report, str) and report in results,
                'Research projected report differs from its durable result event')
        require(all(fact in report.lower() for fact in ('northwind', '47', '12')),
                'Research report lacks the supplied Northwind facts')
        return value
    if name == 'mcp_credential':
        if not denied: read(ctx, f, role)
        if denied:
            return request(ctx, role, 'PATCH', '/mcp/credentials/' + f.id, {'label': 'NW-DENIED-731'}, denied=True)
        result = rpc(ctx, f.data['credential'], 'tools/call', {'name': 'SearchKB', 'arguments': {'fields': ['content', 'meta'], 'kb_id': f.data['kb'], 'query': 'NW-ACCESS-731'}})
        require('NW-ACCESS-731' in json.dumps(result), 'MCP tool returns no grounded content')
        return result
    if name == 'webhook':
        if readonly: return read(ctx, f, role)
        if denied:
            return request(ctx, role, 'PATCH', '/webhooks/' + f.id, {'label': 'NW-DENIED-731'}, denied=True)
        ctx.inst.request('POST', '/api/public/hooks/' + f.id, {'marker': 'NW-ACCESS-731'}, auth=f.data['secret'],
                         expect=(202,), label='webhook-inbound')
        events = request(ctx, role, 'GET', '/threads/' + component_id(f) + '/events')
        require('NW-ACCESS-731' in json.dumps(events), 'Webhook accepts payload without a durable inbound event')
        return events
    if name in ('connection', 'remote_mcp'):
        if denied: return update(ctx, f, role, denied=True)
        result = request(ctx, role, 'POST', '/connections/' + f.id + '/discover_resources', {})
        require(result.get('success') is True and isinstance(result.get('items'), list)
                and isinstance(result.get('resources'), list), 'Connection discovery lacks its declared lists: ' + str(result))
        if f.kind.binding == 'kickside.telegram.connection:telegram_connection' and f.data.get('default_chat_id'):
            require(result.get('success') is True and any(row.get('id') == f.data['default_chat_id']
                    for row in result.get('resources', [])), 'Configured Telegram chat is absent from discovery: ' + str(result))
        elif f.kind.binding == 'kickside.telegram.connection:telegram_connection':
            require(result.get('success') is True and result.get('resources') == [] and result.get('items') == [],
                    'Telegram discovery without a default_chat_id differs from its declared empty result: ' + str(result))
        elif f.kind.name == 'connection':
            require(result['items'] or result['resources'], 'The Jira fixture exposes no declared resources: ' + str(result))
        return result
    if name == 'uploads':
        if denied: return request(ctx, role, 'DELETE', '/uploads/' + f.data['upload'], denied=True)
        result = request(ctx, role, 'GET', '/uploads/' + f.data['upload'] + '/artifacts')
        require((result.get('primary') or {}).get('raw_url'), 'Upload use returns no primary file artifact')
        return result
    if name in ('inbox', 'lifecycle') or f.data.get('generic'):
        if readonly:
            if name == 'inbox':
                events = inbox_event_window(ctx, f, role, f.data['inbox_from_seq'] + 1, inbox_cursor(ctx, f, role))
                require(f.data['inbox_item'] in json.dumps(events) and f.data['notice_marker'] in json.dumps(events),
                        'READ grantee cannot retrieve its shared Inbox notice content')
                return events
            return read(ctx, f, role)
        if name == 'inbox':
            marker = 'NW-ACCESS-731-' + uuid.uuid4().hex[:8]
            item_id = str(uuid.uuid4())
            start_seq = inbox_cursor(ctx, f)
            result = request(ctx, role, 'POST', '/threads/' + component_id(f) + '/events',
                {'type': 'kickside.inbox.events:item.created', 'body': {
                    'id': item_id, 'item_type': 'notice', 'title': marker, 'content': marker,
                    'actor_id': actor_id(ctx, 'owner'), 'created_at': time.time()}}, denied=denied,
                expect=(403, 404) if denied else (201,))
            after = inbox_event_window(ctx, f, 'owner', start_seq + 1, inbox_cursor(ctx, f))
            if denied:
                require(item_id not in json.dumps(after) and marker not in json.dumps(after),
                        'Denied Inbox creation appends its attempted notice to the shared target log')
            else:
                require(item_id in json.dumps(after) and marker in json.dumps(after),
                        'Granted Inbox writer creates no item in the shared target log')
                ctx.inst.wait(lambda: value if item_id in json.dumps(value := request(ctx, 'owner', 'GET',
                              '/inbox?status=all&limit=200')) and marker in json.dumps(value) else None,
                              'Shared Inbox item reaches the owner projection', seconds=35)
            return result
        if name == 'lifecycle':
            marker = 'northwind-absent-' + uuid.uuid4().hex + '@smoke.example'
            before = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
            result = request(ctx, role, 'POST', '/threads/' + component_id(f) + '/events',
                    {'type': 'kickside.lifecycle.events:principal.deleted', 'body': {'principal_id': marker}}, denied=True,
                    expect=(403,))
            after = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
            require(after == before, 'Public lifecycle forgery changes the framework event log')
            return result
        if f.data.get('generic'):
            value = read(ctx, f, role, hidden=denied)
            if not denied:
                events = inbox_event_window(ctx, f, role, 1, inbox_cursor(ctx, f, role))
                require(f.data['marker'] in json.dumps(events), 'Declared conversation read loses its inbound fixture event')
            return value
        raise Failure('Kind has no declared authenticated primary-action adapter: ' + f.kind.binding)
    if name == 'skill':
        if denied: return update(ctx, f, role, denied=True)
        result = read(ctx, f, role)
        require('NW-ACCESS-731' in json.dumps(result), 'Shared skill use lacks its instruction body')
        return result
    if readonly:
        return read(ctx, f, role)
    marker = role + '-' + uuid.uuid4().hex[:8]
    if name == 'graph':
        route = '/workspace/graphs/' + f.id + '/apply'
        body = {'commands': [{'type': 'chestor.graph.events:node.added', 'body': {'id': marker, 'name': marker, 'type': 'concept'}}]}
    elif name == 'database':
        route = '/workspace/db-schema/' + f.id + '/apply'
        body = {'commands': [{'type': 'chestor.db_schema.events:table.added', 'body': {'id': marker, 'name': marker}}]}
    elif name == 'roadmap':
        route = '/workspace/roadmaps/' + f.id + '/apply'
        body = {'commands': [{'type': 'chestor.roadmap.events:layer.added', 'body': {'id': marker, 'title': marker, 'sort_order': 0}}]}
    elif name == 'journal':
        route = '/workspace/journals/' + f.id + '/events'
        body = {'event_type': 'decision', 'title': marker, 'body': 'Northwind NW-ACCESS-731', 'tags': ['northwind']}
    elif name == 'crm':
        route = '/workspace/crm/' + f.id + '/apply'
        body = {'commands': [{'type': 'kickside.crm.events:record.created', 'body': {'record_id': marker,
                 'object_type': 'person', 'values': {'name': marker}}}]}
    elif name == 'folder':
        route = '/workspace/entries'
        body = {'type': 'folder', 'name': marker, 'parent_id': f.id}
    else:
        return update(ctx, f, role, denied=denied)
    before = mutation_snapshot(ctx, f) if denied else None
    result = request(ctx, role, 'POST', route, body, denied=denied)
    if denied:
        require(mutation_snapshot(ctx, f) == before, 'Denied event changes the owner content')
    else:
        def observable():
            if name == 'folder': value = request(ctx, role, 'GET', '/workspace/entries?parent_id=' + f.id)
            elif name == 'crm': value = request(ctx, role, 'POST', '/workspace/crm/' + f.id + '/query', {'kind': 'record', 'record_id': marker})
            else: value = read(ctx, f, role)
            return value if marker in json.dumps(value) else None
        ctx.inst.wait(observable, 'Granted event appears in domain projection', seconds=60)
    return result


def produce_event(ctx, f, role, *, denied=False):
    if f.data.get('generic'):
        marker = 'NW-FORGED-' + uuid.uuid4().hex
        start = inbox_cursor(ctx, f)
        events = inbox_event_window(ctx, f, 'owner', 1, start)
        sample = next(event for event in events if event['type'] == f.data['event_type'])
        body = sample['body']
        if isinstance(body, str):
            body = json.loads(body)
        result = request(ctx, role, 'POST', '/threads/' + f.id + '/events',
                         {'type': sample['type'], 'body': {**body, f.data['identity_field']: marker}}, denied=True, expect=(403,))
        after = inbox_event_window(ctx, f, 'owner', start + 1, inbox_cursor(ctx, f))
        require(marker not in json.dumps(after), 'Public system-only publication appends its forged event')
        return result
    name = f.kind.name
    if name == 'automation':
        before = mutation_snapshot(ctx, f) if denied else None
        result = request(ctx, role, 'POST', '/automations/' + f.id + '/control', {'control': 'pause'}, denied=denied)
        if denied:
            require(mutation_snapshot(ctx, f) == before, 'Denied automation event changes its durable configuration')
        else:
            def paused():
                status = request(ctx, role, 'GET', '/automations/' + f.id + '/status')
                return status if status['drives'] and all(drive['phase'] == 'paused' for drive in status['drives']) else None
            ctx.inst.wait(paused, 'Granted automation pause reaches its declared public state')
            request(ctx, role, 'POST', '/automations/' + f.id + '/control', {'control': 'resume'})
            ctx.inst.wait(lambda: all(drive['phase'] != 'paused' for drive in request(ctx, role, 'GET',
                          '/automations/' + f.id + '/status')['drives']), 'Granted automation resume leaves the paused drive state')
        return result
    if name == 'file':
        return update(ctx, f, role, denied=denied)
    if f.data.get('connection') or f.data.get('oauth') or name in ('connection', 'remote_mcp'):
        if denied:
            return update(ctx, f, role, denied=True)
        before = read(ctx, f, role)
        update(ctx, f, role)
        after = read(ctx, f, role)
        require(after != before and f.title + ' edited by ' + role in json.dumps(after),
                'Authorized connection mutation does not persist its declared metadata')
        return after
    if name.startswith('kb_'):
        marker = 'NW-PRODUCER-' + uuid.uuid4().hex
        result = request(ctx, role, 'POST', '/knowledge/' + f.id + '/ingest',
                         {'title': marker, 'content': marker, 'external_id': marker, 'external_version': '1'}, denied=denied)
        def indexed():
            value = request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/search', {'query': marker, 'limit': 5})
            present = marker in json.dumps(value)
            require(not present, 'Denied ingestion creates indexed content') if denied else None
            return value if present or denied else None
        if denied: indexed()
        else: ctx.inst.wait(indexed, 'Granted ingestion produces new indexed content', seconds=120)
        return result
    if name == 'uploads':
        marker = 'NW-PRODUCER-' + uuid.uuid4().hex
        before = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        result = request(ctx, role, 'POST', '/threads/' + component_id(f) + '/events',
                         {'type': 'kickside.uploads.events:upload.completed', 'body': {
                          'upload_id': f.data['upload'], 'filename': marker}}, denied=denied,
                         expect=(403, 404) if denied else (201,))
        after = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        if denied: require(after == before, 'Denied upload event changes the durable target log')
        else: require(marker in json.dumps(after) and after != before, 'Granted upload event is absent from the target log')
        return result
    if name == 'research':
        marker = 'NW-PRODUCER-' + uuid.uuid4().hex
        before = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        result = request(ctx, role, 'POST', '/threads/' + component_id(f) + '/events',
                         {'type': 'kickside.research.events:progress', 'body': {'status': 'complete',
                          'step': marker, 'findings': [marker], 'eta': '', 'urls': []}}, denied=denied,
                         expect=(403, 404) if denied else (201,))
        after = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        if denied: require(after == before, 'Denied research progress changes the durable event log')
        else: require(after != before, 'Granted research progress creates no durable event')
        return result
    if f.data.get('connection') or name in ('connection', 'remote_mcp'):
        return update(ctx, f, role, denied=denied)
    if name == 'skill':
        if denied: return update(ctx, f, role, denied=True)
        before = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        update(ctx, f, role)
        after = request(ctx, 'owner', 'GET', '/threads/' + component_id(f) + '/events')
        require(after != before, 'Authorized mutation produces no new durable event for ' + name)
        return after
    return primary(ctx, f, role, denied=denied)


def delete(ctx, f, role='owner', *, denied=False):
    name = f.kind.name
    if f.data.get('connection'): name = 'connection'
    route = {'agent': '/user-agents/', 'workflow': '/workflows/', 'connection': '/connections/',
             'skill': '/skills/', 'mcp_credential': '/mcp/credentials/', 'webhook': '/webhooks/',
             'recurring': '/automations/', 'data_sync': '/automations/', 'automation': '/automations/'} .get(name, '/workspace/entries/') + f.id
    if name.startswith('kb_'): route = '/knowledge/' + f.id
    if name == 'remote_mcp': route = '/connections/' + f.id
    if f.data.get('oauth'): route = '/connections/' + f.id
    if name == 'uploads': route = '/uploads/' + f.data['upload']
    if name in ('inbox', 'lifecycle', 'research') or f.data.get('generic'): route = '/threads/' + component_id(f)
    before = mutation_snapshot(ctx, f) if denied else None
    result = request(ctx, role, 'DELETE', route, denied=denied, expect=(403, 404) if denied else (200, 202))
    if denied:
        after = mutation_snapshot(ctx, f, before)
        require(after == before,
                'Denied deletion changes resource content')
    else:
        if name in ('recurring', 'data_sync', 'automation'):
            await_automation_deletion(ctx.inst, f.id, result)
        elif name == 'uploads':
            request(ctx, role, 'GET', '/uploads/' + f.data['upload'], hidden=True)
        else:
            read(ctx, f, role, hidden=True)
    return result


def cascade(ctx, f, role):
    if f.data.get('scope') in ('system', 'user'):
        rows = component_rows(ctx, f.kind.binding)
        target = next(r for r in rows if r['component_id'] == component_id(f))
        require(target['parent_id'] is None, 'Provisioned component is not pinned at its declared root')
        ctx.inst.evidence.save('provisioned-placement', {'binding': f.kind.binding, 'scope': f.data['scope'], 'parent_id': target['parent_id']})
    folder = request(ctx, 'owner', 'POST', '/workspace/entries', {'type': 'folder', 'name': f.title + ' parent'})['entry']
    parent = Fixture(KINDS[0], folder['id'], folder['name'])
    if f.data.get('scope') in ('user', 'system'):
        before = workspace(ctx, f)
        refused = request(ctx, 'owner', 'POST', '/workspace/entries/' + component_id(f) + '/move', {'parent_id': parent.id}, denied=True)
        require('pinned' in json.dumps(refused).lower(), 'Pinned move refusal loses its declared cause: ' + str(refused))
        require(workspace(ctx, f) == before, 'Pinned move refusal changes the workspace placement')
        delete(ctx, parent)
        return
    request(ctx, 'owner', 'POST', '/workspace/entries/' + component_id(f) + '/move', {'parent_id': parent.id})
    entry = workspace(ctx, f)['entry']
    require(entry['parent_id'] == parent.id, 'Folder placement does not persist')
    reader = 'read' if role == 'owner' else role
    share(ctx, parent, reader, 3 if reader == 'write' else 1)
    read(ctx, f, reader)
    primary(ctx, f, reader, readonly=reader != 'write')
    share(ctx, parent, reader, revoke=True)
    read(ctx, f, reader, hidden=True)
    update(ctx, f, reader, denied=True)
    request(ctx, 'owner', 'POST', '/workspace/entries/' + component_id(f) + '/move', {'parent_id': False})
    delete(ctx, parent)


def private_resource_matrix(ctx, kind, role):
    checks = Checks(ctx, kind.name, role)
    f = create(ctx, kind)
    checks.check('owner_management_visibility', lambda: read(ctx, f))
    checks.check('owner_bearer_use', lambda: primary(ctx, f, 'owner'))
    checks.check('owner_edit', lambda: update(ctx, f, 'owner'))
    observer = 'stranger' if role == 'owner' else role
    checks.check('foreign_management_hidden', lambda: read(ctx, f, observer, hidden=True))
    checks.check('foreign_management_write_denied', lambda: update(ctx, f, observer, denied=True))
    checks.check('foreign_management_delete_denied', lambda: delete(ctx, f, observer, denied=True))
    checks.check('owner_delete', lambda: delete(ctx, f))
    checks.finish()


def matrix(ctx, kind, role):
    if kind.name.startswith('binding_'):
        install_package(ctx, kind.package)
        renew(ctx)
        registry(ctx)
        declaration = ctx.state['access_declarations'][kind.binding]
        if (declaration.get('meta', {}).get('component', {}).get('create') and
                declared_contracts(declaration) == {'kickside.contract:component'}):
            from access_ownership import classes
            before = component_rows(ctx, kind.binding)
            require(not before, 'Launcher owns persisted instances despite its creation-only declaration')
            types = request(ctx, role, 'GET', '/automations/types')['types']
            targets = [row for row in types if row['installable'] and row['class'] in classes(declaration)]
            require(targets, 'Launcher exposes no declared installable type')
            for class_name in classes(declaration):
                view = request(ctx, role, 'GET', '/component-kinds/create-view?class=' + quote(class_name))
                require(view.get('success') and view.get('create_view', {}).get('tag_name'),
                        'Launcher declares no resolved creation surface')
            for target in targets:
                native = fixture_kind_for_binding(ctx, target['id'])
                matrix(ctx, native, role)
            require(component_rows(ctx, kind.binding) == before, 'Launcher registers a resource under its own binding')
            return
    if kind.name == 'link':
        install_package(ctx, kind.package)
        renew(ctx)
        registry(ctx)
        declaration = ctx.state['access_declarations'][kind.binding]
        require('create' not in declaration['meta'].get('component', {}), 'Link unexpectedly declares user creation')
        types = request(ctx, role, 'GET', '/workspace/entries/types')['types']
        matches = [row for row in types if row['value'] == 'link']
        require(len(matches) == 1 and matches[0]['creatable'] is False, 'Reference-only link appears in the user-create matrix')
        result = request(ctx, role, 'POST', '/workspace/entries', {'type': 'link', 'name': 'Northwind reference'}, expect=(400,))
        require(result.get('success') is False and 'kind is not creatable: link' in result.get('error', ''),
                'Reference-only link user creation returns no declared validation failure: ' + str(result))
        ctx.inst.evidence.save('link-owning-path', {'declaration': declaration,
            'design_reference': 'chestor/workspace/src/_index.yaml:link; component.create absent; reference-owned kind',
            'user_create_validation': result})
        return
    if kind.binding == 'kickside.oauth:oauth_connection':
        install_package(ctx, kind.package)
        renew(ctx)
        registry(ctx)
        declaration = request(ctx, 'owner', 'GET', '/keeper/registry/entry?id=' + quote(kind.binding, safe=''))['entry']
        require(not declaration.get('meta', {}).get('component', {}).get('create'),
                'OAuth owning-path fixture requires an update when the kind declares user creation')
        require(not declaration['meta'].get('credential_schema'), 'OAuth declares a generic credential form')
        contracts = {c['contract'] for c in declaration['data']['contracts']}
        require('kickside.oauth:oauth_connection_contract' in contracts,
                'OAuth kind lacks its declared authorized-connection contract')
        ctx.inst.evidence.save('oauth-owning-path-declaration', declaration)
        providers = request(ctx, role, 'GET', '/oauth/providers/discover')
        require(isinstance(providers.get('providers'), list), 'OAuth provider discovery lacks its declared list')
        absent = str(uuid.uuid4())
        request(ctx, role, 'GET', '/oauth/connections/' + absent, hidden=True)
        request(ctx, role, 'PUT', '/oauth/connections/' + absent,
                {'connection_name': 'Northwind inaccessible OAuth'}, denied=True)
    if kind.name == 'mcp_credential':
        return private_resource_matrix(ctx, kind, role)
    if kind.name == 'dashboard':
        value = request(ctx, role, 'GET', '/dashboard/overview')
        require(value.get('count') == len(value.get('surfaces', [])),
                'Dashboard overview content differs from its declared count')
        require(value.get('success') is True and isinstance(value.get('surfaces'), list),
                'Dashboard overview lacks its declared surface list')
        return
    checks = Checks(ctx, kind.name, role)
    fixture = checks.check('create', lambda: create(ctx, kind))
    if fixture is None:
        error = checks.rows[-1]['error']
        for action in ('see', 'use', 'edit', 'produce_event', 'share_users', 'share_groups', 'reshare_rule', 'revoke', 'delete_rule', 'folder_cascade'):
            checks.rows.append({'kind': kind.name, 'role': role, 'action': action, 'status': 'fail',
                                'error': 'Fixture creation fails: ' + error})
        checks.finish()
    f = fixture
    scope = f.data.get('scope', 'demand')
    if scope == 'system':
        checks.check('declared_default_access', lambda: default_access(ctx, f))
        checks.check('host_admin_read', lambda: read(ctx, f, 'owner'))
        checks.check('host_admin_write', lambda: update(ctx, f, 'owner'))
        checks.check('framework_publication_denied', lambda: primary(ctx, f, role))
        checks.check('sharing_ADMIN_controls', lambda: system_sharing(ctx, f))
        checks.check('deletion_requires_DELETE', lambda: delete(ctx, f, 'owner', denied=True))
        checks.check('pinned_system_placement', lambda: cascade(ctx, f, role))
        checks.finish()
        return
    elif scope == 'user':
        checks.check('per_user_private_instances', lambda: private_instances(ctx, f))
    if role != 'owner':
        if role != 'stranger': checks.check('grant', lambda: share(ctx, f, role, 3 if role == 'write' else 1))
        checks.check('see', lambda: read(ctx, f, role, hidden=role == 'stranger'))
        checks.check('use', lambda: primary(ctx, f, role, denied=role == 'stranger', readonly=role == 'read'))
        checks.check('edit', lambda: update(ctx, f, role, denied=role != 'write'))
        checks.check('produce_event', lambda: produce_event(ctx, f, role, denied=role == 'stranger' or (role != 'write' and kind.name != 'workflow')))
        checks.check('reshare_ADMIN_rule', lambda: share(ctx, f, 'stranger' if role != 'stranger' else 'read', role=role, denied=True))
        checks.check('delete_DELETE_rule', lambda: delete(ctx, f, role, denied=True))
        if role != 'stranger':
            checks.check('revoke', lambda: share(ctx, f, role, revoke=True))
            checks.check('revoke_hides', lambda: read(ctx, f, role, hidden=True))
            checks.check('revoke_denies_actions', lambda: update(ctx, f, role, denied=True))
    else:
        checks.check('see', lambda: read(ctx, f))
        checks.check('use', lambda: primary(ctx, f, role))
        checks.check('produce_event', lambda: produce_event(ctx, f, role))
        checks.check('edit', lambda: update(ctx, f, role))
        for subject, mask, group in [('read', 1, False), ('write', 3, False), ('app.security:user', 1, True), ('app.security:user', 3, True)]:
            checks.check('share_' + subject + '_' + str(mask), lambda s=subject, m=mask, g=group: share(ctx, f, s, m, group=g))
            observer = ('write' if mask == 3 else 'read') if group else subject
            checks.check('share_observable_' + subject + '_' + str(mask), lambda r=observer: read(ctx, f, r))
            checks.check('share_use_' + subject + '_' + str(mask), lambda r=observer, m=mask: primary(ctx, f, r, readonly=m == 1))
            checks.check('share_event_control_' + subject + '_' + str(mask), lambda r=observer, m=mask: produce_event(ctx, f, r, denied=m == 1 and kind.name != 'workflow'))
            checks.check('revoke_' + subject + '_' + str(mask), lambda s=subject, g=group: share(ctx, f, s, group=g, revoke=True))
            checks.check('revoke_hides_' + subject + '_' + str(mask), lambda r=observer: read(ctx, f, r, hidden=True))
            checks.check('revoke_denies_event_' + subject + '_' + str(mask), lambda r=observer: produce_event(ctx, f, r, denied=True))
    checks.check('folder_cascade', lambda: cascade(ctx, f, 'read' if role == 'stranger' else role))
    if scope == 'demand' or kind.name == 'uploads':
        checks.check('owner_delete', lambda: delete(ctx, f))
    if kind.name == 'file':
        checks.check('upload_parent_cleanup', lambda: delete(ctx, Fixture(kind_named('folder'), f.data['upload_parent'], f.title)))
    checks.finish()


def private_instances(ctx, f):
    ids, checks = {}, Checks(ctx, 'user_private_' + f.kind.name, 'owner')
    for role in ('owner', 'write', 'read', 'stranger'):
        def resolve(r=role):
            actor_browser(ctx, r)
            own = provisioned(ctx, f.kind, r)
            ids[r] = own.id
            return own
        own = checks.check(role + '_deterministic_instance', resolve)
        def visible(r=role, target=own):
            require(target is not None, 'USER instance identity fails for ' + r)
            return read(ctx, target, r)
        checks.check(role + '_private_instance_read', visible)
        for other in ('owner', 'write', 'read', 'stranger'):
            if other != role:
                def hidden(r=other, target=own, actor=role):
                    require(target is not None, 'USER instance identity fails for ' + actor)
                    return read(ctx, target, r, hidden=True)
                checks.check(role + '_hidden_from_' + other, hidden)
    checks.check('four_distinct_instances', lambda: require(len(set(ids.values())) == 4,
                 'USER instances are not private per actor: ' + str(ids)))
    checks.finish()




def system_sharing(ctx, f):
    checks = Checks(ctx, 'system_sharing_' + f.kind.name, 'owner')
    def events(role):
        value = request(ctx, role, 'GET', '/threads/' + f.id + '/events')
        require(isinstance(value.get('events'), list), 'Granted singleton READ does not expose its event list')
    for role, mask in (('read', 1), ('write', 3)):
        checks.check(role + '_grant', lambda r=role, m=mask: share(ctx, f, r, m))
        checks.check(role + '_read', lambda r=role: read(ctx, f, r))
        checks.check(role + '_use', lambda r=role: events(r))
        checks.check(role + '_reshare_denied', lambda r=role: share(ctx, f, 'stranger', role=r, denied=True))
        checks.check(role + '_framework_publication_denied', lambda r=role: primary(ctx, f, r))
        checks.check(role + '_revoke', lambda r=role: share(ctx, f, r, revoke=True))
        checks.check(role + '_revoked_read', lambda r=role: read(ctx, f, r, hidden=True))
    def group_read():
        grants = f.data['declaration']['meta']['provision']['grants']
        original = next(g['access_mask'] for g in grants
                        if g['subject_type'] == 'group' and g['subject_id'] == 'app.security:user')
        share(ctx, f, 'app.security:user', original | 1, group=True)
        try:
            read(ctx, f, 'read')
            events('read')
        finally:
            share(ctx, f, 'app.security:user', original, group=True)
        read(ctx, f, 'read', hidden=not bool(original & 1))
    checks.check('group_read_and_default_restore', group_read)
    checks.check('delegate_ADMIN', lambda: share(ctx, f, 'write', 9))
    checks.check('ADMIN_reshare', lambda: share(ctx, f, 'read', 1, role='write'))
    checks.check('ADMIN_recipient_read', lambda: read(ctx, f, 'read'))
    checks.check('ADMIN_recipient_use', lambda: events('read'))
    checks.check('ADMIN_revoke_recipient', lambda: share(ctx, f, 'read', role='write', revoke=True))
    checks.check('ADMIN_recipient_hidden', lambda: read(ctx, f, 'read', hidden=True))
    checks.check('revoke_ADMIN', lambda: share(ctx, f, 'write', revoke=True))
    checks.check('revoked_ADMIN_denial', lambda: share(ctx, f, 'read', role='write', denied=True))
    checks.finish()


def default_access(ctx, f):
    declaration = f.data['declaration']
    grants = declaration['meta']['provision'].get('grants', [])
    ctx.inst.evidence.save('declared-default-grants', grants)
    checks = Checks(ctx, 'default_' + f.kind.name, 'owner')
    for role in ('write', 'read', 'stranger'):
        mask = 0
        for grant in grants:
            if grant['subject_type'] == 'group' and grant['subject_id'] == 'app.security:user':
                mask |= grant['access_mask']
            if grant['subject_type'] == 'user' and grant['subject_id'] in (actor_id(ctx, role), ctx.state['access_users'][role]['email']):
                mask |= grant['access_mask']
        checks.check(role + '_default_visibility', lambda r=role, m=mask: read(ctx, f, r, hidden=not bool(m & 1)))
        checks.check(role + '_default_use', lambda r=role, m=mask: primary(ctx, f, r, readonly=bool(m & 1) and not bool(m & 2), denied=not bool(m & 3)))
    checks.finish()


def discovery(ctx):
    """Discover live creatable workspace kinds and all class-bearing bindings."""
    inst = ctx.inst
    kinds = request(ctx, 'owner', 'GET', '/workspace/entries/types')
    components = request(ctx, 'owner', 'GET', '/mcp/components')
    registry(ctx)
    bindings = list(ctx.state['access_declarations'].values())
    inst.evidence.save('live-kind-inventory', {'workspace': kinds, 'components': components, 'bindings': bindings})
    known = {k.binding for k in KINDS if k.binding}
    missing = [e['id'] for e in bindings if e['id'] not in known and not supported_installable(e)]
    require(not missing, 'Discovered creatable bindings lack an access matrix adapter: ' + ', '.join(sorted(missing)))


def new_kinds(ctx, role):
    registry(ctx)
    checks = Checks(ctx, 'new_live_bindings', role)
    known = {k.binding for k in KINDS}
    dependencies = request(ctx, 'owner', 'GET', '/keeper/hub/dependencies?entries=1&migrations=0')['dependencies']
    owners = entry_owners(dependencies)
    ctx.inst.evidence.save('live-entry-owners', owners)
    for declaration in ctx.state['access_declarations'].values():
        if declaration['id'] in known:
            continue
        package = owners.get(declaration['id'])
        require(package, 'Live new kind carries no package provenance: ' + declaration['id'])
        kind = Kind('binding_' + declaration['id'].replace('.', '_').replace(':', '_'), package, declaration['id'])
        checks.check(kind.name, lambda k=kind: matrix(ctx, k, role))
    checks.finish()


def matrix_features(kind):
    name = kind.name
    if kind.binding == 'kickside.oauth:oauth_connection':
        return SHARE_FEATURES + http('chestor/workspace', 'POST /api/v1/workspace/entries',
                    'GET /api/v1/workspace/entries/{id}', 'POST /api/v1/workspace/entries/{id}/move',
                    'DELETE /api/v1/workspace/entries/{id}') + http('kickside/oauth',
                    'GET /api/v1/oauth/providers/discover', 'POST /api/v1/oauth/providers/{provider_name}/init',
                    'GET /api/v1/oauth/connections/{component_id}',
                    'PUT /api/v1/oauth/connections/{component_id}') + http('kickside/connection',
                    'DELETE /api/v1/connections/{id}') + http('kickside/mcp',
                    'POST /api/v1/mcp/credentials', 'DELETE /api/v1/mcp/credentials/{id}',
                    'POST /api/public/mcp') + ('kickside/google|mcp_tool|kickside.google.traits:gmail_read_tool',
                    kind.package + '|workspace_kind|' + kind.binding)
    if name.startswith('binding_'):
        if 'connection' not in kind.binding:
            feature = kind.package + '|workspace_kind|' + kind.binding
            require(feature in INVENTORY_IDS, 'Discovered fixture binding is absent from the declared inventory')
            return (feature,) + kind.features
        name = 'connection'
    if name == 'link':
        return http('chestor/workspace', 'GET /api/v1/workspace/entries/types', 'POST /api/v1/workspace/entries') + ('chestor/workspace|workspace_kind|chestor.workspace:link',)
    if name == 'dashboard':
        return http('kickside/dashboard', 'GET /api/v1/dashboard/overview')
    features = [] if name == 'mcp_credential' else list(SHARE_FEATURES)
    if name == 'file':
        features += list(http('kickside/uploads', 'GET /api/v1/uploads/{uuid}/raw'))
    if name not in ('mcp_credential', 'inbox', 'uploads', 'lifecycle'):
        features += list(http('chestor/workspace', 'POST /api/v1/workspace/entries',
                         'GET /api/v1/workspace/entries/{id}',
                         'POST /api/v1/workspace/entries/{id}/move', 'DELETE /api/v1/workspace/entries/{id}'))
    routes = {
      'agent': ('kickside/agents', ('POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}',
                                   'PUT /api/v1/user-agents/{id}', 'DELETE /api/v1/user-agents/{id}')),
      'kb_document': ('kickside/knowledge', ('POST /api/v1/knowledge/create', 'GET /api/v1/knowledge/{id}',
                      'PUT /api/v1/knowledge/{id}', 'POST /api/v1/knowledge/{id}/ingest',
                      'POST /api/v1/knowledge/{id}/search', 'POST /api/v1/knowledge/{id}/ask',
                      'DELETE /api/v1/knowledge/{id}')),
      'workflow': ('kickside/workflows', ('POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply',
                   'POST /api/v1/workflows/{id}/publish', 'POST /api/v1/workflows/{id}/run',
                   'GET /api/v1/workflows/{id}/state', 'GET /api/v1/workflows/{id}/runs/{run_id}',
                   'DELETE /api/v1/workflows/{id}')),
      'recurring': ('kickside/automation', ('POST /api/v1/automations', 'GET /api/v1/automations/{id}/config',
                    'PATCH /api/v1/automations/{id}', 'POST /api/v1/automations/{id}/control',
                    'DELETE /api/v1/automations/{id}', 'GET /api/v1/automations')),
      'connection': ('kickside/connection', ('POST /api/v1/connections', 'GET /api/v1/connections',
                     'PUT /api/v1/connections/{id}', 'DELETE /api/v1/connections/{id}',
                     'POST /api/v1/connections/{id}/discover_resources')),
      'remote_mcp': ('kickside/connection', ('POST /api/v1/connections', 'GET /api/v1/connections',
                     'PUT /api/v1/connections/{id}', 'DELETE /api/v1/connections/{id}',
                     'POST /api/v1/connections/{id}/discover_resources')),
      'uploads': ('kickside/uploads', ('POST /api/v1/uploads', 'GET /api/v1/uploads/{uuid}',
                   'GET /api/v1/uploads/{uuid}/artifacts', 'DELETE /api/v1/uploads/{uuid}')),
      'mcp_credential': ('kickside/mcp', ('POST /api/v1/mcp/credentials', 'GET /api/v1/mcp/credentials',
                        'PATCH /api/v1/mcp/credentials/{id}', 'DELETE /api/v1/mcp/credentials/{id}',
                        'POST /api/public/mcp')),
      'webhook': ('kickside/webhooks', ('POST /api/v1/webhooks', 'GET /api/v1/webhooks',
                  'PATCH /api/v1/webhooks/{id}', 'DELETE /api/v1/webhooks/{id}', 'POST /api/public/hooks/{id}')),
      'research': ('kickside/research', ('POST /api/v1/research', 'GET /api/v1/research/{id}/events')),
      'skill': ('kickside/skills', ('POST /api/v1/skills', 'GET /api/v1/skills/{id}',
                'PATCH /api/v1/skills/{id}', 'DELETE /api/v1/skills/{id}')),
      'crm': ('kickside/crm', ('POST /api/v1/workspace/crm', 'GET /api/v1/workspace/crm/{id}/schema',
              'POST /api/v1/workspace/crm/{id}/schema', 'POST /api/v1/workspace/crm/{id}/apply',
              'POST /api/v1/workspace/crm/{id}/query')),
      'graph': ('chestor/graph', ('GET /api/v1/workspace/graphs/{id}', 'POST /api/v1/workspace/graphs/{id}/apply')),
      'journal': ('chestor/journal', ('GET /api/v1/workspace/journals/{id}/context_pack',
                  'POST /api/v1/workspace/journals/{id}/events')),
      'roadmap': ('chestor/roadmap', ('GET /api/v1/workspace/roadmaps/{id}', 'POST /api/v1/workspace/roadmaps/{id}/apply')),
      'database': ('chestor/db-schema', ('GET /api/v1/workspace/db-schema/{id}', 'POST /api/v1/workspace/db-schema/{id}/apply')),
      'inbox': ('kickside/inbox', ('GET /api/v1/inbox', 'POST /api/v1/inbox')),
      'lifecycle': ('kickside/core', ('GET /api/v1/threads/{id}', 'PATCH /api/v1/threads/{id}',
                    'GET /api/v1/threads/{id}/events', 'POST /api/v1/threads/{id}/events', 'DELETE /api/v1/threads/{id}')),
    }
    if name in ('folder', 'file', 'link', 'graph', 'journal', 'roadmap', 'database', 'workflow', 'crm', 'research'):
        features += list(http('chestor/workspace', 'PUT /api/v1/workspace/entries/{id}'))
    if name == 'kb_wiki': name = 'kb_document'
    if name == 'data_sync': name = 'recurring'
    if name in routes:
        package, paths = routes[name]
        features += list(http(package, *paths))
    if kind.name in ('mcp_credential', 'remote_mcp'):
        features += list(http('kickside/knowledge', 'POST /api/v1/knowledge/create',
                         'POST /api/v1/knowledge/{id}/ingest', 'POST /api/v1/knowledge/{id}/search',
                         'PUT /api/v1/knowledge/{id}'))
    if kind.name == 'data_sync':
        features += list(http('kickside/webhooks', 'POST /api/v1/webhooks', 'POST /api/public/hooks/{id}'))
    if kind.name in ('recurring', 'data_sync'):
        features += list(http('kickside/inbox', 'GET /api/v1/inbox'))
    if kind.name in ('inbox', 'uploads', 'research'):
        features += list(http('kickside/core', 'GET /api/v1/threads/{id}', 'PATCH /api/v1/threads/{id}',
                             'GET /api/v1/threads/{id}/events', 'POST /api/v1/threads/{id}/events'))
        if kind.name != 'uploads':
            features += list(http('kickside/core', 'DELETE /api/v1/threads/{id}'))
    if kind.binding:
        binding_feature = kind.package + '|workspace_kind|' + kind.binding
        if binding_feature in INVENTORY_IDS:
            features.append(binding_feature)
    require(not (set(features) - INVENTORY_IDS), 'Unknown exercised access features: ' + str(set(features) - INVENTORY_IDS))
    return tuple(sorted(set(features)))


REGISTRY = [Scenario('access_users', 'full', access_users, ('first_agent',), issues=('SEC-USERS-MANAGEMENT-POLICY', 'RC2-ICONIFY-MIRROR-ABORT'),
                    features=http('kickside/users', 'POST /api/v1/users', 'GET /api/v1/users/assignable-scopes',
                                  'GET /api/v1/user/me') +
                    http('kickside/component', 'GET /api/v1/components/access-model'), instance_setup=True)]
REGISTRY.append(Scenario('access_catalog', 'full', access_catalog, ('access_users',),
                issues=('R4-HUB-I6', 'SECURITY-SCOPE-PROJECTION-LIVENESS'),
                features=http('chestor/workspace', 'GET /api/v1/workspace/entries/types') +
                         http('kickside/mcp', 'GET /api/v1/mcp/components'), instance_setup=True))

KIND_ISSUES = {
    'binding_kickside_automation_automation_create_kind': ('SMOKE-AUTOMATION-NATIVE-FIXTURE',
                                                            'SMOKE-HUB-INSTALL-RECONCILIATION-FIXTURE',
                                                            'SECURITY-SCOPE-PROJECTION-LIVENESS',
                                                            'KEEPER-PLAN-VERSION-CONFLICT'),
    'binding_kickside_dm_dm_kind': ('SEC-FORGED-SYSTEM-EVENT', 'SEC-PRODUCER-CAP-NOT-GRANTABLE',
                                    'DISCORD-GATEWAY-WAKE', 'IDENTITY-UNLINK-SERIALIZE', 'THREAD-CREATE-GAP'),
    'mcp_credential': ('SEC-MCP-OWNER-CEILING', 'SECURITY-SCOPE-PROJECTION-LIVENESS'),
    'data_sync': ('R3-P01',),
}

ACCESS_ISSUES = {
    ('data_sync', 'stranger'): ('R3-T04',),
    ('binding_kickside_atlassian_confluence_source_content_source', 'owner'): ('R3-T05', 'USERS-DELEGATED-MEMBERSHIP-SCOPE-LOSS'),
    ('binding_kickside_atlassian_jira_source_issues_source', 'owner'): ('R3-T05',),
    ('binding_kickside_crm_sink_crm_sink', 'owner'): ('R3-T05',),
    ('binding_kickside_crm_sink_crm_sink', 'read'): ('SECURITY-SCOPE-PROJECTION-LIVENESS',),
    ('binding_kickside_crm_sink_crm_sink', 'write'): ('SECURITY-SCOPE-PROJECTION-LIVENESS',),
    ('binding_kickside_discord_source_messages_source', 'owner'): ('R3-T05',),
    ('binding_kickside_simplefin_source_balances_source', 'owner'): ('R3-T05',),
    ('binding_kickside_knowledge_injection_injection_kind', 'stranger'): ('R3-T06',),
}

for kind in KINDS:
    for role in ('owner', 'write', 'read', 'stranger'):
        REGISTRY.append(Scenario('access_' + kind.name + '_' + role, 'full',
                        lambda ctx, k=kind, r=role: matrix(ctx, k, r), ('access_users',),
                        model=kind.name in ('agent', 'kb_document', 'kb_wiki', 'mcp_credential', 'remote_mcp', 'research'),
                        features=matrix_features(kind), instance_setup=True, issues=KIND_ISSUES.get(kind.name, ()) + ACCESS_ISSUES.get((kind.name, role), ())))
for role in ('owner', 'write', 'read', 'stranger'):
    REGISTRY.append(Scenario('access_new_live_bindings_' + role, 'full',
                   lambda ctx, r=role: new_kinds(ctx, r), ('access_users',), instance_setup=True,
                   issues=('SMOKE-INSTALLABLE-ACCESS-FIXTURE-CAPABILITY', 'WORKFLOW-UNADMITTED-RUN-READ-PROJECTION')))
REGISTRY.append(Scenario('access_discovery', 'full', discovery, ('access_users',), issues=('SEC-HIDDEN-KIND-LISTING', 'SMOKE-INSTALLABLE-ACCESS-FIXTURE-CAPABILITY'),
                features=http('chestor/workspace', 'GET /api/v1/workspace/entries/types') +
                http('kickside/mcp', 'GET /api/v1/mcp/components')))


def kind_named(name):
    return next(k for k in KINDS if k.name == name)


def regression_agent_read_write_denial(ctx):
    f = create(ctx, kind_named('agent'))
    share(ctx, f, 'read', 1)
    read(ctx, f, 'read')
    update(ctx, f, 'read', denied=True)
    require('edited by read' not in json.dumps(read(ctx, f)), 'READ update persists')


def regression_nested_folder(ctx):
    f = create(ctx, kind_named('folder'))
    child = request(ctx, 'owner', 'POST', '/workspace/entries',
                    {'type': 'folder', 'name': 'Northwind nested regression', 'parent_id': f.id})['entry']
    require(child['parent_id'] == f.id, 'Nested folder creation loses its parent')
    value = request(ctx, 'owner', 'GET', '/workspace/entries/' + child['id'])['entry']
    require(value['parent_id'] == f.id and value['name'] == 'Northwind nested regression', 'Nested folder does not persist')


def regression_graph_events(ctx):
    f = create(ctx, kind_named('graph'))
    request(ctx, 'owner', 'POST', '/workspace/graphs/' + f.id + '/apply', {'commands': [
        {'type': 'chestor.graph.events:node-type.registered', 'body': {'id': 'northwind-task', 'title': 'Northwind Task'}},
        {'type': 'chestor.graph.events:node.added', 'body': {'id': 'northwind-node', 'name': 'Northwind node', 'type': 'northwind-task'}}]})
    value = ctx.inst.wait(lambda: (v if 'northwind-node' in json.dumps(v) and 'northwind-task' in json.dumps(v) else None)
        if (v := read(ctx, f)) else None, 'Declared graph events project', seconds=35)
    require(value['graph']['counts']['nodes'] == 1, 'Graph projection loses its node')


def regression_schema_events(ctx):
    f = create(ctx, kind_named('database'))
    primary(ctx, f, 'owner')
    value = read(ctx, f)
    require(len(value['schema']['tables']) == 1, 'Declared table event does not project exactly once')


def regression_roadmap_projection(ctx):
    f = create(ctx, kind_named('roadmap'))
    request(ctx, 'owner', 'POST', '/workspace/roadmaps/' + f.id + '/apply', {'commands': [
        {'type': 'chestor.roadmap.events:roadmap.settings-updated', 'body': {'origin': '2026-Q4', 'period_kind': 'quarter'}},
        {'type': 'chestor.roadmap.events:layer.added', 'body': {'id': 'northwind-layer', 'title': 'Northwind Delivery', 'sort_order': 0}},
        {'type': 'chestor.roadmap.events:initiative.added', 'body': {'id': 'northwind-initiative', 'title': 'Northwind Coastal Delivery',
                 'layer_id': 'northwind-layer', 'start_period': '2026-Q4', 'span': 1}}]})
    value = ctx.inst.wait(lambda: v if 'northwind-initiative' in json.dumps(v := read(ctx, f)) else None,
                          'Roadmap durable projection', seconds=35)
    require(value['roadmap']['counts']['layers'] == 1 and value['roadmap']['counts']['initiatives'] == 1,
            'Roadmap durable counts differ from appended events')


def regression_foreign_automation_config(ctx):
    f = create(ctx, kind_named('recurring'))
    read(ctx, f)
    request(ctx, 'stranger', 'GET', '/automations/' + f.id + '/config', denied=True)
    require(f.id in json.dumps(read(ctx, f)), 'Foreign config request damages its owner resource')


def regression_workflow_member_apply(ctx):
    install_package(ctx, 'kickside/workflows')
    renew(ctx)
    fid = request(ctx, 'write', 'POST', '/workflows', {'title': 'Northwind member Start regression'})['flow_id']
    request(ctx, 'write', 'POST', '/workflows/' + fid + '/apply', {'commands': [
        {'type': 'node.add', 'body': {'id': 'start', 'kind': 'start', 'config': {}}}]})
    state = request(ctx, 'write', 'GET', '/workflows/' + fid + '/state')['state']
    require(any(n['id'] == 'start' for n in state['nodes']), 'Member Start apply has no durable node')


def regression_unsupported_inbox_type(ctx):
    from data_contracts import inbox_items, event_window
    owner_request = lambda *args, **kwargs: request(ctx, 'owner', *args, **kwargs)
    before, listing = inbox_items(owner_request, '/inbox')
    base = '/threads/' + listing['component_id'] + '/events'
    start = owner_request('GET', base + '?ascending=false&limit=1')['events']
    from_seq = max((event['seq'] for event in start), default=0) + 1
    title = 'Northwind unsupported type ' + uuid.uuid4().hex
    body = {'item_type': 'notification', 'title': title}
    result = owner_request('POST', '/inbox', body, expect=(400,))
    require(result.get('success') is False and result.get('error'), 'Unsupported inbox kind has no validation error')
    after, _ = inbox_items(owner_request, '/inbox')
    require([item for item in after if item['title'] == title] == [item for item in before if item['title'] == title],
            'Rejected inbox kind produces a durable item')
    tail = owner_request('GET', base + '?ascending=false&limit=1')['events']
    to_seq = max((event['seq'] for event in tail), default=0)
    events = event_window(owner_request, base, from_seq, to_seq)
    require(not any(title in json.dumps(event.get('body')) for event in events),
            'Rejected inbox kind appends a durable creation event')


def regression_knowledge_model_catalog(ctx):
    result = request(ctx, 'owner', 'GET', '/knowledge/models')
    require(ctx.state['embed'] in json.dumps(result), 'Configured embedding model is missing from knowledge/models: ' + str(result))


def regression_workspace_deleted(ctx):
    f = create(ctx, kind_named('folder'))
    delete(ctx, f)
    listed = request(ctx, 'owner', 'GET', '/workspace/entries')
    require(f.id not in json.dumps(listed), 'Deleted workspace entry remains in its parent listing')


def regression_workspace_granted_parent(ctx):
    parent = create(ctx, kind_named('folder'))
    install_package(ctx, 'chestor/graph')
    renew(ctx)
    child = request(ctx, 'owner', 'POST', '/workspace/entries',
          {'type': 'graph', 'name': 'Northwind shared Planning graph', 'parent_id': parent.id})['entry']
    f = Fixture(kind_named('graph'), child['id'], child['name'])
    checks = Checks(ctx, 'regression_R2-002_apps', 'read')
    email = ctx.state['access_users']['read']['email']
    checks.check('qa_email_workspace_grant', lambda: request(ctx, 'owner', 'POST', '/workspace/entries/' + parent.id + '/grant',
                 {'subject_id': email, 'subject_type': 'user', 'access_mask': 1}))
    checks.check('qa_email_granted_folder_read', lambda: read(ctx, parent, 'read'))
    checks.check('qa_email_granted_child_read', lambda: read(ctx, f, 'read'))
    checks.check('qa_email_workspace_revoke', lambda: request(ctx, 'owner', 'POST', '/workspace/entries/' + parent.id + '/revoke',
                 {'subject_id': email, 'subject_type': 'user'}))
    checks.check('canonical_actor_grant', lambda: share(ctx, parent, 'read', 1))
    checks.check('granted_folder', lambda: read(ctx, parent, 'read'))
    checks.check('cascade_child', lambda: read(ctx, f, 'read'))
    checks.check('read_write_denied', lambda: primary(ctx, f, 'read', denied=True))
    checks.check('revoke_parent', lambda: share(ctx, parent, 'read', revoke=True))
    checks.check('revoked_folder_hidden', lambda: read(ctx, parent, 'read', hidden=True))
    checks.check('revoked_child_hidden', lambda: read(ctx, f, 'read', hidden=True))
    checks.finish()


def regression_foreign_chestor_reads(ctx):
    checks = Checks(ctx, 'regression_N06', 'stranger')
    for name in ('graph', 'database'):
        f = create(ctx, kind_named(name))
        read(ctx, f)
        checks.check(name, lambda fixture=f: read(ctx, fixture, 'stranger', hidden=True))
    checks.finish()


def regression_member_research(ctx):
    f = create(ctx, kind_named('research'), 'write')
    value = primary(ctx, f, 'write')
    require(value['run']['result_md'], 'Member research lacks a persisted report')
    read(ctx, f, 'stranger', hidden=True)


def regression_member_upload_delete(ctx):
    from scenarios import pdf_fixture
    from PIL import Image, ImageDraw
    from docx import Document
    paths = [pdf_fixture(ctx.inst, 'northwind-member-delete')]
    image = ctx.inst.root / 'temp' / 'northwind-member-delete.png'
    im = Image.new('RGB', (320, 200), 'white')
    ImageDraw.Draw(im).text((20, 20), 'Northwind NW-ACCESS-731', fill='black')
    im.save(image)
    paths.append(image)
    doc = Document()
    doc.add_paragraph('Northwind NW-ACCESS-731')
    path = ctx.inst.root / 'temp' / 'northwind-member-delete.docx'
    doc.save(path)
    paths.append(path)
    checks = Checks(ctx, 'regression_R2-001', 'write')
    for path in paths:
        def lifecycle(p=path):
            upload_id = upload_path(ctx, 'write', p)
            request(ctx, 'write', 'DELETE', '/uploads/' + upload_id)
            request(ctx, 'write', 'GET', '/uploads/' + upload_id, hidden=True)
            listed = request(ctx, 'write', 'GET', '/uploads/list')
            require(upload_id not in json.dumps(listed), 'Deleted member upload remains listed')
        checks.check(path.suffix, lifecycle)
    checks.finish()


def upload_path(ctx, role, path, *, failed=False, target_parent_id=None):
    import mimetypes
    boundary = 'northwind_' + uuid.uuid4().hex
    raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
           f'Content-Type: {mimetypes.guess_type(path.name)[0] or "application/octet-stream"}\r\n\r\n').encode()
    raw += path.read_bytes() + f'\r\n--{boundary}--\r\n'.encode()
    if target_parent_id is not None:
        metadata = json.dumps({'target_parent_id': target_parent_id})
        raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="metadata"\r\n\r\n'
               f'{metadata}\r\n').encode() + raw
    result = request(ctx, role, 'POST', '/uploads', raw=raw,
              headers={'Content-Type': 'multipart/form-data; boundary=' + boundary}, timeout=120)
    uid = result['uuid']
    def ready():
        row = request(ctx, role, 'GET', '/uploads/' + uid)['upload']
        if failed:
            return row if row['status'] in ('failed', 'error') else None
        require(row['status'] not in ('failed', 'error'), 'Upload processing fails: ' + str(row))
        return row if row['status'] == 'completed' else None
    ctx.inst.wait(ready, 'Regression upload processing', seconds=120)
    return uid


def regression_png_ingest(ctx):
    from PIL import Image, ImageDraw, ImageFont
    checks = Checks(ctx, 'regression_R2-002', 'write')
    for role in ('owner', 'write'):
        for name in ('kb_document', 'kb_wiki'):
            def ingest(r=role, n=name):
                f = create(ctx, kind_named(n), r)
                path = ctx.inst.root / 'temp' / ('northwind-ingest-' + r + '-' + n + '.png')
                im = Image.new('RGB', (600, 180), 'white')
                ImageDraw.Draw(im).multiline_text((20, 35), 'NORTHWIND PNG\nNW-PNG-731 17 DAYS', fill='black',
                    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 28), spacing=12)
                im.save(path)
                uid = upload_path(ctx, r, path)
                result = request(ctx, r, 'POST', '/knowledge/' + f.id + '/ingest-upload', {'upload_id': uid})
                require(result.get('queued', 0) > 0 and not result.get('errors'), 'PNG returns no queued artifact: ' + str(result))
                ctx.inst.wait(lambda: 'NW-PNG-731' in json.dumps(request(ctx, r, 'POST', '/knowledge/' + f.id + '/search',
                     {'query': 'NW-PNG-731', 'limit': 5})), 'PNG ingestion produces searchable content', seconds=120)
            checks.check(role + '_' + name, ingest)
    checks.finish()


def regression_mcp_last_trait(ctx):
    f = create(ctx, kind_named('kb_document'))
    credential = request(ctx, 'owner', 'POST', '/mcp/credentials', {'label': 'Northwind last active trait',
              'whitelist_component_ids': [f.id], 'allowed_trait_ids': ['kickside.knowledge.traits:kb_reader'],
              'default_active_trait_ids': ['kickside.knowledge.traits:kb_reader']})['credential']
    ctx.inst.redactor.add(credential['token'])
    rpc(ctx, credential, 'initialize', {'protocolVersion': '2025-03-26', 'capabilities': {},
                               'clientInfo': {'name': 'Northwind access regression', 'version': '1'}})
    tools = rpc(ctx, credential, 'tools/list', {})['tools']
    require(any(t['name'] == 'QueryKB' for t in tools), 'Default KB reader does not activate')
    dropped = rpc(ctx, credential, 'tools/call', {'name': 'drop_trait', 'arguments': {'id': 'kickside.knowledge.traits:kb_reader'}})
    require('QueryKB' not in json.dumps(rpc(ctx, credential, 'tools/list', {})), 'Dropping last default trait reactivates QueryKB: ' + str(dropped))
    require('QueryKB' not in json.dumps(rpc(ctx, credential, 'tools/list', {})), 'Empty active trait set does not persist across tool listings')


def regression_mcp_profile_context(ctx):
    f = create(ctx, kind_named('kb_document'), 'write')
    credential = request(ctx, 'write', 'POST', '/mcp/credentials', {'label': 'Northwind profile QueryKB',
              'whitelist_component_ids': [f.id], 'allowed_trait_ids': ['kickside.knowledge.traits:kb_reader'],
              'default_active_trait_ids': ['kickside.knowledge.traits:kb_reader'],
              'allowed_tool_ids': ['kickside.knowledge.traits:kb_search_tool'],
              'agent_profile': {'trait_contexts': {'kickside.knowledge.traits:kb_reader': {'kb_ids': [f.id]}}}})['credential']
    ctx.inst.redactor.add(credential['token'])
    rpc(ctx, credential, 'initialize', {'protocolVersion': '2025-03-26', 'capabilities': {},
                               'clientInfo': {'name': 'Northwind profile regression', 'version': '1'}})
    scoped = rpc(ctx, credential, 'tools/call', {'name': 'SearchKB', 'arguments': {'fields': ['content', 'meta'], 'kb_id': f.id, 'query': 'NW-ACCESS-731'}})
    require('NW-ACCESS-731' in json.dumps(scoped), 'Explicit own SearchKB misses its facts')
    value = rpc(ctx, credential, 'tools/call', {'name': 'QueryKB', 'arguments': {'fields': ['content', 'meta'], 'query': 'What is the return window and access marker?'}})
    require('17' in json.dumps(value) and 'NW-ACCESS-731' in json.dumps(value), 'Stored reader profile does not reach QueryKB')


def regression_mcp_crm_error(ctx):
    f = create(ctx, kind_named('crm'))
    credential = request(ctx, 'owner', 'POST', '/mcp/credentials', {'label': 'Northwind CRM denial protocol',
                  'whitelist_component_ids': [f.id], 'allowed_tool_ids': ['kickside.crm.traits:get_record']})['credential']
    ctx.inst.redactor.add(credential['token'])
    value = ctx.inst.request('POST', '/api/public/mcp', {'jsonrpc': '2.0', 'id': uuid.uuid4().hex,
             'method': 'tools/call', 'params': {'name': 'CRMGetRecord', 'arguments': {
             'crm_id': str(uuid.uuid4()), 'object_type': 'person', 'record_id': 'northwind-outside'}}},
             auth=credential['token'], headers={'Accept': 'application/json, text/event-stream'},
             label='crm-outside-component', timeout=120)
    require(value.get('error') or value.get('result', {}).get('isError') is True,
            'MCP denied CRM call reports success: ' + str(value))
    require('NW-ACCESS-731' not in json.dumps(value), 'MCP denied CRM call leaks scoped facts')


def regression_run_code_bounded(ctx):
    install_package(ctx, 'kickside/code-execution')
    renew(ctx)
    credential = request(ctx, 'write', 'POST', '/mcp/credentials', {'label': 'Northwind bounded RunCode',
                  'allowed_tool_ids': ['kickside.code_execution.traits:run_code_tool']})['credential']
    ctx.inst.redactor.add(credential['token'])
    value = ctx.inst.request('POST', '/api/public/mcp', {'jsonrpc': '2.0', 'id': uuid.uuid4().hex,
             'method': 'tools/call', 'params': {'name': 'RunCode', 'arguments': {'code': 'while true do end'}}},
             auth=credential['token'], headers={'Accept': 'application/json, text/event-stream'},
             label='bounded-run-code', timeout=120)
    require(value.get('error') or value.get('result', {}).get('isError') or 'limit' in json.dumps(value).lower(),
            'Unbounded code has no observable bounded execution failure: ' + str(value))


def regression_member_code_authoring(ctx):
    install_package(ctx, 'kickside/code-execution')
    renew(ctx)
    reference = request(ctx, 'write', 'GET', '/code-execution/sdk-reference')
    require('sdk' in json.dumps(reference).lower(), 'Member SDK reference contains no SDK contract')
    tools = request(ctx, 'write', 'POST', '/code-execution/recommend-tools', {'need': 'Calculate fictional Northwind total 24*12'})
    require(tools.get('success') and ('tools' in tools or 'recommendations' in tools), 'Tool recommendation returns no observable result: ' + str(tools))


def regression_member_code_assist(ctx):
    install_package(ctx, 'kickside/code-execution')
    renew(ctx)
    value = request(ctx, 'write', 'POST', '/code-execution/assist',
                    {'intent': 'Return a Lua table with fictional Northwind company and total 42. No tool calls.',
                     'model': ctx.state['model']}, timeout=120)
    require(value.get('ok') is True and value.get('validated') is True and 'Northwind' in str(value.get('code')),
            'Member code assist returns no validated Northwind Lua: ' + str(value))


def regression_member_renderer_catalog(ctx):
    renew(ctx)
    result = request(ctx, 'write', 'GET', '/renderers')
    require(result.get('renderers') and result.get('success'), 'Member renderer catalog contains no renderers: ' + str(result))


def regression_member_sync_create(ctx, *, delivery=False):
    install_package(ctx, 'kickside/sync')
    renew(ctx)
    f = create(ctx, kind_named('kb_document'), 'write')
    declaration = ctx.state['access_declarations']['kickside.inbox:inbox_binding']
    class_name = declaration['meta']['class'][0]
    inbox = provision_id(actor_id(ctx, 'write'), class_name)
    marker = 'NW-SYNC-621-MEMBER-' + uuid.uuid4().hex[:6]
    payload = {'title': 'Northwind member inbox to KB', 'enabled': True,
       'trigger': {'v': 2, 'source': 'kickside.inbox:item_created', 'config': {'thread_id': inbox}, 'backfill': 'none'},
       'sink_port': 'kickside.knowledge.sink:kb_embed', 'sink_config': {'kb_id': f.id},
       'transform': {'kind': 'expr', 'expr': json.dumps({'title': 'Northwind delivered', 'content': marker,
                              'external_id': marker, 'external_version': '1'})}}
    result = request(ctx, 'write', 'POST', '/automations', {'type': 'kickside.sync:sync', 'input': payload})
    aid = result['id']
    require(aid in json.dumps(request(ctx, 'write', 'GET', '/automations/' + aid + '/config')), 'Member Data Sync creation does not persist')
    if delivery:
        request(ctx, 'write', 'POST', '/inbox', {'item_type': 'notice', 'title': marker, 'content': marker})
        ctx.inst.wait(lambda: marker in json.dumps(request(ctx, 'write', 'POST', '/knowledge/' + f.id + '/search',
                      {'query': marker, 'limit': 5})), 'Member sync produces searchable KB content', seconds=120)
        status = request(ctx, 'write', 'GET', '/automations/' + aid + '/status')
        require('retry' not in json.dumps(status).lower(), 'Delivered member sync still retries: ' + str(status))


def regression_journal_background(ctx):
    f = create(ctx, kind_named('journal'))
    primary(ctx, f, 'owner')
    def ready():
        health = request(ctx, 'owner', 'GET', '/workspace/journals/' + f.id + '/health')
        events = request(ctx, 'owner', 'GET', '/workspace/journals/' + f.id + '/events')
        ctx.inst.evidence.save('journal-background-state', {'health': health, 'events': events})
        return (health, events) if 'pending' not in json.dumps(events) and 'unknown' not in json.dumps(health) else None
    value = ctx.inst.wait(ready, 'Journal health and embeddings complete', seconds=120)
    recall = request(ctx, 'owner', 'POST', '/workspace/journals/' + f.id + '/search',
                     {'query': 'Northwind NW-ACCESS-731', 'limit': 5})
    require('NW-ACCESS-731' in json.dumps(recall), 'Journal semantic recall loses appended facts: ' + str(value))


def regression_wiki_enrichment(ctx):
    f = create(ctx, kind_named('kb_wiki'))
    request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/ingest', {'title': 'Northwind Tea relationships',
                'content': 'Northwind Tea is stored in East Warehouse. Ada manages East Warehouse. '
                           'Northwind Tea supplies Delivery Team. NW-WIKI-731.',
                'external_id': 'northwind-tea-relations', 'external_version': '1'})
    request(ctx, 'owner', 'POST', '/kb10/' + f.id + '/enrichment', {})
    def populated():
        value = request(ctx, 'owner', 'GET', '/kb10/' + f.id + '/graph')
        graph = value.get('graph', value)
        return value if len(graph.get('nodes', [])) > 1 and graph.get('edges') else None
    graph = ctx.inst.wait(populated, 'Wiki enrichment materializes concepts and edges', seconds=120)
    require('Northwind' in json.dumps(graph) or 'Tea' in json.dumps(graph), 'Wiki graph loses ingested relationship content')


REGRESSION_CASES = (
    ('QP-006', regression_nested_folder, False, http('chestor/workspace', 'POST /api/v1/workspace/entries', 'GET /api/v1/workspace/entries/{id}'), True),
    ('QA-002', regression_graph_events, False, http('chestor/graph', 'GET /api/v1/workspace/graphs/{id}', 'POST /api/v1/workspace/graphs/{id}/apply'), True),
    ('QA-003', regression_schema_events, False, http('chestor/db-schema', 'GET /api/v1/workspace/db-schema/{id}', 'POST /api/v1/workspace/db-schema/{id}/apply'), True),
    ('QA-004', regression_roadmap_projection, False, http('chestor/roadmap', 'GET /api/v1/workspace/roadmaps/{id}', 'POST /api/v1/workspace/roadmaps/{id}/apply'), True),
    ('QP-009', regression_agent_read_write_denial, True, SHARE_FEATURES + http('kickside/agents', 'POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}', 'PUT /api/v1/user-agents/{id}'), True),
    ('QP-010', regression_member_sync_create, True, http('kickside/automation', 'POST /api/v1/automations', 'GET /api/v1/automations/{id}/config'), True),
    ('QP-011', regression_workspace_deleted, False, http('chestor/workspace', 'POST /api/v1/workspace/entries', 'DELETE /api/v1/workspace/entries/{id}', 'GET /api/v1/workspace/entries/{id}', 'GET /api/v1/workspace/entries'), True),
    ('QP-012', regression_knowledge_model_catalog, True, http('kickside/knowledge', 'GET /api/v1/knowledge/models'), False),
    ('QP-013', regression_unsupported_inbox_type, False, http('kickside/inbox', 'GET /api/v1/inbox', 'POST /api/v1/inbox'), False),
    ('QP-016', regression_foreign_automation_config, False, http('kickside/automation', 'POST /api/v1/automations', 'GET /api/v1/automations/{id}/config'), True),
    ('QP-017', regression_workflow_member_apply, False, http('kickside/workflows', 'POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply', 'GET /api/v1/workflows/{id}/state'), True),
    ('QP-019', regression_mcp_last_trait, True, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp'), True),
    ('QA-007', regression_wiki_enrichment, True, http('kickside/knowledge', 'POST /api/v1/knowledge/create', 'POST /api/v1/knowledge/{id}/ingest') + http('kickside/kb10', 'POST /api/v1/kb10/{id}/enrichment', 'GET /api/v1/kb10/{id}/graph'), True),
    ('QA-008', regression_journal_background, True, http('chestor/journal', 'POST /api/v1/workspace/journals/{id}/events', 'GET /api/v1/workspace/journals/{id}/health', 'POST /api/v1/workspace/journals/{id}/search'), True),
    ('QA-012', regression_mcp_crm_error, False, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp'), True),
    ('QA-P03', regression_run_code_bounded, False, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp'), True),
    ('QA-P04', regression_member_code_authoring, False, http('kickside/code-execution', 'GET /api/v1/code-execution/sdk-reference', 'POST /api/v1/code-execution/recommend-tools'), True),
    ('QA-P05', lambda ctx: delete(ctx, create(ctx, kind_named('kb_wiki'), 'write'), 'write'), True, http('kickside/knowledge', 'POST /api/v1/knowledge/create', 'POST /api/v1/knowledge/{id}/ingest', 'DELETE /api/v1/knowledge/{id}'), True),
    ('QA-P06', lambda ctx: delete(ctx, create(ctx, kind_named('remote_mcp')), 'stranger', denied=True), True, http('kickside/connection', 'POST /api/v1/connections', 'DELETE /api/v1/connections/{id}'), True),
    ('R2-001', regression_member_upload_delete, True, http('kickside/uploads', 'POST /api/v1/uploads', 'DELETE /api/v1/uploads/{uuid}', 'GET /api/v1/uploads/{uuid}', 'GET /api/v1/uploads/list'), False),
    ('R2-002', regression_png_ingest, True, http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/ingest-upload', 'POST /api/v1/knowledge/{id}/search') + http('kickside/uploads', 'POST /api/v1/uploads'), True),
    ('R2-003', lambda ctx: regression_member_sync_create(ctx, delivery=True), True, http('kickside/automation', 'POST /api/v1/automations', 'GET /api/v1/automations/{id}/status') + http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/search'), True),
    ('R2-005', regression_mcp_profile_context, True, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp'), True),
    ('R2-015', regression_member_renderer_catalog, False, http('kickside/renderers', 'GET /api/v1/renderers'), False),
    ('R2-002_apps', regression_workspace_granted_parent, False, SHARE_FEATURES + http('chestor/workspace', 'POST /api/v1/workspace/entries', 'GET /api/v1/workspace/entries/{id}') + http('chestor/graph', 'GET /api/v1/workspace/graphs/{id}', 'POST /api/v1/workspace/graphs/{id}/apply'), True),
    ('R2-003_apps', regression_member_research, True, http('kickside/research', 'POST /api/v1/research', 'GET /api/v1/research/{id}/events'), True),
    ('QA-R2-01', regression_member_code_assist, True, http('kickside/code-execution', 'POST /api/v1/code-execution/assist'), True),
    ('N06', regression_foreign_chestor_reads, False, http('chestor/graph', 'GET /api/v1/workspace/graphs/{id}') + http('chestor/db-schema', 'GET /api/v1/workspace/db-schema/{id}'), True),
)
REGRESSION_ISSUES = {
    'QA-002': ('SECURITY-SCOPE-PROJECTION-LIVENESS',),
    'QA-004': ('SECURITY-SCOPE-PROJECTION-LIVENESS',),
    'R2-008': ('SECURITY-SCOPE-PROJECTION-LIVENESS',),
    'N01': ('SMOKE-KNOWLEDGE-SEARCH-MODEL-OUTPUT',),
}
for regression_id, operation, model, features, instance_setup in REGRESSION_CASES:
    require(not (set(features) - INVENTORY_IDS), 'Unknown regression features: ' + str(set(features) - INVENTORY_IDS))
    REGISTRY.append(Scenario('regression_' + regression_id, 'full', operation, ('access_users',), model=model,
                    features=features, instance_setup=instance_setup, issues=REGRESSION_ISSUES.get(regression_id, ())))


def open_knowledge(ctx, f, role='owner'):
    browser = actor_browser(ctx, role)
    browser.frame.locator('.sidebar').get_by_text('Knowledge', exact=True).click()
    browser.frame.get_by_text(f.title, exact=True).first.click()
    browser.frame.locator('kx-kb-wiki' if f.kind.name == 'kb_wiki' else 'kx-kb-document').wait_for(timeout=30000)
    return browser


def regression_manual_document_title(ctx):
    f = create(ctx, kind_named('kb_document'))
    title = 'Northwind Shipping Policy'
    request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/ingest', {'title': title,
        'content': '# Northwind Shipping Policy\nNW-TITLE-731 ships in 17 days.',
        'external_id': 'northwind-policy-title', 'external_version': '1'})
    ctx.inst.wait(lambda: title in json.dumps(request(ctx, 'owner', 'GET', '/kb10/' + f.id + '/browse')),
                  'Manual document browse title', seconds=120)
    browser = open_knowledge(ctx, f)
    browser.frame.get_by_text(title, exact=True).first.click()
    heading = browser.frame.locator('.dr-title')
    heading.wait_for(timeout=30000)
    require(heading.inner_text() == title, 'Document reader renders an ID instead of the supplied title: ' + heading.inner_text())
    browser.clean(ctx.state['access_browser_error_start'])


def regression_updated_document_ui_delete(ctx):
    f = create(ctx, kind_named('kb_document'))
    title = 'Northwind Dispatch Revision'
    for revision, marker in [('1', 'SILVER OTTER'), ('2', 'GOLDEN BADGER')]:
        request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/ingest', {'title': title,
            'content': '# ' + title + '\nDispatch marker ' + marker,
            'external_id': 'northwind-revision', 'external_version': revision})
        ctx.inst.wait(lambda m=marker: m in json.dumps(request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/search',
                      {'query': m, 'limit': 5})), 'Updated dispatch revision indexes ' + revision, seconds=120)
    browser = open_knowledge(ctx, f)
    browser.frame.get_by_text(title, exact=True).first.click()
    browser.frame.get_by_role('button', name='Delete ' + title, exact=True).click()
    dialog = browser.frame.get_by_role('dialog')
    dialog.get_by_role('button', name='Delete', exact=True).click()
    def removed():
        browse = request(ctx, 'owner', 'GET', '/kb10/' + f.id + '/browse')
        search = request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/search', {'query': 'GOLDEN BADGER', 'limit': 5})
        return (browse, search) if title not in json.dumps(browse) and 'GOLDEN BADGER' not in json.dumps(search) else None
    ctx.inst.wait(removed, 'UI deletion removes the updated document and indexed content', seconds=120)
    answer = request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/ask', {'query': 'What is the dispatch marker?'}, timeout=120)
    require('GOLDEN BADGER' not in json.dumps(answer) and 'SILVER OTTER' not in json.dumps(answer), 'Deleted revision still appears in Ask')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_wiki_keyboard(ctx):
    f = create(ctx, kind_named('kb_wiki'))
    title = 'Northwind Tea'
    request(ctx, 'owner', 'POST', '/knowledge/' + f.id + '/ingest',
            {'title': 'Northwind Tea relationships',
             'content': 'Northwind Tea is stored in East Warehouse. Ada manages East Warehouse. '
                        'Northwind Tea supplies Delivery Team. NW-KEYBOARD-731.',
             'external_id': 'northwind-keyboard-tea', 'external_version': '1'})
    ctx.inst.wait(lambda: 'NW-KEYBOARD-731' in json.dumps(request(ctx, 'owner', 'POST',
        '/knowledge/' + f.id + '/search', {'query': 'NW-KEYBOARD-731', 'limit': 5})),
        'Tea Wiki source is indexed before taxonomy changes', seconds=120)
    topic = request(ctx, 'owner', 'POST', '/kb10/' + f.id + '/concepts',
                    {'label': title, 'description': 'Northwind Tea navigation NW-KEYBOARD-731.'})
    require(topic.get('success') is True, 'Declared topic creation fails')
    ctx.inst.wait(lambda: title in json.dumps(request(ctx, 'owner', 'GET', '/kb10/' + f.id + '/graph')),
                  'Tea Wiki materializes the explicitly named navigation topic', seconds=120)
    browser = open_knowledge(ctx, f)
    browser.frame.get_by_role('tab', name='Wiki', exact=True).click()
    browser.frame.get_by_label('Filter the contents', exact=True).fill(title)
    option = browser.frame.get_by_role('option', name=title, exact=False).first
    option.wait_for(timeout=30000)
    option.focus()
    require(option.evaluate('el => el === el.getRootNode().activeElement'), 'Wiki navigation option cannot receive keyboard focus')
    option.press('Enter')
    browser.frame.get_by_role('heading', name=title, exact=True).wait_for(timeout=30000)
    browser.frame.locator('kx-kb-wiki').get_by_text('NW-KEYBOARD-731', exact=False).first.wait_for(timeout=30000)
    require(browser.frame.locator('kx-kb-wiki').get_by_text('NW-KEYBOARD-731', exact=False).first.is_visible(),
            'Enter does not open the indexed Wiki page')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_graph_legend_label(ctx):
    f = create(ctx, kind_named('kb_wiki'))
    browser = open_knowledge(ctx, f)
    browser.frame.get_by_role('tab', name='Graph', exact=True).click()
    toggle = browser.frame.locator('.gx-legend-toggle')
    toggle.wait_for(timeout=30000)
    toggle.click()
    require(toggle.get_attribute('aria-label') and toggle.get_attribute('aria-label').strip(),
            'Expanded legend toggle lacks an explicit aria-label')
    toggle.press('Enter')
    require(toggle.get_attribute('aria-label') and toggle.get_attribute('aria-label').strip(),
            'Collapsed legend toggle loses its explicit aria-label')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_kb_keyboard_controls(ctx):
    f = create(ctx, kind_named('kb_document'))
    browser = open_knowledge(ctx, f)
    browser.frame.get_by_role('button', name='Add document', exact=True).click()
    browse = browser.frame.locator('.ap-drop')
    browse.wait_for(timeout=30000)
    require(browse.evaluate("el => el.tagName === 'BUTTON'"), 'Upload browse action is not a native keyboard button')
    browse.focus()
    require(browse.evaluate('el => el === el.getRootNode().activeElement'), 'Upload browse action cannot receive keyboard focus')
    with browser.page.expect_file_chooser() as chooser:
        browse.press('Enter')
    require(chooser.value.is_multiple, 'Enter does not open the declared multi-file upload picker')
    chooser.value.set_files([])
    browser.frame.get_by_role('dialog', name='Add to library', exact=True).get_by_role('button', name='Close', exact=True).click()
    browser.frame.locator('.kd-row').filter(has=browser.frame.get_by_text(f.title, exact=True)).first.click()
    browser.frame.get_by_role('tab', name='Chunks', exact=False).click()
    chunk = browser.frame.locator('.dr-chunk-h').first
    chunk.wait_for(timeout=30000)
    require(chunk.evaluate("el => el.tagName === 'BUTTON'"), 'Document chunk is not a native keyboard button')
    before = chunk.get_attribute('aria-expanded')
    chunk.focus()
    chunk.press('Enter')
    require(chunk.get_attribute('aria-expanded') != before, 'Enter does not expand/collapse the document chunk')
    browser.clean(ctx.state['access_browser_error_start'])
    wiki = create(ctx, kind_named('kb_wiki'))
    def attached_document():
        graph = request(ctx, 'owner', 'GET', '/kb10/' + wiki.id + '/graph')
        documents = [node for node in graph['nodes'] if node.get('node_type') == 'document']
        if not documents:
            return None
        require(len(documents) == 1 and documents[0].get('label'), 'Wiki fixture has no unique labeled document: ' + str(graph))
        return documents[0]
    document = ctx.inst.wait(attached_document, 'Wiki chunk fixture attaches its indexed document', seconds=120)
    browser = open_knowledge(ctx, wiki)
    browser.frame.get_by_role('tab', name='Wiki', exact=True).click()
    browser.frame.get_by_label('Filter the contents', exact=True).fill(document['label'])
    browser.frame.get_by_role('option').filter(has_text=document['label']).last.click()
    browser.frame.get_by_role('tab', name='Chunks', exact=False).click()
    chunk = browser.frame.locator('.wr-chunk-h').first
    chunk.wait_for(timeout=30000)
    require(chunk.evaluate("el => el.tagName === 'BUTTON'"), 'Wiki chunk is not a native keyboard button')
    before = chunk.get_attribute('aria-expanded')
    chunk.focus()
    chunk.press('Enter')
    require(chunk.get_attribute('aria-expanded') != before, 'Enter does not expand/collapse the Wiki chunk')
    browser.frame.get_by_role('tab', name='Graph', exact=True).click()
    controls = browser.frame.locator('.gx-controls')
    controls.wait_for(timeout=30000)
    for label in ('Zoom in', 'Zoom out', 'Fit to view'):
        control = controls.get_by_role('button', name=label, exact=True)
        require(control.get_attribute('aria-label') == label, 'Graph icon lacks its explicit label: ' + label)
        control.focus()
        require(control.evaluate('el => el === el.getRootNode().activeElement'), 'Graph control cannot receive focus: ' + label)
        before = controls.locator('.gx-ctrl-read').inner_text()
        control.press('Enter')
        if label != 'Fit to view':
            require(controls.locator('.gx-ctrl-read').inner_text() != before, 'Enter does not change graph zoom: ' + label)
    browser.clean(ctx.state['access_browser_error_start'])


def regression_upload_error_presentation(ctx):
    import re
    import zipfile
    pdf = ctx.inst.root / 'temp' / 'northwind-corrupt.pdf'
    pdf.write_bytes(b'%PDF-1.4\nNorthwind corrupt PDF without a cross-reference table\n')
    archive = ctx.inst.root / 'temp' / 'northwind-over-limit-members.zip'
    with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as zip:
        for index in range(10001): zip.writestr('northwind-' + str(index) + '.txt', 'Northwind')
    checks = Checks(ctx, 'regression_QA-P07', 'owner')
    for path in (pdf, archive):
        def inspect(p=path):
            uid = upload_path(ctx, 'owner', p, failed=True)
            row = request(ctx, 'owner', 'GET', '/uploads/' + uid)['upload']
            message = str(row.get('error') or row.get('error_message') or row.get('meta', {}).get('error'))
            require(message and message != 'None', 'Failed upload lacks a user-facing error')
            require(not re.search(r'(?:kickside\.|:[0-9]+:|stack traceback)', message), 'Upload error exposes internal source details: ' + message)
            browser = actor_browser(ctx, 'owner')
            browser.frame.locator('.sidebar').get_by_text('Uploads', exact=True).click()
            browser.frame.get_by_text(p.name, exact=True).first.wait_for(timeout=30000)
            require(not re.search(r'(?:kickside\.|:[0-9]+:)', browser.frame.locator('kickside-uploads').inner_text()),
                    'Uploads UI exposes internal registry/source details')
            browser.clean(ctx.state['access_browser_error_start'])
        checks.check(path.suffix, inspect)
    checks.finish()


def regression_recurring_json_ui(ctx):
    install_package(ctx, 'kickside/automation')
    renew(ctx)
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('Automations', exact=True).click()
    panel = browser.frame.locator('kickside-automation-manager')
    panel.get_by_role('button', name='Add automation', exact=False).first.click()
    dialog = browser.frame.get_by_role('dialog')
    dialog.get_by_role('button', name='Recurring task', exact=False).click()
    dialog.get_by_label('Title', exact=False).fill('Northwind JSON recurring regression')
    dialog.get_by_label('Enabled', exact=True).check()
    dialog.get_by_label('Schedule (UTC)', exact=False).fill(json.dumps({'type': 'interval', 'expression': '2s'}))
    dialog.get_by_label('Target', exact=False).fill(json.dumps({'kind': 'action', 'port': 'kickside.inbox.sink:inbox_item',
           'config': {}, 'input': {'title': 'Northwind JSON notice', 'content': 'NW-JSON-731'}}))
    with browser.page.expect_response(lambda r: '/api/v1/automations' in r.url and r.request.method == 'POST') as response:
        dialog.get_by_role('button', name='Create', exact=True).click()
    value = response.value.json()
    require(response.value.status == 200 and value.get('id'), 'Recurring UI does not submit structured objects: ' + str(value))
    aid = value['id']
    config = request(ctx, 'write', 'GET', '/automations/' + aid + '/config')
    require('interval' in json.dumps(config) and 'Northwind JSON notice' in json.dumps(config), 'Recurring UI config does not persist')
    ctx.inst.wait(lambda: 'NW-JSON-731' in json.dumps(request(ctx, 'write', 'GET', '/inbox?status=all&limit=200')),
                  'UI recurring task delivers its event', seconds=60)
    browser.clean(ctx.state['access_browser_error_start'])


def regression_mcp_setup_mounted(ctx):
    install_package(ctx, 'kickside/mcp-client')
    renew(ctx)
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('Connections', exact=True).click()
    panel = browser.frame.locator('kickside-connections')
    panel.locator('.setup-card').filter(has_text='Remote MCP').click()
    dialog = browser.frame.get_by_role('dialog')
    dialog.get_by_label('Connection name', exact=True).wait_for(timeout=30000)
    dialog.get_by_label('Streamable HTTP endpoint', exact=False).wait_for(timeout=30000)
    dialog.get_by_role('button', name='Verify and connect', exact=False).wait_for(timeout=30000)
    require(dialog.locator('.kx-engine-host').evaluate('el => el.children.length') > 0, 'Remote MCP configure host is empty')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_provider_scroll(ctx):
    install_package(ctx, 'kickside/atlassian')
    renew(ctx)
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('Connections', exact=True).click()
    panel = browser.frame.locator('kickside-connections')
    panel.get_by_role('button', name='Connect app', exact=False).first.click()
    dialog = browser.frame.get_by_role('dialog')
    tile = dialog.get_by_role('button', name='Atlassian', exact=False)
    tile.scroll_into_view_if_needed()
    tile.click()
    field = dialog.get_by_label('Connection name', exact=True)
    field.wait_for(timeout=30000)
    geometry = field.bounding_box()
    ctx.inst.evidence.save('provider-setup-geometry', geometry)
    require(geometry and geometry['y'] >= 0 and geometry['y'] + geometry['height'] <= 1000,
            'Provider selection clips the setup name input: ' + str(geometry))
    browser.clean(ctx.state['access_browser_error_start'])


def regression_mcp_reconnect_identity(ctx):
    f = create(ctx, kind_named('remote_mcp'), 'write')
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('Connections', exact=True).click()
    panel = browser.frame.locator('kickside-connections')
    before = request(ctx, 'write', 'GET', '/connections')['connections']
    panel.locator('.live-card').filter(has_text=f.title).get_by_role('button', name='Reconnect ' + f.title, exact=True).click()
    dialog = browser.frame.get_by_role('dialog')
    dialog.get_by_label('Connection name', exact=True).fill(f.title)
    dialog.get_by_label('Streamable HTTP endpoint', exact=False).fill(ctx.inst.base + '/api/public/mcp')
    dialog.get_by_role('radio', name='Access token', exact=False).click()
    dialog.get_by_label('Access token', exact=False).fill(f.data['credential']['token'])
    with browser.page.expect_response(lambda r: ('/api/v1/connections' in r.url or '/api/v1/mcp-client/connections/' in r.url)
                                      and r.request.method in ('POST', 'PUT')) as response:
        dialog.get_by_role('button', name='Verify and', exact=False).click()
    value = response.value.json()
    require(response.value.status == 200 and value.get('success'), 'MCP reconnect fails: ' + str(value))
    after = request(ctx, 'write', 'GET', '/connections')['connections']
    require({r['component_id'] for r in before} == {r['component_id'] for r in after},
            'Reconnect creates a second connection instead of preserving the selected ID: ' + str(after))
    require(response.value.request.method == 'PUT' and f.id in response.value.url,
            'Reconnect does not update the selected component: ' + response.value.request.method + ' ' + response.value.url)
    browser.clean(ctx.state['access_browser_error_start'])


def regression_mcp_nested_trait_selection(ctx):
    f = create(ctx, kind_named('kb_document'), 'write')
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('MCP', exact=True).click()
    panel = browser.frame.locator('kickside-mcp')
    panel.get_by_role('button', name='New credential', exact=True).first.click()
    dialog = browser.frame.get_by_role('dialog', name='New MCP credential', exact=True)
    dialog.get_by_label('Label', exact=True).fill('Northwind UI KB Reader')
    dialog.locator('.kx-browse-row').filter(has_text=f.title).get_by_role('button', name='Add', exact=True).click()
    picker = dialog.locator('wc-agent-trait-picker')
    picker.get_by_role('combobox', name='Search capabilities to add', exact=True).click()
    browser.frame.get_by_role('option').filter(has=browser.frame.get_by_text('Knowledge Base Reader', exact=True)).click()
    picker.locator('wc-kb-picker input.ms-input').fill(f.title)
    browser.frame.locator('.ms-row').filter(has_text=f.title).click()
    require('0 active' not in picker.inner_text(), 'Nested KB change clears the active capability selection')
    with browser.page.expect_response(lambda r: '/api/v1/mcp/credentials' in r.url and r.request.method == 'POST') as response:
        dialog.get_by_role('button', name='Create credential', exact=True).click()
    body = response.value.request.post_data_json
    value = response.value.json()
    require(response.value.status == 200 and value.get('credential'), 'UI credential does not mint')
    credential = value['credential']
    ctx.inst.redactor.add(credential['token'])
    ctx.inst.evidence.save('ui-mint-request', body)
    trait = 'kickside.knowledge.traits:kb_reader'
    require(trait in body['allowed_trait_ids'] and trait in body['default_active_trait_ids'], 'UI mint loses reader capability')
    require(f.id in json.dumps(body['agent_profile']['trait_contexts'][trait]), 'UI mint loses nested KB context')
    result = rpc(ctx, credential, 'tools/call', {'name': 'QueryKB', 'arguments': {'fields': ['content', 'meta'], 'query': 'Northwind return window'}})
    require('17' in json.dumps(result), 'UI minted credential cannot query its selected KB')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_workflow_history_ui(ctx):
    install_package(ctx, 'kickside/workflows')
    renew(ctx)
    title = 'Northwind Member Review history'
    fid = request(ctx, 'write', 'POST', '/workflows', {'title': title})['flow_id']
    browser = actor_browser(ctx, 'write')
    browser.frame.locator('.sidebar').get_by_text('Workflows', exact=True).click()
    browser.frame.get_by_text(title, exact=True).first.click()
    start = browser.frame.locator('.wf-palette-item').filter(has=browser.frame.get_by_text('Start', exact=True)).first
    start.drag_to(browser.frame.locator('.vue-flow__pane'), target_position={'x': 300, 'y': 180})
    def count_starts():
        state = request(ctx, 'write', 'GET', '/workflows/' + fid + '/state')['state']
        return sum(n['kind'] == 'start' for n in state['nodes'])
    ctx.inst.wait(lambda: count_starts() == 1, 'UI Start placement persists', seconds=35)
    browser.frame.get_by_role('button', name='Undo', exact=True).click(timeout=30000)
    ctx.inst.wait(lambda: count_starts() == 0, 'UI Undo removes the added Start', seconds=35)
    browser.frame.get_by_role('button', name='Redo', exact=True).click(timeout=30000)
    ctx.inst.wait(lambda: count_starts() == 1, 'UI Redo restores exactly one Start', seconds=35)
    browser.clean(ctx.state['access_browser_error_start'])


UI_REGRESSIONS = (
    ('QA-014', regression_manual_document_title, True, http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/ingest') + http('kickside/kb10', 'GET /api/v1/kb10/{id}/browse'), True),
    ('N01', regression_updated_document_ui_delete, True, http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/ingest', 'DELETE /api/v1/knowledge/{id}/content', 'POST /api/v1/knowledge/{id}/search', 'POST /api/v1/knowledge/{id}/ask'), True),
    ('QA-P07', regression_upload_error_presentation, False, http('kickside/uploads', 'POST /api/v1/uploads', 'GET /api/v1/uploads/{uuid}'), False),
    ('QA-P08', regression_kb_keyboard_controls, True, http('kickside/kb10', 'GET /api/v1/kb10/{id}/browse', 'GET /api/v1/kb10/{id}/graph') +
     http('kickside/knowledge', 'POST /api/v1/knowledge/create', 'POST /api/v1/knowledge/{id}/ingest', 'POST /api/v1/knowledge/{id}/search'), True),
    ('QP-003', regression_recurring_json_ui, False, http('kickside/automation', 'POST /api/v1/automations', 'GET /api/v1/automations/{id}/config') + http('kickside/inbox', 'GET /api/v1/inbox'), True),
    ('QP-018', regression_mcp_setup_mounted, False, http('kickside/connection', 'GET /api/v1/connections/providers/discover'), True),
    ('R2-004', regression_mcp_nested_trait_selection, True, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp', 'GET /api/v1/mcp/components'), True),
    ('QA-R2-02', regression_mcp_reconnect_identity, True, http('kickside/connection', 'POST /api/v1/connections', 'GET /api/v1/connections') + http('kickside/mcp-client', 'PUT /api/v1/mcp-client/connections/{id}'), True),
    ('QA-R2-03', regression_wiki_keyboard, True, http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/ingest') + http('kickside/kb10', 'POST /api/v1/kb10/{id}/concepts', 'GET /api/v1/kb10/{id}/graph'), True),
    ('QA-R2-04', regression_graph_legend_label, True, http('kickside/kb10', 'GET /api/v1/kb10/{id}/graph'), True),
    ('QA-R2-06', regression_provider_scroll, False, http('kickside/connection', 'GET /api/v1/connections/providers/discover'), True),
    ('R2-010', regression_workflow_history_ui, False, http('kickside/workflows', 'POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply', 'POST /api/v1/workflows/{id}/undo', 'POST /api/v1/workflows/{id}/redo', 'GET /api/v1/workflows/{id}/state'), True),
)
for regression_id, operation, model, features, instance_setup in UI_REGRESSIONS:
    require(not (set(features) - INVENTORY_IDS), 'Unknown UI regression features: ' + str(set(features) - INVENTORY_IDS))
    REGISTRY.append(Scenario('regression_' + regression_id, 'full', operation, ('access_users',), model=model,
                    features=features, instance_setup=instance_setup, issues=REGRESSION_ISSUES.get(regression_id, ())))


def configured_agent(ctx, role, title, *, traits=(), contexts=None, delegates=None, options=None, agent_options=None, prompt='Use only fictional Northwind data.'):
    catalog = request(ctx, role, 'GET', '/agents/traits')['traits']
    missing = set(traits) - {t['id'] for t in catalog}
    require(not missing, 'Agent fixture traits are absent from the live capability catalog: ' + str(sorted(missing)))
    body = {'title': title, 'model': ctx.state['model'], 'traits': list(traits),
            'trait_contexts': contexts or {}, 'prompt': prompt}
    if delegates is not None: body['delegates'] = delegates
    if options is not None: body['trait_agent_options'] = options
    if agent_options is not None: body['agent_options'] = agent_options
    aid = request(ctx, role, 'POST', '/user-agents', body)['component_id']
    agent = request(ctx, role, 'GET', '/user-agents/' + aid)['agent']
    require(agent['title'] == title and agent['traits'] == list(traits), 'Configured agent does not persist its declared capabilities')
    if agent_options is not None:
        require(agent.get('agent_options') == agent_options, 'Configured agent does not persist its checkpoint policy')
    return aid


def regression_gemma_tool_names(ctx):
    install_package(ctx, 'kickside/code-execution')
    install_package(ctx, 'kickside/webscout')
    install_package(ctx, 'kickside/crm')
    renew(ctx)
    crm = create(ctx, kind_named('crm'))
    catalog = request(ctx, 'owner', 'GET', '/agents/traits')['traits']
    def declared(title):
        matches = [t['id'] for t in catalog if t.get('title') == title]
        require(len(matches) == 1, 'Live capability catalog has no unique ' + title + ': ' + str(matches))
        return matches[0]
    code_trait, research_trait, crm_trait = [declared(t) for t in ('Code Execution', 'Web Researcher', 'CRM')]
    aid = configured_agent(ctx, 'owner', 'Northwind Tool Runner regression',
            traits=(code_trait, research_trait, crm_trait),
            contexts={crm_trait: {'crm_id': crm.id}},
            prompt='Use the advertised tools. Execute arithmetic with RunCode and use CRM tools for records.')
    messages = chat(ctx, aid, 'owner', 'Use RunCode to calculate 24*12. Then use CRMListRecords for the fictional person object '
                       'in CRM ' + crm.id + '. Explain the arithmetic and CRM tool results in ten numbered points, '
                       'one full sentence per point. End with NW-TOOLS-731.', 'NW-TOOLS-731')
    text = json.dumps(messages)
    require('288' in text and 'RunCode' in text and 'CRMListRecords' in text, 'Tool Runner does not execute its advertised code and CRM tools')
    require('Failed to get tool schema' not in text and 'tool is not allowed' not in text and 'unmarshal' not in text,
            'Gemma tool-name/schema integration fails: ' + text)


def regression_artifact_bar_chart(ctx):
    aid = configured_agent(ctx, 'write', 'Northwind chart guidance regression', traits=('kickside.agents.traits:artifacts',))
    messages = chat(ctx, aid, 'write', 'Use CreateArtifact to create one standalone text/html bar-chart artifact comparing '
                    'Demo A=47 and Demo B=12. Use self-contained SVG with two bars and visible labels. '
                    'Do not use code fences or external scripts. End your reply with NW-CHART-731.', 'NW-CHART-731')
    require('CreateArtifact' in json.dumps(messages), 'Chart request never creates a durable artifact')
    sid = ctx.state['access_last_chat']['session']
    artifacts = request(ctx, 'write', 'GET', '/artifacts?session_id=' + sid)['artifacts']
    require(artifacts, 'Chart tool does not persist an artifact')
    browser = ctx.state['access_last_chat']['browser']
    require(len(artifacts) == 1 and artifacts[0]['content_type'] == 'text/html', 'Chart output is not one standalone HTML artifact')
    artifact = artifacts[0]
    response = browser.context.request.get(ctx.inst.base + '/api/v1/artifact/' + artifact['artifact_id'] + '/content',
                    headers={'Authorization': 'Bearer ' + identity(ctx, 'write')}, timeout=30000)
    require(response.status == 200, 'Persisted chart content is unavailable: HTTP ' + str(response.status))
    content = response.text()
    ctx.inst.evidence.save('persisted-chart-content', {'artifact': artifact, 'content': content})
    require('```' not in content, 'HTML chart artifact contains a literal code fence')
    preview = browser.open_page()
    try:
        preview.set_content(content, wait_until='networkidle', timeout=30000)
        plotted = preview.locator('svg')
        plotted.wait_for(timeout=30000)
        labels = ' '.join(plotted.locator('text').all_text_contents())
        require(plotted.locator('rect').count() >= 2 and '47' in labels and '12' in labels,
                'Persisted HTML chart does not render both requested bars and data labels')
    finally:
        preview.close()
    browser.clean(ctx.state['access_browser_error_start'])


def regression_delegation(ctx):
    suffix = uuid.uuid4().hex[:8]
    specialist_title = 'Northwind Calculation Specialist ' + suffix
    specialist = configured_agent(ctx, 'write', specialist_title,
                    prompt='Compute fictional arithmetic, show the calculation, and include NW-DELEGATE-144.')
    router = configured_agent(ctx, 'write', 'Northwind Routing Assistant ' + suffix, delegates=[specialist],
                    prompt='Delegate arithmetic to ' + specialist_title + ' using the delegation tool. Do not calculate yourself.')
    messages = chat(ctx, router, 'write', 'Delegate 24 times 12 to ' + specialist_title + '. '
                     'Use the advertised delegate tool. Explain its result and calculation in ten numbered points, '
                     'one full sentence per point. End with NW-DELEGATE-144.', 'NW-DELEGATE-144')
    text = json.dumps(messages)
    require('delegat' in text.lower() and '288' in text, 'Routing chat produces no delegated calculation result')
    require('Failed to execute delegation workflow' not in text and 'not allowed to register name' not in text,
            'Delegation fails its orchestrator admission: ' + text)


def regression_durable_memory(ctx):
    install_package(ctx, 'kickside/memory')
    renew(ctx)
    f = create(ctx, kind_named('kb_document'), 'write')
    trait = 'kickside.memory:knowledge_based_memory'
    aid = configured_agent(ctx, 'write', 'Northwind Durable Memory Assistant', traits=(trait,),
              contexts={trait: {'kb_ids': [f.id], 'write_kb_id': f.id, 'read_enabled': True,
                        'write_enabled': True, 'checkpoint_model': ctx.state['model']}},
              agent_options={'checkpoint': {'model': ctx.state['model'], 'token_threshold': 2000}},
              prompt='Preserve stable fictional Northwind preferences in configured memory. Recall prior chats using MemoryRecall.')
    ledger = '\n'.join(f'Northwind dispatch {number:03d}: the delivery team consistently prefers green tea for every '
              'scheduled shipment and retains the agreed return window of 29 days under marker NW-LONGTERM-144-921.'
              for number in range(100))
    messages = chat(ctx, aid, 'write', 'Remember durably that the fictional Northwind delivery team prefers green tea; '
              'marker NW-LONGTERM-144-921 and return window 29 days. Give a detailed policy in ten numbered points '
              'and end with NW-MEMORY-SOURCE. Use this fictional ledger as context; keep the answer to ten points:\n'
              + ledger, 'NW-MEMORY-SOURCE')
    require(any(m['type'] == 'assistant' and m.get('metadata', {}).get('tokens', {}).get('context_tokens', 0) > 2000
                for m in messages), 'Durable memory chat does not cross its configured checkpoint threshold')
    browser = ctx.state['access_last_chat']['browser']
    browser.stream('Preserve that green tea preference, marker NW-LONGTERM-144-921 and return window in long-term memory. '
                   'Give twenty numbered dispatch checks with two full sentences per check and end with NW-MEMORY-WRITTEN.',
                   'NW-MEMORY-WRITTEN', 'memory-write-followup')
    ctx.inst.wait(lambda: 'NW-LONGTERM-144-921' in json.dumps(request(ctx, 'write', 'POST', '/knowledge/' + f.id + '/search',
                        {'query': 'NW-LONGTERM-144-921', 'limit': 5})), 'Durable memory writes into the owned KB', seconds=120)
    recall = chat(ctx, aid, 'write', 'Use persisted memory from prior chats to state Northwind delivery team preference '
                  'and return window. Do not guess. Explain twenty numbered dispatch checks with two full sentences per check '
                  'and end with NW-MEMORY-RECALL.', 'NW-MEMORY-RECALL')
    text = '\n'.join(str(m.get('data', '')) for m in recall if m['type'] == 'assistant').lower()
    require('green tea' in text and '29' in text and 'nw-longterm-144-921' in text,
            'Independent recall does not retrieve the persisted preference and source marker: ' + text)
    require('no relevant long-term memory' not in text, 'Independent recall reports no durable memory')


def regression_crm_sdk_schema(ctx):
    install_package(ctx, 'kickside/code-execution')
    renew(ctx)
    f = create(ctx, kind_named('crm'))
    credential = request(ctx, 'owner', 'POST', '/mcp/credentials', {'label': 'Northwind CRM SDK schema regression',
          'whitelist_component_ids': [f.id], 'allowed_tool_ids': [
          'kickside.code_execution.traits:run_code_tool', 'kickside.crm.traits:list_records']})['credential']
    ctx.inst.redactor.add(credential['token'])
    direct = rpc(ctx, credential, 'tools/call', {'name': 'CRMListRecords', 'arguments': {'crm_id': f.id, 'object_type': 'person'}})
    direct_text = json.dumps(direct)
    require('rows' in direct_text and 'total' in direct_text and 'limit' in direct_text and 'offset' in direct_text,
            'Direct CRM list produces no declared paginated row collection: ' + direct_text)
    code = 'local result = sdk.tools["kickside.crm.traits:list_records"]({crm_id="' + f.id + '",object_type="person"}); return result'
    value = rpc(ctx, credential, 'tools/call', {'name': 'RunCode', 'arguments': {'code': code}})
    text = json.dumps(value)
    require('rows' in text and 'total' in text and 'limit' in text and 'offset' in text and
            'unmarshal' not in text and 'invalid arguments' not in text and 'success": false' not in text,
            'CRM SDK list does not return the actual typed record collection: ' + text)


def regression_cross_user_review(ctx):
    install_package(ctx, 'kickside/workflows')
    renew(ctx)
    reviewer = create(ctx, kind_named('inbox'), 'write')
    request(ctx, 'write', 'POST', '/component-access/share/' + reviewer.id,
            {'subject_id': ctx.state['access_owner_id'], 'subject_type': 'user', 'access_mask': 2})
    fid = request(ctx, 'owner', 'POST', '/workflows', {'title': 'Northwind cross-user review'})['flow_id']
    schema = {'type': 'object', 'additionalProperties': True}
    commands = [
      {'type': 'node.add', 'body': {'id': 'start', 'kind': 'start', 'config': {'input_schema': schema}}},
      {'type': 'node.add', 'body': {'id': 'review', 'kind': 'block', 'ref': 'kickside.inbox:block.review',
       'config': {'assignee': actor_id(ctx, 'write'), 'prompt': 'Approve Northwind 47+12 crates.',
                  'form': {'type': 'object', 'properties': {'reason': {'type': 'string'}}}}}},
      {'type': 'node.add', 'body': {'id': 'finish', 'kind': 'finish', 'config': {'output_schema': schema}}},
      {'type': 'edge.add', 'body': {'id': 'start-review', 'source_id': 'start', 'source_port': 'out', 'target_id': 'review', 'target_port': 'in'}},
      {'type': 'edge.add', 'body': {'id': 'review-finish', 'source_id': 'review', 'source_port': 'then', 'target_id': 'finish', 'target_port': 'in'}}]
    request(ctx, 'owner', 'POST', '/workflows/' + fid + '/apply', {'commands': commands})
    run = request(ctx, 'owner', 'POST', '/workflows/' + fid + '/simulate', {'input': {'marker': 'NW-REVIEW-731'}})
    rid = run['run_id']
    def assigned():
        detail = request(ctx, 'owner', 'GET', '/workflows/' + fid + '/runs/' + rid)
        require(detail['run']['status'] != 'failed', 'Cross-user draft review fails: ' + str(detail))
        items = request(ctx, 'write', 'GET', '/inbox?status=all&limit=200')['items']
        matches = [i for i in items if i.get('metadata', {}).get('run_id') == rid]
        return matches[0] if len(matches) == 1 else None
    item = ctx.inst.wait(assigned, 'Assigned cross-user Inbox review', seconds=60)
    request(ctx, 'write', 'POST', '/inbox/' + item['id'] + '/approve', {'form': {'reason': 'Northwind review approved'}})
    value = ctx.inst.wait(lambda: v if (v := request(ctx, 'owner', 'GET', '/workflows/' + fid + '/runs/' + rid))['run']['status'] == 'completed' else None,
                          'Cross-user approval reaches Finish', seconds=60)
    require('Northwind review approved' in json.dumps(value['run']['output']), 'Cross-user approval loses the form result')


def regression_member_oauth(ctx):
    import base64
    from urllib.parse import urlencode, urlsplit, parse_qs
    f = create(ctx, kind_named('kb_document'), 'write')
    redirect = 'https://northwind.invalid/mcp/callback'
    client = ctx.inst.request('POST', '/api/public/mcp/oauth/register', {'client_name': 'Northwind member PKCE',
                      'redirect_uris': [redirect], 'token_endpoint_auth_method': 'none'}, auth='',
                      expect=(201,), label='member-oauth-register')
    verifier = uuid.uuid4().hex + uuid.uuid4().hex
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).decode().rstrip('=')
    params = {'response_type': 'code', 'client_id': client['client_id'], 'redirect_uri': redirect,
              'code_challenge': challenge, 'code_challenge_method': 'S256', 'state': 'NW-OAUTH-731',
              'resource': ctx.inst.base + '/api/public/mcp'}
    context = request(ctx, 'write', 'GET', '/mcp/oauth/context?' + urlencode(params))
    require(context['client_name'] == 'Northwind member PKCE', 'Member consent context loses client identity')
    consent = request(ctx, 'write', 'POST', '/mcp/oauth/consent', {'params': params,
                    'mint': {'label': 'Northwind member OAuth reader', 'whitelist_component_ids': [f.id],
                    'policy': {'allowed_tool_ids': ['kickside.knowledge.traits:kb_search_tool']}}})
    redirect_url = consent['redirect_to']
    query = parse_qs(urlsplit(redirect_url).query)
    require(query.get('state') == ['NW-OAUTH-731'] and query.get('code'), 'Member OAuth consent does not issue its PKCE code')
    code = query['code'][0]
    ctx.inst.redactor.add(code)
    form = urlencode({'grant_type': 'authorization_code', 'client_id': client['client_id'], 'redirect_uri': redirect,
                      'code': code, 'code_verifier': verifier}).encode()
    value = ctx.inst.request('POST', '/api/public/mcp/oauth/token', raw=form, auth='',
                headers={'Content-Type': 'application/x-www-form-urlencoded'}, label='member-oauth-exchange')
    ctx.inst.redactor.add(value['access_token'])
    result = rpc(ctx, {'token': value['access_token']}, 'tools/call', {'name': 'SearchKB',
                                  'arguments': {'fields': ['content', 'meta'], 'kb_id': f.id, 'query': 'NW-ACCESS-731'}})
    require('NW-ACCESS-731' in json.dumps(result), 'Member OAuth credential cannot retrieve scoped content')
    replay = ctx.inst.request('POST', '/api/public/mcp/oauth/token', raw=form, auth='',
             headers={'Content-Type': 'application/x-www-form-urlencoded'}, expect=(400,), label='member-oauth-code-replay')
    require(replay.get('error'), 'OAuth authorization code permits replay')


def regression_lifecycle_reaper(ctx):
    email = 'northwind-reap-' + uuid.uuid4().hex[:8] + '@smoke.example'
    user = request(ctx, 'owner', 'POST', '/users', {'email': email, 'password': ctx.inst.password,
                       'full_name': 'Northwind reaper fixture', 'groups': ['app.security:members']}, expect=(201,))['user']
    uid = user['user_id']
    value = request(ctx, 'owner', 'DELETE', '/users/' + uid, expect=(202,))
    require(value.get('result', {}).get('deletion_queued') is True
            and value['result'].get('user_id') == uid, 'User deletion does not enqueue its lifecycle work')
    def gone():
        users = request(ctx, 'owner', 'GET', '/users')
        log = (ctx.inst.root / 'runtime.log').read_text()
        require('lifecycle reaper requires a framework caller' not in log, 'Lifecycle reaper rejects its framework worker context')
        return users if uid not in json.dumps(users) and email not in json.dumps(users) else None
    ctx.inst.wait(gone, 'Lifecycle physically removes the queued fictional user', seconds=60)


def regression_retention_worker(ctx):
    request(ctx, 'owner', 'GET', '/settings?namespace=kickside.retention')
    log = (ctx.inst.root / 'runtime.log').read_text()
    require('marked thread state unavailable: missing result' not in log,
            'Component retention owner reports unavailable marked-thread state')
    jobs = request(ctx, 'owner', 'GET', '/jobs')
    ctx.inst.evidence.save('retention-health-observation', {'jobs': jobs, 'runtime_log': str(ctx.inst.root / 'runtime.log')})
    require('retention' in log.lower(), 'Fresh runtime produces no retention worker observation')


EXTRA_REGRESSIONS = (
    ('QP-004', regression_gemma_tool_names, True, http('kickside/agents', 'POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}'), True),
    ('QP-008', regression_cross_user_review, False, http('kickside/workflows', 'POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply', 'POST /api/v1/workflows/{id}/simulate', 'GET /api/v1/workflows/{id}/runs/{run_id}') + http('kickside/inbox', 'GET /api/v1/inbox', 'POST /api/v1/inbox/{id}/{action}'), True),
    ('QA-010', regression_artifact_bar_chart, True, http('kickside/agents', 'POST /api/v1/user-agents'), False),
    ('R2-007', regression_delegation, True, http('kickside/agents', 'POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}'), False),
    ('R2-008', regression_durable_memory, True, http('kickside/agents', 'POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}') + http('kickside/knowledge', 'POST /api/v1/knowledge/{id}/search'), True),
    ('R2-009', regression_crm_sdk_schema, False, http('kickside/mcp', 'POST /api/v1/mcp/credentials', 'POST /api/public/mcp'), True),
    ('R2-011', regression_lifecycle_reaper, False, http('kickside/users', 'POST /api/v1/users', 'DELETE /api/v1/users/{id}', 'GET /api/v1/users'), False),
    ('R2-012', regression_retention_worker, False, http('kickside/settings', 'GET /api/v1/settings') + http('kickside/core', 'GET /api/v1/jobs'), False),
    ('R2-013', regression_member_oauth, True, http('kickside/mcp', 'POST /api/public/mcp/oauth/register', 'GET /api/v1/mcp/oauth/context', 'POST /api/v1/mcp/oauth/consent', 'POST /api/public/mcp/oauth/token', 'POST /api/public/mcp'), True),
)
for regression_id, operation, model, features, instance_setup in EXTRA_REGRESSIONS:
    require(not (set(features) - INVENTORY_IDS), 'Unknown extra regression features: ' + str(set(features) - INVENTORY_IDS))
    REGISTRY.append(Scenario('regression_' + regression_id, 'full', operation, ('access_users',), model=model,
                    features=features, instance_setup=instance_setup, issues=REGRESSION_ISSUES.get(regression_id, ())))


def regression_workflow_picker(ctx, mode):
    install_package(ctx, 'kickside/workflows')
    install_package(ctx, 'kickside/research' if mode == 'research' else 'kickside/crm')
    renew(ctx)
    target = create(ctx, kind_named('agent' if mode == 'research' else 'crm'))
    if mode == 'crm':
        base = '/workspace/crm/' + target.id
        current = request(ctx, 'owner', 'GET', base + '/schema')
        current['declaration']['attributes'].append({'object_type': 'person', 'attr': 'email',
             'label': 'Email', 'data_type': 'text', 'config': {'alias': 'email'}})
        request(ctx, 'owner', 'POST', base + '/schema', {'declaration': current['declaration'],
             'expected_revision': current['revision']})
    title = 'Northwind picker ' + mode + ' regression'
    fid = request(ctx, 'owner', 'POST', '/workflows', {'title': title})['flow_id']
    ref = 'kickside.research:block.research' if mode == 'research' else 'kickside.crm.blocks:resolve'
    schema = {'type': 'object', 'additionalProperties': True}
    commands = [
      {'type': 'node.add', 'body': {'id': 'start', 'kind': 'start', 'config': {'input_schema': schema}}},
      {'type': 'node.add', 'body': {'id': 'target', 'kind': 'block', 'ref': ref,
       'title': 'Research' if mode == 'research' else 'Resolve CRM record',
       'config': {'prompt': 'Analyze fictional Northwind 47+12 crates. NW-BLOCK-731.'} if mode == 'research' else {}}},
      {'type': 'node.add', 'body': {'id': 'finish', 'kind': 'finish', 'config': {'output_schema': schema}}},
      {'type': 'edge.add', 'body': {'id': 'start-target', 'source_id': 'start', 'source_port': 'out', 'target_id': 'target', 'target_port': 'in'}},
      {'type': 'edge.add', 'body': {'id': 'target-finish', 'source_id': 'target', 'source_port': 'result', 'target_id': 'finish', 'target_port': 'in'}}]
    request(ctx, 'owner', 'POST', '/workflows/' + fid + '/apply', {'commands': commands})
    browser = actor_browser(ctx, 'owner')
    browser.frame.locator('.sidebar').get_by_text('Workflows', exact=True).click()
    browser.frame.get_by_text(title, exact=True).first.click()
    browser.frame.locator('.vue-flow__node').filter(has_text='Research' if mode == 'research' else 'Resolve CRM record').click()
    picker = browser.frame.locator('wc-agent-picker' if mode == 'research' else 'wc-component-picker')
    if mode == 'research':
        picker.locator('input.ms-input').fill(target.title)
        browser.frame.locator('.ms-row').filter(has_text=target.title).click()
    else:
        picker.locator('.cp').click()
        browser.frame.locator('.cp-row-main').filter(has=browser.frame.get_by_text(target.title, exact=True)).click()
    key = 'agent_id' if mode == 'research' else 'crm_id'
    def persisted():
        state = request(ctx, 'owner', 'GET', '/workflows/' + fid + '/state')['state']
        node = next(n for n in state['nodes'] if n['id'] == 'target')
        return node if node['config'].get(key) else None
    node = ctx.inst.wait(persisted, 'Workflow picker commits canonical ' + key, seconds=35)
    require(node['config'][key].removeprefix('user_agent:') == target.id, 'Workflow picker stores the wrong component identifier: ' + str(node['config']))
    require('thread_id' not in node['config'] and 'selected' not in node['config'], 'Workflow picker retains non-contract selector fields: ' + str(node['config']))
    input = {} if mode == 'research' else {'object_type': 'person', 'record_id': 'northwind-picker-person',
                'match': {'alias_type': 'email', 'alias_value': 'picker@northwind.example'}, 'values': {'name': 'Northwind picker person'}}
    run = request(ctx, 'owner', 'POST', '/workflows/' + fid + '/simulate', {'input': input})
    def finished():
        value = request(ctx, 'owner', 'GET', '/workflows/' + fid + '/runs/' + run['run_id'])
        require(value['run']['status'] != 'failed', 'UI-configured workflow run fails: ' + str(value))
        return value if value['run']['status'] == 'completed' else None
    value = ctx.inst.wait(finished, 'UI-configured block completes', seconds=180 if mode == 'research' else 60)
    require('research_id' in json.dumps(value['run']['output']) if mode == 'research' else
            'Northwind picker person' in json.dumps(value['run']['output']), 'UI-configured block returns no observable domain result')
    browser.clean(ctx.state['access_browser_error_start'])


def regression_crm_member_create_affordance(ctx):
    install_package(ctx, 'kickside/crm')
    renew(ctx)
    browser = actor_browser(ctx, 'read')
    browser.frame.locator('.sidebar').get_by_text('CRM', exact=True).click()
    panel = browser.frame.locator('kickside-crm-surface')
    panel.wait_for(timeout=30000)
    button = panel.get_by_role('button', name='Create CRM', exact=True)
    require(button.count() == 0 or button.is_disabled(), 'Member UI offers an enabled Create CRM action while schema administration denies it')
    browser.clean(ctx.state['access_browser_error_start'])


for regression_id, operation, features, instance_setup in (
    ('N04', lambda ctx: regression_workflow_picker(ctx, 'research'),
        http('kickside/workflows', 'POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply',
             'GET /api/v1/workflows/{id}/state', 'POST /api/v1/workflows/{id}/simulate', 'GET /api/v1/workflows/{id}/runs/{run_id}'), True),
    ('N05', lambda ctx: regression_workflow_picker(ctx, 'crm'),
        http('kickside/workflows', 'POST /api/v1/workflows', 'POST /api/v1/workflows/{id}/apply',
             'GET /api/v1/workflows/{id}/state', 'POST /api/v1/workflows/{id}/simulate', 'GET /api/v1/workflows/{id}/runs/{run_id}'), True),
    ('QP-021', regression_crm_member_create_affordance, (), True)):
    require(not (set(features) - INVENTORY_IDS), 'Unknown picker regression features: ' + str(set(features) - INVENTORY_IDS))
    REGISTRY.append(Scenario('regression_' + regression_id, 'full', operation, ('access_users',), model=True, features=features, instance_setup=instance_setup))


def regression_member_upload_resources(ctx):
    from docx import Document
    from module_harness import run_focused_module_test
    run_focused_module_test(ctx, 'kickside/uploads', 'module_access_policy')
    from scenarios_security_model import install_probe
    install_probe(ctx)
    versions = ctx.inst.request('GET', '/api/v1/security-model/requirement-choices', label='converter-version-range')
    require(versions['converter_older_selection'] != '0.8.10',
            'Upload dependency retains a converter release without its member policy: ' + str(versions))
    require(versions['converter_newer_selection'] == '0.9.0',
            'Upload dependency excludes a newer converter release: ' + str(versions))
    doc = Document()
    doc.add_paragraph('Fictional Northwind member upload NW-MEMBER-UPLOAD-731')
    path = ctx.inst.root / 'temp' / 'northwind-member-resource.docx'
    doc.save(path)
    text = ctx.inst.root / 'temp' / 'northwind-member-resource.txt'
    text.write_text('Fictional Northwind member upload NW-MEMBER-UPLOAD-731. Returns take 17 days.')
    for document in (path, text):
        uid = upload_path(ctx, 'write', document)
        listed = request(ctx, 'write', 'GET', '/uploads/list')
        require(uid in json.dumps(listed), 'Owned member upload does not appear in resource listing')
        artifacts = request(ctx, 'write', 'GET', '/uploads/' + uid + '/artifacts')
        require(artifacts.get('artifacts') or artifacts.get('primary'), 'Owned member upload exposes no converted artifacts')
        request(ctx, 'stranger', 'GET', '/uploads/' + uid, hidden=True)
    install_package(ctx, 'kickside/knowledge')
    renew(ctx)
    created = request(ctx, 'write', 'POST', '/knowledge/create', {
        'name': 'Northwind member uploaded knowledge', 'kb_type': 'document',
        'embedding_model': ctx.state['embed'], 'config': {'analysis_model': ctx.state['model']}})
    base = '/knowledge/' + created['id']
    ctx.inst.cleanup.add('member upload KB', lambda: request(ctx, 'write', 'DELETE', base))
    queued = request(ctx, 'write', 'POST', base + '/ingest-upload', {'upload_id': uid})
    require(queued.get('queued') == 1, 'Member upload does not enter its KB: ' + str(queued))
    ctx.inst.wait(lambda: 'NW-MEMBER-UPLOAD-731' in json.dumps(request(ctx, 'write', 'POST',
        base + '/search', {'query': 'NW-MEMBER-UPLOAD-731', 'limit': 5})),
        'Member converted document becomes searchable', seconds=120)
    answer = request(ctx, 'write', 'POST', base + '/ask', {
        'query': 'What is the exact Northwind member upload marker? Cite the document.'}, timeout=120)
    require('NW-MEMBER-UPLOAD-731' in str(answer.get('answer', '')),
            'Member KB answer omits the uploaded marker: ' + str(answer))
    require(answer.get('citations') or answer.get('sources') or re.search(r'\[\d+\]', str(answer.get('answer', ''))),
            'Member KB answer omits its uploaded source: ' + str(answer))


REGISTRY.append(Scenario('regression_QP-002', 'full', regression_member_upload_resources, ('access_users',), model=True,
    issues=('MEMBER-DOCUMENT-CONVERSION-DENIED',), features=http('kickside/uploads', 'POST /api/v1/uploads', 'GET /api/v1/uploads/list',
                  'GET /api/v1/uploads/{uuid}', 'GET /api/v1/uploads/{uuid}/artifacts') +
             http('kickside/knowledge', 'POST /api/v1/knowledge/create',
                  'POST /api/v1/knowledge/{id}/ingest-upload', 'POST /api/v1/knowledge/{id}/search',
                  'POST /api/v1/knowledge/{id}/ask', 'DELETE /api/v1/knowledge/{id}')))

# Findings use source-lane keys because the round-two reports reuse R2 IDs.
QA_MAP = {
 'qa-platform': {
   'QP-001': ('platform_onboarding',), 'QP-002': ('regression_QP-002',), 'QP-003': ('regression_QP-003',),
   'QP-004': ('regression_QP-004',), 'QP-005': ('regression_QA-007',), 'QP-006': ('regression_QP-006',),
   'QP-007': ('upstream:session',), 'QP-008': ('regression_QP-008',), 'QP-009': ('regression_QP-009',),
   'QP-010': ('regression_QP-010',), 'QP-011': ('regression_QP-011',), 'QP-012': ('regression_QP-012',),
   'QP-013': ('regression_QP-013',), 'QP-014': ('upstream:session',), 'QP-015': ('regression_QA-010',),
   'QP-016': ('regression_QP-016',), 'QP-017': ('regression_QP-017',), 'QP-018': ('regression_QP-018',),
   'QP-019': ('regression_QP-019',), 'QP-020': ('upstream:session',), 'QP-021': ('regression_QP-021',)},
 'qa-apps': {
   'QA-001': ('chestor_entries', 'regression_QP-006'), 'QA-002': ('regression_QA-002',),
   'QA-003': ('regression_QA-003',), 'QA-004': ('regression_QA-004',), 'QA-005': ('research',),
   'QA-006': ('upstream:web-host',), 'QA-007': ('regression_QA-007',), 'QA-008': ('regression_QA-008',),
   'QA-009': ('upstream:runtime',), 'QA-010': ('regression_QA-010',), 'QA-011': ('first_agent',),
   'QA-012': ('regression_QA-012',), 'QA-013': ('upstream:session',), 'QA-014': ('regression_QA-014',),
   'QA-015': ('regression_QP-011',)},
 'qa-flows': {
   'D01': ('first_agent',), 'D02': ('regression_QP-012',), 'D03': ('regression_QP-003',),
   'D04': ('research',), 'D05': ('regression_QA-002',), 'D06': ('regression_QA-003',),
   'D07': ('regression_QA-004',), 'D08': ('regression_QA-007',), 'D09': ('upload_attach_send',),
   'D10': ('upstream:session',), 'D11': ('regression_QP-004',), 'D12': ('regression_QA-014',),
   'D13': ('regression_QP-002',), 'D14': ('upstream:keeper-export',), 'D15': ('regression_QA-010',)},
 'qa-providers': {
   'QA-P01': ('converters', 'regression_QP-002'), 'QA-P02': ('regression_QP-018',),
   'QA-P03': ('regression_QA-P03',), 'QA-P04': ('regression_QA-P04',), 'QA-P05': ('regression_QA-P05',),
   'QA-P06': ('regression_QA-P06',), 'QA-P07': ('regression_QA-P07',), 'QA-P08': ('regression_QA-P08',),
   'QA-P09': ('upstream:web-host',)},
 'qa-platform-r2': {
   'QP-019': ('regression_QP-019',), 'R2-001': ('regression_R2-001',), 'R2-002': ('regression_R2-002',),
   'R2-003': ('regression_R2-003',), 'R2-004': ('regression_R2-004',), 'R2-005': ('regression_R2-005',),
   'R2-006': ('regression_QA-R2-02',), 'R2-007': ('regression_R2-007',), 'R2-008': ('regression_R2-008',),
   'R2-009': ('regression_R2-009',), 'R2-010': ('regression_R2-010',), 'R2-011': ('regression_R2-011',),
   'R2-012': ('regression_R2-012',), 'R2-013': ('regression_R2-013',), 'R2-014': ('upstream:web-host',),
   'R2-015': ('regression_R2-015',)},
 'qa-apps-r2': {'R2-001': ('boot',), 'R2-002': ('regression_R2-002_apps',), 'R2-003': ('regression_R2-003_apps',)},
 'qa-flows-r2': {'N01': ('regression_N01',), 'N02': ('regression_R2-009',), 'N03': ('regression_R2-015',),
                  'N04': ('regression_N04',), 'N05': ('regression_N05',), 'N06': ('regression_N06',)},
 'qa-providers-r2': {'QA-R2-01': ('regression_QA-R2-01',), 'QA-R2-02': ('regression_QA-R2-02',),
                     'QA-R2-03': ('regression_QA-R2-03',), 'QA-R2-04': ('regression_QA-R2-04',),
                     'QA-R2-05': ('regression_R2-008',), 'QA-R2-06': ('regression_QA-R2-06',)},
}


from mcp_contracts import ACCESS_SCENARIOS
REGISTRY.extend(ACCESS_SCENARIOS)
