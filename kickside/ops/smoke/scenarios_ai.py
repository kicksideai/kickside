"""Northwind agent, knowledge, skill, model, and file user journeys."""
import base64
import hashlib
import json
from pathlib import Path
import re
import urllib.request
from urllib.parse import quote, urlsplit
import uuid
import zipfile
import zlib
from datetime import datetime
from contextlib import contextmanager

from browser import ASSISTANT_TEXT, CHAT_STATE
from harness import Scenario, require
from runtime import Instance
from scenarios import Context, first_agent, platform_onboarding, nav


def http(component, *routes):
    return tuple('kickside/' + component + '|http|' + route for route in routes)


def feature(component, category, entry):
    return 'kickside/' + component + '|' + category + '|' + entry


def shot(ctx, label):
    browser, inst = ctx.browser, ctx.inst
    browser.settle()
    for frame in browser.page.frames:
        frame.evaluate('window.scrollTo(0, window.scrollY)')
    dark = browser.page.evaluate("matchMedia('(prefers-color-scheme: dark)').matches")
    browser.screenshot(label, full_page=False, mask_color='#18181b' if dark else '#f8f9fa')



def open_named_session(ctx, sid):
    ctx.browser.settle()
    title = 'Northwind ' + ctx.inst.evidence.label.removeprefix('ai_').replace('_', ' ')
    renamed = ctx.inst.request('POST', '/api/v1/sessions/rename', {'session_id': sid, 'title': title}, label='northwind-session-name')
    require(renamed.get('success'), 'Session rename does not confirm the Northwind title')
    saved = ctx.inst.request('GET', '/api/v1/sessions/get?session_id=' + sid, label='northwind-session-title')['session']
    require(saved['title'] == title, 'Session rename does not persist the Northwind title')
    ctx.browser.settle()
    ctx.browser.frame.locator('.sidebar').get_by_text('Sessions', exact=True).click()
    ctx.browser.settle()
    ctx.browser.frame.locator('kickside-sessions').first.locator('.kx-row-main').filter(has_text=title).click()
    ctx.browser.frame.locator('wippy-chat').get_by_placeholder('Type a message', exact=True).wait_for(timeout=30000)
    ctx.browser.settle()


def chat_shot(ctx, label):
    open_named_session(ctx, ctx.state['ai_current_chat'])
    latest = next(m['data'] for m in reversed(messages(ctx, ctx.state['ai_current_chat'])) if m['type'] == 'assistant' and isinstance(m['data'], str) and m['data'].strip())
    markers = re.findall(r'\bNW-[A-Z0-9-]+\b', latest)
    require(markers, 'Chat screenshot has no completed Northwind answer marker')
    ctx.inst.wait(lambda: markers[-1] in ctx.browser.frame.locator('wippy-chat').evaluate(ASSISTANT_TEXT), 'Chat evidence restores the latest completed assistant answer', seconds=30, pump=ctx.browser.page.wait_for_timeout)
    ctx.browser.frame.locator('wippy-chat .chat-message:not(.chat-message--user) .chat-message__body').last.evaluate("el => el.scrollIntoView({block:'start',inline:'nearest'})")
    shot(ctx, label)


@contextmanager
def dark_ui(ctx):
    browser = ctx.browser
    browser.settle()
    with browser.presentation(viewport={'width': 1600, 'height': 1000}, device_scale_factor=2, color_scheme='dark'):
        browser.page.goto(ctx.inst.base + '/', wait_until='domcontentloaded', timeout=30000)
        browser.frame.locator('.sidebar').wait_for(timeout=30000)
        browser.settle()
        yield


def hero_shot(ctx, label):
    with dark_ui(ctx):
        chat_shot(ctx, label + '-dark')


def ai_boot(ctx: Context):
    inst: Instance = ctx.inst
    inst.boot()
    browser = ctx.browser_check(viewport={'width': 1600, 'height': 1000}, device_scale_factor=2, color_scheme='light')
    browser.page.goto(inst.base + '/login', wait_until='domcontentloaded', timeout=30000)
    browser.page.get_by_role('button', name='Sign In', exact=True).wait_for(timeout=30000)


def ai_ready(ctx):
    platform_onboarding(ctx)
    first_agent(ctx)
    shot(ctx, 'northwind-workspace')


def install_module(component, parameters=()):
    def run(ctx):
        ctx.browser.settle()
        ctx.browser.page.goto('about:blank')
        if parameters:
            plan = ctx.inst.request('POST', '/api/v1/keeper/hub/dependencies/plan', {
                'component': component, 'run_migrations': True,
                'parameters': [{'name': name, 'value': value} for name, value in parameters],
            }, label='bound-install-plan', timeout=120)
            unresolved = [r['parameter_name'] for r in plan['requirements'] if r.get('missing') or r.get('invalid')]
            require(not unresolved, 'Hub bindings remain unresolved: ' + str(unresolved))
            body = dict(plan['install_payload'])
            body.pop('version', None)
            body['run_migrations'] = True
            installed = ctx.inst.request('POST', '/api/v1/keeper/hub/dependencies/install', body, label='bound-install', timeout=300)
            require(not [r for r in installed.get('execution', []) if r.get('status') in ('failed', 'rollback_failed')], 'Hub installation has failed steps: ' + str(installed))
        else:
            ctx.inst.install(component)
        dependencies = ctx.inst.request('GET', '/api/v1/keeper/hub/dependencies?entries=0&migrations=1', label='installed-module')['dependencies']
        require(any(row.get('component') == component for row in dependencies), 'Installed module is absent from Hub dependencies: ' + component)
        ctx.browser.page.goto(ctx.inst.base + '/', wait_until='domcontentloaded', timeout=30000)
        ctx.browser.frame.locator('.sidebar').wait_for(timeout=30000)
    return run


def ocr_config(ctx):
    for key, value in [('ocr_enabled', True), ('ocr_images', True), ('ocr_model', ctx.state['model'])]:
        ctx.inst.request('PUT', '/api/v1/settings', {'namespace': 'kickside.llm_ocr', 'key': key, 'value': value}, label='ocr-config-' + key)
    values = ctx.inst.request('GET', '/api/v1/settings?namespace=kickside.llm_ocr', label='ocr-config-persisted')['settings']
    require(any(s['key'] == 'ocr_model' and s['value'] == ctx.state['model'] for s in values), 'OCR does not persist requested Gemma model')


def ocr_settings(ctx):
    fresh_ui(ctx)
    nav(ctx, 'System')
    ctx.browser.frame.get_by_text('Settings', exact=True).first.click()
    ctx.browser.frame.get_by_text('LLM OCR', exact=True).first.click()
    ctx.browser.frame.get_by_text(re.compile(r'^Ocr model', re.I)).first.wait_for(timeout=30000)
    require(ctx.state['model'] in ctx.browser.frame.locator('kickside-settings-provider').evaluate("el => [...el.shadowRoot.querySelectorAll('input')].map(i => i.value).join('\\n')"), 'Settings UI omits configured Gemma OCR model')
    shot(ctx, 'northwind-ocr-settings')


def fresh_ui(ctx):
    ctx.browser.reload()
    ctx.browser.frame.locator('.sidebar').wait_for(timeout=30000)
    ctx.browser.settle()


def raw_bytes(ctx, path, label):
    inst = ctx.inst
    require(path.startswith('/'), 'Raw asset is not a product-relative URL')
    request = urllib.request.Request(inst.base + path, headers={'Authorization': 'Bearer ' + inst.token})
    with urllib.request.urlopen(request, timeout=35) as response:
        data, mime = response.read(), response.headers.get('Content-Type', '')
    require(data, label + ': empty asset')
    inst.evidence.save(label, {'path': path, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(), 'mime': mime})
    return data, mime


def completed_upload(ctx, path, mime=None):
    inst = ctx.inst
    if mime:
        boundary = 'northwind_' + uuid.uuid4().hex
        raw = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\nContent-Type: {mime}\r\n\r\n'.encode()
            + path.read_bytes() + f'\r\n--{boundary}--\r\n'.encode())
        uploaded = inst.request('POST', '/api/v1/uploads', raw=raw, headers={'Content-Type': 'multipart/form-data; boundary=' + boundary}, label='typed-upload', timeout=120)
    else:
        uploaded = inst.upload(path, 'upload')
    uid = uploaded['uuid']
    ctx.state.setdefault('ai_uploads', {})[path.name] = uid
    def complete():
        upload = inst.request('GET', '/api/v1/uploads/' + uid, label='upload-state')['upload']
        require(upload['status'] not in ('error', 'failed'), path.name + ': processing fails: ' + str(upload))
        return upload if upload['status'] == 'completed' else None
    inst.wait(complete, path.name + ' conversion', seconds=120, pump=ctx.browser.page.wait_for_timeout)
    artifacts = inst.request('GET', '/api/v1/uploads/' + uid + '/artifacts', label='upload-artifacts')
    if path.suffix not in ('.zip', '.bin', '.dxf', '.svg'):
        require(artifacts.get('primary'), path.name + ': no primary content artifact')
    original, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/raw', 'original-file')
    require(original == path.read_bytes(), path.name + ': download differs from original bytes')
    return uid, artifacts


POLICY = '''# Northwind delivery handbook

All companies, people, and figures in this document are fictional.
Northwind's returns window is 17 days. The warehouse marker is NW-AI-731.
Mira Chen leads the Harbor warehouse. Jonah Reed leads the Orchard warehouse.
Returns reach Harbor before 16:00. Orchard dispatches orders at 09:30.

| Product | Units | Revenue |
| --- | ---: | ---: |
| Harbor tea | 120 | 2400 |
| Orchard coffee | 80 | 3200 |
| Meadow cocoa | 95 | 2850 |

Inspect packaging, record the batch, and confirm dispatch with the customer.
'''


LEGACY_FIXTURES = {
    'doc': 'eNrtWl1sk1UYfk/7rbQFtlIGzIHwMeoGg/0xfoaAbF0Zbed+YMMh/+s6pJO2cxtxGi/40UwSSTAmmhATY4JXGoJwY7yBmMiNMS5GLjCa4JUx8aIaL8SE1ec93+moZbivHVGRvsvznfOdfue85+c573nPORv/at6t9z8u/YEyZDNZaSLpIFtamgBcqRdELCptIplMppKTeXmo5I4KeQw1jF8BwGM+C7ADDsAJzAbmAHOBQqBIjfs8Febl4ZRdFMffCOm0nWIIh+glykYWgjHp5ZnJM2HyO7OS15+7/pT9zmX+81rA898NzAeKgQWSE0SLgBLgMaAUWAwsAR4HliqdOsLlQBmwAvCo9HKEFSq+CmElsBpYA1QB1UANUAvUAWuBemAdsB7YAGwEGoBNwJNyPSPaAmwFngK2AY1AE+AFmgEfsB1oAXYAfiAABIFW4GmgDWgHOoBOYCewC+gCuoHdwDNAD7AHeBbYC+wD9gMHgIPAIeCwamPoX7ShApqtToNDNodFcuKaQY0WHr+2SN9QfDh+ZETviQ+Fq3zx549H+2MjkhNtXZzmi/dJJnC8Gi/y9+oG+m3T5Rem56Iw3IicxQ3WOVHKPslIQ9C/rl+SFtnPNoxWHFYtSr10TPG4zkOVHtHkYdZ4K6kjaKWdQHOwhKJ+hzYMtPot1BnUKObX7CPAwWAB9eL3Q/5N2t/Wp4UmlgmyiBY5h/zUD71hisC2PidnTxHNv5Cg4gujZPMIzJCOoA3KbVC+BApnkaFoCbjsdXE56+Qc9KINYVhmHRzrp1HYaZ59LnKHi4VAiZh7F95Amk41LiHcgmeWBr5GaFh+a8E81VC2HZwvQ7llolHWrw+1G8QXEZQfk/Vzo7RRWb85qJ9bdsoGu+Cs3P4DdlEuW7hWlEtrEUC+sKwTv82XuQxNpGaqZhgMyTA2KXvthpXh+H47TXqY1Sq0uNISOg0rJM3QjoXnhBlKdKM5UVRpGN3RjvBFhLzMMQW4ketRjhnOdaHDoxRCTiZOfbk57U0gW0SRLQJNYt70eXigQqjpEPKlBkOH/n5Z1hETbWq1QZOJ2rXLxT7+l9J1GL8gDBynsRncnENZXIth7jP0knDl3uJUOem9GMi5Pve0zJN7y7wgsrCfhU38HOQ5KqYwZGrZrJRvVkn90+es96RaVY4KK+kyolFeZrIE1D74MjPsgy4HaXpeZ12TEzOtaJNgvyVM3+AvU0bxN7hocNF0ZRSRlmAjc1lzJtjaCs1w+q7Q7AT7RiflV6fk87R8vmrwlpWTTTJ/m3yeFGzkT8v4VcErFBwMacFdCcXyEyGalRjjfLbPxsfPctLLb/eM2WX39k1OJCWxpQn2nDK/TJ7pbN8x11nKXsjcgN+3S3mrdjsPXGt01XfsmgZ8Td2j3/Z9sqJibEHdnCoHnfm+RRTsH+DFhALb230XvYdPzbD/J5JcY8tUfg3deu29X293HHV9+KadVldcucnseEX510L5pw7lhzqVfzlb+Y3sa4eVvz2oCPnTHcN3tiimNabpMxOfSn7+QIg23YHB+d19NZ3EanUupObe+LHeWMMU1HZoC6mm8O579bT7lCK58gsVt6XFM+V1+TyvyHDehNXmb4qzGLmLwujlT1WYkh/xfsoycxNyXeSWL4KxCKXp/1oYu6q8/D/legb//imx0M185+clL4+0GKcjI3QUW8qI3MvzBrOHqrAFC+C5keqpDmk6PJEWfNE3uV3jzZmOvP1IOY7Q2HDVIUcY3wrpueTlPy/w8faoUxn6I98dj4I4dtuFdUA5ARj3NYoD7O/fTv/w/i95eWgFuy6rh3o9xAfg1iraUkuNnUTeTgstuzReo1+6tXX5pZhWBqw4F9M8QD1+X7mWCmvN7385xXLjyxvvVi92vfUO9r9rbl/k+5WCjDS+EylR+8DU/xek9rr3S8/Lg5MHef8rT4My7pCmysMHoNdKUuckzfIYe5A6KEQDWdffDVawxgJSNwwmZWDynKYDfkx/zv3nhHbWa81CP9c3dVJRR93wpEI516EQpWl09w7YTB6u6xfq2sUnD+B7Z9z+bPQvI+PckuS4dcF7jMqLDOYe3ycdmbyxu3szdT9ZCf2p+2qz+p8APlLxHqkrjH6Iw7c9Li9ZYvLOzJyU5tD+lUDq3KzgHs3Z9UdDDvp5tFMt/BOshvTz',
    'ppt': 'eNrt3QdcVfX/x/FzAQEFBRX3wr1w4t4o7gFuzYUpOBFEcM8ytaFprtx7pGZ7+qtsW85Ky0zTxq9hS/v1+5Ul8H+fAVwRSK1/D7PX0/OGy71n33M+3++5o44dzX9u02PFPzEyaWp4GimpuQ1vt/t8lVmezh+BhuGhXy4lJTU1NW2c2coc5Q7lTmWucpd5v6ZLxU0n2fltPocz9Rzl0nNlPuc+zvOdW8mj+Cn+Sl4lnxJgHwJGfqWAUlAJUgophZ1joqh+F1OKKyWUkkoppbRSRglWyirllPJKBaWiUkmprFRRqirVlOpKiFJDqanUUmordZS6SqhST6mvNFAaKo2UxkoT63g2jGZKc6WF0lJppYQprZU2SrjSVmmntFc6KB2VTkpnpYvSVemmRCiRSnelh9JT6aX0VvoofZV+Sn/lNmWAMlAZpAxWhihRylDldmWYMlyJVmKUEcpIZZQyWhmjjFVilXFKnBKvjFcSlAlKopKkTFQmKZOVKcpUZZoyXZmhzDTPY2Wete5x+peo56Kt5puoeU0xrkdhHTHux9O1TJNyjeNdK5Z/48tPcc59M17O35nP/xTK5C3LpWfXM499DAUunDN8TcTR/PPOG7uM56p+Zdb2rj2DI+MmRSdExo0alxjcuGHmY8lldwP+gP8FVHDajiH7l3x2oYLZzvzkaXYz2iQlJERrqb0nRCf4Zjf5uNIXzNr6yrFjC80/p67oO1/jBqbeHdmtfd48xc22K2/HDuE9nOPZ19dc3c6xVU+bjVfH8LBek08Ne7Z85fmF6vjXyG3cfaadK9fA0eZRb3Rs2y18b+uoO37nXDL3wdW7wLzH48ThE2trlghcutLXqB5yaa/ZvuTKdJ/ZJhR1+lEup001p/062TD267eXM78qzuO9nMfN6cx99GTaspx2ekhkx/Ah7Xt3DLce9U+7p0OXjt0693Sbd5jTjgUYX3m+mMuc99ee5jKOaGV+UiMfrAZ9f1F7mR7Osuzn2+oRugKMHz3v8zDb4QMB/lZb/G5A2roGGCd8apsTGs8EhFvrPErtVrTaqGAtM1ptU7DV7sSq9Rt3TceId2CgM68wtVCjNN3YGz7evAODnXl1s1q+OK1NT2tNzPVro9a+s3VPm2uZl7bbnle4tmu05tJHLXDG/K5rvZStAYFWHzrGqeoFDLumPxSQx6nJHkZuX7PnsiUgn8vsraUmt1R/rFyw3XswUsweRlo7EGw9B57WvH4w7OPUnqs9ZuqMYE1b8Kppj7vMKXOa5urlHfNo5XG9y7ngGeWZ8zT7jMzTBOaa7ZXdNAGGn9ezdt/0wqP2Pr1g9hE9A+1+bUHnfHKl70u7Jng696eljOG6EGL1y3zMcmR4e2euLnZ/zeX27A0z/Kxx57iMZ1ONO62HXnIFatT9Lv3hm+ry1bJcHlqmbyuj9MVAe1m+5mPmfT9rrmWMCwElrbn9xzPIuKqwzrbPrWM+LVzmOqamjePlf2U91ji/eo6wq5NVR607kz3rOd3T4Tn83lPEME7p9P/pcmrqHqt5MLuZv3mmld9434y/zcda+ZmL8/FqZj1qL2OXciCH31kto0ymVXZlrHJTZ9Wy+23Ozl9PeZVLqanPZrHK5iq6r7K5CS2NDwPsvXbKqlp5dd0RcdV9zbVpVz4jXm4VOe3pr5VpnFxZjOMyXF6Rzk2t4qw/o9dnz/q8p31U/OLZKr95NPzgmd6aZNEuRxkXfILt2mIfUrNn2w8cWmWEhh4yDh1KNR5//PH08TRm+rjmr9RV9g3znhUrVrjN75A1TkxMqGHOMnRVqH24ap6ph2KuWm5oaKj1+8yZM9Y6dO3a1Th//ny262fOw1wXraGxf//+bMczHzNi7PUzVmW/vaGrUo1VqYes+Znba2RbT13XUE9DsqlzrhzqaUg29dSVQz39veVkVU9Dsqmnrhzqado0gTm2Mos05skc9kpwDtV/WzZ7pWQOrcy+bPZK0RyW8242eyUoh1bm62z2Sto0wTnsleHWaxsZx4pH+l5Ju0rzz3Zt7Wnd90rW01y9PPe9cq3Lcd8rWU/jvlfsadz3SuZp6JFkNU2k9srHTpuQOsPI8kw/ae0NH2vqYOcRe9ysK8mV43pkO19z+68c1zPb+ZrbfeW4Xs64WVeMK8eNuoatNNy2MuQ6tjLkOrYy5Dq2MuQ6ttIeNz7TVgZkOhPN+pDxXBpXPJcBmc7AIOdMTxu3pNtWZjXfjK20rxbTtjKr+WZspX1VnbaVAZnO6CDnjL5y3Nn/iK0MMPy9vvIxx/a4cM7HbOl8L/g612DeXub9nhcSrMe9LthX3rkvBGbRB/Mw8ljTmR2+XIYzvst8Rdy+3+VlX5YXdPldiLFedXVpvNlGhdpTrXdJ7PdJ7HdK7PdK7HdLDGO+1Yd7STP20nWEYV1H1IvearzsMtfmFdfB1QOM960u3gfWz1T1mHM71xca21r6IqO5a1GRTzwCjPzO2r/q53vFdULeC1OtvzakX7sbxsaAds61cYKujkfqan2UrmWH65mO0a1hum+UHhtnXYMH6/5Rus6N11+Jemyk7jEfT9Rj0daxsTnA23ll37TLuo6tkN57NfdXB4+M/eWRzf763G1/Bcdlvb+MP7y/wj3N/XWbX+b9lfYaiLm/Vrsy9pfL2V+tXX9kf8UZSdb+Mqcy91leo6d+D7OmMecSkenxLvo5URmrMXtpDqO03N8br521FHv9fm/MUVrv3x+vp+6dfE3j2bfH/e645rFSzXn3x34fyWXYV7iFrD6Dff1l/yzsXGPZx9PmK46nBm7nn2c2x9PWOr9//v3x4+lp/+YuT5/debM6/3yc42mB2/nn5RxPHs57XpsCKqW/CuCKz2Ok9TjSXm+sbO2XfE6t+8XtHN4VkM+Z0m7BclsPbQ8o7iwll9UTM/tHZn/H7L9szGXvvWZue88rm703vO5fs/em+/YLzHrv5Xb23qIc916VjL03OY+1H3Laf79ew/4r7SzHO8v9t9/b3oeN3PZhruxagNC/Zh++lf/+IlnvQ19nH95zXfvQI8d9ePJPOQa7GRn7z2yNzZexlxp+F8z3dWdbJ/oc81UU51W2pa7x76wylrlWzOuj/VRAU5ovAfkaXp7mUdLS06VlB17j6x4B6lfM9jHX9mLa60XOq27+Xq2c/kIjH7PXY/cXzC3z8bX7C8e8r6O/4JvRX3jebWtdvtm0f6F/RX/hE1ek60j+Kn5ZHS25nKPF3+1o8chiCzyy2YKt9f6K4/1lv/OuBvk/K5zVFnjnsAUt3NuMbLZgeP2/YgtMXXybubLagjzOFiz+09qMS9dwvpZzluKTbc075/SUW7m3HdnsxQoN/pq9+F6eMgWz3otpLe8Df2Ld++VP3Y/uR2Ou7CrCNexHjz+8H+8sYh6NLxXNqQ3+847GX//fjkbv7KpSw79mL5pHY9Z70ff/4Wg8+afuxw5u1d1s68zW+E61xrWvaoc/fmGk2uFFG+51a4eNG2yHv/ds4ZXxfoYpb6DbO2Eue3tUwTwDjYw3XDysz6OZ633leznm58rMNjzKw27D+3mY71jbbbj5aqSnv92GX3ZdRxvun9GGt9N0m/2dNtzfPtbK61gba62sVyPzKFs8xKWjrNP35t4zb881DjySVc2b40rYcFT7dGFU7fR3EHM+4sx3Flt62M+yfcS18tA/v9le5pG1K4sjq1o2V/jdjL5GDSPM6KifDY1Qo47uC7aOwGDnc32Zz2P7Pi/nyLOPtsJXHW2VnHF8czzazJJi7k3z82uPezn9CWdvvqcjrpt15s7Wv7usbbX7PF5hLqtf6Gm943otZ6i5oCv3137vVh5Rnud8rCsgt+PdPC7M5znt3d20vqb584/3NRcUvM063u4p2MM6U5633v0fYv2L1L9e2v/mWXZvwbrW49/UNNfp0qDXXLVLtDc/sVs97TwvY4Rf9LbGqXwxONPx6pnpk2IFrGO2+UXzk5ytXK0vmsewK7czk5eNq2fiyjSTtJfnrpqJvTZhF4PcPs/ivhYtjYoXQ+xPs1jvf7ucVsRI/9xIs4v2Weuqnvb6SkQW0+T5nWl0hle3H7vg+a0rozdvPuiZZXWob1WHB112dVhsHqPu1SHQrg5Vrqc6BGZUB/NzoOnVIdDegDY6ooparUzGkfy+9a73lUem2Sc/mr9qln1yP2f+i93m7+HMf4vm3yzT/G/07DB71Q3zf144q9fF/HPoVbu3HFaZzrblMNuMZVb78UfPKsNIKGK/A26/3/6GnoR86sWs0KX7abPRM/5rv1uvJfwrr31cH6tgb4t1EFufZv7nyO7za+b+OHfXhh8vRYwM3L3E16he+YkPzeftXyrWPmkf/nDO3w7OOdnLOSajnONirPN58cnO58Tnun3+rLRzu7YznzTXctv8bHZtrYdZzXydz8aZMwu2jmmnv7DfO8fPyQIAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAgL9OcjaZbxjGAuVu5R7lXuU+ZaGySLlfWawsUR5QlirLlOXKCmWl8qCySlmtrFHWKuuU9coGZaOySdmsbFG2KtuU7coOZafykLJL2a3sUR5W9iqPKI8qjymPK08oTypPKU8rzyjPKs8pzyv7lH8pLygvKi8p+5WXlVeUV5XXlNeVN5Q3lQPKW8rbykHlkHJYOaIcVY4p7yjvKu8px5UTyvvKB8pJ5UPllPKRclo5o3ysnFXOKZ8onyqfKZ8r/1a+UL5UvlK+Vs4r3yjfKt8p3ys/KBeUi8qPyn+Un5T/Kv9TflZ+US4pvyq/KZeVZCVFSVUMlwbFQ/FUvJRcirfio/gquZU8ip/ir+RV8ikBSqCSXymgFFSClEJKYaWIUlQpphRXSigllVJKaaWMEqyUVcop5ZUKSkWlklJZqaJUVaop1ZUQpYZSU6ml1FbqKHWVUKWeUl9poDRUGimNlSZKU6WZ0lxpobRUWilhSmuljRKutFXaKe2VDkpHpZPSWemidFW6KRFKpNJd6aH0VHopvZU+Sl+ln9JfuU0ZoAxUBimDlSFKlDJUuV0ZpgxXopUYZYQyUhmljFbGKGOVWGWcEqfEK+OVBGWCkqgkKROVScpkZYoyVZmmTFdmKDOVWcpsZY5yh3KnMle5S5mnzFcWKHcr9yj3KvcpC5VFyv3KYmWJ8oCyVFmmLFdWKCuVB5VVympljbJWWaesVzYoG5VNymZli7JV2aZsV3YoO5WHlF3KbmWP8rCyV3lEeVR5THlceUJ5UnlKeVp5RnlWeU55Xtmn/Et5QXlReUnZr7ysvKK8qrymvK68obypHFDeUt5WDiqHlMPKEeWockx5R3lXeU85rpxQ3lc+UE4qHyqnlI+U08oZ5WPlrHJO+UT5VPlM+Vz5t/KF8qXylfK1cl75RvlW+U75XvlBuaBcVH5U/qP8pPxX+Z/ys/KLckn5VflNuawkKylKqmJ46PxXPBRPxUvJpXgrPoqvklvJo/gp/kpeJZ8SoAQq+ZUCSkElSCmkFFaKKEWVYkpxpYRSUimllFbKKMFKWaWcUl6poFRUKimVlSpKVaWaUl0JUWooNZVaSm2ljlJXCVXqKfWVBkpDpZHSWGmiNFWaKc2VFkpLpZUSprRW2ijhSlulndJe6aB0VDopnZUuSlelmxKhRCrdlR5KT6WX0lvpo/RV+in9lduUAcpAZZAyWBmiRClDlduVYcpwJVqJUUYoI5VRymhljDJWiVXGKXFKvDJeSVAmKIlKkjJRmaRMVqYoU5VpynRlhjJTmaXMVuYodyh3KnOVu5R5ynxlgXK3co9yr3KfslBZpNyvLFaWKA8oS5VlynJlhbJSeVBZpaxW1ihrlXXKemWDslHZpGxWtihblW3KdmWHslN5SNml7Fb2KA8re5VHlEeVx5THlSeUJ5WnlKeVZ5RnleeU55V9yr+UF5QXlZeU/crLyivKq8pryuvKG8qbygHlLeVt5aBySDmsHFGOKseUd5R3lfeU48oJ5X3lA+Wk8qFySvlIOa2cUT5WzirnlE+UT5XPlM+VfytfKF8qXylfK+eVb5Rvle+U75UflAvKReVH5T/KT8p/lf8pPyu/KJeUX5XflMtKspKipCqGp85/xUPxVLyUXIq34qP4KrmVPIqf4q/kVfIpAUqgkl8poBRUgpRCSmGliFJUKaYUV0ooJZVSSmmljBKslFXKKeWVCkpFpZJSWamiVFWqKdWVEKWGUlOppdRW6ih1lVClnlJfaaA0VBopjZUmSlOlmdJcaaG0VFopYUprpY0SrrRV2intlQ5KR6WT0lnponRVuikRSqTSXemh9FR6Kb2VPkpfpZ/SX7lNGaAMVAYpg5UhSpQyVLldGaYMV6KVGGWEMlIZpYxWxihjlVhlnBKnxCvjlQRlgpKoJCkTlUnKZGWKMlWZpkxXZigp6t+bmaPbXAEBAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAPD/LyUlxf55+XKyhuTklBsb7MnT58aOBQAAAG7a/v+lSymnT/+i4cyZSzc22JNrPvT/AeBvITnZrNWnT18qWfKUhlKlPrqxwZ5c80mfJwCA+g8AuDnr/5kzl+wyXrr06Rsb7Mk1H+o/AFD/AQDUfwAA9R8AQP0HAFD/AQDUfwAA9R8AQP0HAFD/AQDUfwAA9R8AQP0HAFD/AQDUfwCg/lP/AYD6T/0HAOo/9R8AqP/UfwC4Jes///9fAKD+U/8B4J8gJcWs1ZcupZw+/YuGM2cu3dhgT675pM8TAAAAwM3Z/9fPy5eTNSQnp9zYYE+ePjd2LAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAwN9USkqK/TP58uWUZDF/MjBkOUjK5cvmQZJ22HAGAQAA/B37///94fzJfZs/emn7qRe36ScDQ3bDqRe2nHn9EV0t0v8H/vb1PzlZP7858+6rCwccWj367ZUjDj44koEhu+HtFcPfWj3+119+pv4Dt0b9/+7ciQPLRxzbnHRkQ8LRjRMYGLIbjqyPP7x52q+XfqH+A7dG/f/23Ik3lsUc3ZR4eP14NQEMDNkNh9fFHdo0lfoP3Er9/zeXxRzblHhk/fijugRgYMhmOLIu7jD1H6D+M1D/AVD/Gaj/AKj/DNR/ANR/Buo/AOo/A/UfAPWfgfoPgPrPQP0HQP1noP4DoP4zUP8BUP8ZqP8AqP8M1H8A1H8G6j8A6j8D9R8A9Z+B+g+A+s9A/QdA/Weg/gOg/jNQ/wHqP/WfgfoPUP+pcgzUf4D6z8BA/Qeo/wwM1H+A+s9A/af+A9R/Buo/AOo/A/UfAPX/lql41j7Rz8Pr4g+tG394fXz6PTlMcmS9eePwOnOq3x2f+g+A+n/TFv9D6+KPb0k8uX3SsQ0Jup1W5MfnPO37W5OOb048tPaaxqf+A6D+34RNwKmdk5+8Z9im6f3eWjNOVV0d+xy69Hapf2fjhIfnDnl2UbRajcNpTQb1H8BNW//tGqUuq3vJUrlL78TqxsE1cRp0wx7fvmG9PJI+jkaIP7g2Lr1OZnSkr5zKrioa7eDajHketfrYhzPmlsWdV9/j/LReeLHWP3OJzrQOV0x11bqlb+Ph9eOPb0nqFVazbLECrUMrVChR8KE5A9/fOtFeYQ3mhmv8tJ1g36lH39ucOKZHk0KBfvePi/hwx6RDt0QTQP0Hbu3+v35+uH3SiS1Jh+3ibxbARN1j17eT2yd+vGvqmV1TPto52R7/ox2T9VPjn9xmVrljGxI0wumHpny8e+o7mybYdfjoBqdin9oxSR1j3U6/oXm+u2mCPYlmazUH8Se3TTRXYJ3Zi9aY9mp8sG2i2fdObzWsez7YOvHwlUXVnrnK79XFVvcc22iusH5mahqsBU22l64bWh8NH1ov+LzxYOyIiMZvPhj728v3jO3RtElIsLbdXmetsFZJ63DmoSlHrB11dvdUbbgmVBOgG1tm9FeTcWBVrL2x1H8AN239V8lVUR3Xq9my8d3Va1W39uT2SWsn9R4Z2VgV9f1tE9dM6t21abVOjavMjemocvfayrGD2oe+vSZu6fjuEweEnd455a3V49Tv7diwclSn+i8vG62pDlt1TyVXow3uWO/5+2NUPAd3CH3mvuF6VMMLS0YObF+3Y8Mqcb2bq2x+tHPK1EGtF47tqgKrkaM61lMFVsmdERV+35guWivrfdXxuqF77hje3r4n/bpAMx/WpcHWmbdpQ9JaH6fI6yGt0oC2dV5aOupdt/bu+ObE5xZFD+loboi2V4vu3Lhql6bVHoiP1CTq7T+2YGjHRlWGd224YGTn6mWLaBzdo/0Q2Tzk4buG3B/bLaJZdc3qibtv79EypH2Dyube2z5Jc9NqVyldWCup1u3QunjqP4Cbtv6rRqlS1atSMr5Pc5VfdVzVxZ02uE3l0oU+3Ttdxb9gvjzzRnRaldSrUqmgJXERKnGFAv1VnzV+i9rlzj85p3XdCmF1y6sq9mldK6RsUZX99A72m6vGFQ/Kp4eOb0kqEZRv55yBKpKH1sZXKV1oUPu6e+YOblojuEPDyt88NadNvYox3Rp9+dgsjVyyUMAry0d/tndGp8ZVVdjP7ZmmSTTohqp0jxY1zEZn1Tjr5aZ4u5iXLVZgwchOmrkaI/s1HKfOb0l6dmG0NuGpe4dlXOBYTd6uOwYVLZBXpX7+iE6FAv2Wxkcuiu2q+WyfPeDA6nEF8uVJ6Nti7eQ+2iLVc10g1ChXdHCHepum9esfXjuPb66ZUW1fWzFGWxfbs6lajaCAPIvjIj7ZM00b5TKM8PoVdd10C7wLQP0Hbvn636BaqYR+Lc7tnvb26rizu6bOGtq2WtkiXz8xu239SqN7NLn00t3/fnSm+drIxoT9S0erPr+5KnZ8nxaqz6q3Kn0q1588PF1FVQ+pZtqvftv1v3SRwEfmRakOlyoSqNKqqq4Lh+Ci+dVPVvuy7/4RBQPyqM50a1Z9RESjLx6btXdeVJmi+dVp16PdmlWL7tbwrFv9Vw98SMd6Pz43V+ustdWlhBakwl6xZJBaqO+evkONl1qxY1aRt1+feW5RTJECeZ++b/jxK+u/CnWpwoEntibVqVRCbcfPLy744tGZakpO7Zy8KLabmj/N5/wTs2dEtS1TJL/6/9q0BxN7aldotnnz+Ly+cuzc6A61KxTXVJ8+PH3brAGbpvfT7X7htbu3CGlWs6waMvvFK+o/gJu2/p95aEr9KqXiejdXv/qNlbEf7Zg8fXCbasGFP39kRota5coXL9ixURV18lvVKT9pQNjrK2OLFcyr+q/xuzSppuqqP3U5cHBt/LubE6sFF1F32nzpY61Z/3W/yubDc4eoUOvG9lkDPt87Y/qQ8NoVi6toa5L9S0cVzu/34gMjI5uHjIhorCWqLAdb9V/lvWuTasO7Nkzv/6sg9w6rpQZFVwq3ta0zKrLxM/cNV419d9OEKmUKNw4pM7xLQ92vlVSrZL4MZX2GU2tYOL+/6n+m/v/uO836r83XOqtv375BpdahFVrUKqvOvDZBa6hpNZPVE3upcdEW6c7yJQpGNKveq1VNP99c+xaPGNerWXi9CnpIbcEZ872Mqbpw0DyfvOd2NVLRXRue3T3Vfhud+g/g5qz/6vGqtvvn9qlVoZgKac3yxQL9fetWKvHlY7PaN6isnvm++2M0qFvboWFl1QR1p1X/4/s0V7ugiqcO/KvLx6j/rwuEUoUDVk/q/eGOyXb//+j68SrXqtK6IggK9FuZ2PN/L8xXgS1brIDaGk3ywpKRBQP8VJbV1e8fXvunffOeXxRTIF8e1f8fn79LxTbavf7vnqraW7VM4RlDwhP7t+zRIkSXHo/Ojzq1Y3KlUoW0MlMHtUnq36phtdKq5/bHb97bnFP9L1EoQBcFoZVLainPLYrWNjaqXnpMj6abpvVTqVdJP//knJlD2+p65P1tE19dMea5hdGPL7j9oTkDtYaa7Z3DO9SrUuqrx2ed3T3toTsGqRVbEhdRsVTQJ3umqdWoXraIlniE/j+Am/v1H9VM893b5WP2zovav2y0aqDK12d7Z6xK6qUO7ZpJvTdM7atKrhtvr4kratX/8X1a6KLg6ydmt6xdPqxu+T1zhwxoV1d12HznN+291x2zBwYF+M0b0em1lWNVRRtULb1gZKcDq8ZVNl//D9UkLWuX07SaSZcmVTWCyubQzvX9fL1HRjbWslS3F4zsbL+Rar/+oyuOge3q/vD0HR9sTfr2qTkdGlbp26b2vx+dqUuGFQk9zj8x+6OdkzV+6SL574/tdnKb+RGmp+8dXijQ74m7b1d7dHBNnPvrPyWC8qmw3xXTsULJoB2zByyJi9TGPjp/qBanbRzbs+naSb3V3Kgt0KLH922h1X5kXtTE21rpWkb74ZVlozW+WqLF4yLy++feNL1/s5plx/Vupm6/FqE2bun47mpb7XcBqP8Abs7XfxpULaWy9sWjM9Vj/3zvjDuHt1cX2vr45cRVE3upOHduXFXVTNX1xSUjSxTKZ9b/vi3ahFb4dO/0Nx+MHd29ibrfA9vX3bd4xHHrY5yas/reoyIbe+fyDK9XUQ3BhH4t54/opC69OuHqKqux0CSjuzfWrDTbaYPb3Demy5geTTTPR+ZH6UohrG4F1dL0T/LY9V8XI4M6hH7/9B0nt03U5Ymaj9va1lH9L1e8wIoJZv0/vXPy8c2JxQrmWxwX8cVjMz9/ZIauUIoF5Xt95dgvH5upKwXzI0/rnfqvFk2P6rLl3tFdOjWqoplvnt7f/A7vlsTddw7q1Liq2poHk3pq66w3tcdPGhCmC6KI5iG6dtA42lJdffRoWUN3aoU1575tau1bHKOGRvOcPqSNWhYtkfoP4Gbu/2f6/M/UwW3U77U+0+58/l/Dye3mB/hfemCU/YJ/vNX/t/q68c7n/3dNfc/8/H98+gfvT2xJKlEoYN3kPirXKrmaocY5bF0aaEJ7Euu92vgPrBF0j+5XXTUXao2Q/j0C+/X/ni1rqNN+e+f6aj7U2S5aMO9T9w7TbCuUCFITFtWpnu6vXKpQ/aql1PpMHBDWvUVI16bVvL08OzepqkI9oX/Ld633BdJf/3lLVwTW1wfsN46tT5CaXyXTmmtN1DLaq22vg/mFBetrDvZLSfZoH6d9/l9/qiF7Z6P93w4af/LKLy9Q/wHcbPXf/pDM8oTuqoeqZtZ/90bd2qFL4iLsD8ynfzf24No4++P0qq7qNs8a2rZfeO0zu6bq/iy//2uXx9ieTdXN3jF7gLrl6V/4tb//e8iaJO0bx/H2i/zOF2yt79W6f2/3sPU5z60zb0vs30qXEgl9W8wa2u6V5aO1CI3wQHyk7tSgq5L7xnbV+LpfvfrYns2S+rfSmPoZ26vZPaM6299QO7E16bEFQxtWK/32mrhjVuPifP93nf2N4IT0FbBXLP0a5FDGaM5/Jsj9y9Hu33HO9JVq6j+Am63+5/z93z9QOsbbrUBaj/pPqIT2f2BN3fKze6ad2z1NTY+uR+w5a211/7k9U8/unqZO+FHrW8a6cW7PNLsbbz06Lf0rzAzUf+CfW/+XRh/dMMGuzCqJB9fG211f+0+zW5v237exX8xP69Jb/Vvzdrz7PVcPh9e5/zeCchrzGof0FXvbvBKJt69H0uZprq3uf3ttRqfd3ihn5LXx6bftqw/3lfmDK3ZrD4fXxh7aSP0Hbp36/+3ZEweWj3hn80Q1Acc2JjIwZDeYV1Kbpv36y8/Uf+CWqf8v3zvgrRUj3lwWfWB5DANDdsObD9x+YHXCb5cuUf+Bv339t07h3y79/N259zV8/8kHDAw5DN+dPXHhi49T3A4eAAAA/O36//qZfPlySnIyA0POg44T/Ug/bDiDAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAArp2Bq0wzwo3WRh0jzGhjNDYaGPWMGkZbI1S32+hWHf0LN+rqVphutTXa6VZt/WugNDYa6rEwo77+aqS/ZljzC1O82K1/Gz2MOP1LNIL17I7T7wRjynVNX9jIZaSdXy79HbhwzvA1EUfzzztv7DKeq/pVVtN4KrM97NsuHWdxRqwRb0QYtxujr3v9Cxgehoczz+s5xxsbacuPMMYa0Te8//Jo6ddbY8z9FOTcbmMkaZ8naA3GWc9Cb2OCbidc8/KLa/nm+ZbrOrbf3F8hzu1kTZNC0/CPFWmMMobpyEuyjsEJN3T+Xe/xb56rw5zbkTr7J1lHvHlrlHMWhOv2MK1TrHNeZK+Klu+t3z7Xcfyb50tQRft2LqOntZxYY6hV+zpqeTFaeoJ1T6LWKE735Lz8693+Isq3+72d5Wfe1utbn0Y3sHyzVpR02bf/D/QBL44=',
    'xls': 'eNrtWE9IFFEY/73Z2XZddF1tDbSQQdDKlDAPRqCbQlkHzYQwIqhxZ8RFmYlxVYwo07wEQdCp6CJ46WJ16Q8V1K1DYNQhCAK1Y6egoIM6fe/bWdPagwslFfMb3ve+973vN++beX9nXs+VzE/fr1jAT2hBACtuAbassQlKBdlCDFTvulLN5mFKro9/CgVh6sgtQTwpehWSfSj7ewEK7qkvSAKLlE7jHDpty9Q2EW0cgy5kDM0kBW6TJYpyjqqUZZLlVpZ32fMpy4NsucaymXznxSnMJTpr93uj+KRSxXVRyPs+ZM57tjSgDC/lKL50XWR8g2h1Uvrg31lRqRZiBtRv7aZlOvrgPOLUgTP46mrAl+xMfa759s21C5D923p7KIf9hqIC43DP8gCfwg7sVWVNnKabk+4fTVmGZpiDqRHTSZlDS7jAy6+XVl2k3lPferS+qbGBJnO3mR52rKEDWkOTZuhjQzQ8uhzbGE6mg8AJK5WWlm5zxLSGzQhwRHd6bUdLm3oUOOYk+3XH0JJ2X59pFgIdpm7Yo1RO2npEtsvLRWzdclHE06iQpIFi1kt4MsVoQ1i68/lNR29X4gxbxnmLyGwk1fLZ4eKyZBA5yjVBlgFKtczYw3KC77qd9QqWcfKkvKarzFMOT7LPFa6toXYaGW8TO9fou0if+nT8UeXUx8Ru0mfbF87HZ98lplFFb9QgvrwmUSfqxK2bEo8T2Vx4i84HluW/LEBhJebF7nq7ZTGWEWG1hGWmJN+OWC0p3ruSbJGDLZgd8PwFs9XVkmQHPbaSg60wW/bRM45esheFtE1Uy7YlL5CDF2BeyOMFmNfGPGVflqfm4KnMC3s8lXkXmXe1TvJalFI8kEHREv0DEfjw4cOHDx8+fPjYNAg+hGaO3Kp3AJcHxpD3X2eZ0or/m+S/RTdsutL0YXoIFuUOxvIaP9sQFNl7iQ1ysv8LJXqodQcD6OU4BvIev/RZI9Y+z4aJsd83hfJtfyWfOP9w+98Bo8XyQA==',
    'xlsb': 'eNqNegVcVc3W/iGlG6QbpENASkC6Qw4hLUgj3Y0IIt0lIkgfWqQOCNIgHYfm0HDokpTww/u+wH259/v/P/jts9c8e61nz6xZE3tm1BQREAkAAAAKwKBcyqBRtTbUDw4AwEQCAPCvUT1JezsXMzsXIw1PBzNnAw4P29e5GoqKbVw4jehXIXFFvJNH1LTTR/MI3woyFz2VzGuaeVFTC6D+u5UG0tJF7YpUP05Szg5ekNdEsn7ijsVsxHMWxEfyoog42lSYgX5KDsnKxJgNQKdk/SKO19kDTW/fZqQl7D3WCBWCeeQQrJOSUZPQen3SCYY8W52r1MI9pCt+vj1g/kDhQcYQ7LE7qQlYoJ3wKVj/Y9XTpxoHLxuYVJD1LTRR3w/pYKbmj37lr/iMpxepoh6ZwoBPbWuKIBTMg42NRQJfnNU4Gu13KT33xMAXenoSERTlx7zuPLdSbfGjl1KyKrKyM4CKvC41A7P6G6thsSkm0XjMAQNvhDrj2iYBrWM56kptr1w6RX0EWxCznk+ZCyQ2N5/s3Wi3tWlk/jtmvHoUjA9iXjUqDUGHE5zCoCu+ZPzjHuEze+TWjrp5+Gn5Jxjb8ie8i/FnTfN7H38LdxquJi8z2ZMPJ4gid6SYYqL2Kav4WCgBr+TIublQDYW6pjt/x0+hNZjl/up5fYam9o/qkgxcat+5TinBAwDo13cjJ7PXzpwcf35L4vqsW7kwAnbfNrwdang+1UycrcJpNjhNPIlQUH0gidczMZvUeDkzQrSmgC2smyTqize1s2qlkUFuwqnSWsck/oXHxHvc2qPMSEs104TWwoqgk1SF3FgbOmc755GTEG4yfl7wuF0a00DhioT7Q9SjsGAPt7P8WK7n3QODWY941HpqmXtkMb6TqLGnbLzkC3rrC9DMeDrXkpS3jrnb7qO0NRxM1OVsYjFZ6tN1RHPkSaoJtBS0KjcxdBag9TQy5Z03y+7blvFXDcd/VJGAEonYuOPcX8tYeTVAyM7j93oV9zs9hmzsZ4Hjna3AcLKCRrh/umXj6Gfy7nWq7totpNd3j9ecf3nG3d7JxsTe3obDxMruLy+p96i0cuF8o9vHHpsNjQ/EkXD1EpdrkDgF4NubFIXaeibLXlFKIinnBgosm/a9EOY5f2E/2yn+zYT24bfAQibxTPu3jo3ae5AZo5kA6rHZNRoPDS6a5jxRMT2x+UzBh7GusVJa+zlNjzv1mEFVAtYsElIZ/nUxMn0keYYy0B2XaCfNpPHHfXyZosMEetnUtn0M783DTqm2B11L2+sHGyPWzoKC6T+9mwmwdNmIJdxy3N962jVyXMSwLTi+2HyJvvOjsMvku1lBUu8e0XAUGy2V04Wa57MM1aPxZ6VE0JBR1D66fB+7n1W967tcB2uLG9j/dFOvbi33wnVjD0MAALD/ctO/OygyTtkGKE7wbi61o5QNnPUS9QVkTF/c8kO1lQnh8uqSRHtYm5YeE4GNGO+YzPD7dk0p8Yfq0PcEJTbt3+fL9qiY3ucJnwSKOk9rdDGmL3xz70/9uf+J5wT0vBWjyFbHVLLosc9Xbw8Uci0VOoW3eqt78hxehlNM80o1z7X1I2sMOd7oDtau8pa4soOQ+0EtpRs0U1l2VpiKsfWuUCDFOcKhaYLUJ1JVL6qa9x8lh7c8g4XKMMRJCoLQBsmRSRcTcl8SqXsmZqZ06gaZU9oUcEcGCm7B1SAN6igqxAqOzvKzgmCMJibGw2tH1TUq9v32jamnE/wUVPIGV3GQB85ZK0/1YrtJrDquVMv3W9YSh7x55nn564nVRxDT80f7zlgoxtfOoyjXq4b4i6i0hMd4bDUa9c4WQElEsjPEOwccR5al0ow6skQ4/FYwR+OTfbezmfd7nDbT9DLur9vMcc9+rp+M/wi2Sf5oOCT8CkK1dHY4yTO4E73+/nBEe0RsE8+jy3djNs88UUGTHeiO0xDmq2+WwntJA/MqfrvpujNWKYGu+KCPclhmhMqlVFa4loh4OycvBmEY3BLN1a5hr567Hg2OWaKQhD+AWQlOYH9bg8engEdY1UC2AL0Y7vWDGjUiH/auUf0zGuiHFUZJriOBBeu6+/8rGpxdPF+bOf+JhRJtRUv85wQiP1f7TUlthQyYpAvYIdhMVbYsiy3SCugS5bm5dmx2xYy6akHVYTt17iPBTIGsyWEQ1SSORQmcksFEai2cYqTzSabTbLy5NCPD1RVtRoTUb4KCsJm5ZW7nkAX05xhNC297Lp6W7eMjmCySMXHtj1D3LzzO4wq18ivNmDjSpbVmrg1pFkPAEo9Bkm9F13qbUNTJrB0d/zE+w+d9pcSpHKTlyw7NvrCiJx+Wx6EtJVSyBNMRd3JDWb8bW4Twhxl/zACMrq6tV8nN2ZknK/y5pZR0qeITa6NjlOR2meU4WYi+V63LviGERCd1K//I6LlBhJKbzYqvpDlBFe14Li6PbsR8hdLpht1H2V81Ra7SpBcTK1rDWETFLO+7n9uEFZKNRmgANaJ0xTXSdVUqFU5VpHxGXz3MGZqsjuZCGip08sNpxM2lkKXeKviMLQSzSXk2s4OCM95aofdUypR+0COAGK0FPkfwWUpIpWK/NGF8TOXyKw01QsHFbjUc9qC2pxkmBspoeNoxFGgmXl+EJPqt2gVLBPIJYbQtgXVhFD81t/GygnHUBZY9Emk1AtUGR00OV9UySPDQTASGFHHzqAk+PGROJFXBVtQ0l6sCrqcIuaubi0epVpUmfdfrxVk80flEDcRRzwrRzCmQe65DuJO0cfLZK6Q1nIGuctsfzjmPe6EP3tBe6YG83O9Ah9w8dRBIo1C7gLZUA1e7MAYXCvzRryTOFFT+eOgohodZlOC0GHcMeunxHfSmysSutMSUjLAT7rC6mVWR1w5R+wz/hBAWqhfHfqSC1ejDldGYoS/MzfNkL6bn+Je3O1Bu7MfPtcT20FWeVyB9n1PuEo4ZUW+bN+HCl7aipe/jBaJeVXxzOBhNw1a0l6VsOL/w5uN3lmwbY+0QDfN7jXO1luVwMfolNzos0WNCrE8w/mNis4C544gO5s8BpKswx0hheooxpu91/SIWDGPMUQZtDH7Qzvjz/OzwF/AvavZGhJano4XLdGCCvnWBh8nBmOjntbowzseqr59+jZdo6JG3C8uPAbWugCPzn7iDe/butSFOtp9dNdcp+utele6vNmTqZOxuZWfh/PcI9HeS+89M6l8jUOTAF2sEbgzk+YZBd64pG6grSu5HTkVUU/YfCCh6MB6Rng43iYNjPherTElZ+Yu56LxLdruN7tx39MgyIgEmX34g7SP357f2kBVMyZHbyk80pbNuwWD56/H81iIa4zFFG2EaCDY2ZxRfphLpNNMfK35nA0lZ+7JR8kxyHkSUtZa3HkwHnDVryFpkkkNRNlwGLqiE2RXLL2K3GCoqxVi42mah4iyf0uBKuCvQD79SuRJhC6ie7k1fYP6z3CGf+Kunr1OFd+X+M5I4W5qZudyU/F8J7tuRNzKmWeF6GvluN3AXZ9VueEdc6okBDq5X4hFch+kUcpqSXFa1vT2iq0HB/H6GceOZ+XFI+8BipLobOW9lBZx4c7mrMWTPxae2Di4WHTmc58Fwufo7SXfPL/45JIoU3fHhOhoKXaR56osl6pNKy9yKUQTWFvPLXyApMoRD3+blrMd5x1kuzZ+5dxA0jMZ41SM46tNdhvS1zmNAv3oeEXR3oJnhN473sOCNyEm/TZiryjBnZSWCVCufh0S6MzqufpZrvVI0VXyZuqvfVHncPOw07CFKqKfTCvNyPV7YQfinb5of4jTGIQMAHdR/TamvfeNiaWZr9tfvvwJhW6fPfvoJwXfsid/UazsK1Qz+iK5oetXnibziS25rsUlW4dwMpPSJbUxoGNtWyXAIdcbJcPXbXBis5fCsbpNU+JntY5e43nHfpdKs45+wC7FhvYyH7Ipcnc4uLYV7HE3Jy0XS9SIZP29lM/yijCArShZjCeHaPaoL/mCFEpoFl8XWgukcFuxpe/5zoDSLnWQeh4yaKpQdwoXA/EgvicGHWGYRLW0e5RMmCXZsEAeVDkJUlLi9XEmcrXYxTmwXLy1B4FTDQszpxBoDS9928AFWcJlcmGlT/kWTqkWvQs5Ec23KxXvQHhN/FI5tCtdAAMXJXigvZlhQ0APMERWxvCrfUV2056l0abGbeRngGf2XXBy5Tj4QIzShqE2ND4qXualXyvr5UyAXFkVfZsCZaHFlwZAOu2c/eRpc3wdS6LZ1vb7GHsnlCqbPiPIVgY9nsPuTPWhpTYUxMq9NHZlHjaZ9leblDprRZJP2tvk0iOWUgEANL5J/sFqP2Tt8lINGUsbecsF1fF1L0fkUzcwsy4aR6se3F+hDNEsGgJx+6Tipdx9V9nDOAn7iMPiqRFA8RqPeY2QygX+3HQet6tlLvxC0Yj1immphohMhO2jvUGvpxBF9hD/JZf9GiwfB82yxfX0Jlh/tL+Z7sHy4X/YTzL9il2mZkOh9OcL/+2Ivd//EXqzR71cJJWfFqfZ3n6ONaLG5HEOWykbv442zr8lz/pfjA6p8e4yyl2zfZCXYIoRPJ2p7JF5kmsPPwGDxa0jvmfCYADnIE7Fv4L/xSGuE2BiPBp0ImqJR7Hn5FDiIoHiwhSJ7CKMQznO8qmyvdROsiQdX9IRmx/o5UVo4uuMKDfdtlwjW4qtSL4eq5D58GbwqimSfjRtVzhtn8Trii8wWHE2YgXvoQOpLSaun1DLouCXyXjmcRxL8cih+qRsogUWE01uz/LoRG5Ca/diwINs3BSoMeWEkZdkvstR6iIN5w2fEHw/GPk1xMuKujZlfx83rXcMoxxFKFRHVXJt9h5Htfs6R/o4lTvHxGI07zisd07TIBkAwPrsFHpLOljL3OxzpDtF4XTwlKI081UVAlr773KMXn+z15t5i5GDHLMv3uurB87iKqazjGKcfMGInbk0tS5FRlHjkBB1jdRBQZMN1dDSkCT1N4xJZ14MzTSR7YZGmJxc5mkNlDGemrj5uSa4YhFfxwKPN6XRLZRxvgECaZtHzw27EaMx6E3O+cq43hIoVlLBnOrIhvxWKky6/6ReQUK7wS4KqDEiZ7h79pndBuGveueeZidhXW6/3Jr5NhzgExHUu6LRfDOknVUAvihPJCtPcVpBtSmZR5Q/HFi1KthzFVCZ/NqN9Cud0PVDv9noTH3KOncRnSyB7TqwS11O4WlF4JqbF5+l9VpEELjbjsEmN1uop7n7Iep70aBRvtfGABE8BezlhkSKWSBT++CoT/a3VI0LdRuMPiIMfVfPnZchzbd/ynCHFNosHMHIpc6cb068WyrakcMHlxqcWCbM3LDokY5pnty8u5jKiUINJQ4ZtuLGCP3yM1CosqnFXPLKyp69AnNFYn5Q3ZQTKJ/DRF658jWWaetieVWr8nLUxNDhxH41cBVlhNqS+PwI3iCzokHp/RMlvAYp+4vKJe9QsmbenSPjhqlWwFOOMBl4qF8b83guqxpdG45te2smVvIVV2As2Njz8g3LbJJ3lsG6Ogv2hHGOshQxGy3f4g4/ilILXmKxTGj+S/TJ43BOEPilDm9I8LPv1KETKxE4NT6LYAvosm7DGBGTtX7e8cZxvIWgD/eEh0NbYS6LjujCpFHq87h0YjxX+5lCcS+DJL95Vd1zi7hIifASs9zGKNlgoA4gn3WE269aNhIk9SaNE0i54mWehccqrPB5ThsM/iunM6hpZFC1argq8cfN3poPH4sz74dT7+kyodmIGylk2dcBT+tplk/aiDLO/n3WpFa2DEo8gB/1R6Gwoj0c0nba0uo9WvT8hbz63MdvorqxxLOr7FYJzZlug9YqlfWtlAjVhbnZR1CkJdXhI2k0wVqZXvtXHsl65kxVs2FwKTg8Eum1J9NkmoSuPsz8Uthuyq1GIbmJ6jBK+83qH4+Krav5PxXMXn6VFUf2EKW0fIl22krclGgINtZxJB5w7mPBJni+YXTdXvIrUC5LToIEjdqNormXDLJWshlVQW2+7kllfU1JZ5gmrJJH8YRAhi/QAqwCh8GejooeFKiCIS2GfbqeV6IdpUq3GU4MU/pWqS85/DmErMaJ0CBgAQIIFHID4P4b3u4F9O73HnkKesEP13MvHfsvmnDHqUWP1kvtFBNU+ncpFGI/jYs0rYcdVkafmbVGY/s0BbdTm/HCoQRNN8KEU1G1xIqxEbBYxMTKZk1/MWpf8NPP7tw9Q2gp3VoXVwd93XdxdXTUvC8+Fqn9uK/nFQE8lnYHV7BxXJXCXrcWBfed+0d8KEnfNDy5fz38NXIa38Iv5KdbfUf69QesDXyIIWzAhmueMKf47UXetvrzdEvFZATXZuIT0K6jftEaPmdfYrtAmWNt6mtZpxatUYzorogrirG2tl3AZ5VVaVS09aV8Gm7KePshdkigy0toQ3HILqRx1ttSC5f5y6Nb8mGZDZYe1kga2cRmthiZHeVdVG9S6QIWr2Lt7Ir3svlQbtF9WjzmfWE/jHypynFRWS58dKoi4605nfeulcNea/mQgtPOtWmQAaq/IfqJ0EYRDlr1LuBw/xaQp3YppmYccYk78MPmHQqowUuEPjWaNr8wYjs0gIqC0ZvtjlkpmLYlwI8fCUY00UAXIDUixY57J5SKdz1L9UFLbaeMDiGPaG6pRU2hC8va4aWdfxiGXEqgJbMj7pXSUVmdl3QTJBNEAW/NklCI/kPK49kBqQajA+Lw8JeW0D9yu05BikCDwZx6RBmFPSun0kw8pxdjpuo+rXlu3QwyAKVNl4A9cIKpJ8lqj87olyAeQK5As74mybarXS+tpiCSwLNXL1LoT8gqoqrT9YYqnyn6jCfIZRAfszANsvXLPbyKo4q+/ANWZkLzYUUhxhW1U4KnNOG+AQehAauySWOuX+0bJzWx1rg8oJhfJj43xCNNt1LWdSjN3wpz7NnTTJVu5GqY+2jzXdistKDWAr3Wg6P/MgaGTj+jBz9C0qTlJ0D/nZLGgrzQc5pyuu6QJi2l1qWOEW67He3Zc2ANxBQI8NeG7UhwM8J41PoIzLSvM81I+ArEkwT1mPHeFAXgYX9cgZu4qAqaK/uWfaw9k604X9cNZzztLwk8o4jUdW7dAckDMmu0HlWNjS+TzyIGaqAGcioCZor9c5tnkuGIIIxo9EqzJKZ0Dm7nyp+sNDZE57JT2x4vDtFxlMOUBVNukSYCrYsS9Bl3+KmPrEUgS6EFetY0RTH8bQHXCRwt3hYS65131auML6FGem/J26hS/68uNQUgs6BUw4HcU4il3XtT3A1AdNyp6U6152Iqz61K/MW7KoDqbmDbf+kzw7FyCMyfk80kuWh3ubxfWyg+/RrygIKW0+YwxQ0IlvjQ09AB/XW9NXsLMkyaEbhKY7ijUrVQZWuxs5upXPjXaV/9LqBL59+ExH8JFnJvNF/JP8402rUCfITCHKyVaQBjruTy0xE66xQSHKMqEZOIzhd/ntSFfyvNCoWjbsjo+2C/IGeV2jRvgQaVZq+B8wPmUi53ayRmSvVy6fwVW0bSbnVR6mS/EurN2fpZSKCT0beOBKsQRetIFRnUV7lcu/Kmaw+zaeGxVRNEQ3Xgw8CiC6pf/OFgHWgOb+VjnC1sdpRCoyewfpyQi+F09T9hhqdAyHThQMnHRXPsyAouBjdJBFo8IO2Lii/PJaLpATd4OyHlzQyV9Ws19dSavzhlWDRFxxsp+o7M3A7aFutMYOzZ/3MPWORRabdvsFxn2tReoye6vwWiIOjvIc1Z1ybDvKa2r+TgNhOlVl/XXbD1Lt5/8LDRXQedqL+xevEMBU4E0bP4SZJO7Um4rKsc5q8Epc+iTR7WnMQwO7vwe/Eyez8flYrgT4TdRYO6CCGp4PyFqeJQ2ErAPb1a8fO1wU+eqJHsnLJZVGstijnecFFzdlUg3TgmO0k+71JODatK68zmN044TgveGWup3uoytsVPqccGfujTpxsnAjV1HstrfSbbJ3s+8XH5Hy08Oju2KlNEmDj+Ghs14LDPS8pOAQV3KstopocfQqJnXy6u0/FTg4C5oaGk5MTi3i13WjsHJna6GFOwdW08FruoqkbWjc0qLnQ4jJyoSd54eagyJbTXhM915Xmn/yOlF8bs+kdDjbRk7TqfzWOH3M8bLUbTu5OD4LpHIGavlflp3SkE2Q9e8DNkdWidOJ5XChTTs3yveFH6gX7+JHv8+hnAubP8oygLFgcodzqIVDKtSWdKs+UeNx9d4NbldRyF8QIO8xgvfqGICIprfQ11Ftx2WXh8/O/9L60FIDMgYSJcnU5LE3QN9zGf3V0MEjmj1RKzYpX3g47OwnoUEgoS1erJADMCePC0lKF+V9cYn0HnemFJJ2hRXlZV1jwZgxr9tHxySBw+5Wooc2E5Boz6P5hetsbEegDSCEIAheVFKFGlC3EcQCAjEA4Tlwcdc2bYUoQPhJEAmorQo1HD0AXGUzy4tiODlAKgv6QCIVCh/pMycB4PGjIgxADS414gBAqTiAFRquIRDwEZ0NO6zawUyuMwrfcZrY8RyXMAGth3+DTiGdAOW8dxqGtLfgJwbeH+B1KscgYBrmWkdDsO/9u3fYttTaZw/zxlEAlLSbzKi9puU5G97E8bbLPeKIf4B1YYR3uuEUAdcZ1XBC9BqlMn4t8jdinMt4D2Tx6ZeJbx5VdtriRv+naJrielfpIi3mjuENwVdtLsFp8Jv6KH+p7c+SXlwkz+c7OgbBe7VW5Au8hZsvNMMuwWT7jTf34J2d5rBt6DMnebbW5DyTjPgBhw6uQVFom7BgVuw4bYgQ/l3mqG3oP+d5rtbUPtOM+gW5LzTfHML3jnE/fbtjt/vwNsXOdrfgbecjreBh1Mie1NJljw3NSeliHtbH7fVPcRyGwPV1DeBUQm4rcO6W0ptyRvKWuRbyr7bx0c0N9aQwNsM3QZoG8mtYuet4tZtpHbdYkw32TmtO72hmcOh+jvj82mLN3Huo+hzY20fjv03+R7F/k3OG1lF/+ZsOml7cENKePvcT/PW3gj/ln+8/Va18e/3WwOMzA+yzS9VnX4iXpXV+mBclCRNjYf5nBqr5u71iF0d//w80pBx4euzypno7xLZM1CBVms7vOS4nWFxsFlM8nNNW4dB41DJfVV93pRGxbiiuIM54kL+owR/egfz4TBYNVhP+mEEBPg5GC/fcChOscryR26YrhMoLmmGkFtge/85xXVChR2M70gGdURLH3mYMVHGkOxVyJBMOhpXhB7epOEDqUyaec5ztIAiC70mYLf+kWsICXmSI5EO1mACa4grcowWMyQLlUS7yaVb/cjtX0INgPCwgyHO11dwkr61vCJH1bWC7lDcpkTv6LZA8ArSyAl/6gV4UCN2AA19zmG/5oExK0dj+XZcuNjk8cM+M5cfNaZGogVHvlidw/lGxb6NT46W8zNcLhq9yw73a30exf8G/PPDQVjD9IPDdcof/nbty9bM1MqY08rW2MKMm8PBzmITY88ZthK3aPLy5fzRoQPq9R0CV6HqwAX3o9674Q1BfFp++Cu59o81yZieJATEWMTkE1Qs9NrJB+bia+8OF887MGFa3GJM2japzaCUVQg3kNoqQbC77aXfqwJLC3mF85pe+JRy4qcGl/pjggy8cPcySGeg3MgPBwCkEwEABH9velgaO5mZAl2c/qza/vmqyR1mV3krjhOwAGodJ2xFkvvV9xs5uO1lLB/DZ3g5FNTOz7B5ly2SY8J64dStzU9z0L6M9x9XJ7EFe2E6oiuKjSvKfsu0VE5J9ddjq7dnQr5ncsprZ/dUf8MVX5XOIxc9VVeRqiPI+MgLVWeRkqPK8RGa8wdgA3pH+n0ZE4YvsRPy5PQqeA9D09qepeAZKwQx4+ZJ/VMgwTWXSsE1lrkn4BurBAnirt4TkHCJpF7gGIvfE/CM5YMYcJPuCYS4WlIvcI2l7wn4xspBfLhj9wRM2tc4xm9RyKQmUlpKhc8e/NPR0TtDzn92qb8hAm4+IW9Xxv99TTxdY9ruIxeObxnEv/xCpYnNWIfgiQmnRmzhk8IpbIcIVK1mmeqv9c8T/HeJW7WKcqxHH3XwLAnCeiPAYJEjCiEH26A4JEZtXBn9x5roZy/mqeaKqFu7XrDYTDH/QL/oxWw08L9wcC509eMVjswM4BUkR5/RFV42tnzR421dCmlikCLL9bZeKKBq1R2hS9ip54rK4uV9u1r5hB3mBqPzKJT9na2rGty2GZsqUsi+W6deoWXVZmOsCHypNv8NqJKtucueV+fqgt+/unU+xbtphpY7QAMjydalGG+rTxc813/6hnVdkHcWhXOh+BcrG00Rk575LkKnPwB5ecW7i2eWOoV/ecGw3FuExk9wxoZ+A76TWH6qNYbGOoNXFlFI2tmvIiMyne0wR/qNUSHLrP5uEX1BE1qcZTximm8FWqFH7PaSWVXkw/WieuXhj80R7MXjBiIDSTn2edPoeYnPIxtMCrE1R1A6YGctLvKL+cIHPqEseEb2wwZqdt3SMQnjPdSs2ucBMdu0kx2TT2Y7ywQQz+hgZ8KFKAMf/HtO+jJLxTMIz61iNKaTumUNP7zhQd81YXNLpWxcX5Z89E7BJZ11ZUt+R0wEo8wjfwI4kHt0XuzJ/6S/0oSLeP3jCTb54Wa9F72U8yu+mmJXc4SmuYBibV+soOS2mR8YhtugoYZzeF4hQo8Zl4hPxthEJKHA6HmZfmoFFPlhhgsMYnq2pXY0JxGIvvlgh99ioWw/IK7JTJ/KK6DxHCJolrUSGvSw6IcQ3C93yDB46aGq5wGYauDXw39G4MP8/C8x16nu6wvn+jK1f6XmZO/gzGns4PAn8pSjvmC1cOFI+ud2q5RxcWXyErKYSF1IJBha5axyKGAx+asSB0Qdf329aYfqfNSUIAIsqnbNFsouC6Vj3E8pCmSuscYLfRCwxYjxZBwJSEOX/zK3ywWBUnZ3r22HMgbRTfnrA31EqyZjfP/g7rVXEV/dtbqzwA5x6+tWg+XVrgM0ul6pX92pWvtHWc6CSa9k+8j+mft3Dz+5UV13VFfXbYjyP5ZgrnspYydPeTtTM49/LcRABwxY3xoDOvZ//9x6VsLHGfkpJdLW8NjyyyCbp/xzNGZ6LQlm65WuNzywuJzcFVvDeHr5LosJUPR8r/u8l1NlFU+Uvvu4wckP6NFX9gTkqqSWytVITrOWC0/S70/hLvcXA31IDj2bDFCnRfcsm1QoPN6WkwWGmU+ES2hiHAs3jSOlIkIOW8TFXkjGUIzNPBcgR/hNT64q/pJIBSatE/CVAaWkqZ1FatihCqt6hfqBZvJzPKGafX41HQG8SsNBfH7Ehhj86i+k/cGY5wvqE3lc+ujTJq2R5UMCM2jTYwjuLR3I84YTx7geOE9EsaJb3FRa48uhzB5ET75glLUwjs+zbz+jWM76pUNad7ZvlInlpOjLH54Z6wSdJMlu0vnG/+rg3uZVURnG96o/QXHtXtx/D4xX9k5mfyLDLLrerpULQ9pfepByZGewWS2B5lk1wSVSoMmrz15vTKgyjBgsG1zodxfgyoR9YaaiOl7FK5bJVtWPdgtKNIKNcp8QRc70ZSoYgSi286a0AlTILLZGqLtR5gpNIdzdmKTDwrPEn1bPSB6zImsE1e8NumVfaiJPJW5LDkvioAg4rmZRJFnBbRpkamYzHifLCRm2mxoWrh0TKnK9AZITvhttxppVlSOuCvsuI+48W/v04HvJW3+RS2w1RTh4AoT//WTXzV8AHOC/n/O6T3D/rNEdQRncP04e3Te8fxrnzrAD/v91Nuc+z/3jKnc82wj/cXjlvvH90w13xudI98463De9v6l7ZyqG+n/a4r1PeH+39I6QH+3/tHd6n/D+FuMdIRD9v2443ie4v8B7R0CE978v995nuT/bu2NRpfyvc7/7BPdnY3cEe5T/fW52n+H+NOOOQYf6f5903Ge5P1TcsUzT/ufAcd/6fld9Z51K9//ruO9z3e+b7rgO6f9LT6WmiIT85zH29f/RdbRPMPxJ/Q/VY2wL',
    'sav': 'eNqNUk1Lw0AQTVpLG8TiwauwoAd7qKnxk0IhwbZY7BdJSvEibDdbG2y7ZbtpqRf/mv9GPPkLNO7GTQ0VxAePTJg3b3Zm97DeNMyjgwJwuo4DqpZrgXqjWQNFMGJsNi/r+oPPRsHgGJGJ3vefIPVaEOk2hp7DIAMpRVEyyg/S6+jOLJ2BDmLAuDCMcskon1/VfcR8MoVj0CaUjZb+1AMeHvsLTFdgHtAFXoG/wF1FPy3RT9HUiC3Lvq3ZIDpPyBGn1X01Ymwh8vlkfV6N2LU71d61+6/6JFK5TMReu+E6PJ+VO9jizHHuJbTCVxWeco63D0XJJrRpqQnD91B+nz/lv9BtS208fy7hvRmLHqJmR8YFzu8dVSaQPmKqyYkrM0q8ADEtmqASTH02F4W70iglDZKX+yqDdr9oNYqXpydyOTeQDggFDMP1fXHcm5u6DkUj/pAAIsMhxrHO/aVr8VdGllyGCFz7vfTNL8uFgjA=',
}


def fixture(ctx, suffix):
    path = ctx.inst.root / 'temp' / ('Northwind-delivery.' + suffix)
    marker = 'NW-AI-731'
    aliases = {'dot': 'doc', 'pps': 'ppt', 'markdown': 'md', 'htm': 'html', 'jpeg': 'jpg', 'tif': 'tiff'}
    if suffix in aliases:
        data = fixture(ctx, aliases[suffix]).read_bytes()
        if suffix in ('dot', 'pps'):
            old = 'Northwind NW-AI-731 '.encode('utf-16le')
            require(old in data, 'Legacy Northwind fixture omits its title')
            data = data.replace(old, ('Northwind NW-AI-731' + ('D' if suffix == 'dot' else 'P')).encode('utf-16le'))
        elif suffix in ('markdown', 'htm'):
            data += ('\n<!-- Fictional Northwind ' + suffix + ' source -->\n').encode()
        elif suffix in ('jpeg', 'tif'):
            from PIL import Image, ImageDraw
            image = Image.open(fixture(ctx, aliases[suffix]))
            ImageDraw.Draw(image).text((80, 950), 'Fictional Northwind ' + suffix + ' edition', fill='#18483f')
            if suffix == 'jpeg':
                image.save(path, format='JPEG', comment=b'Fictional Northwind alternate JPEG')
            else:
                image.save(path, format='TIFF', tiffinfo={270: 'Fictional Northwind alternate TIFF'})
            return path
        path.write_bytes(data)
    elif suffix in LEGACY_FIXTURES:
        path.write_bytes(zlib.decompress(base64.b64decode(LEGACY_FIXTURES[suffix])))
    elif suffix in ('css', 'py', 'lua', 'rb', 'php', 'java', 'c', 'cpp', 'cs', 'go', 'sh', 'bat', 'ps1', 'yml', 'yaml', 'toml', 'sql', 'tsv'):
        prefix = '--' if suffix in ('lua', 'sql') else '#' if suffix in ('py', 'rb', 'sh', 'yml', 'yaml', 'toml') else '//'
        text = 'Product\tUnits\tMarker\nHarbor tea\t120\tNW-AI-731\n' if suffix == 'tsv' else prefix + ' Fictional Northwind ' + suffix + ' source NW-AI-731: returns within 17 days.\n'
        path.write_text(text)
    elif suffix == 'svg':
        path.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1500" height="1000"><rect width="1500" height="1000" fill="white"/><g font-family="sans-serif" font-size="64" fill="#18483f"><text x="80" y="160">Fictional Northwind</text><text x="80" y="320">Warehouse NW-AI-731</text><text x="80" y="480">Returns within 17 days</text><text x="80" y="640">Harbor tea: 120 units</text></g></svg>')
    elif suffix == 'ico':
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (256, 256), '#fffdf7')
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=26)
        for y, line in zip((25, 80, 135, 190), ('NORTHWIND', 'NW-AI-731', 'Returns: 17 days', 'Fictional icon')):
            draw.text((10, y), line, font=font, fill='#18483f')
        image.save(path, format='ICO', sizes=[(256, 256)])
    elif suffix == 'dxf':
        import ezdxf
        drawing = ezdxf.new()
        drawing.modelspace().add_text('Fictional Northwind warehouse NW-AI-731')
        drawing.modelspace().add_line((0, 0), (120, 80))
        drawing.saveas(path)
    elif suffix in ('txt', 'md', 'csv', 'html', 'json', 'xml', 'js', 'rtf', 'bin'):
        text = {
            'txt': POLICY, 'md': POLICY + '\nFictional Northwind Markdown reference edition.\n',
            'csv': 'Product,Units,Marker\nHarbor tea,120,NW-AI-731\nOrchard coffee,80,NW-AI-731\n',
            'html': '<!doctype html><html><head><title>Northwind delivery</title></head><body><h1>Northwind delivery handbook</h1><p>Fictional warehouse NW-AI-731: returns within 17 days.</p><table><tr><th>Product</th><th>Units</th></tr><tr><td>Harbor tea</td><td>120</td></tr><tr><td>Orchard coffee</td><td>80</td></tr></table></body></html>',
            'json': json.dumps({'company': 'Fictional Northwind', 'marker': marker, 'returns_days': 17, 'orders': [{'product': 'Harbor tea', 'units': 120}]}),
            'xml': '<northwind><marker>NW-AI-731</marker><returns days="17"/><product units="120">Harbor tea</product></northwind>',
            'js': 'const northwind = {company: "Fictional Northwind", marker: "NW-AI-731", returnsDays: 17};\n',
            'rtf': r'{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs28 Fictional Northwind delivery\par NW-AI-731: returns within 17 days.\par Harbor tea: 120 units.}',
            'bin': 'Fictional Northwind binary shipment NW-AI-731\x00\x01',
        }[suffix]
        path.write_bytes(text.encode())
    elif suffix == 'pdf':
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(str(path))
        for i, line in enumerate(POLICY.splitlines()):
            canvas.drawString(45, 790 - i * 24, line)
        canvas.save()
    elif suffix == 'docx':
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.add_heading('Northwind delivery handbook', 0)
        for line in POLICY.splitlines()[2:8]:
            doc.add_paragraph(line)
        table = doc.add_table(rows=1, cols=3)
        for cell, title in zip(table.rows[0].cells, ['Product', 'Units', 'Revenue']):
            cell.text = title
        for row in [('Harbor tea', '120', '2400'), ('Orchard coffee', '80', '3200'), ('Meadow cocoa', '95', '2850')]:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
        doc.add_picture(str(fixture(ctx, 'png')), width=Inches(5))
        doc.save(path)
    elif suffix in ('xlsx', 'xlsm'):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = 'Northwind deliveries'
        for row in [('Fictional Northwind ' + suffix + ' workbook', marker, 'Returns: 17 days'), ('Product', 'Units', 'Revenue'), ('Harbor tea', 120, 2400), ('Orchard coffee', 80, 3200), ('Meadow cocoa', 95, 2850)]:
            sheet.append(row)
        book.save(path)
        if suffix == 'xlsm':
            import io
            output = io.BytesIO()
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
                for item in source.infolist():
                    content = source.read(item.filename)
                    if item.filename == '[Content_Types].xml':
                        content = content.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                    target.writestr(item, content)
            path.write_bytes(output.getvalue())
    elif suffix == 'pptx':
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'Northwind delivery handbook'
        slide.placeholders[1].text = 'Fictional Northwind NW-AI-731\nReturns within 17 days\nHarbor tea: 120 units\nOrchard coffee: 80 units\nMeadow cocoa: 95 units'
        deck.save(path)
    elif suffix in ('png', 'jpg', 'webp', 'tiff', 'bmp', 'gif'):
        from PIL import Image, ImageDraw, ImageFont
        im = Image.new('RGB', (1500, 1000), '#fffdf7')
        draw = ImageDraw.Draw(im)
        font = ImageFont.load_default(size=44)
        draw.rectangle((0, 0, 1500, 155), fill='#18483f')
        draw.text((80, 50), 'NORTHWIND  /  DELIVERY HANDBOOK', font=font, fill='white')
        for i, line in enumerate(['Fictional training document', 'Warehouse marker: NW-AI-731', 'Return window: 17 days', 'Harbor tea: 120 units', 'Orchard coffee: 80 units', 'Meadow cocoa: 95 units', 'Dispatch confirmed by Mira Chen']):
            draw.text((80, 215 + 95 * i), line, font=font, fill='#18483f')
        im.save(path)
    elif suffix == 'scan.pdf':
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.utils import ImageReader
        from PIL import Image
        im = Image.open(fixture(ctx, 'png'))
        canvas = Canvas(str(path), pagesize=(750, 500))
        canvas.drawImage(ImageReader(im), 0, 0, width=750, height=500)
        canvas.showPage()
        canvas.drawImage(ImageReader(im), 0, 0, width=750, height=500)
        canvas.save()
    elif suffix == 'zip':
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('Northwind/handbook.md', POLICY)
            archive.writestr('Northwind/shipments.csv', 'Product,Units,Marker\nHarbor tea,120,NW-AI-731\n')
    else:
        require(False, 'No Northwind fixture generator for ' + suffix)
    return path


def preview(ctx, filename, markdown=False):
    from playwright.sync_api import expect
    nav(ctx, 'Uploads')
    uploads = ctx.browser.frame.locator('kickside-uploads')
    uploads.get_by_label('Search files', exact=True).fill(filename)
    uid = ctx.state['ai_uploads'][filename]
    uploads.locator('[data-testid=uploads-row-preview][data-upload-id="' + uid + '"]').click()
    drawer = ctx.browser.frame.get_by_role('dialog', name='File preview')
    expect(drawer).to_be_visible(timeout=30000)
    expect(drawer).to_contain_text(filename)
    if markdown:
        drawer.get_by_test_id('uploads-view-markdown').click()
        expect(drawer).to_contain_text('NW-AI-731', timeout=30000)
    shot(ctx, 'northwind-file-preview')
    drawer.get_by_role('button', name='Close preview', exact=True).click()


def upload_format(suffix):
    def run(ctx):
        fresh_ui(ctx)
        path = fixture(ctx, suffix)
        uid, artifacts = completed_upload(ctx, path)
        if suffix in DOCUMENT_FORMATS:
            require(artifacts['primary']['mime_type'] == 'text/markdown', suffix + ': document does not publish converted Markdown: ' + str(artifacts['primary']))
        if suffix not in ('zip', 'bin', 'dxf', 'svg'):
            primary = artifacts['primary']
            data, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + primary['artifact_id'] + '/raw', 'converted-content')
        image = (suffix in IMAGE_FORMATS and suffix != 'svg') or suffix == 'scan.pdf'
        textual = suffix not in IMAGE_FORMATS + ('zip', 'bin', 'dxf')
        if textual:
            require('NW-AI-731' in data.decode(), suffix + ': converted content omits Northwind marker')
        if image:
            ocr = [a for a in [artifacts['primary'], *artifacts['artifacts']] if a.get('kind') in ('markdown', 'text') or a.get('mime_type') == 'text/markdown']
            require(ocr, suffix + ': no OCR Markdown artifact')
            contents = []
            for a in ocr:
                content, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + a['artifact_id'] + '/raw', 'ocr-' + a['artifact_id'])
                contents.append(content.decode())
            require(any('NW-AI-731' in text and '17' in text for text in contents), suffix + ': OCR omits warehouse marker or returns window')
        if suffix == 'svg':
            require(not artifacts.get('primary'), 'Source-only SVG unexpectedly publishes a text primary')
            upload = ctx.inst.request('GET', '/api/v1/uploads/' + uid, label='svg-source-state')['upload']
            image_id = upload.get('meta', {}).get('primary_artifact_id')
            require(image_id and upload['mime_type'] == 'image/svg+xml', 'SVG omits its original image artifact')
            data, mime = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + image_id + '/raw', 'svg-source-artifact')
            require(data == path.read_bytes() and mime.startswith('image/svg+xml'), 'SVG source artifact differs from the original')
            require(b'<svg' in data and b'NW-AI-731' in data and b'17' in data, 'SVG markup omits its declared source facts')
            require(not any(a.get('kind') in ('markdown', 'text') for a in artifacts['artifacts']), 'Source-only SVG unexpectedly publishes OCR text')
        if suffix == 'zip':
            parent = ctx.inst.request('GET', '/api/v1/uploads/' + uid, label='archive-container')['upload']
            require(parent['meta']['archive_member_count'] == 2 and len(parent['meta']['child_upload_ids']) == 2, 'Archive omits extracted members')
            for child in parent['meta']['child_upload_ids']:
                record = ctx.inst.request('GET', '/api/v1/uploads/' + child, label='archive-member-' + child)['upload']
                require(record['meta'].get('parent_upload_id') == uid, 'Archive child loses parent provenance: ' + str(record))
                content, _ = raw_bytes(ctx, '/api/v1/uploads/' + child + '/raw', 'archive-member-bytes-' + child)
                require(b'NW-AI-731' in content, 'Archive child bytes omit Northwind marker')
        listed = ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='upload-list')
        require(uid in json.dumps(listed), 'Uploaded file is absent from list')
        preview(ctx, path.name, markdown=suffix not in ('zip', 'bin', 'dxf', 'svg') and artifacts['primary']['mime_type'] == 'text/markdown')
    return run


def docx_primary_body(ctx):
    uid, artifacts = completed_upload(ctx, fixture(ctx, 'docx'))
    primary = artifacts['primary']
    require(primary['mime_type'] == 'text/markdown' and primary['metadata'].get('converter') == 'doc2md', 'Word body does not publish doc2md Markdown')
    content, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + primary['artifact_id'] + '/raw', 'word-primary-body')
    text = content.decode()
    required = ('Revenue', '2400', 'Jonah Reed', '09:30')
    ctx.inst.evidence.save('word-body-facts', {'expected': required, 'markdown': text})
    require(all(value in text for value in required), 'Converted Word Markdown loses original body facts: ' + str([value for value in required if value not in text]))


def models_catalog(ctx):
    inst = ctx.inst
    admin = inst.request('GET', '/api/v1/models/admin', label='model-admin')
    require(any(m['name'] == ctx.state['model'] for m in admin['models']), 'Gemma is absent from model admin')
    models = inst.request('GET', '/api/v1/models/list?capability=generate', label='model-list')['models']
    require(any(m['name'] == ctx.state['model'] for m in models), 'Gemma is absent from generation picker')
    byname = inst.request('GET', '/api/v1/models/by-name?models=' + quote(ctx.state['model']), label='model-by-name')['models']
    require(byname[ctx.state['model']]['name'] == ctx.state['model'], 'Model lookup resolves a different card')
    classes = inst.request('GET', '/api/v1/models/classes', label='model-classes')['classes']
    require(any(ctx.state['model'] in json.dumps(c['models']) for c in classes), 'Gemma is absent from model classes')
    tested = inst.request('POST', '/api/v1/models/cards/test', {'name': ctx.state['model']}, label='gemma-health', timeout=120)
    require(tested['output'].strip().lower().strip('.!') == 'ok', 'Gemma generation probe does not return ok: ' + str(tested))


def provider_preset(entry):
    def run(ctx):
        inst = ctx.inst
        preset = next(p for p in inst.request('GET', '/api/v1/models/admin', label='preset')['presets'] if p['id'] == entry)
        title = 'Northwind ' + preset['title'] + ' configuration'
        created = inst.request('POST', '/api/v1/models/provider-profiles', {
            'name': title, 'provider_id': preset['provider_id'], 'config': preset['config'], 'enabled': False,
        }, expect=(201,), label='preset-create')['profile']
        pid = created['id']
        require(created['name'] == title and created['provider_id'] == preset['provider_id'], 'Provider preset creates a different adapter')
        current = next(p for p in inst.request('GET', '/api/v1/models/admin', label='preset-persisted')['profiles'] if p['id'] == pid)
        require(current['enabled'] is False and current['config'] == created['config'], 'Preset defaults do not persist')
        updated = inst.request('PUT', '/api/v1/models/provider-profiles/' + pid, {'name': title + ' revised'}, label='preset-update')['profile']
        require(updated['name'] == title + ' revised', 'Provider rename does not persist')
        inst.request('DELETE', '/api/v1/models/provider-profiles/' + pid, label='preset-delete')
        require(not any(p['id'] == pid for p in inst.request('GET', '/api/v1/models/admin', label='preset-absent')['profiles']), 'Deleted provider remains in admin')
    return run


def provider_live(ctx):
    inst = ctx.inst
    profile = next(p for p in inst.request('GET', '/api/v1/models/admin')['profiles'] if p['name'] == 'On-prem GPU server')
    pid = profile['id']
    test = inst.request('POST', '/api/v1/models/provider-profiles/' + pid + '/test', {'model': inst.config['model']['llm']}, label='provider-live-test', timeout=120)
    require(test.get('message') and test.get('success'), 'Provider probe supplies no observable result')
    original = next(m for m in inst.request('GET', '/api/v1/models/admin')['models'] if m['name'] == ctx.state['model'])
    discovered = inst.request('POST', '/api/v1/models/provider-profiles/' + pid + '/discover', {}, label='provider-discover', timeout=120)
    require(inst.config['model']['llm'] in json.dumps(discovered), 'Gemma is absent from discovery')
    card = {**original, 'name': 'northwind/gemma-review', 'title': 'Northwind Gemma review', 'class': []}
    inst.request('PUT', '/api/v1/models/cards', card, label='card-create')
    require(any(m['name'] == card['name'] and m['title'] == card['title'] for m in inst.request('GET', '/api/v1/models/admin', label='card-read')['models']), 'Model card does not persist')
    inst.request('DELETE', '/api/v1/models/cards?name=' + quote(card['name']), label='card-delete')
    require(not any(m['name'] == card['name'] for m in inst.request('GET', '/api/v1/models/admin', label='card-absent')['models']), 'Deleted card remains listed')
    retained = next(m for m in inst.request('GET', '/api/v1/models/admin', label='discovery-class-assignments')['models'] if m['name'] == ctx.state['model'])
    require(retained['class'] == original['class'], 'Provider rediscovery changes Gemma class assignments: before=' + str(original['class']) + '; after=' + str(retained['class']))


def kb_create(ctx, kind, title, *, analysis=False):
    config = {'extraction_model': ctx.state['model']}
    if analysis:
        config['analysis_model'] = ctx.state['model']
    created = ctx.inst.request('POST', '/api/v1/knowledge/create', {'name': title, 'kb_type': kind,
        'embedding_model': ctx.state['embed'], 'description': 'Fictional Northwind delivery and returns knowledge',
        'config': config}, label='kb-create')
    kid = created['id']
    current = ctx.inst.request('GET', '/api/v1/knowledge/' + kid, label='kb-created')
    require(current['name'] == title and current['kb_type'] == kind and current['embedding_model'] == ctx.state['embed'], 'Created KB config differs')
    return kid


def kb_ingest(ctx, kid, marker='NW-AI-731', content=POLICY):
    inst = ctx.inst
    base = '/api/v1/knowledge/' + kid
    receipt = inst.request('POST', base + '/ingest', {'title': 'Northwind delivery handbook',
        'content': content.replace('NW-AI-731', marker), 'content_type': 'text/markdown',
        'external_id': 'northwind-delivery', 'external_version': '1'}, label='kb-ingest')
    require(receipt.get('thread_id') and receipt.get('projection_id'), 'Ingest supplies no durable projection receipt')
    return kb_indexed(ctx, kid, marker)


def kb_indexed(ctx, kid, marker='NW-AI-731'):
    inst = ctx.inst
    base = '/api/v1/knowledge/' + kid
    def indexed():
        result = inst.request('POST', base + '/search', {'query': marker + ' Northwind delivery returns', 'limit': 10}, label='kb-indexed')
        return result if marker in json.dumps(result) else None
    return inst.wait(indexed, 'Northwind knowledge indexed', seconds=120, pump=ctx.browser.page.wait_for_timeout)


def knowledge_catalog(ctx):
    inst = ctx.inst
    engines = inst.request('GET', '/api/v1/knowledge/engines', label='engines')['engines']
    require({'document', 'wiki'} <= {e['kb_type'] for e in engines}, 'Document/wiki engines are missing')
    models = inst.request('GET', '/api/v1/knowledge/models', label='knowledge-models')
    require(ctx.state['embed'] in json.dumps(models), 'Embedding model is absent from KB model picker')
    renderers = inst.request('GET', '/api/v1/renderers', label='renderers')
    require(renderers.get('renderers'), 'Renderer catalog is empty')


def knowledge_document(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'document', 'Northwind delivery library', analysis=True)
    ctx.state['ai_document_kb'] = kid
    kb_ingest(ctx, kid)
    require(kid in json.dumps(inst.request('GET', '/api/v1/knowledge/list', label='kb-list')), 'KB list omits new library')
    base = '/api/v1/knowledge/' + kid
    inst.request('PUT', base, {'name': 'Northwind delivery library', 'description': 'Harbor and Orchard returns, shipment, and product references'}, label='kb-update')
    require('Harbor and Orchard' in inst.request('GET', base, label='kb-updated')['description'], 'KB update does not persist')
    token = inst.request('POST', base + '/upload-token', {}, label='kb-upload-token')
    require(token.get('upload_token') and token.get('upload_id'), 'KB upload token is empty')
    for suffix in ('txt', 'docx', 'xlsx'):
        path = fixture(ctx, suffix)
        uid, _ = completed_upload(ctx, path)
        queued = inst.request('POST', base + '/ingest-upload', {'upload_id': uid}, label='kb-file-' + suffix)
        require(queued['queued'] == 1 and queued['total'] == 1, suffix + ': file is not queued for KB ingestion')
    inst.wait(lambda: (lambda v: v if v['document_count'] >= 4 else None)(inst.request('GET', base, label='kb-file-count')), 'Several document types indexed', seconds=120, pump=ctx.browser.page.wait_for_timeout)
    nav(ctx, 'Knowledge')
    ctx.browser.frame.get_by_text('Northwind delivery library', exact=True).first.click()
    ctx.browser.frame.get_by_text('Northwind delivery handbook', exact=False).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-document-library')


def require_cited_sources(answer, expected_facts):
    references = [int(value) for value in re.findall(r'\[(\d+)\]', str(answer.get('answer')))]
    results = answer.get('results', [])
    require(references and all(1 <= index <= len(results) for index in references), 'Answer citations do not identify returned source passages')
    cited = [results[index - 1] for index in references]
    require(all(row.get('id') and row.get('content') for row in cited), 'Cited passages omit their source identity or text')
    evidence = '\n'.join(row['content'] for row in cited)
    require(all(fact in evidence for fact in expected_facts), 'Cited passages do not support the expected source facts')


def knowledge_answer(ctx):
    base = '/api/v1/knowledge/' + ctx.state['ai_document_kb']
    answer = ctx.inst.request('POST', base + '/ask', {'query': 'What is the Northwind return window and warehouse marker? Cite the handbook.'}, label='cited-answer', timeout=120)
    require('17' in str(answer['answer']) and 'NW-AI-731' in str(answer['answer']), 'Knowledge answer omits source facts: ' + str(answer))
    require(answer.get('citations') or answer.get('sources') or re.search(r'\[\d+\]', str(answer['answer'])), 'Knowledge answer omits citations: ' + str(answer))
    require_cited_sources(answer, ('NW-AI-731', '17'))


def renderer_docx_grounded_answer_settles(ctx):
    inst, browser = ctx.inst, ctx.browser
    title = 'Northwind DOCX grounded answer'
    kid = kb_create(ctx, 'document', title, analysis=True)
    inst.cleanup.add('renderer knowledge', lambda: inst.request('DELETE', '/api/v1/knowledge/' + kid))
    uid, artifacts = completed_upload(ctx, fixture(ctx, 'docx'))
    primary = artifacts['primary']
    content, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + primary['artifact_id'] + '/raw',
                           'renderer-docx-converted')
    require('NW-AI-731' in content.decode('utf-8') and '17' in content.decode('utf-8'),
            'Converted DOCX omits source facts')
    base = '/api/v1/knowledge/' + kid
    queued = inst.request('POST', base + '/ingest-upload', {'upload_id': uid}, label='renderer-docx-ingest')
    require(queued['queued'] == 1 and queued['total'] == 1, 'DOCX is not queued for ingestion')
    kb_indexed(ctx, kid)
    nav(ctx, 'Knowledge')
    browser.frame.get_by_text(title, exact=True).first.click()
    browser.frame.get_by_role('button', name='Search', exact=True).click()
    browser.frame.get_by_role('button', name='Ask', exact=True).first.click()
    question = browser.frame.get_by_placeholder('Ask this knowledge base a question…')
    question.fill('What is the Northwind warehouse marker and return window? Cite the handbook.')
    with browser.page.expect_response(lambda r: r.request.method == 'POST' and urlsplit(r.url).path == base + '/ask',
                                      timeout=120000) as response:
        question.press('Enter')
    answer = response.value.json()
    inst.evidence.save('renderer-docx-grounded-answer', {'status': response.value.status, 'response': answer})
    require(response.value.status == 200 and answer['results'], 'DOCX knowledge UI ask fails')
    require('NW-AI-731' in str(answer.get('answer')) and '17' in str(answer.get('answer')),
            'DOCX knowledge answer omits source facts')
    require_cited_sources(answer, ('NW-AI-731', '17'))
    browser.settle()
    browser.clean(0)
    shot(ctx, 'renderer-docx-grounded-answer')


def knowledge_delete_content(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'document', 'Northwind temporary records')
    kb_ingest(ctx, kid, 'NW-DELETE-811')
    base = '/api/v1/knowledge/' + kid
    deleted = inst.request('DELETE', base + '/content', {'external_id': 'northwind-delivery'}, label='content-delete')
    require(deleted.get('thread_id'), 'Content delete supplies no durable receipt')
    inst.wait(lambda: (lambda r: r if 'NW-DELETE-811' not in json.dumps(r) else None)(inst.request('POST', base + '/search', {'query': 'NW-DELETE-811', 'limit': 10}, label='content-absent')), 'Deleted content leaves search', seconds=120)
    inst.request('DELETE', base, label='kb-delete')
    require(kid not in json.dumps(inst.request('GET', '/api/v1/knowledge/list', label='kb-absent')), 'Deleted KB remains listed')


def agent_crud(ctx):
    from playwright.sync_api import expect
    inst, browser = ctx.inst, ctx.browser
    fresh_ui(ctx)
    nav(ctx, 'Agents')
    agents = browser.frame.locator('kickside-agents')
    agents.get_by_role('button', name='New agent', exact=True).first.click()
    form = agents.get_by_role('dialog', name='New agent')
    form.get_by_label('Agent name', exact=True).fill('Northwind logistics analyst')
    form.get_by_placeholder('What this agent is for.').fill('Analyzes fictional Harbor and Orchard delivery records.')
    form.get_by_placeholder('You are a helpful assistant that…').fill('You are Northwind logistics analyst Mira. Use only fictional Northwind records. Cite handbook sources.')
    form.get_by_role('tab', name='Model', exact=True).click()
    tokens = form.get_by_role('spinbutton', name=re.compile(r'^Max tokens'))
    ceiling = tokens.get_attribute('max')
    require(ceiling and int(ceiling) > 0, 'Selected agent model has no output budget')
    tokens.fill(ceiling)
    form.get_by_role('tab', name='Capabilities', exact=False).click()
    form.get_by_role('button', name='Remove Artifacts capability', exact=True).click()
    picker = form.get_by_role('combobox', name='Search capabilities to add')
    picker.fill('Artifacts')
    form.get_by_role('option', name='Artifacts', exact=False).click()
    expect(form.get_by_role('button', name='Remove Artifacts capability')).to_be_visible()
    shot(ctx, 'northwind-agent-capabilities')
    with browser.page.expect_response(lambda r: '/api/v1/user-agents' in r.url and r.request.method == 'POST') as response:
        form.get_by_role('button', name='Create agent', exact=True).click()
    created = response.value.json()
    require(created.get('component_id'), 'Agent UI creation supplies no component ID')
    aid = created['component_id']
    ctx.state['ai_agent'] = aid
    kb = kb_create(ctx, 'document', 'Northwind agent-scoped handbook')
    kb_ingest(ctx, kb)
    traits = ['kickside.agents.traits:artifacts', 'kickside.knowledge.traits:kb_reader']
    scopes = {'kickside.knowledge.traits:kb_reader': {'kb_ids': [kb]}}
    inst.request('PUT', '/api/v1/user-agents/' + aid, {'model': ctx.state['model'], 'traits': traits, 'trait_contexts': scopes}, label='agent-model-and-scope')
    current = inst.request('GET', '/api/v1/user-agents/' + aid, label='agent-persisted')['agent']
    require(current['title'] == 'Northwind logistics analyst' and current['model'] == ctx.state['model'] and current['traits'] == traits and current['trait_contexts'] == scopes, 'Agent model/KB scope differs after save')
    listed = inst.request('GET', '/api/v1/user-agents', label='agent-list')['agents']
    require(any(a['id'] == aid for a in listed), 'Agent list omits created analyst')
    agents.get_by_role('button', name='Edit Northwind logistics analyst', exact=True).click()
    edit = agents.get_by_role('dialog', name='Edit agent')
    edit.get_by_role('tab', name='Capabilities', exact=False).click()
    edit.get_by_role('button', name='Remove Artifacts capability', exact=True).click()
    with browser.page.expect_response(lambda r: '/api/v1/user-agents/' + aid in r.url and r.request.method == 'PUT') as response:
        edit.get_by_role('button', name='Save changes', exact=True).click()
    require(response.value.json().get('success'), 'Agent capability removal does not save')
    require('kickside.agents.traits:artifacts' not in inst.request('GET', '/api/v1/user-agents/' + aid, label='agent-capability-removed')['agent']['traits'], 'Removed agent capability remains persisted')
    shot(ctx, 'northwind-agent-roster')
    inst.request('DELETE', '/api/v1/user-agents/' + aid, label='agent-delete')
    require(not any(a['id'] == aid for a in inst.request('GET', '/api/v1/user-agents', label='agent-absent')['agents']), 'Deleted agent remains listed')


def agent_catalog(ctx):
    inst = ctx.inst
    agents = inst.request('GET', '/api/v1/agents/list', label='agents-catalog')
    require(agents.get('agents'), 'Agent picker is empty')
    name = agents['agents'][0]['name']
    byname = inst.request('GET', '/api/v1/agents/by-name?agents=' + quote(name), label='agent-by-name')
    require(name in json.dumps(byname), 'Agent name resolution omits requested agent')
    namespaces = inst.request('GET', '/api/v1/agents/namespaces', label='agent-namespaces')
    require(namespaces.get('namespaces'), 'Namespace picker is empty')
    traits = inst.request('GET', '/api/v1/agents/traits', label='agent-traits')['traits']
    require(any(t['id'] == 'kickside.agents.traits:artifacts' for t in traits), 'Artifacts is absent from capability picker')
    tools = inst.request('GET', '/api/v1/agents/tools?q=ViewFiles', label='agent-tools')['tools']
    require(any(t['llm_alias'] == 'ViewFiles' for t in tools), 'Public ViewFiles tool is absent from tool picker')
    memory = inst.request('GET', '/api/v1/agents/memory-options', label='memory-options')
    require(memory.get('providers') and isinstance(memory.get('items'), list), 'Memory target picker has no provider')


def open_agent(ctx, aid):
    nav(ctx, 'Agents')
    agent = ctx.inst.request('GET', '/api/v1/user-agents/' + aid)['agent']
    card = ctx.browser.frame.locator('kickside-agents .kx-card').filter(has_text=agent['title'])
    card.locator('.kx-card-chat').click()
    ctx.browser.frame.locator('wippy-chat').get_by_placeholder('Type a message', exact=True).wait_for(timeout=30000)


def messages(ctx, sid):
    stored, seen, cursors = [], set(), set()
    cursor = None
    while True:
        path = '/api/v1/sessions/messages?session_id=' + quote(sid, safe='') + '&limit=100&direction=before'
        if cursor is not None:
            path += '&cursor=' + quote(cursor, safe='')
        page = ctx.inst.request('GET', path, label='session-messages-' + str(len(cursors)))
        older = []
        for row in page['messages']:
            message_id = row['message_id']
            if message_id not in seen:
                older.append(row)
                seen.add(message_id)
        stored = older + stored
        pagination = page['pagination']
        if not pagination['has_more']:
            return stored
        cursor = pagination.get('next_cursor')
        require(isinstance(cursor, str) and cursor and cursor not in cursors,
                'Session history pagination does not advance')
        cursors.add(cursor)


def function_messages(stored, alias):
    return [row for row in stored if isinstance(row.get('metadata'), dict)
            and row['metadata'].get('function_name') == alias]


def create_agent(ctx, title, traits=(), contexts=None, **options):
    body = {'title': title, 'description': 'Fictional Northwind logistics, returns, and product planning.', 'model': ctx.state['model'], 'traits': list(traits),
        'trait_contexts': contexts or {}, 'prompt': 'You work with fictional Northwind data. Use the requested tools and report their actual results.'}
    body.update(options)
    aid = ctx.inst.request('POST', '/api/v1/user-agents', body, label='chat-agent-create')['component_id']
    require(ctx.inst.request('GET', '/api/v1/user-agents/' + aid, label='chat-agent-config')['agent']['model'] == ctx.state['model'], 'Chat agent does not use Gemma')
    return aid


def start_chat(ctx, aid):
    fresh_ui(ctx)
    before = {s['session_id'] for s in ctx.inst.request('GET', '/api/v1/sessions?limit=100')['sessions']}
    open_agent(ctx, aid)
    sessions = ctx.inst.request('GET', '/api/v1/sessions?limit=100', label='independent-session')['sessions']
    new = [s for s in sessions if s['session_id'] not in before]
    require(len(new) == 1, 'Starting an agent chat does not create exactly one session')
    ctx.state['ai_current_chat'] = new[0]['session_id']
    return new[0]['session_id']


def tool_events(ctx, start, alias):
    frames = ctx.browser.frames[start:]
    events = [f for f in frames if isinstance(f.get('data'), dict) and 'tool' in str(f['data'].get('type', '')).lower()]
    ctx.inst.evidence.save('chat-tool-events', events)
    require(any(alias in json.dumps(f) for f in events), 'Chat has no observable ' + alias + ' tool invocation')
    functions = function_messages(messages(ctx, ctx.state['ai_current_chat']), alias)
    require(functions and all(m['metadata'].get('status') == 'success' for m in functions), alias + ': persisted tool invocation does not succeed')
    return functions


def scoped_knowledge_chat(alias):
    def run(ctx):
        kid = kb_create(ctx, 'document', 'Northwind scoped ' + alias + ' handbook')
        kb_ingest(ctx, kid)
        if alias == 'RetrieveKB':
            ctx.inst.request('PUT', '/api/v1/knowledge/' + kid, {'config': {'query_driver': 'agentic', 'analysis_model': ctx.state['model'], 'extra_prompt': 'Always call RetrieveKB before answering. Search for the Northwind delivery handbook warehouse marker and return window.'}}, label='agentic-retrieval-config')
            saved = ctx.inst.request('GET', '/api/v1/knowledge/' + kid, label='agentic-retrieval-persisted')
            require(saved['config']['query_driver'] == 'agentic' and saved['config']['analysis_model'] == ctx.state['model'], 'KB does not persist its agentic retrieval driver')
            answer = ctx.inst.request('POST', '/api/v1/knowledge/' + kid + '/ask', {'query': 'Use RetrieveKB to retrieve the fictional Northwind delivery handbook. What is its exact warehouse marker and returns window?', 'model': ctx.state['model']}, label='agentic-retrieve-answer', timeout=120)
            require('NW-AI-731' in answer['answer'] and '17' in answer['answer'], 'Agentic retrieval omits source facts')
            require(answer['results'] and 'NW-AI-731' in json.dumps(answer['results']), 'Agentic retrieval supplies no source passages')
            return
        trait = 'kickside.knowledge.traits:kb_reader'
        aid = create_agent(ctx, 'Northwind ' + alias + ' reader', (trait,), {trait: {'kb_ids': [kid]}})
        sid = start_chat(ctx, aid)
        start = len(ctx.browser.frames)
        ctx.browser.stream('Call ' + alias + ' to retrieve Northwind delivery handbook passages about the returns window and warehouse marker. Report the exact marker and days from the tool. Explain ten delivery checks and end with NW-SCOPED-' + alias.upper() + '.', 'NW-SCOPED-' + alias.upper(), 'scoped-' + alias)
        tool_events(ctx, start, alias)
        answers = '\n'.join(str(m['data']) for m in messages(ctx, sid) if m['type'] == 'assistant')
        require('NW-AI-731' in answers and '17' in answers, alias + ': scoped Gemma answer omits handbook facts')
        chat_shot(ctx, 'northwind-scoped-knowledge')
    return run


def memory_chat(ctx):
    kid = kb_create(ctx, 'document', 'Northwind dispatch memories')
    kb_ingest(ctx, kid, 'NW-MEMORY-731')
    trait = 'kickside.memory:knowledge_based_memory'
    aid = create_agent(ctx, 'Northwind memory analyst', (trait,), {trait: {'kb_ids': [kid], 'read_enabled': True, 'write_enabled': False}})
    sid = start_chat(ctx, aid)
    start = len(ctx.browser.frames)
    ctx.browser.stream('Call SearchMemory with fields ["content", "text", "summary"] for Northwind warehouse marker and returns policy. Report exact marker and days. Explain ten warehouse checks and end with NW-MEMORY-DONE.', 'NW-MEMORY-DONE', 'memory-recall')
    tool_events(ctx, start, 'SearchMemory')
    require(any(m['type'] == 'assistant' and 'NW-MEMORY-731' in str(m['data']) and '17' in str(m['data']) for m in messages(ctx, sid)), 'Long-term memory recall omits stored facts')
    chat_shot(ctx, 'northwind-memory-chat')


def capabilities_chat(master):
    def run(ctx):
        uid, _ = completed_upload(ctx, fixture(ctx, 'txt'))
        trait = 'kickside.agents.traits:' + ('trait_master' if master else 'trait_curator')
        files = 'kickside.uploads.traits:file_manager'
        aid = create_agent(ctx, 'Northwind ' + ('all' if master else 'selected') + ' capabilities', (trait,), {trait: {'inject_trait_catalog': True}})
        sid = start_chat(ctx, aid)
        chat = ctx.browser.frame.locator('wippy-chat')
        ctx.inst.evidence.save('chat-controls', chat.aria_snapshot())
        start = len(ctx.browser.frames)
        ctx.browser.stream('Call Capabilities list, then activate ' + files + '. Use the newly available ViewFiles tool on upload_ids ["' + uid + '"] to read the Northwind handbook. Report the exact warehouse marker and returns window. Explain ten dispatch checks, keep the capability active, and end with NW-CAP-ACTIVE.', 'NW-CAP-ACTIVE', 'chat-capability-activate')
        tool_events(ctx, start, 'Capabilities')
        tool_events(ctx, start, 'ViewFiles')
        config = ctx.inst.request('GET', '/api/v1/sessions/get?session_id=' + sid, label='chat-active-capability')['session']['config']
        require(files in json.dumps(config['active_traits']), 'Activated File Manager does not persist in session capabilities')
        require(any(m['type'] == 'assistant' and 'NW-AI-731' in str(m['data']) for m in messages(ctx, sid)), 'Activated capability does not read actual file contents')
        start = len(ctx.browser.frames)
        ctx.browser.stream('Call Capabilities deactivate for ' + files + '. Then call Capabilities list and report that File Manager is inactive. Explain ten short reasons to finish a dispatch review, then end with NW-CAP-INACTIVE.', 'NW-CAP-INACTIVE', 'chat-capability-deactivate')
        events = tool_events(ctx, start, 'Capabilities')
        require('deactivat' in json.dumps(events).lower(), 'Chat does not deactivate File Manager')
        config = ctx.inst.request('GET', '/api/v1/sessions/get?session_id=' + sid, label='chat-inactive-capability')['session']['config']
        require(files not in json.dumps(config['active_traits']), 'Deactivated File Manager remains in session capabilities')
        chat_shot(ctx, 'northwind-capability-chat')
    return run


def platform_tools(ctx):
    trait = 'kickside.agents.traits:platform_manager'
    rpc = mcp(ctx, trait_ids=(trait,))
    tool_value(rpc, 'use_trait', {'id': trait}, 'platform-activate')
    listed = json.loads(tool_value(rpc, 'Platform', {'action': 'list_agents'}, 'platform-agents'))
    require(ctx.state['agent'] in json.dumps(listed), 'Platform agent list omits Northwind assistant')
    created = json.loads(tool_value(rpc, 'Platform', {'action': 'upsert_agent', 'name': 'Northwind platform planner', 'model': ctx.state['model'], 'prompt': 'Fictional Northwind planning only.', 'traits': []}, 'platform-create'))
    aid = created['agent']['component_id']
    tool_value(rpc, 'Platform', {'action': 'update_agent', 'component_id': aid, 'description': 'Northwind NW-PLATFORM-731 planner'}, 'platform-update')
    require(ctx.inst.request('GET', '/api/v1/user-agents/' + aid, label='platform-agent-persisted')['agent']['description'] == 'Northwind NW-PLATFORM-731 planner', 'Platform agent update loses description')
    from agent_memory_binding import check_memory_binding
    check_memory_binding(ctx, rpc, aid)
    tool_value(rpc, 'Platform', {'action': 'delete_agent', 'component_id': aid}, 'platform-delete')
    require(aid not in json.dumps(ctx.inst.request('GET', '/api/v1/user-agents', label='platform-agent-absent')), 'Platform deletion retains agent')


def roster_chat(ctx):
    target = create_agent(ctx, 'Northwind carton specialist', prompt='You are a fictional Northwind carton specialist. A Harbor pallet holds 12 cartons and Orchard holds 8. Return the sum with marker NW-DELEGATE-731.')
    trait = 'kickside.agents.traits:agent_roster'
    aid = create_agent(ctx, 'Northwind roster coordinator', (trait,))
    sid = start_chat(ctx, aid)
    start = len(ctx.browser.frames)
    ctx.browser.stream('Call Agents list, then Agents delegate to user_agent:' + target + ' with task "Compute Harbor 12 plus Orchard 8 cartons and return NW-DELEGATE-731". Report the actual delegated sum. Explain ten checks and end with NW-ROSTER-DONE.', 'NW-ROSTER-DONE', 'roster-delegate')
    tool_events(ctx, start, 'Agents')
    require(any(m['type'] == 'assistant' and 'NW-DELEGATE-731' in str(m['data']) and '20' in str(m['data']) for m in messages(ctx, sid)), 'Roster delegation omits specialist result')
    chat_shot(ctx, 'northwind-agent-delegation')


def time_chat(ctx):
    aid = create_agent(ctx, 'Northwind time-aware planner', ('kickside.agents.traits:time_aware',))
    sid = start_chat(ctx, aid)
    ctx.browser.stream('Report the current date from your injected time context in YYYY-MM-DD format. Do not use a tool or guess. Explain ten fictional Northwind scheduling checks and end with NW-TIME-DONE.', 'NW-TIME-DONE', 'time-awareness')
    require(any(m['type'] == 'assistant' and datetime.now().date().isoformat() in str(m['data']) for m in messages(ctx, sid)), 'Time-aware capability supplies a different current date')
    chat_shot(ctx, 'northwind-time-planning')


def delegation_chat(ctx):
    target = create_agent(ctx, 'Northwind sealed cartons specialist', prompt='You only calculate fictional Northwind carton totals. Harbor has 12 cartons and Orchard has 8. Return the total 20 and marker NW-PRIVATE-DELEGATE-731.')
    aid = create_agent(ctx, 'Northwind delegate planner', delegates=[target])
    sid = start_chat(ctx, aid)
    start = len(ctx.browser.frames)
    ctx.browser.stream('Call your delegate tool for Northwind sealed cartons specialist with a task to compute Harbor 12 plus Orchard 8 cartons. Report the actual delegated total and its marker. Explain ten dispatch checks and end with NW-DELEGATION-DONE.', 'NW-DELEGATION-DONE', 'configured-delegation')
    events = [f for f in ctx.browser.frames[start:] if isinstance(f.get('data'), dict) and 'tool' in str(f['data'].get('type', '')).lower()]
    ctx.inst.evidence.save('configured-delegate-events', events)
    functions = [m for m in messages(ctx, sid) if isinstance(m.get('metadata'), dict) and m['metadata'].get('function_name')]
    require(events and any(m['metadata'].get('status') == 'success' and 'NW-PRIVATE-DELEGATE-731' in json.dumps(m['metadata'].get('result')) for m in functions), 'Configured delegate has no observable specialist result')
    require(any(m['type'] == 'assistant' and 'NW-PRIVATE-DELEGATE-731' in str(m['data']) and '20' in str(m['data']) for m in messages(ctx, sid)), 'Configured delegation loses specialist answer')
    hero_shot(ctx, 'northwind-delegate-planner')


def redirect_chat(ctx):
    target = create_agent(ctx, 'Northwind returns specialist', prompt='You are the fictional Northwind returns specialist. The return window is 17 days and the marker is NW-REDIRECT-731. On handoff report both, give ten returns handling steps, and end with NW-REDIRECT-DONE.')
    trait = 'kickside.agents.traits:redirect'
    aid = create_agent(ctx, 'Northwind returns routing', (trait,), {trait: {'targets': [{'agent_id': 'user_agent:' + target, 'comment': 'Handles fictional Northwind returns'}]}})
    sid = start_chat(ctx, aid)
    start = len(ctx.browser.frames)
    chat = ctx.browser.frame.locator('wippy-chat')
    box = chat.get_by_placeholder('Type a message', exact=True)
    box.fill('Call Redirect switch to user_agent:' + target + ' now. Hand this session to the returns specialist.')
    box.press('Enter')
    def handed_off():
        updates = [f['data'] for f in ctx.browser.frames[start:] if isinstance(f.get('data'), dict)]
        state = chat.evaluate(CHAT_STATE)
        return (any(d.get('agent') == 'user_agent:' + target for d in updates)
                and any(d.get('status') == 'idle' and d.get('interaction', {}).get('can_send') is True for d in updates)
                and state['enabled'] and not state['busy'])
    ctx.inst.wait(handed_off, 'Redirect selects the target and completes its handoff turn', seconds=180,
                  pump=ctx.browser.page.wait_for_timeout)
    ctx.inst.evidence.save('configured-redirect-handoff-websocket', ctx.browser.frames[start:])
    tool_events(ctx, start, 'Redirect')
    ctx.browser.stream('Report your warehouse marker and return window. Give ten returns handling steps in short paragraphs and end with NW-REDIRECT-DONE.', 'NW-REDIRECT-DONE', 'configured-redirect')
    persisted = messages(ctx, sid)
    require(any(m['type'] == 'system' and isinstance(m.get('metadata'), dict)
                and m['metadata'].get('system_action') == 'config_change'
                and m['metadata'].get('new_agent') == 'user_agent:' + target for m in persisted),
            'Redirect does not persist target agent handoff')
    require(any(m['type'] == 'assistant' and 'NW-REDIRECT-731' in str(m['data']) and '17' in str(m['data']) for m in persisted), 'Receiving agent omits its returns facts')
    hero_shot(ctx, 'northwind-returns-handoff')


def chat_multiturn(ctx):
    inst, browser = ctx.inst, ctx.browser
    fresh_ui(ctx)
    inst.request('PUT', '/api/v1/user-agents/' + ctx.state['agent'], {'model': ctx.state['model']}, label='chat-gemma')
    before = {s['session_id'] for s in inst.request('GET', '/api/v1/sessions?limit=100')['sessions']}
    open_agent(ctx, ctx.state['agent'])
    browser.stream('Remember our fictional Northwind project is Harbor Bloom and its dispatch target is 73 cartons. Give ten two-sentence logistics steps. End with NW-FIRST-TURN.', 'NW-FIRST-TURN', 'gemma-first-turn')
    sessions = inst.request('GET', '/api/v1/sessions?limit=100', label='chat-session')['sessions']
    new = [s for s in sessions if s['session_id'] not in before]
    require(len(new) == 1, 'One chat creates a different number of durable sessions')
    sid = new[0]['session_id']
    ctx.state['ai_chat_session'] = sid
    ctx.state['ai_current_chat'] = sid
    browser.stream('What project name and dispatch target did I ask you to remember? Repeat them exactly. Give ten numbered checks, then end with NW-SECOND-TURN.', 'NW-SECOND-TURN', 'gemma-second-turn')
    persisted = messages(ctx, sid)
    answers = [json.dumps(m.get('data')) for m in persisted if m['type'] == 'assistant']
    require(any('NW-SECOND-TURN' in a and 'Harbor Bloom' in a and '73' in a for a in answers), 'Gemma loses first-turn project facts')
    require(sum(m['type'] == 'user' for m in persisted) == 2, 'Two chat turns persist a different number of user messages')
    browser.reload()
    open_named_session(ctx, sid)
    inst.wait(lambda: 'Harbor Bloom' in browser.frame.locator('wippy-chat').evaluate(ASSISTANT_TEXT), 'Reload restores visible assistant history', seconds=30, pump=browser.page.wait_for_timeout)
    browser.stream('Repeat the project name and 73-carton target from our earlier discussion. Give ten delivery checks and end with NW-RELOAD-TURN.', 'NW-RELOAD-TURN', 'gemma-after-reload')
    require(any(m['type'] == 'assistant' and 'Harbor Bloom' in json.dumps(m.get('data')) and 'NW-RELOAD-TURN' in json.dumps(m.get('data')) for m in messages(ctx, sid)), 'Reload continuation does not persist with original context')
    chat_shot(ctx, 'northwind-gemma-multiturn')
    hero_shot(ctx, 'northwind-gemma-multiturn')



def generated_session_title(ctx):
    sid = start_chat(ctx, create_agent(ctx, 'Northwind delivery title analyst'))
    ctx.browser.stream('Draft ten fictional Northwind Harbor Bloom shipment checks for 73 cartons and end with NW-TITLE-DONE.', 'NW-TITLE-DONE', 'generated-session-title')
    def titled():
        session = ctx.inst.request('GET', '/api/v1/sessions/get?session_id=' + sid, label='generated-title-persisted')['session']
        return session if session.get('title', '').strip() else None
    session = ctx.inst.wait(titled, 'Completed Gemma chat receives a generated session title', seconds=30, pump=ctx.browser.page.wait_for_timeout)
    require(session['title'] != 'Untitled chat', 'Generated title remains the empty-chat placeholder')
    ctx.browser.open_session(sid)
    shot(ctx, 'northwind-generated-session-title')


def session_controls(ctx):
    inst = ctx.inst
    sid = start_chat(ctx, create_agent(ctx, 'Northwind session controls'))
    ctx.browser.stream('Write ten fictional Harbor Bloom dispatch checklist items and end with NW-SESSION-DONE.', 'NW-SESSION-DONE', 'session-control-chat')
    for favorite in (True, False):
        result = inst.request('POST', '/api/v1/sessions/favorite', {'session_id': sid, 'favorite': favorite}, label='favorite-' + str(favorite))
        require(result['favorite'] is favorite, 'Favorite toggle returns a different value')
        inst.wait(lambda: (lambda s: s if bool(s.get('public_meta', {}).get('favorite')) is favorite else None)(next(s for s in inst.request('GET', '/api/v1/sessions?limit=100', label='favorite-persisted')['sessions'] if s['session_id'] == sid)), 'Favorite flag persists')
    token = inst.request('POST', '/api/v1/sessions/start-token', {'agent': ctx.state['agent'], 'model': ctx.state['model']}, label='session-start-token')['token']
    require(token, 'Session restart supplies no start token')
    open_named_session(ctx, sid)
    ctx.inst.wait(lambda: 'NW-SESSION-DONE' in ctx.browser.frame.locator('wippy-chat').evaluate(ASSISTANT_TEXT), 'Named session restores its completed assistant answer', seconds=30, pump=ctx.browser.page.wait_for_timeout)
    shot(ctx, 'northwind-session-history')


def artifact_chat(kind):
    def run(ctx):
        browser, inst = ctx.browser, ctx.inst
        if kind == 'chart':
            inst.request('GET', '/api/public/facade/config', label='artifact-facade-config')
        sid = start_chat(ctx, create_agent(ctx, 'Northwind ' + kind + ' designer', ('kickside.agents.traits:artifacts',)))
        title = 'Northwind ' + kind + ' review'
        instructions = {
            'document': 'Create a Markdown document with headings, a returns policy of 17 days, a dispatch checklist, and the warehouse marker NW-AI-731.',
            'table': 'Create a Markdown table with Product, Units, and Revenue columns. Include Harbor tea 120 2400, Orchard coffee 80 3200, and Meadow cocoa 95 2850.',
            'chart': 'Create a complete self-contained HTML document with an inline SVG bar chart, axes, data labels, and a legend. Plot Harbor tea 120, Orchard coffee 80, and Meadow cocoa 95. Use SVG directly.',
        }[kind]
        browser.stream('Use CreateArtifact, with title "' + title + '", instructions true and inline render. ' + instructions + ' Include the resulting embed instruction in your reply. Explain ten useful observations in short paragraphs and end with NW-ARTIFACT-' + kind.upper() + '.', 'NW-ARTIFACT-' + kind.upper(), 'gemma-artifact-' + kind)
        persisted = messages(ctx, sid)
        require(title in json.dumps(persisted), kind + ': artifact title is absent from persisted messages')
        start = [m for m in persisted if isinstance(m.get('metadata'), dict) and m['metadata'].get('function_name') == 'CreateArtifact']
        require(start and all(m['metadata'].get('status') == 'success' for m in start), kind + ': CreateArtifact does not succeed')
        catalog = inst.request('GET', '/api/v1/artifacts?session_id=' + sid, label='chat-artifact-catalog')['artifacts']
        artifact = next((a for a in catalog if a['title'] == title), None)
        require(artifact, kind + ': durable artifact metadata loses its exact title')
        require(artifact['content_type'] == ('text/html' if kind == 'chart' else 'text/markdown'), kind + ': artifact metadata loses its declared content type')
        reference = '<artifact id="' + artifact['artifact_id'] + '"/>'
        require(any(m['type'] == 'developer' and reference in str(m['data']) for m in persisted), kind + ': session omits returned artifact embed instruction')
        require(any(m['type'] == 'assistant' and reference in str(m['data']) for m in persisted), kind + ': model omits returned artifact embed tag')
        chat = browser.frame.locator('wippy-chat')
        rendered = chat.evaluate(ASSISTANT_TEXT)
        chat_shot(ctx, 'northwind-' + kind + '-artifact')
        if kind == 'chart':
            inst.evidence.save('inline-chart-configuration', browser.frame.locator('body').evaluate("el => ({allowAdditionalTags: el.ownerDocument.defaultView.__WIPPY_APP_CONFIG__?.hostConfig?.allowAdditionalTags ?? null})"))
            inst.evidence.save('inline-chart-body', chat.locator('.chat-message:not(.chat-message--user) .chat-message__body').all_inner_texts())
            hero_shot(ctx, 'northwind-chart-artifact')
        if kind == 'document':
            require('NW-AI-731' in rendered, 'Document artifact omits warehouse marker')
        else:
            require('harbor tea' in rendered.casefold() and 'orchard coffee' in rendered.casefold(), kind + ': artifact omits product labels')
        body = chat.locator('.chat-message:not(.chat-message--user) .chat-message__body')
        require(body.locator('table' if kind == 'table' else 'svg' if kind == 'chart' else 'h1, h2').count() > 0, kind + ': rich artifact does not render its expected structure')
        if kind == 'chart':
            require(body.locator('svg rect, svg path').count() >= 3, 'Inline chart omits its three rendered product bars')
    return run



NORTHWIND_CHART = """<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Northwind dispatch review</title><style>body{margin:0;padding:24px;font-family:system-ui,sans-serif;background:#f5f8f6;color:#173f35}main{max-width:760px;margin:auto;background:white;padding:24px;border-radius:18px}h1{font-size:26px;margin:0 0 8px}p{line-height:1.5}svg{display:block;width:100%;height:auto}footer{font-size:14px}</style></head><body><main><h1>Northwind dispatch review</h1><p>Harbor and Orchard warehouses · Fictional planning data</p><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 390" role="img" aria-label="Northwind product dispatch units"><g stroke="#dbe7df"><line x1="65" y1="95" x2="565" y2="95"/><line x1="65" y1="195" x2="565" y2="195"/><line x1="65" y1="295" x2="565" y2="295"/></g><g fill="#567167" font-size="14"><text x="25" y="100">120</text><text x="35" y="200">60</text><text x="45" y="300">0</text><text x="20" y="60">Units</text></g><rect data-series="Harbor tea" x="110" y="95" width="95" height="200" rx="6" fill="#25795a"/><rect data-series="Orchard coffee" x="270" y="162" width="95" height="133" rx="6" fill="#75aa71"/><rect data-series="Meadow cocoa" x="430" y="137" width="95" height="158" rx="6" fill="#c1b17c"/><g fill="#173f35" text-anchor="middle" font-size="17" font-weight="600"><text x="157" y="80">120</text><text x="317" y="147">80</text><text x="477" y="122">95</text></g><g fill="#173f35" text-anchor="middle" font-size="15"><text x="157" y="325">Harbor tea</text><text x="317" y="325">Orchard coffee</text><text x="477" y="325">Meadow cocoa</text></g><text x="315" y="370" text-anchor="middle" fill="#567167" font-size="14">295 units prepared for dispatch</text></svg><footer>Fictional Northwind · Dispatch reviewed by Mira Chen · Warehouse NW-AI-731</footer></main></body></html>"""

def standalone_chart(ctx):
    aid = create_agent(ctx, 'Northwind chart presenter', ('kickside.agents.traits:artifacts',))
    sid = start_chat(ctx, aid)
    title = 'Northwind dispatch chart'
    ctx.browser.stream('Call CreateArtifact with title "' + title + '", content_type text/html, render standalone, and instructions true. Copy this complete fictional Northwind chart HTML exactly into content. Preserve its SVG data-series attributes, all labels, and values. Return the embed instruction and a 150-word inventory analysis, ending with NW-STANDALONE-CHART.\n' + NORTHWIND_CHART, 'NW-STANDALONE-CHART', 'standalone-chart')
    persisted = messages(ctx, sid)
    functions = [m for m in persisted if isinstance(m.get('metadata'), dict) and m['metadata'].get('function_name') == 'CreateArtifact']
    require(functions and all(m['metadata'].get('status') == 'success' for m in functions), 'Chart artifact tool does not succeed')
    artifacts = ctx.inst.request('GET', '/api/v1/artifacts?session_id=' + sid, label='standalone-chart-catalog')['artifacts']
    artifact = next((a for a in artifacts if a['title'] == title), None)
    require(artifact and artifact['content_type'] == 'text/html', 'Standalone chart lacks durable HTML artifact')
    content, _ = raw_bytes(ctx, '/api/v1/artifact/' + artifact['artifact_id'] + '/content', 'standalone-chart-html')
    require(b'<svg' in content and all(label.encode() in content.lower() for label in ('harbor tea', 'orchard coffee', 'meadow cocoa')), 'Durable chart omits SVG or product series')
    open_named_session(ctx, sid)
    ctx.browser.frame.get_by_text('Artifacts', exact=True).first.click()
    ctx.browser.settle()
    ctx.browser.frame.get_by_text(title, exact=True).first.click()
    def chart_frame():
        for frame in ctx.browser.page.frames:
            if frame.locator('svg').count() and 'harbor tea' in frame.locator('body').inner_text().casefold():
                return frame
        return None
    rendered = ctx.inst.wait(chart_frame, 'Standalone chart renders SVG series', seconds=30, pump=ctx.browser.page.wait_for_timeout)
    require(rendered.locator('svg rect[data-series]').count() == 3, 'Chart does not render its three product bars')
    require(all(label in rendered.locator('body').inner_text().casefold() for label in ('harbor tea', 'orchard coffee', 'meadow cocoa', '120', '80', '95')), 'Chart view omits series labels or counts')
    shot(ctx, 'northwind-standalone-chart')
    with dark_ui(ctx):
        open_named_session(ctx, sid)
        ctx.browser.frame.get_by_text('Artifacts', exact=True).first.click()
        ctx.browser.settle()
        ctx.browser.frame.get_by_text(title, exact=True).first.click()
        dark_chart = ctx.inst.wait(chart_frame, 'Dark chart renders SVG series', seconds=30, pump=ctx.browser.page.wait_for_timeout)
        require(dark_chart.locator('svg rect[data-series]').count() == 3, 'Dark chart omits its product bars')
        shot(ctx, 'northwind-standalone-chart-dark')


def skills_lifecycle(ctx):
    from playwright.sync_api import expect
    inst, browser = ctx.inst, ctx.browser
    fresh_ui(ctx)
    nav(ctx, 'Skills')
    skills = browser.frame.locator('kickside-skills')
    skills.get_by_role('button', name='New skill', exact=True).first.click()
    dialog = skills.get_by_role('dialog', name='Edit skill')
    dialog.get_by_placeholder('e.g. Triage a bug report').fill('Northwind dispatch review')
    dialog.get_by_placeholder('Summarize a long chat into action items.').fill('Reviews fictional Harbor and Orchard dispatch readiness.')
    body = '1. Read the Northwind handbook.\n2. Verify NW-SKILL-731 and the 17-day return window.\n3. Confirm product totals.\n4. Prepare the dispatch checklist.'
    dialog.get_by_label('Skill instructions', exact=True).fill(body)
    with browser.page.expect_response(lambda r: '/api/v1/skills' in r.url and r.request.method == 'POST') as response:
        dialog.get_by_role('button', name='Create skill', exact=True).click()
    sid = response.value.json()['skill_id']
    ctx.state['ai_skill'] = sid
    current = inst.request('GET', '/api/v1/skills/' + sid, label='skill-created')['skill']
    require(current['body'] == body and current['name'] == 'Northwind dispatch review', 'Created skill instructions differ')
    require(sid in json.dumps(inst.request('GET', '/api/v1/skills', label='skill-list')), 'Skills list omits created skill')
    skills.get_by_role('button', name='Edit Northwind dispatch review', exact=True).click()
    expect(dialog.get_by_label('Skill instructions', exact=True)).to_have_value(body, timeout=30000)
    dialog.get_by_label('Skill instructions', exact=True).fill(body + '\n5. Record Mira Chen as dispatch reviewer.')
    with browser.page.expect_response(lambda r: '/api/v1/skills/' + sid in r.url and r.request.method == 'PATCH') as response:
        dialog.get_by_role('button', name='Save new version', exact=True).click()
    require(response.value.json().get('success'), 'Skill refinement does not save')
    updated = inst.request('GET', '/api/v1/skills/' + sid, label='skill-refined')['skill']
    require('Mira Chen' in updated['body'] and updated['version'] > current['version'], 'Skill refinement fails to advance content/version')
    versions = inst.request('GET', '/api/v1/skills/' + sid + '/versions', label='skill-versions')['versions']
    require(len(versions) == 2, 'Skill version history differs from two saved versions')
    previous = inst.request('GET', '/api/v1/skills/' + sid + '?version=' + str(current['version']), label='skill-original')['skill']
    require(previous['body'] == body, 'Skill version history mutates the original instructions')
    shot(ctx, 'northwind-skills-library')


def skills_delete(ctx):
    inst = ctx.inst
    sid = inst.request('POST', '/api/v1/skills', {'name': 'Northwind temporary checklist', 'description': 'Fictional scratch skill', 'body': 'Inspect Northwind carton NW-REMOVE-731.'}, label='skill-delete-fixture')['skill_id']
    inst.request('DELETE', '/api/v1/skills/' + sid, label='skill-delete')
    require(sid not in json.dumps(inst.request('GET', '/api/v1/skills', label='skill-absent')), 'Deleted skill remains listed')


def mcp(ctx, tool_ids=(), trait_ids=(), components=()):
    inst = ctx.inst
    credential = inst.request('POST', '/api/v1/mcp/credentials', {
        'label': 'Northwind AI smoke client', 'allowed_tool_ids': list(tool_ids),
        'allowed_trait_ids': list(trait_ids), 'whitelist_component_ids': list(components),
    }, label='mcp-credential')['credential']
    token, cid = credential['token'], credential['credential_id']
    inst.redactor.add(token)
    inst.cleanup.add('Northwind MCP credential ' + str(cid), lambda: inst.request('POST', '/api/v1/mcp/credentials/' + str(cid) + '/revoke', {}))
    def rpc(method, params, label):
        result = inst.request('POST', '/api/public/mcp', {'jsonrpc': '2.0', 'id': uuid.uuid4().hex,
            'method': method, 'params': params}, auth=token, label=label,
            headers={'Accept': 'application/json, text/event-stream'}, timeout=120)
        require('error' not in result, 'MCP protocol error: ' + str(result))
        require(not result.get('result', {}).get('isError'), 'MCP tool error: ' + str(result))
        return result['result']
    initialized = rpc('initialize', {'protocolVersion': '2025-03-26', 'capabilities': {},
        'clientInfo': {'name': 'Northwind AI Smoke', 'version': '1'}}, 'mcp-initialize')
    require(initialized.get('serverInfo'), 'MCP initialize omits server identity')
    return rpc


def tool_value(rpc, alias, args, label):
    result = rpc('tools/call', {'name': alias, 'arguments': args}, label)
    require(result.get('content'), alias + ': tool supplies no observable content')
    text = '\n'.join(item.get('text', '') for item in result['content'])
    require(text and not text.startswith(('Error:', 'tool error:', '_Error')), alias + ': tool reports failure: ' + text)
    return text


def knowledge_tool(alias, tool_id, *, analysis=False):
    def run(ctx):
        kid = kb_create(ctx, 'document', 'Northwind ' + alias + ' reference', analysis=analysis)
        content = 'Fictional Northwind NW-AI-731: returns within 17 days.\n' + POLICY if alias == 'SearchKB' else POLICY
        kb_ingest(ctx, kid, content=content)
        rpc = mcp(ctx, (tool_id,), components=(kid,))
        args = {'query': 'Northwind warehouse marker and 17-day return window', 'kb_id': kid}
        if alias in ('QueryKB', 'RetrieveKB', 'SearchKB'):
            args['fields'] = ['content', 'meta']
        if alias == 'FindKnowledgeBases':
            args = {'query': 'Northwind'}
        if alias == 'Knowledge':
            args = {'action': 'list_knowledge_bases'}
        result = tool_value(rpc, alias, args, 'knowledge-tool')
        require(kid in result if alias in ('FindKnowledgeBases', 'Knowledge') else 'NW-AI-731' in result, alias + ': tool omits the known Northwind source: ' + result)
    return run


def knowledge_embed_tool(alias, tool_id):
    def run(ctx):
        inst = ctx.inst
        kid = kb_create(ctx, 'document', 'Northwind ' + alias + ' intake')
        rpc = mcp(ctx, (tool_id,), components=(kid,))
        result = tool_value(rpc, alias, {'kb_id': kid, 'content': 'Fictional Northwind intake marker NW-TOOL-842. Returns within 17 days.',
            'title': 'Northwind MCP intake', 'external_id': 'northwind-tool-intake', 'external_version': '1'}, 'kb-tool-embed')
        require('queued' in result.lower() or 'thread_id' in result, 'Tool does not supply an ingestion receipt')
        inst.wait(lambda: (lambda r: r if 'NW-TOOL-842' in json.dumps(r) else None)(inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-TOOL-842'}, label='tool-content-indexed')), 'Tool content reaches search', seconds=120)
    return run


def wiki_lifecycle(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'wiki', 'Northwind operations wiki', analysis=True)
    ctx.state['ai_wiki_kb'] = kid
    kb_ingest(ctx, kid)
    base = '/api/v1/kb10/' + kid
    concepts = []
    for label, description in [('Harbor warehouse', 'Northwind Harbor return processing and afternoon shipments.'), ('Orchard warehouse', 'Northwind Orchard morning dispatch and coffee inventory.'), ('Returns policy', 'Customers return Northwind products within 17 days.'), ('Product catalog', 'Harbor tea, Orchard coffee, and Meadow cocoa.'), ('Dispatch checklist', 'Inspect packaging, record batches, and confirm shipment.')]:
        created = inst.request('POST', base + '/concepts', {'label': label, 'description': description}, label='concept-' + str(len(concepts)))
        require(created.get('node_id'), 'Create concept omits durable node ID')
        concepts.append(created['node_id'])
    ctx.state['ai_wiki_concepts'] = concepts
    graph = inst.request('GET', base + '/graph', label='wiki-graph')
    require(set(concepts) <= {n['id'] for n in graph['nodes']} and len(graph['edges']) >= len(concepts), 'Wiki graph omits concepts or parent relationships')
    browse = inst.request('GET', base + '/browse?node_type=document&include_content=true&limit=100', label='wiki-documents')
    require(browse['nodes'], 'Wiki browse omits ingested handbook')
    ctx.state['ai_wiki_document'] = browse['nodes'][0]['id']
    searched = inst.request('GET', base + '/search?q=' + quote('Northwind NW-AI-731'), label='wiki-search')
    require('NW-AI-731' in json.dumps(searched), 'Wiki search omits handbook marker')
    fresh_ui(ctx)
    nav(ctx, 'Knowledge')
    ctx.browser.frame.get_by_text('Northwind operations wiki', exact=True).first.click()
    ctx.browser.frame.get_by_role('tab', name='Graph', exact=True).click()
    ctx.browser.frame.get_by_text('Harbor warehouse', exact=True).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-knowledge-graph')
    ctx.browser.frame.get_by_role('tab', name='Wiki', exact=True).click()
    ctx.browser.frame.get_by_text('Orchard warehouse', exact=True).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-operations-wiki')



def wiki_file_ingestion(ctx):
    kid = kb_create(ctx, 'wiki', 'Northwind warehouse file library')
    base = '/api/v1/knowledge/' + kid
    for suffix in ('txt', 'docx', 'xlsx'):
        uid, _ = completed_upload(ctx, fixture(ctx, suffix))
        queued = ctx.inst.request('POST', base + '/ingest-upload', {'upload_id': uid}, label='wiki-file-' + suffix)
        require(queued['queued'] == 1 and queued['total'] == 1, suffix + ': wiki does not queue source file')
    ctx.inst.wait(lambda: (lambda result: result if result['document_count'] >= 3 else None)(ctx.inst.request('GET', base, label='wiki-file-count')), 'Wiki indexes three public file formats', seconds=120, pump=ctx.browser.page.wait_for_timeout)
    source = ctx.inst.request('POST', base + '/search', {'query': 'Northwind warehouse NW-AI-731 returns', 'limit': 10}, label='wiki-file-search')
    require('NW-AI-731' in json.dumps(source), 'Wiki file search omits source warehouse marker')
    filenames = {item.get('meta', {}).get('metadata', {}).get('original_filename') for item in source['items']}
    require({'Northwind-delivery.txt', 'Northwind-delivery.docx', 'Northwind-delivery.xlsx'} <= filenames, 'Wiki search loses an ingested public file format: ' + str(filenames))
    fresh_ui(ctx)
    nav(ctx, 'Knowledge')
    ctx.browser.frame.get_by_text('Northwind warehouse file library', exact=True).click()
    ctx.browser.frame.get_by_role('button', name='Search', exact=True).click()
    search = ctx.browser.frame.get_by_placeholder('Search every indexed passage…')
    search.fill('NW-AI-731')
    search.press('Enter')
    ctx.browser.frame.get_by_text('NW-AI-731', exact=False).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-wiki-file-library')


def wiki_cited_answer(ctx):
    answer = ctx.inst.request('POST', '/api/v1/knowledge/' + ctx.state['ai_wiki_kb'] + '/ask', {'query': 'What is the Northwind warehouse marker and return window? Cite the handbook.'}, label='wiki-cited-answer', timeout=120)
    require('17' in str(answer.get('answer')) and 'NW-AI-731' in str(answer.get('answer')), 'Wiki answer omits its handbook facts: ' + str(answer))
    require(answer.get('citations') or answer.get('sources') or re.search(r'\[\d+\]', str(answer.get('answer'))), 'Wiki answer omits citations: ' + str(answer))
    require_cited_sources(answer, ('NW-AI-731', '17'))


def wiki_update_delete(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'wiki', 'Northwind taxonomy revisions')
    base = '/api/v1/kb10/' + kid
    parent = inst.request('POST', base + '/concepts', {'label': 'Northwind delivery', 'description': 'Fictional Northwind taxonomy'}, label='parent-concept')['node_id']
    child = inst.request('POST', base + '/concepts', {'label': 'Northwind returns'}, label='child-concept')['node_id']
    inst.request('PATCH', base + '/concepts/' + child, {'label': 'Northwind customer returns', 'description': '17-day returns at Harbor'}, label='concept-update')
    updated = inst.request('GET', base + '/graph', label='concept-updated')
    require(any(n['id'] == child and n['label'] == 'Northwind customer returns' and '17-day' in n['description'] for n in updated['nodes']), 'Concept update does not persist label/description')
    inst.request('PUT', base + '/concepts/' + child + '/parent', {'parent_id': parent}, label='concept-move')
    graph = inst.request('GET', base + '/graph', label='concept-moved')
    require(any(e['from_node_id'] == child and e['to_node_id'] == parent for e in graph['edges']), 'Moving a concept does not change its parent edge')
    inst.request('DELETE', base + '/concepts/' + child, {}, label='concept-delete')
    require(child not in {n['id'] for n in inst.request('GET', base + '/graph', label='concept-absent')['nodes']}, 'Deleted concept remains in graph')
    inst.request('DELETE', '/api/v1/knowledge/' + kid, label='wiki-delete')
    require(kid not in json.dumps(inst.request('GET', '/api/v1/knowledge/list', label='wiki-absent')), 'Deleted wiki remains listed')


def wiki_node_delete(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'document', 'Northwind removable handbook')
    kb_ingest(ctx, kid, 'NW-NODE-936')
    base = '/api/v1/kb10/' + kid
    docs = inst.request('GET', base + '/browse?node_type=document&limit=100', label='nodes-before')['nodes']
    require(len(docs) == 1, 'One handbook produces a different number of documents')
    refused = inst.request('DELETE', base + '/nodes/' + docs[0]['id'], expect=(409,), label='raw-node-delete-guard')
    require(refused['success'] is False and refused['canonical'] == 'knowledge.content.delete' and refused['error'] == 'raw node delete is disabled; delete source content through /api/v1/knowledge/' + kid + '/content using external_id, upload_id, source_url, or source_ref', 'Raw node delete does not identify its canonical deletion path')
    retained = inst.request('GET', base + '/browse?node_type=document&limit=100', label='guard-retains-source')['nodes']
    require([n['id'] for n in retained] == [n['id'] for n in docs], 'Rejected raw node deletion changes source documents')
    deleted = inst.request('DELETE', '/api/v1/knowledge/' + kid + '/content', {'external_id': 'northwind-delivery'}, label='node-source-delete')
    require(deleted.get('thread_id'), 'Canonical deletion omits its durable source receipt')
    inst.wait(lambda: (lambda r: r if not r['nodes'] else None)(inst.request('GET', base + '/browse?node_type=document&limit=100', label='nodes-after')), 'Deleted document leaves browse', seconds=120)
    require('NW-NODE-936' not in json.dumps(inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-NODE-936'}, label='node-search-absent')), 'Deleted node remains searchable')


def wiki_source(ctx):
    inst = ctx.inst
    kid = kb_create(ctx, 'document', 'Northwind original source library')
    path = fixture(ctx, 'pdf')
    uid, _ = completed_upload(ctx, path)
    inst.request('POST', '/api/v1/knowledge/' + kid + '/ingest-upload', {'upload_id': uid}, label='source-ingest')
    def document():
        docs = inst.request('GET', '/api/v1/kb10/' + kid + '/browse?node_type=document&limit=100', label='source-documents')['nodes']
        return docs[0] if docs else None
    doc = inst.wait(document, 'Original source document indexed', seconds=120)
    source = inst.request('GET', '/api/v1/kb10/' + kid + '/nodes/' + doc['id'] + '/source', label='source-link')
    require(source['filename'] == path.name and source['content_type'] == 'application/pdf', 'Source metadata differs from uploaded PDF')
    data, _ = raw_bytes(ctx, source['url'], 'kb-original-source')
    require(data == path.read_bytes(), 'KB source URL retrieves different bytes')


def wiki_maintenance(operation):
    def run(ctx):
        inst = ctx.inst
        kid = kb_create(ctx, 'document', 'Northwind ' + operation + ' maintenance')
        kb_ingest(ctx, kid)
        base = '/api/v1/kb10/' + kid
        receipt = inst.request('POST', base + '/' + operation, {}, label='maintenance-schedule', timeout=120)
        require(receipt.get('scheduled') and receipt.get('job_id'), operation + ': maintenance supplies no scheduled job receipt')
        def finished():
            state = inst.request('GET', base + '/maintenance', label='maintenance-state')
            run = state.get(operation)
            if not run:
                return None
            require(run.get('status') not in ('failed', 'error'), operation + ': maintenance fails: ' + str(run))
            return run if run.get('in_flight') is False else None
        run = inst.wait(finished, 'Knowledge ' + operation + ' completes', seconds=180)
        require(run.get('finished_at') and run.get('last_error') in (None, ''), 'Maintenance completion lacks finished timestamp or records an error: ' + str(run))
        searched = inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-AI-731'}, label='maintenance-search')
        require('NW-AI-731' in json.dumps(searched), 'Maintenance loses indexed source facts')
    return run


def wiki_tool(alias, tool_id):
    def run(ctx):
        kid = kb_create(ctx, 'document', 'Northwind ' + alias + ' library')
        kb_ingest(ctx, kid)
        docs = ctx.inst.request('GET', '/api/v1/kb10/' + kid + '/browse?node_type=document&limit=100', label='tool-document')['nodes']
        require(docs, 'Knowledge tool fixture has no document')
        rpc = mcp(ctx, (tool_id,), components=(kid,))
        args = {'kb_id': kid}
        if alias in ('GetKBNode', 'TraverseKBGraph'):
            args['node_id'] = docs[0]['id']
        result = tool_value(rpc, alias, args, 'wiki-tool')
        require('NW-AI-731' in result if alias == 'GetKBNode' else docs[0]['id'] in result, alias + ': knowledge tool omits source node: ' + result)
        if alias in ('BrowseKB', 'SurveyKB', 'TraverseKBGraph'):
            require('NW-AI-731' not in result, alias + ': collection includes full source content')
    return run


def mcp_capability_toggle(ctx, trait, alias):
    rpc = mcp(ctx, trait_ids=(trait,))
    before = rpc('tools/list', {}, 'tools-before')['tools']
    require(alias not in {t['name'] for t in before}, 'Inactive capability exposes its tool')
    activated = tool_value(rpc, 'use_trait', {'id': trait}, 'capability-activate')
    require(trait in activated, 'Activation receipt omits capability')
    active = rpc('tools/list', {}, 'tools-active')['tools']
    require(alias in {t['name'] for t in active}, 'Activated capability does not expose its tool')
    result = tool_value(rpc, alias, {'action': 'overview'} if alias == 'SelfReflection' else {'limit': 50} if alias == 'ListFiles' else {'action': 'list'}, 'active-tool')
    require(result, 'Activated capability tool supplies no content')
    dropped = tool_value(rpc, 'drop_trait', {'id': trait}, 'capability-deactivate')
    require(trait not in json.dumps(json.loads(dropped)['active_traits']), 'Deactivation retains active capability')
    after = rpc('tools/list', {}, 'tools-after')['tools']
    require(alias not in {t['name'] for t in after}, 'Deactivated capability retains its tool')


def reflection(ctx):
    trait = 'kickside.self_reflection.traits:self_reflection'
    mcp_capability_toggle(ctx, trait, 'SelfReflection')
    rpc = mcp(ctx, ('kickside.self_reflection.traits:self_reflection_tool',))
    result = tool_value(rpc, 'SelfReflection', {'action': 'search_registry', 'namespace': 'kickside.agents', 'q': 'artifacts'}, 'reflection-registry')
    require('kickside.agents.traits:artifacts' in result, 'Self Reflection omits the known artifact capability')


def files_tools(ctx):
    uid, _ = completed_upload(ctx, fixture(ctx, 'txt'))
    ids = ['kickside.uploads.traits:' + name for name in ('list_files', 'count_files', 'view_files', 'upload_file')]
    rpc = mcp(ctx, ids)
    listed = tool_value(rpc, 'ListFiles', {'limit': 50}, 'files-list')
    require(uid in listed and 'Northwind-delivery.txt' in listed, 'ListFiles omits uploaded Northwind handbook')
    counted = tool_value(rpc, 'CountFiles', {}, 'files-count')
    count = json.loads(counted).get('count')
    require(type(count) is int and count > 0, 'CountFiles does not report nonzero uploaded files: ' + counted)
    viewed = tool_value(rpc, 'ViewFiles', {'upload_ids': [uid], 'format': 'text'}, 'files-view')
    require('NW-AI-731' in viewed and '17 days' in viewed, 'ViewFiles omits actual handbook contents')
    uploaded = tool_value(rpc, 'UploadFile', {'content_base64': base64.b64encode(b'Fictional Northwind tool upload NW-FILE-965').decode(), 'filename': 'Northwind-tool-note.txt', 'content_type': 'text/plain'}, 'files-upload')
    upload_receipt = json.loads(uploaded)
    require(upload_receipt.get('upload_id') and upload_receipt['size'] == 43, 'UploadFile omits uploaded file result')
    require('Northwind-tool-note.txt' in json.dumps(ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='tool-upload-persisted')), 'Tool upload does not persist in file list')
    mcp_capability_toggle(ctx, 'kickside.uploads.traits:file_manager', 'ListFiles')


def skill_tools(ctx):
    ids = ('kickside.skills.traits:skill_tool', 'kickside.skills.traits:skill_file_tool')
    rpc = mcp(ctx, ids)
    saved = tool_value(rpc, 'Skill', {'action': 'save', 'name': 'Northwind bundled dispatch', 'description': 'Fictional Northwind bundle',
        'body': 'Read references/dispatch.md and return NW-BUNDLE-731.',
        'files': [{'path': 'references/dispatch.md', 'kind': 'reference', 'content': 'Northwind dispatch marker NW-BUNDLE-731. Returns within 17 days.'}]}, 'skill-bundle-save')
    skills = ctx.inst.request('GET', '/api/v1/skills?q=' + quote('Northwind bundled dispatch'),
                              label='skill-bundle-persisted')['skills']
    matches = [skill for skill in skills if skill['name'] == 'Northwind bundled dispatch']
    require(len(matches) == 1, 'Skill save does not persist one named bundle: ' + str(skills))
    sid = matches[0]['skill_id']
    require(sid in saved, 'Skill save omits its persisted identifier: ' + saved)
    used = tool_value(rpc, 'Skill', {'action': 'use', 'skill_id': sid}, 'skill-bundle-use')
    require('references/dispatch.md' in used and 'NW-BUNDLE-731' in used, 'Skill use omits bundle manifest/instructions')
    listed = tool_value(rpc, 'SkillFile', {'action': 'list', 'skill_id': sid}, 'skill-files')
    require('references/dispatch.md' in listed, 'SkillFile omits saved reference')
    read = tool_value(rpc, 'SkillFile', {'action': 'read', 'skill_id': sid, 'path': 'references/dispatch.md'}, 'skill-file-read')
    require('NW-BUNDLE-731' in read, 'SkillFile read omits reference content')
    mcp_capability_toggle(ctx, 'kickside.skills.traits:skills', 'Skill')
    return sid


def organize_files(ctx):
    ids = []
    for suffix in ('txt', 'md', 'pdf'):
        uid, _ = completed_upload(ctx, fixture(ctx, suffix))
        ids.append(uid)
    result = ctx.inst.request('POST', '/api/v1/knowledge/organize', {'upload_ids': ids, 'kb_type': 'document', 'embedding_model': ctx.state['embed']}, label='organize-files', timeout=120)
    require(result.get('kb_id') and result['queued'] == 3 and result['total'] == 3, 'Organize omits KB or loses selected uploads: ' + str(result))
    ctx.inst.wait(lambda: (lambda r: r if r['document_count'] >= 3 else None)(ctx.inst.request('GET', '/api/v1/knowledge/' + result['kb_id'], label='organized-library')), 'Organized uploads reach knowledge library', seconds=120)


def upload_delete(ctx):
    path = ctx.inst.root / 'temp' / 'Northwind-temporary-note.txt'
    path.write_text('Fictional Northwind scratch marker NW-DELETE-UPLOAD-731')
    uid, _ = completed_upload(ctx, path)
    removed = ctx.inst.request('DELETE', '/api/v1/uploads/' + uid, label='upload-delete')
    require(removed.get('success'), 'Upload deletion does not confirm removal')
    require(uid not in json.dumps(ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='upload-absent')), 'Deleted upload remains listed')
    missing = ctx.inst.request('GET', '/api/v1/uploads/' + uid, expect=(404,), label='upload-deleted-detail')
    require(missing.get('error') and not missing.get('upload'), 'Deleted upload remains readable')


def upload_workspace(ctx):
    from playwright.sync_api import expect
    fresh_ui(ctx)
    nav(ctx, 'Uploads')
    uid, _ = completed_upload(ctx, fixture(ctx, 'docx'))
    uploads = ctx.browser.frame.locator('kickside-uploads')
    uploads.get_by_label('Search files', exact=True).fill('Northwind-delivery.docx')
    expect(uploads).to_contain_text('Northwind-delivery.docx', timeout=30000)
    require(uid in json.dumps(ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='workspace-upload')), 'Uploads workspace omits created document')
    shot(ctx, 'northwind-uploads-workspace')


def storage_settings(ctx):
    fresh_ui(ctx)
    nav(ctx, 'System')
    ctx.browser.frame.get_by_text('Settings', exact=True).first.click()
    ctx.browser.frame.get_by_text('Files & Uploads', exact=True).first.click()
    ctx.browser.frame.get_by_text('Maximum document size', exact=True).first.wait_for(timeout=30000)
    ctx.browser.frame.get_by_text('Processing time limit', exact=True).first.wait_for(timeout=30000)
    require(ctx.browser.frame.get_by_role('switch', name='Rasterize empty documents', exact=True).is_checked(), 'Files settings disable rasterization by default')
    shot(ctx, 'northwind-files-settings')


def pdf_rasterization(ctx):
    from reportlab.pdfgen.canvas import Canvas
    from PIL import Image
    import io
    path = ctx.inst.root / 'temp' / 'Northwind-vector-scan.pdf'
    canvas = Canvas(str(path), pageCompression=0)
    for warehouse in ('Harbor', 'Orchard'):
        canvas.setFont('Helvetica-Bold', 28)
        canvas.drawString(50, 740, 'Fictional Northwind ' + warehouse)
        canvas.drawString(50, 690, 'NW-AI-731: Returns within 17 days')
        canvas.setStrokeColorRGB(.86, .91, .87)
        for i in range(180):
            canvas.rect(40 + i % 15 * 34, 60 + i // 15 * 44, 32, 42)
        canvas.showPage()
    canvas.save()
    uid, artifacts = completed_upload(ctx, path)
    pages = [a for a in artifacts['artifacts'] if a.get('role') == 'rendered-page']
    require(len(pages) == 2, 'Two sparse PDF pages do not produce two rendered-page artifacts')
    for a in pages:
        data, mime = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + a['artifact_id'] + '/raw', 'raster-page-' + a['artifact_id'])
        image = Image.open(io.BytesIO(data))
        require(mime.startswith('image/') and image.width > 1000 and image.height > 1500, 'PDF page does not render at the configured resolution')
    content, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + artifacts['primary']['artifact_id'] + '/raw', 'raster-ocr-content')
    require(all(s in content.decode() for s in ('NW-AI-731', 'Harbor', 'Orchard', '17')), 'Rasterized PDF OCR omits page facts')
    fresh_ui(ctx)
    preview(ctx, path.name, markdown=True)


def wiki_enrichment(ctx):
    kid = kb_create(ctx, 'wiki', 'Northwind automatic concept placement')
    prompt = 'Use BrowseTree first. Create a concept named Northwind returns operations with CreateConcept, then call AttachDocument with the provided document node id and that concept id. Finish only after attachment.'
    ctx.inst.request('PUT', '/api/v1/knowledge/' + kid, {'config': {'extraction_model': ctx.state['model'], 'concept_prompt': prompt}}, label='enrichment-config')
    saved = ctx.inst.request('GET', '/api/v1/knowledge/' + kid, label='enrichment-config-persisted')
    require(saved['config'].get('concept_prompt') == prompt, 'Knowledge update loses concept placement instructions')
    kb_ingest(ctx, kid)
    retried = ctx.inst.request('POST', '/api/v1/kb10/' + kid + '/enrichment', {}, label='enrichment-retry')
    require(retried['documents'] == 1 and 0 <= retried['requested'] <= 1, 'Enrichment retry loses source document')
    def placed():
        graph = ctx.inst.request('GET', '/api/v1/kb10/' + kid + '/graph?limit=100', label='automatic-concept-graph')
        return graph if any(n.get('label') == 'Northwind returns operations' for n in graph['nodes']) and any(e.get('edge_type') == 'categorized_as' for e in graph['edges']) else None
    graph = ctx.inst.wait(placed, 'Gemma creates and attaches Northwind concept', seconds=120, pump=ctx.browser.page.wait_for_timeout)
    concept = next(n for n in graph['nodes'] if n.get('label') == 'Northwind returns operations')
    require(not concept.get('metadata', {}).get('manual'), 'Placement creates a manual concept instead of automatic taxonomy')
    require(any(e['to_node_id'] == concept['id'] and e['edge_type'] == 'categorized_as' for e in graph['edges']), 'Automatic concept does not contain source document')
    fresh_ui(ctx)
    nav(ctx, 'Knowledge')
    ctx.browser.frame.get_by_text('Northwind automatic concept placement', exact=True).click()
    ctx.browser.frame.get_by_role('tab', name='Graph', exact=True).click()
    ctx.browser.frame.get_by_text('Northwind returns operations', exact=True).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-automatic-knowledge-graph')


def wiki_default_placement(ctx):
    kid = kb_create(ctx, 'wiki', 'Northwind delivery taxonomy')
    kb_ingest(ctx, kid)
    retried = ctx.inst.request('POST', '/api/v1/kb10/' + kid + '/enrichment', {}, label='default-placement-retry')
    require(retried['documents'] == 1 and 0 <= retried['requested'] <= 1, 'Default enrichment retry loses source document')
    def placed():
        graph = ctx.inst.request('GET', '/api/v1/kb10/' + kid + '/graph?limit=100', label='default-placement-graph')
        return graph if any(e['edge_type'] == 'categorized_as' for e in graph['edges']) else None
    graph = ctx.inst.wait(placed, 'Gemma places Northwind delivery handbook', seconds=120, pump=ctx.browser.page.wait_for_timeout)
    concepts = [n for n in graph['nodes'] if n['node_type'] == 'concept' and not n['is_root']]
    require(concepts and any(n['document_count'] > 0 for n in concepts), 'Automatic taxonomy creates no populated topic')
    docs = {n['id'] for n in graph['nodes'] if n['node_type'] == 'document'}
    require(any(e['from_node_id'] in docs and e['to_node_id'] in {n['id'] for n in concepts} and e['edge_type'] == 'categorized_as' for e in graph['edges']), 'Automatic taxonomy does not attach the source document')
    fresh_ui(ctx)
    nav(ctx, 'Knowledge')
    ctx.browser.frame.get_by_text('Northwind delivery taxonomy', exact=True).click()
    ctx.browser.frame.get_by_role('tab', name='Graph', exact=True).click()
    ctx.browser.frame.get_by_text(concepts[0]['label'], exact=True).first.wait_for(timeout=30000)
    shot(ctx, 'northwind-delivery-taxonomy')


def sav_typed_upload(ctx):
    path = ctx.inst.root / 'temp' / 'Northwind-typed-survey.sav'
    content = fixture(ctx, 'sav').read_bytes()
    old = b'Fictional Northwind delivery survey'
    new = b'Fictional Northwind typed survey'.ljust(len(old))
    require(old in content, 'Northwind SPSS fixture lacks its fictional label')
    path.write_bytes(content.replace(old, new))
    uid, artifacts = completed_upload(ctx, path, 'application/x-spss-sav')
    data, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/artifacts/' + artifacts['primary']['artifact_id'] + '/raw', 'spss-markdown')
    require('NW-AI-731' in data.decode() and 'Harbor tea' in data.decode(), 'Typed SPSS conversion loses survey records')
    fresh_ui(ctx)
    preview(ctx, path.name, markdown=True)


def knowledge_capability(trait, alias, args):
    def run(ctx):
        kid = kb_create(ctx, 'document', 'Northwind ' + alias + ' capability library')
        kb_ingest(ctx, kid)
        rpc = mcp(ctx, trait_ids=(trait,), components=(kid,))
        tool_value(rpc, 'use_trait', {'id': trait}, 'knowledge-capability-activate')
        tools = rpc('tools/list', {}, 'knowledge-capability-tools')['tools']
        require(alias in {t['name'] for t in tools}, 'Activated knowledge capability omits ' + alias)
        result = tool_value(rpc, alias, args(kid), 'knowledge-capability-call')
        require(kid in result or 'NW-AI-731' in result or 'Northwind' in result, 'Knowledge capability omits actual library: ' + result)
        tool_value(rpc, 'drop_trait', {'id': trait}, 'knowledge-capability-deactivate')
        require(alias not in {t['name'] for t in rpc('tools/list', {}, 'knowledge-capability-removed')['tools']}, 'Deactivated knowledge capability retains ' + alias)
    return run


def knowledge_ingest_tools(ctx):
    kid = kb_create(ctx, 'document', 'Northwind MCP ingestion lifecycle')
    trait = 'kickside.knowledge.traits.kb_ingest:kb_ingest'
    rpc = mcp(ctx, trait_ids=(trait,), components=(kid,))
    tool_value(rpc, 'use_trait', {'id': trait}, 'ingest-capability-activate')
    selected = json.loads(tool_value(rpc, 'SelectKB', {'kb_id': kid}, 'select-active-kb'))
    require(selected['kb_id'] == kid and selected['success'], 'SelectKB selects a different knowledge base')
    queued = tool_value(rpc, 'EmbedContent', {'kb_id': kid, 'content': 'Fictional Northwind intake NW-INGEST-731, returns in 17 days.', 'title': 'Northwind tool intake', 'external_id': 'northwind-ingest-tool', 'external_version': '1'}, 'ingest-content')
    require('queued' in queued.lower() or 'thread_id' in queued, 'EmbedContent omits durable ingestion receipt')
    ctx.inst.wait(lambda: 'NW-INGEST-731' in json.dumps(ctx.inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-INGEST-731'}, label='ingest-search')), 'MCP ingestion reaches search', seconds=120)
    deleted = tool_value(rpc, 'DeleteContent', {'kb_id': kid, 'external_id': 'northwind-ingest-tool', 'external_version': '2'}, 'ingest-delete')
    require('queued' in deleted.lower() or 'thread_id' in deleted, 'DeleteContent omits durable deletion receipt')
    ctx.inst.wait(lambda: 'NW-INGEST-731' not in json.dumps(ctx.inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-INGEST-731'}, label='ingest-deleted-search')), 'MCP content deletion removes search passage', seconds=120)
    tool_value(rpc, 'drop_trait', {'id': trait}, 'ingest-capability-deactivate')
    require('SelectKB' not in {t['name'] for t in rpc('tools/list', {}, 'ingest-tools-removed')['tools']}, 'Deactivated ingestion capability retains tools')


def knowledge_writer(ctx):
    kid = kb_create(ctx, 'document', 'Northwind writable capability library')
    trait = 'kickside.knowledge.traits:kb_writer'
    rpc = mcp(ctx, trait_ids=(trait,), components=(kid,))
    tool_value(rpc, 'use_trait', {'id': trait}, 'writer-capability-activate')
    result = tool_value(rpc, 'EmbedToKB', {'kb_id': kid, 'content': 'Fictional Northwind writer marker NW-WRITER-731.', 'title': 'Northwind writer note'}, 'writer-content')
    require('queued' in result.lower() or 'thread_id' in result, 'Writer does not queue content')
    ctx.inst.wait(lambda: 'NW-WRITER-731' in json.dumps(ctx.inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NW-WRITER-731'}, label='writer-search')), 'Activated writer persists searchable content', seconds=120)


def models_ui(ctx):
    fresh_ui(ctx)
    nav(ctx, 'System')
    ctx.browser.frame.get_by_text('Models', exact=True).first.click()
    ctx.browser.frame.get_by_text('Gemma', exact=False).first.wait_for(timeout=30000)
    require(ctx.state['model'] in json.dumps(ctx.inst.request('GET', '/api/v1/models/admin', label='model-ui-card')), 'Model UI lacks configured Gemma card')
    shot(ctx, 'northwind-gemma-models')


def upload_direct_validation(ctx):
    before = ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='direct-upload-before')
    missing_filename = ctx.inst.request('POST', '/api/v1/uploads/presigned', {'size': 32}, expect=(400,), label='presigned-invalid-input')
    require(missing_filename['success'] is False and missing_filename['error'] == 'Missing required field: filename', 'Presigned upload does not explain missing filename')
    missing_id = ctx.inst.request('POST', '/api/v1/uploads/complete', {}, expect=(400,), label='complete-invalid-input')
    require(missing_id['success'] is False and missing_id['error'] == 'Missing required field: upload_id', 'Upload completion does not explain missing upload id')
    after = ctx.inst.request('GET', '/api/v1/uploads/list?limit=100', label='direct-upload-after')
    require(before == after, 'Invalid direct-upload requests create file records')


from renderer_math_mermaid import markdown_math_mermaid


REGISTRY = [
    Scenario('ai_boot', 'full', ai_boot),
    Scenario('ai_ready', 'full', ai_ready, ('ai_boot',), True, instance_setup=True),
    Scenario('ai_install_skills', 'full', install_module('kickside/skills', (('kickside.skills:upload_storage', 'app:uploads'),)), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_reflection', 'full', install_module('kickside/self-reflection', (('kickside.self_reflection.security:admin_scope', 'app.security:admin'),)), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_memory', 'full', install_module('kickside/memory'), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_mcp', 'full', install_module('kickside/mcp'), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_ocr', 'full', install_module('kickside/llm-ocr'), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_pdf2image', 'full', install_module('kickside/pdf2image'), ('ai_ready',), instance_setup=True),
    Scenario('ai_install_html2pdf', 'full', install_module('kickside/html2pdf'), ('ai_ready',), instance_setup=True),
    Scenario('ai_ocr_config', 'full', ocr_config, ('ai_install_ocr',), True),
    Scenario('ai_ocr_settings', 'full', ocr_settings, ('ai_ocr_config',), True, features=(feature('llm-ocr', 'navigation', 'kickside.llm_ocr:settings_block'),)),
    Scenario('ai_model_catalog', 'full', models_catalog, ('ai_ready',), True, features=http('models', 'GET /api/v1/models/admin', 'GET /api/v1/models/list', 'GET /api/v1/models/by-name', 'GET /api/v1/models/classes', 'POST /api/v1/models/cards/test')),
    Scenario('ai_markdown_math_mermaid', 'full', markdown_math_mermaid, ('ai_ready',), features=http('renderers', 'GET /api/v1/renderers')),
    Scenario('ai_models_ui', 'full', models_ui, ('ai_ready',), True, features=(feature('models', 'navigation', 'kickside.models:tab.models'),)),
    Scenario('ai_knowledge_catalog', 'full', knowledge_catalog, ('ai_ready',), True, features=http('knowledge', 'GET /api/v1/knowledge/engines', 'GET /api/v1/knowledge/models') + http('renderers', 'GET /api/v1/renderers')),
    Scenario('ai_agent_catalog', 'full', agent_catalog, ('ai_ready',), features=http('agents', 'GET /api/v1/agents/list', 'GET /api/v1/agents/by-name', 'GET /api/v1/agents/namespaces', 'GET /api/v1/agents/tools', 'GET /api/v1/agents/traits', 'GET /api/v1/agents/memory-options')),
    Scenario('ai_agent_crud', 'full', agent_crud, ('ai_ready',), True, features=http('agents', 'POST /api/v1/user-agents', 'GET /api/v1/user-agents/{id}', 'PUT /api/v1/user-agents/{id}', 'GET /api/v1/user-agents', 'DELETE /api/v1/user-agents/{id}') + (feature('agents', 'navigation', 'kickside.agents:nav_item'), feature('agents', 'workspace_kind', 'kickside.agents:user_agent'))),
    Scenario('ai_gemma_multiturn', 'full', chat_multiturn, ('ai_ready',), True, features=(feature('sessions', 'navigation', 'kickside.sessions:nav_item'),)),
    Scenario('ai_generated_session_title', 'full', generated_session_title, ('ai_ready',), True, features=(feature('sessions', 'navigation', 'kickside.sessions:nav_item'),)),
    Scenario('ai_session_controls', 'full', session_controls, ('ai_ready',), True, features=http('sessions', 'POST /api/v1/sessions/favorite', 'POST /api/v1/sessions/start-token') + (feature('sessions', 'navigation', 'kickside.sessions:nav_item'),)),
    Scenario('ai_memory_chat', 'full', memory_chat, ('ai_install_memory',), True, issues=('PROVIDER-MEMORY-COMPACT-LISTS',), features=(feature('memory', 'capability', 'kickside.memory:knowledge_based_memory'), feature('memory', 'mcp_tool', 'kickside.memory:memory_recall_tool'))),
    Scenario('ai_selected_capabilities_chat', 'full', capabilities_chat(False), ('ai_ready',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:trait_curator'), feature('agents', 'mcp_tool', 'kickside.agents.traits:trait_curator_tool'))),
    Scenario('ai_all_capabilities_chat', 'full', capabilities_chat(True), ('ai_ready',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:trait_master'), feature('agents', 'mcp_tool', 'kickside.agents.traits:trait_curator_tool'))),
    Scenario('ai_platform_tools', 'full', platform_tools, ('ai_install_mcp',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:platform_manager'), feature('agents', 'mcp_tool', 'kickside.agents.traits:platform_manager_tool'))),
    Scenario('ai_roster_chat', 'full', roster_chat, ('ai_ready',), True, issues=('R4-I13',), features=(feature('agents', 'capability', 'kickside.agents.traits:agent_roster'), feature('agents', 'mcp_tool', 'kickside.agents.traits:agents_action'))),
    Scenario('ai_time_chat', 'full', time_chat, ('ai_ready',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:time_aware'),)),
    Scenario('ai_delegation_chat', 'full', delegation_chat, ('ai_ready',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:delegate'),)),
    Scenario('ai_redirect_chat', 'full', redirect_chat, ('ai_ready',), True, features=(feature('agents', 'capability', 'kickside.agents.traits:redirect'), feature('agents', 'mcp_tool', 'kickside.agents.traits:redirect_action'))),
    Scenario('ai_skills_lifecycle', 'full', skills_lifecycle, ('ai_install_skills',), features=http('skills', 'POST /api/v1/skills', 'GET /api/v1/skills', 'GET /api/v1/skills/{id}', 'PATCH /api/v1/skills/{id}', 'GET /api/v1/skills/{id}/versions') + (feature('skills', 'navigation', 'kickside.skills:nav_item'), feature('skills', 'workspace_kind', 'kickside.skills:skill'))),
    Scenario('ai_skills_delete', 'full', skills_delete, ('ai_install_skills',), features=http('skills', 'DELETE /api/v1/skills/{id}')),
    Scenario('ai_document_library', 'full', knowledge_document, ('ai_ready',), True, issues=('SMOKE-KNOWLEDGE-SEARCH-MODEL-OUTPUT',), features=http('knowledge', 'POST /api/v1/knowledge/create', 'GET /api/v1/knowledge/{id}', 'GET /api/v1/knowledge/list', 'PUT /api/v1/knowledge/{id}', 'POST /api/v1/knowledge/{id}/ingest', 'POST /api/v1/knowledge/{id}/search', 'POST /api/v1/knowledge/{id}/ingest-upload', 'POST /api/v1/knowledge/{id}/upload-token') + (feature('knowledge', 'navigation', 'kickside.knowledge:nav_item'), feature('knowledge', 'workspace_kind', 'kickside.knowledge.injection:injection_kind'), feature('kb10', 'workspace_kind', 'kickside.kb10:embeddable'))),
    Scenario('ai_document_cited_answer', 'full', knowledge_answer, ('ai_document_library',), True, issues=('KB-CITED-PASSAGE-REFERENCES',), features=http('knowledge', 'POST /api/v1/knowledge/{id}/ask')),
    Scenario('ai_knowledge_delete', 'full', knowledge_delete_content, ('ai_ready',), True, features=http('knowledge', 'DELETE /api/v1/knowledge/{id}/content', 'DELETE /api/v1/knowledge/{id}')),
    Scenario('ai_wiki_lifecycle', 'full', wiki_lifecycle, ('ai_ready',), True, features=http('kb10', 'POST /api/v1/kb10/{id}/concepts', 'GET /api/v1/kb10/{id}/graph', 'GET /api/v1/kb10/{id}/browse', 'GET /api/v1/kb10/{id}/search')),
    Scenario('ai_wiki_file_ingestion', 'full', wiki_file_ingestion, ('ai_ready',), True, features=http('knowledge', 'POST /api/v1/knowledge/{id}/ingest-upload', 'POST /api/v1/knowledge/{id}/search') + (feature('kb10', 'workspace_kind', 'kickside.kb10:embeddable'),)),
    Scenario('ai_wiki_cited_answer', 'full', wiki_cited_answer, ('ai_wiki_lifecycle',), True, issues=('R5-T04',), features=http('knowledge', 'POST /api/v1/knowledge/{id}/ask')),
    Scenario('ai_wiki_update_delete', 'full', wiki_update_delete, ('ai_ready',), True, features=http('kb10', 'PATCH /api/v1/kb10/{id}/concepts/{node_id}', 'PUT /api/v1/kb10/{id}/concepts/{node_id}/parent', 'DELETE /api/v1/kb10/{id}/concepts/{node_id}')),
    Scenario('ai_wiki_node_delete', 'full', wiki_node_delete, ('ai_ready',), True, features=http('kb10', 'DELETE /api/v1/kb10/{id}/nodes/{node_id}') + http('knowledge', 'DELETE /api/v1/knowledge/{id}/content')),
    Scenario('ai_wiki_source', 'full', wiki_source, ('ai_ready',), True, features=http('kb10', 'GET /api/v1/kb10/{id}/nodes/{node_id}/source')),
    Scenario('ai_wiki_reembed', 'full', wiki_maintenance('reembed'), ('ai_ready',), True, features=http('kb10', 'POST /api/v1/kb10/{id}/reembed', 'GET /api/v1/kb10/{id}/maintenance')),
    Scenario('ai_wiki_resummarize', 'full', wiki_maintenance('resummarize'), ('ai_ready',), True, features=http('kb10', 'POST /api/v1/kb10/{id}/resummarize', 'GET /api/v1/kb10/{id}/maintenance')),
    Scenario('ai_self_reflection', 'full', reflection, ('ai_install_reflection', 'ai_install_mcp'), issues=('B02',), features=(feature('self-reflection', 'capability', 'kickside.self_reflection.traits:self_reflection'), feature('self-reflection', 'mcp_tool', 'kickside.self_reflection.traits:self_reflection_tool'))),
    Scenario('ai_file_tools', 'full', files_tools, ('ai_install_mcp',), features=tuple(feature('uploads', 'mcp_tool', 'kickside.uploads.traits:' + n) for n in ('list_files', 'count_files', 'view_files', 'upload_file')) + (feature('uploads', 'capability', 'kickside.uploads.traits:file_manager'),)),
    Scenario('ai_skill_tools', 'full', skill_tools, ('ai_install_skills', 'ai_install_mcp'), features=(feature('skills', 'capability', 'kickside.skills.traits:skills'), feature('skills', 'mcp_tool', 'kickside.skills.traits:skill_tool'), feature('skills', 'mcp_tool', 'kickside.skills.traits:skill_file_tool'))),
    Scenario('ai_organize_files', 'full', organize_files, ('ai_ready',), True, features=http('knowledge', 'POST /api/v1/knowledge/organize')),
    Scenario('ai_upload_delete', 'full', upload_delete, ('ai_ready',), features=http('uploads', 'DELETE /api/v1/uploads/{uuid}')),
    Scenario('ai_upload_direct_validation', 'full', upload_direct_validation, ('ai_ready',), features=http('uploads', 'POST /api/v1/uploads/presigned', 'POST /api/v1/uploads/complete')),
    Scenario('ai_upload_workspace', 'full', upload_workspace, ('ai_ready',), features=(feature('uploads', 'workspace_kind', 'kickside.uploads:uploads_component'),), issues=('R4-P02', 'R4-I14')),
    Scenario('ai_storage_settings', 'full', storage_settings, ('ai_ready',), features=(feature('uploads', 'navigation', 'kickside.uploads:settings.block.storage'),)),
    Scenario('ai_pdf_rasterization', 'full', pdf_rasterization, ('ai_install_pdf2image', 'ai_ocr_config'), True, features=tuple(feature('pdf2image', 'converter', 'kickside.convert.pdf2image:' + name) for name in ('capability', 'pages', 'render')) + (feature('doc2md', 'converter', 'kickside.convert.doc2md:page_resources'),)),
    Scenario('ai_wiki_enrichment', 'full', wiki_enrichment, ('ai_ready',), True, issues=('R4-I7',), features=http('kb10', 'POST /api/v1/kb10/{id}/enrichment') + tuple(feature('kb10', 'mcp_tool', 'kickside.kb10.graph:' + name) for name in ('browse_tree', 'create_concept', 'attach_document'))),
    Scenario('ai_wiki_default_placement', 'full', wiki_default_placement, ('ai_ready',), True, features=http('kb10', 'POST /api/v1/kb10/{id}/enrichment') + tuple(feature('kb10', 'mcp_tool', 'kickside.kb10.graph:' + name) for name in ('browse_tree', 'create_concept', 'attach_document'))),
    Scenario('ai_upload_sav_typed', 'full', sav_typed_upload, ('ai_ready',), features=(feature('doc2md', 'converter', 'kickside.convert.doc2md:converter'), feature('uploads', 'upload_kind', 'kickside.uploads.documents:types.documents'))),
    Scenario('ai_knowledge_ingest_tools', 'full', knowledge_ingest_tools, ('ai_install_mcp',), True, features=(feature('knowledge', 'capability', 'kickside.knowledge.traits.kb_ingest:kb_ingest'),) + tuple(feature('knowledge', 'mcp_tool', 'kickside.knowledge.traits.kb_ingest:' + name) for name in ('select_kb', 'embed_content', 'delete_content'))),
    Scenario('ai_knowledge_writer', 'full', knowledge_writer, ('ai_install_mcp',), True, features=(feature('knowledge', 'capability', 'kickside.knowledge.traits:kb_writer'), feature('knowledge', 'mcp_tool', 'kickside.knowledge.traits:kb_embed_tool'))),
]

for entry, alias, args, issues in [
        ('kb_explorer', 'SearchKB', lambda kid: {'query': 'Northwind returns', 'kb_id': kid, 'fields': ['content', 'meta']}, ('R4-I7',)),
        ('knowledge_manager', 'Knowledge', lambda kid: {'action': 'list_knowledge_bases', 'component_id': kid}, ('TS-KNOWLEDGE-CAPABILITY-ARGUMENTS',))]:
    trait = 'kickside.knowledge.traits:' + entry
    REGISTRY.append(Scenario('ai_knowledge_capability_' + entry, 'full', knowledge_capability(trait, alias, args), ('ai_install_mcp',), True, issues=issues, features=(feature('knowledge', 'capability', trait),)))

for alias, entry, analysis in [('QueryKB', 'kb_query_tool', False), ('RetrieveKB', 'kb_retrieve_tool', False), ('AskKB', 'kb_ask_tool', True), ('SearchKB', 'kb_search_tool', False), ('FindKnowledgeBases', 'kb_find_tool', False), ('Knowledge', 'knowledge_manager_tool', False)]:
    tool_id = 'kickside.knowledge.traits:' + entry
    REGISTRY.append(Scenario('ai_knowledge_tool_' + entry, 'full', scoped_knowledge_chat(alias) if alias in ('QueryKB', 'RetrieveKB') else knowledge_tool(alias, tool_id, analysis=analysis), ('ai_install_mcp',), True, features=(feature('knowledge', 'mcp_tool', tool_id),) + ((feature('knowledge', 'capability', 'kickside.knowledge.traits:kb_reader'),) if alias == 'QueryKB' else ())))

for alias, tool_id in [('EmbedToKB', 'kickside.knowledge.traits:kb_embed_tool'), ('EmbedContent', 'kickside.knowledge.traits.kb_ingest:embed_content')]:
    REGISTRY.append(Scenario('ai_knowledge_embed_' + alias.lower(), 'full', knowledge_embed_tool(alias, tool_id), ('ai_install_mcp',), True, features=(feature('knowledge', 'mcp_tool', tool_id),)))

for alias, entry in [('BrowseKB', 'kb_browse_tool'), ('GetKBNode', 'kb_get_node_tool'), ('TraverseKBGraph', 'kb_graph_tool'), ('SurveyKB', 'kb_summaries_tool')]:
    tool_id = 'kickside.kb10.traits:' + entry
    REGISTRY.append(Scenario('ai_wiki_tool_' + entry, 'full', wiki_tool(alias, tool_id), ('ai_install_mcp',), True, issues=('PROVIDER-KB10-COMPACT-LISTS',), features=(feature('kb10', 'mcp_tool', tool_id),)))

for kind, issues in [('document', ('ARTIFACT-EMBED-PROMPT-CONFLICT',)), ('table', ('ARTIFACT-EMBED-PROMPT-CONFLICT',)), ('chart', ('R4-I15', 'R4-U03', 'ARTIFACT-EMBED-PROMPT-CONFLICT'))]:
    REGISTRY.append(Scenario('ai_artifact_' + kind, 'full', artifact_chat(kind), ('ai_ready',), True, issues=issues,
        features=(feature('agents', 'capability', 'kickside.agents.traits:artifacts'), feature('agents', 'mcp_tool', 'kickside.agents.traits:create_artifact'))))

for preset in ('anthropic', 'bedrock', 'cerebras', 'deepinfra', 'fireworks', 'google_gemini', 'google_vertex', 'groq', 'lmstudio', 'mistral', 'ollama', 'openai', 'openai_compatible', 'openai_compatible_plain', 'openrouter', 'together', 'vllm', 'xai'):
    entry = 'kickside.models:preset.' + preset
    REGISTRY.append(Scenario('ai_provider_preset_' + preset, 'full', provider_preset(entry), ('ai_ready',),
        features=(feature('models', 'model_provider', entry),) + http('models', 'POST /api/v1/models/provider-profiles', 'PUT /api/v1/models/provider-profiles/{id}', 'DELETE /api/v1/models/provider-profiles/{id}')))

DOCUMENT_FORMATS = ('txt', 'md', 'markdown', 'rtf', 'pdf', 'docx', 'doc', 'dot', 'xlsx', 'xlsm', 'xlsb', 'xls', 'pptx', 'ppt', 'pps', 'sav', 'scan.pdf')
IMAGE_FORMATS = ('png', 'jpg', 'jpeg', 'webp', 'tiff', 'tif', 'bmp', 'gif', 'svg', 'ico')
TEXT_FORMATS = ('csv', 'tsv', 'html', 'htm', 'json', 'xml', 'js', 'css', 'py', 'lua', 'rb', 'php', 'java', 'c', 'cpp', 'cs', 'go', 'sh', 'bat', 'ps1', 'yml', 'yaml', 'toml', 'sql')
for suffix in DOCUMENT_FORMATS + IMAGE_FORMATS + TEXT_FORMATS + ('dxf', 'zip', 'bin'):
    kind = 'documents' if suffix in DOCUMENT_FORMATS else 'images' if suffix in IMAGE_FORMATS else 'text' if suffix in TEXT_FORMATS else 'generic' if suffix in ('bin', 'dxf') else 'archive'
    entry = 'kickside.uploads.documents:types.documents' if kind == 'documents' else 'kickside.uploads.archive:types.archive' if kind == 'archive' else 'kickside.uploads:types.' + kind
    claims = http('uploads', 'POST /api/v1/uploads', 'GET /api/v1/uploads/{uuid}', 'GET /api/v1/uploads/{uuid}/artifacts', 'GET /api/v1/uploads/{uuid}/raw', 'GET /api/v1/uploads/list') + (feature('uploads', 'upload_kind', entry), feature('uploads', 'navigation', 'kickside.uploads:nav_item'))
    if suffix not in ('zip', 'bin', 'dxf'):
        claims += http('uploads', 'GET /api/v1/uploads/{uuid}/artifacts/{artifact_id}/raw')
    if kind == 'documents':
        claims += (feature('doc2md', 'converter', 'kickside.convert.doc2md:converter'),)
    prerequisites = ('ai_ocr_config',) if suffix in IMAGE_FORMATS + ('scan.pdf',) else ('ai_ready',)
    REGISTRY.append(Scenario('ai_upload_' + suffix.replace('.', '_'), 'full', upload_format(suffix), prerequisites, True, features=claims))

REGISTRY.append(Scenario('ai_artifact_standalone_chart', 'full', standalone_chart, ('ai_ready',), True, issues=('R4-I15', 'R4-U03'), features=(feature('agents', 'capability', 'kickside.agents.traits:artifacts'), feature('agents', 'mcp_tool', 'kickside.agents.traits:create_artifact'))))

REGISTRY.append(Scenario('ai_docx_primary_body', 'full', docx_primary_body, ('ai_ocr_config',), True,
    features=http('uploads', 'POST /api/v1/uploads', 'GET /api/v1/uploads/{uuid}', 'GET /api/v1/uploads/{uuid}/artifacts', 'GET /api/v1/uploads/{uuid}/raw', 'GET /api/v1/uploads/{uuid}/artifacts/{artifact_id}/raw') + (feature('uploads', 'upload_kind', 'kickside.uploads.documents:types.documents'), feature('doc2md', 'converter', 'kickside.convert.doc2md:converter'))))

REGISTRY.append(Scenario('renderer_docx_grounded_answer_settles', 'full', renderer_docx_grounded_answer_settles,
    ('ai_ready',), True, issues=('RENDERER-WORKER-REQUEST-SETTLE',)))

REGISTRY.append(Scenario('ai_provider_live', 'full', provider_live, ('ai_ready',), True, features=http('models', 'POST /api/v1/models/provider-profiles/{id}/test', 'POST /api/v1/models/provider-profiles/{id}/discover', 'PUT /api/v1/models/cards', 'DELETE /api/v1/models/cards')))


def ai_event_contract_setup(ctx):
    from scenarios_providers import install_probe
    for module in ('kickside/sync', 'kickside/crm'):
        ctx.inst.install(module)
    ctx.state['probe_http_resources'] = ['https://example.com/*']
    install_probe(ctx)


def ai_skill_event_delivery(ctx):
    from scenarios_crm_chestor import crm_destination, sync_trigger, delivered
    inst = ctx.inst
    sid = inst.request('POST', '/api/v1/skills', {'name': 'Northwind event checklist', 'description': 'Dispatch checks',
        'body': 'Inspect 48 crates.'}, label='skill-event-create')['skill_id']
    inst.cleanup.add('Skill event source', lambda: inst.request('DELETE', '/api/v1/skills/' + sid), scope='scenario')
    marker = 'skill-updated-' + uuid.uuid4().hex[:8]
    cid, destination = crm_destination(ctx, marker)
    inst.cleanup.add('Skill event destination', lambda: inst.request('DELETE', '/api/v1/threads/' + cid), scope='scenario')
    sync_trigger(ctx, 'kickside.skills:skill_updated', sid, cid, marker, scope='scenario')
    inst.request('PATCH', '/api/v1/skills/' + sid, {'body': 'Inspect 48 crates and record NW-SKILL-EVENT.'}, label='skill-event-update')
    require('NW-SKILL-EVENT' in inst.request('GET', '/api/v1/skills/' + sid)['skill']['body'], 'Skill event source does not save its revision')
    delivered(ctx, destination, marker)


def ai_knowledge_event_delivery(ctx):
    from scenarios_crm_chestor import crm_destination, sync_trigger, delivered
    kid = kb_create(ctx, 'document', 'Northwind event knowledge')
    marker = 'knowledge-updated-' + uuid.uuid4().hex[:8]
    cid, destination = crm_destination(ctx, marker)
    ctx.inst.cleanup.add('Knowledge event source', lambda: ctx.inst.request('DELETE', '/api/v1/knowledge/' + kid), scope='scenario')
    ctx.inst.cleanup.add('Knowledge event destination', lambda: ctx.inst.request('DELETE', '/api/v1/threads/' + cid), scope='scenario')
    sync_trigger(ctx, 'kickside.knowledge.injection:content_updated', kid, cid, marker, scope='scenario')
    kb_ingest(ctx, kid)
    delivered(ctx, destination, marker)


def ai_upload_event_delivery(ctx):
    from scenarios_crm_chestor import crm_destination, sync_trigger, delivered
    completed_upload(ctx, fixture(ctx, 'txt'))
    marker = 'upload-completed-' + uuid.uuid4().hex[:8]
    cid, destination = crm_destination(ctx, marker)
    ctx.inst.cleanup.add('Upload event destination', lambda: ctx.inst.request('DELETE', '/api/v1/threads/' + cid), scope='scenario')
    sync_trigger(ctx, 'kickside.uploads:upload_completed', None, cid, marker, scope='scenario')
    path = ctx.inst.app.parent / 'Northwind-event-upload.txt'
    path.write_text('NW-UPLOAD-EVENT: 48 fictional dispatch crates.')
    uid, _ = completed_upload(ctx, path)
    require(ctx.inst.request('GET', '/api/v1/uploads/' + uid)['upload']['status'] == 'completed', 'Upload event source does not finish processing')
    delivered(ctx, destination, marker)


def ai_fetch_tools(ctx):
    from scenarios_providers import probe
    url = 'https://example.com/'
    kid = kb_create(ctx, 'document', 'Northwind URL ingestion')
    ctx.inst.cleanup.add('URL ingestion KB', lambda: ctx.inst.request('DELETE', '/api/v1/knowledge/' + kid), scope='scenario')
    fetched = probe(ctx, tool='kickside.knowledge.traits.kb_ingest:fetch_url',
        context={'kb_id': kid}, arguments={'url': url, 'kb_id': kid, 'embed': True}, label='knowledge-fetch-url')
    fetched = json.loads(fetched) if isinstance(fetched, str) else fetched
    require(fetched['success'] and fetched['status'] == 'queued' and fetched['thread_id'], 'FetchURL does not queue a durable KB event: ' + str(fetched))
    def indexed():
        result = ctx.inst.request('POST', '/api/v1/knowledge/' + kid + '/search', {'query': 'NORTHWIND-PAGE-731', 'limit': 10}, label='fetched-kb-source')
        return result if 'Example Domain' in json.dumps(result) else None
    ctx.inst.wait(indexed, 'Fetched URL indexes its fixture content', seconds=120)
    uploaded = probe(ctx, tool='kickside.uploads.traits:fetch_url', arguments={'url': url, 'wait': True}, label='uploads-fetch-url')
    uploaded = json.loads(uploaded) if isinstance(uploaded, str) else uploaded
    ctx.inst.evidence.save('fetched-upload-result', uploaded)
    uid = uploaded['upload_id']
    require(ctx.inst.request('GET', '/api/v1/uploads/' + uid)['upload']['status'] == 'completed', 'Uploads FetchURL does not finish processing')
    raw, _ = raw_bytes(ctx, '/api/v1/uploads/' + uid + '/raw', 'fetched-upload-source')
    require(b'Example Domain' in raw, 'Uploads FetchURL loses fixture bytes')


def ai_wiki_placement_tools(ctx):
    from scenarios_providers import probe
    kid = kb_create(ctx, 'wiki', 'Northwind placement tools')
    ctx.inst.cleanup.add('Placement tools KB', lambda: ctx.inst.request('DELETE', '/api/v1/knowledge/' + kid), scope='scenario')
    kb_ingest(ctx, kid)
    base = '/api/v1/kb10/' + kid
    documents = ctx.inst.request('GET', base + '/browse?node_type=document&limit=100')['nodes']
    require(len(documents) == 1, 'Placement fixture does not have one indexed document')
    document = documents[0]['id']
    context = {'graph_component_id': kid, 'graph_store_id': kid, 'graph_embedding_model': ctx.state['embed']}
    def call(name, args):
        return probe(ctx, tool='kickside.kb10.graph:' + name, context=context, arguments=args, label='placement-' + name)
    def documents_of(concept):
        return {row['id'] for row in call('get_concept_documents', {'concept_id': concept})['documents']}
    def graph():
        return ctx.inst.request('GET', base + '/graph?limit=100', label='placement-graph')
    nodes = graph()['nodes']
    root = next(n['id'] for n in nodes if n.get('concept_type') == 'root')
    for label in ('Northwind dispatch duplicate', 'Northwind canonical dispatch'):
        call('create_concept', {'label': label, 'parent_id': root})
    nodes = graph()['nodes']
    source = next(n['id'] for n in nodes if n['label'] == 'Northwind dispatch duplicate')
    target = next(n['id'] for n in nodes if n['label'] == 'Northwind canonical dispatch')
    call('attach_document', {'document_node_id': document, 'concept_id': source})
    require(document in documents_of(source), 'GetConceptDocuments omits the attached document')
    call('move_document', {'document_node_id': document, 'concept_id': target})
    require(document in documents_of(target) and document not in documents_of(source),
        'MoveDocument keeps the former categorization')
    call('move_document', {'document_node_id': document, 'concept_id': source})
    call('merge_concept', {'from_id': source, 'into_id': target})
    require(source not in {n['id'] for n in graph()['nodes']} and document in documents_of(target),
        'MergeConcept loses the document or keeps the duplicate')


REGISTRY += [
    Scenario('ai_event_contract_setup', 'full', ai_event_contract_setup, ('ai_ready',), instance_setup=True),
    Scenario('ai_skill_event_delivery', 'full', ai_skill_event_delivery, ('ai_event_contract_setup', 'ai_install_skills'),
        features=(feature('skills', 'trigger', 'kickside.skills:skill_updated'),)),
    Scenario('ai_knowledge_event_delivery', 'full', ai_knowledge_event_delivery, ('ai_event_contract_setup',), True,
        features=(feature('knowledge', 'trigger', 'kickside.knowledge.injection:content_updated'),)),
    Scenario('ai_upload_event_delivery', 'full', ai_upload_event_delivery, ('ai_event_contract_setup',),
        features=(feature('uploads', 'trigger', 'kickside.uploads:upload_completed'),)),
    Scenario('ai_fetch_tools', 'full', ai_fetch_tools, ('ai_event_contract_setup',), True,
        features=(feature('uploads', 'mcp_tool', 'kickside.uploads.traits:fetch_url'),
            feature('knowledge', 'mcp_tool', 'kickside.knowledge.traits.kb_ingest:fetch_url'))),
    Scenario('ai_wiki_placement_tools', 'full', ai_wiki_placement_tools, ('ai_event_contract_setup',), True, issues=('R4-I7',),
        features=tuple(feature('kb10', 'mcp_tool', 'kickside.kb10.graph:' + name)
            for name in ('get_concept_documents', 'move_document', 'merge_concept'))),
]


def ai_user_aware_sink(ctx):
    from scenarios_providers import probe
    user = ctx.inst.request('GET', '/api/v1/user/me', label='sink-execution-user')['user']
    name = user['full_name']
    require(name and name != 'Unknown', 'User fixture has no full name')
    aid = create_agent(ctx, 'Northwind user-aware receipt agent', ('kickside.users.traits:user_aware',),
        prompt='Return only JSON with user_name equal to the quoted User name in your injected context and crates equal to the sum in the task. Do not guess a name.')
    ctx.inst.cleanup.add('User-aware sink agent', lambda: ctx.inst.request('DELETE', '/api/v1/user-agents/' + aid), scope='scenario')
    result = probe(ctx, port='kickside.agents:invoke_sink', write=True, arguments={
        'origin': 'app.smoke.providers:probe', 'identity_key': 'northwind-user-sink', 'idempotency_key': uuid.uuid4().hex, 'sink_op': 'upsert',
        'config': {'agent_id': aid}, 'input': {'content': 'Harbor has 31 crates and Orchard has 17. Report their total and the injected user name.'}}, label='user-aware-agent-sink')
    require(result['result'] == 'created' and result['dest_ref']['agent_id'] == 'user_agent:' + aid,
        'Agent sink does not return the selected agent receipt: ' + str(result))
    output = result['output']
    require(re.search(r'"user_name"\s*:\s*' + re.escape(json.dumps(name)), output)
        and re.search(r'"crates"\s*:\s*48\b', output),
        'Agent sink loses user context or computed output: ' + str(output))


REGISTRY.append(Scenario('ai_user_aware_sink', 'full', ai_user_aware_sink, ('ai_event_contract_setup',), True,
    features=(feature('agents', 'sink', 'kickside.agents:invoke_sink'),
        feature('users', 'capability', 'kickside.users.traits:user_aware'))))


def regression_skills_upgrade(ctx):
    from playwright.sync_api import expect
    from module_lifecycle_contract import published_fixture
    from upgrade_regressions import upgrade_application
    module = 'kickside/skills'
    with published_fixture(ctx, module, published_app=True) as fixture:
        fixture.inst.install(module, [{'name': 'kickside.skills:upload_storage', 'value': 'app:uploads'}])
        fixture.inst.install('kickside/mcp')
        sid = skill_tools(fixture)
        fixture.inst.cleanup.add('upgrade skill', lambda: fixture.inst.request('DELETE', '/api/v1/skills/' + sid))
        upgrade_application(fixture, ctx.inst.config['source'], ctx.inst.config['replace'])
        nav(fixture, 'Skills')
        skills = fixture.browser.frame.locator('kickside-skills')
        skills.get_by_role('button', name='Open Northwind bundled dispatch', exact=True).click()
        dialog = skills.get_by_role('dialog', name='Edit skill', exact=True)
        fixture.browser.screenshot('upgraded-skill-dialog')
        file = dialog.locator('.kx-file').filter(has_text='references/dispatch.md')
        expect(file).to_be_visible(timeout=30000)
        file.get_by_text('references/dispatch.md', exact=True).click()
        expect(skills.get_by_text(re.compile('Northwind dispatch marker NW-BUNDLE-731. Returns within 17 days.'))).to_be_visible(timeout=30000)
        fixture.browser.screenshot('upgraded-skill-file')


REGISTRY.append(Scenario('regression_skills_upgrade', 'full', regression_skills_upgrade, ('ai_boot',),
    features=(feature('skills', 'http', 'GET /api/v1/skills/{id}'),
        feature('skills', 'mcp_tool', 'kickside.skills.traits:skill_file_tool')),
    issues=('SKILLS-BUNDLE-UPGRADE', 'SMOKE-UPGRADE-MANIFEST-SOURCE-ROOT', 'SMOKE-UPGRADE-SKILL-RECEIPT')))


def regression_reflection_chat(ctx):
    trait = 'kickside.self_reflection.traits:self_reflection'
    aid = create_agent(ctx, 'Northwind registry reflection assistant', (trait,),
        prompt='Use SelfReflection when asked to inspect registry declarations. Report actual results without guessing.')
    ctx.inst.cleanup.add('reflection chat agent', lambda: ctx.inst.request(
        'DELETE', '/api/v1/user-agents/' + aid), scope='scenario')
    browser = ctx.browser_check()
    browser.login()
    browser.frame.locator('.sidebar').wait_for(timeout=30000)
    browser.open_chat(aid)
    browser.stream('Use SelfReflection search_registry with namespace kickside.agents and q artifacts. '
        'Quote the matching capability ID, explain the result in ten numbered sentences, and end with NW-REFLECTION-READY.',
        'NW-REFLECTION-READY', 'reflection-chat')
    require('kickside.agents.traits:artifacts' in browser.frame.locator('wippy-chat').evaluate(ASSISTANT_TEXT),
        'Self Reflection chat omits the actual registry capability')
    sid = browser.current_session()
    stored = ctx.inst.request('GET', '/api/v1/sessions/messages?session_id=' + sid)['messages']
    require(any(row['type'] == 'function' and row.get('metadata', {}).get('function_name') == 'SelfReflection'
        and row['metadata'].get('status') == 'success' for row in stored), 'Chat does not execute SelfReflection successfully')


REGISTRY.append(Scenario('regression_reflection_chat', 'full', regression_reflection_chat,
    ('ai_ready', 'ai_install_reflection'), model=True, issues=('B02',),
    features=(feature('self-reflection', 'mcp_tool', 'kickside.self_reflection.traits:self_reflection_tool'),)))


from scenarios_knowledge_artifact_events import large_upload

REGISTRY.append(Scenario('knowledge_large_upload', 'full', large_upload, ('ai_ready',), True,
             issues=('KNOWLEDGE-BOUNDED-CONTENT-EVENT', 'SMOKE-KNOWLEDGE-SETUP-OWNERSHIP'), features=(
                 'kickside/knowledge|http|POST /api/v1/knowledge/{id}/ingest-upload',
                 'kickside/knowledge|http|POST /api/v1/knowledge/{id}/search',
                 'kickside/uploads|http|POST /api/v1/uploads',
                 'kickside/core|http|GET /api/v1/threads/{id}/events',
             )))
