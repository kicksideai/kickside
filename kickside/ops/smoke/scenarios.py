"""Product scenarios consolidated from the onboarding, release, QA, and browser proofs."""
import json
import re
import time
from urllib.parse import quote
import uuid

from browser import Browser
from harness import Failure, Scenario, require


class Context:
    def __init__(self, inst, browser_session=None):
        self.inst = inst
        self.browser_session = browser_session
        from mocks import Fixtures
        self.fixtures = Fixtures()
        self.browser = None
        self.browser_navigation = None
        self.state = {}

    def browser_check(self, **presentation):
        if self.browser is None:
            self.browser = Browser(self.inst, session=self.browser_session, **presentation)
            self.inst.cleanup.add('browser', self.browser.close)
            if self.browser_navigation is not None:
                try:
                    self.browser.restore(self.browser_navigation)
                except Exception:
                    self.close_browser()
                    raise
                self.browser_navigation = None
        return self.browser

    def resume_browser(self):
        if self.browser_navigation is not None:
            self.browser_check()

    def suspend_browser(self):
        if self.browser is None:
            return
        try:
            self.browser_navigation = self.browser.navigation()
        finally:
            self.close_browser()

    def close_browser(self):
        browser, self.browser = self.browser, None
        if browser is None:
            return
        self.inst.cleanup.remove(browser.close)
        browser.close()

    def restart(self, hard=False):
        return self.restart_runtime(self, hard=hard)

    def paused(self, hard=False):
        return self.pause_runtime(self, hard=hard)

    def navigate(self, navigation):
        surface, agent, session = navigation
        require(surface == 'chat', 'Unknown declared browser navigation: ' + surface)
        self.browser_check().open_chat(self.state[agent], self.state.get(session))


def boot(ctx):
    inst = ctx.inst
    inst.boot()
    browser = ctx.browser_check()
    browser.page.goto(inst.base + '/login', wait_until='domcontentloaded', timeout=30000)
    browser.page.get_by_role('button', name='Sign In', exact=True).wait_for(timeout=30000)


def select_onboarding_models(ctx, models, generation, embedding):
    expected = {}
    for select in models.get_by_role('combobox').all():
        label = select.get_attribute('aria-label') or ''
        card = embedding if 'embed' in label.lower() else generation
        select.click()
        ctx.browser.frame.get_by_role('option', name=card['title'], exact=True).click()
        ctx.browser.settle()
        expected[label.removeprefix('Model for ')] = card['name']
    classes = ctx.inst.request('GET', '/api/v1/models/classes', label='onboarding-model-winners')['classes']
    actual = {c.get('title') or c['name']: c['models'][0]['name'] for c in classes if c.get('models')}
    require(expected and all(actual.get(title) == name for title, name in expected.items()),
            'Onboarding model classes differ from the requested selections: ' + str(actual))


def platform_onboarding(ctx, *, embedding_dimensions=512):
    from playwright.sync_api import expect
    inst, browser = ctx.inst, ctx.browser_check()
    browser.login()
    frame = browser.frame
    frame.get_by_role('heading', name='Name your workspace', exact=True).wait_for(timeout=30000)
    expect(frame.locator('.sp-screen')).to_be_visible()
    expect(frame.locator('dialog.sp-dialog')).to_have_count(0)
    for width, height in ((1440, 900), (390, 844)):
        browser.page.set_viewport_size({'width': width, 'height': height})
        browser.screenshot(f'instance-wizard-{width}x{height}', full_page=False, mask_color='#e5e7eb')
    browser.page.set_viewport_size({'width': 1440, 'height': 900})
    frame.get_by_label('Workspace name', exact=True).fill('Northwind smoke workspace')
    frame.get_by_role('button', name='Continue ›', exact=True).click()
    frame.get_by_role('heading', name='Add your logo', exact=True).wait_for(timeout=30000)
    frame.get_by_role('button', name='Continue ›', exact=True).click()
    models = frame.locator('kickside-onboard-models')
    model = inst.config['model']
    if model:
        models.get_by_role('button', name=re.compile('vLLM', re.I)).click()
        models.get_by_label('Name', exact=True).fill('On-prem GPU server')
        models.get_by_label(re.compile('^Base URL')).fill(model['base'])
        models.get_by_label(re.compile('^API key')).fill(model['key'])
        browser.screenshot('provider-setup-masked')
        models.get_by_role('button', name='Save provider', exact=True).click()
        models.get_by_role('button', name='Auto-select', exact=True).wait_for(timeout=120000)
        from model_fixture import configure_cards
        generation, embedding = configure_cards(inst, inst.request('GET', '/api/v1/models/admin')['models'],
                                                embedding_dimensions=embedding_dimensions)
        models.get_by_role('button', name='‹ Back', exact=True).click()
        frame.get_by_role('heading', name='Add your logo', exact=True).wait_for(timeout=30000)
        frame.get_by_role('button', name='Continue ›', exact=True).click()
        models.get_by_role('button', name='Auto-select', exact=True).wait_for(timeout=120000)
        ctx.state.update(model=generation['name'], embed=embedding['name'])
        select_onboarding_models(ctx, models, generation, embedding)
        models.get_by_role('button', name='Continue ›', exact=True).click()
    else:
        models.get_by_role('button', name='Skip for now', exact=True).click()
    modules = frame.locator('kickside-onboard-modules')
    modules.wait_for(state='visible', timeout=30000)
    require('Cron' not in modules.inner_text(), 'Retired Cron appears in onboarding')
    modules.get_by_role('list', name='Included capabilities').get_by_text('Knowledge base', exact=False).wait_for(timeout=30000)
    action = modules.get_by_role('button', name=re.compile(r'^(Install \d+ packages?|Continue with this setup)')).first
    action.wait_for(state='visible', timeout=300000)
    if action.inner_text().strip().startswith('Install '):
        with browser.successful_response('POST', '/api/v1/keeper/hub/dependencies/install',
                                        'onboarding-module-install'):
            action.click()
            modules.get_by_role('button', name=re.compile('^Continue with this setup')).wait_for(state='visible', timeout=300000)
    modules.get_by_role('button', name=re.compile('^Continue with this setup')).click(timeout=300000)
    frame.get_by_role('heading', name='Connect an external agent', exact=True).wait_for(timeout=30000)
    with browser.page.expect_response(lambda r: '/api/v1/keeper/mcp/tokens' in r.url and r.request.method == 'POST') as response:
        frame.get_by_role('button', name='Generate access token', exact=True).click()
    issued = response.value.json()
    require(response.value.status == 200 and issued.get('success') and issued.get('token', {}).get('token'),
            'Onboarding does not create a Keeper MCP token')
    token = issued['token']['token']
    inst.redactor.add(token)
    ctx.state['keeper_ui_token'] = token
    require(frame.locator('.sp-mcp-row').filter(has=frame.get_by_text('Token', exact=True)).locator('code').inner_text() == token,
            'Onboarding does not reveal the issued token')
    import hashlib
    inst.cleanup.add('UI Keeper token', lambda: inst.request('POST', '/api/v1/keeper/mcp/tokens/revoke',
        {'token': hashlib.sha256(token.encode()).hexdigest()}, label='keeper-ui-token-revoke'))
    frame.get_by_role('button', name='Continue ›', exact=True).click()
    frame.get_by_role('heading', name="You're ready to go", exact=True).wait_for(timeout=30000)
    with browser.page.expect_response(lambda r: '/api/v1/ui/splash/welcome/seen' in r.url and r.request.method == 'POST') as response:
        frame.get_by_role('button', name='Enter workspace', exact=True).click()
    require(response.value.status == 200 and response.value.json().get('success'), 'Wizard completion does not persist')
    nav = inst.request('GET', '/api/v1/nav/list?manage=1', label='completed-splash')
    require(nav.get('config', {}).get('completed_splashes', []).count('welcome') == 1, 'Welcome completion count differs')
    require((nav.get('splash') or {}).get('id') == 'first_agent', 'First-agent prompt does not follow platform onboarding')
    frame.locator('kickside-agent-onboarding').get_by_label('Agent name', exact=True).wait_for(timeout=30000)


def first_agent(ctx):
    from playwright.sync_api import expect
    browser, inst = ctx.browser, ctx.inst
    dialog = browser.frame.locator('dialog.sp-dialog[open]')
    expect(dialog).to_be_visible()
    expect(dialog).to_have_attribute('aria-modal', 'true')
    expect(dialog.locator('.sp-aside, .sp-eyebrow')).to_have_count(0)
    personal = dialog.locator('kickside-agent-onboarding')
    catalog = inst.request('GET', '/api/v1/agents/traits', label='installed-capabilities')['traits']
    require(catalog, 'Installed capability catalog is empty')
    picker = personal.locator('wc-agent-trait-picker')
    picker.get_by_role('combobox', name='Search capabilities to add').click()
    preselected = {'kickside.agents.traits:trait_master', 'kickside.agents.traits:artifacts', 'kickside.uploads.traits:file_manager'}
    require(preselected <= {trait['id'] for trait in catalog}, 'Default first-agent capabilities are not installed')
    expect(picker.get_by_role('option')).to_have_count(len(catalog) - len(preselected))
    for trait in catalog:
        expect(picker.locator('.kx-trait-name').get_by_text(trait.get('title') or trait['id'], exact=True)).to_have_count(1)
    picker.get_by_role('combobox', name='Search capabilities to add').press('Escape')
    expect(personal.get_by_label('Agent name', exact=True)).to_have_value('My Assistant')
    for title in ('All Capabilities', 'Artifacts', 'File Manager'):
        expect(picker.locator('.kx-trait-name').get_by_text(title, exact=True)).to_be_visible()
    personal.get_by_label('Agent name', exact=True).fill('Northwind smoke assistant')
    with browser.page.expect_response(lambda r: '/api/v1/ui/splash/first_agent/seen' in r.url and r.request.method == 'POST') as seen_response:
        with browser.page.expect_response(lambda r: '/api/v1/user-agents' in r.url and r.request.method == 'POST') as response:
            personal.get_by_role('button', name='Create agent', exact=True).click()
        created = response.value.json()
        require(response.value.status == 200 and created.get('success') and created.get('component_id'), 'First-agent creation fails: ' + str(created))
        ctx.state['agent'] = created['component_id']
        launch_catalog = inst.request('GET', '/api/v1/agents/list', label='first-agent-launch-catalog')
        launch_agent = next((agent for agent in launch_catalog.get('agents', [])
                             if agent.get('name') == 'user_agent:' + created['component_id']), None)
        require(launch_agent and launch_agent.get('start_token'),
                'Created agent has no chat start token in the launch catalog: ' + str(launch_catalog))
    require(seen_response.value.status == 200 and seen_response.value.json().get('success') is True,
            'Agent completion does not persist personal onboarding seen')
    browser.frame.locator('wippy-chat').wait_for(timeout=30000)
    expect(browser.frame.locator('.sp-screen, dialog.sp-dialog')).to_have_count(0, timeout=30000)
    require((inst.request('GET', '/api/v1/nav/list', label='first-agent-completed').get('splash') or {}).get('id') != 'first_agent',
            'Agent completion does not mark personal onboarding seen')
    agent = inst.request('GET', '/api/v1/user-agents/' + created['component_id'], label='first-agent')['agent']
    require(agent['title'] == 'Northwind smoke assistant', 'Agent name does not persist')
    require(agent['traits'] == ['kickside.agents.traits:trait_master', 'kickside.agents.traits:artifacts', 'kickside.uploads.traits:file_manager'],
            'First agent does not start with All Capabilities, Artifacts and File Manager')


def first_agent_dialog(ctx):
    from playwright.sync_api import expect
    inst, browser = ctx.inst, ctx.browser
    for width, height, action in ((1440, 900, 'close'), (390, 844, 'escape')):
        with browser.presentation(viewport={'width': width, 'height': height}):
            email = f'dialog-{width}-' + uuid.uuid4().hex[:8] + '@smoke.example'
            inst.request('POST', '/api/v1/users', {'email': email, 'password': inst.password,
                         'full_name': 'First agent dialog member', 'groups': ['app.security:members']},
                         expect=(201,), label=f'dialog-member-{width}')
            token = inst.request('POST', '/api/public/user/token', {'email': email, 'password': inst.password})['token']
            inst.redactor.add(token)
            browser.login(email, inst.password)
            dialog = browser.frame.locator('dialog.sp-dialog[open]')
            expect(dialog).to_be_visible(timeout=30000)
            expect(dialog.locator('kickside-agent-onboarding').get_by_label('Agent name', exact=True)).to_be_visible(timeout=30000)
            expect(dialog).to_have_attribute('aria-modal', 'true')
            expect(dialog.locator('#sp-title')).to_have_text('Create your first agent')
            expect(browser.frame.locator('.sp-screen')).to_have_count(0)
            expect(dialog.locator('.sp-aside, .sp-eyebrow')).to_have_count(0)
            require((inst.request('GET', '/api/v1/nav/list', auth=token, label=f'dialog-nav-{width}').get('splash') or {}).get('presentation') == 'dialog',
                    'Member onboarding does not declare dialog presentation')
            inst.evidence.save(f'dialog-layout-{width}', dialog.evaluate("""el => {
                const out=[];
                function walk(node, depth=0) {
                    if (depth > 9) return;
                    if (node instanceof Element) {
                        const css=getComputedStyle(node), rect=node.getBoundingClientRect();
                        out.push({tag:node.tagName, class:node.className, width:rect.width, height:rect.height,
                            cssHeight:css.height, minHeight:css.minHeight, maxHeight:css.maxHeight,
                            display:css.display, position:css.position, inset:css.inset, flex:css.flex});
                        if (node.shadowRoot) walk(node.shadowRoot, depth+1);
                    }
                    for(const child of node.children || []) walk(child, depth+1);
                }
                walk(el); return out;
            }"""))
            bounds = dialog.bounding_box()
            require(bounds and bounds['width'] <= 560 and bounds['height'] < height - 32,
                    'Dialog does not size to its content: ' + str(bounds))
            require(abs(bounds['x'] + bounds['width'] / 2 - width / 2) <= 1,
                    'Dialog is not horizontally centered')
            require(abs(bounds['y'] + bounds['height'] / 2 - height / 2) <= 1,
                    'Dialog is not vertically centered')
            browser.screenshot(f'first-agent-dialog-{width}x{height}', full_page=False, mask_color='#e5e7eb')
            dialog.get_by_role('button', name='Close onboarding', exact=True).focus()
            for key in ['Shift+Tab'] + ['Tab'] * 12:
                browser.page.keyboard.press(key)
                require(dialog.evaluate("el => {let active=el.getRootNode().activeElement; while(active?.shadowRoot?.activeElement) active=active.shadowRoot.activeElement; for(let node=active;node;node=node.parentNode||node.host) if(node===el) return true; return false}"),
                        'Keyboard focus escapes the modal')
            with browser.page.expect_response(lambda r: '/api/v1/ui/splash/first_agent/seen' in r.url and r.request.method == 'POST') as response:
                if action == 'close':
                    dialog.get_by_role('button', name='Close onboarding', exact=True).click()
                else:
                    dialog.get_by_role('button', name='Close onboarding', exact=True).focus()
                    browser.page.keyboard.press('Escape')
            require(response.value.status == 200 and response.value.json().get('success') is True,
                    'Dialog dismissal does not persist')
            expect(browser.frame.locator('dialog.sp-dialog')).to_have_count(0, timeout=30000)
            require(not inst.request('GET', '/api/v1/user-agents', auth=token, label=f'dialog-no-agent-{width}').get('agents'),
                    'Dismissing onboarding creates an agent')
            require((inst.request('GET', '/api/v1/nav/list', auth=token, label=f'dialog-seen-{width}').get('splash') or {}).get('id') != 'first_agent',
                    'Dismissed dialog resolves again')
            browser.screenshot(f'first-agent-after-close-{width}x{height}', full_page=False, mask_color='#e5e7eb')
            browser.reload()
            browser.frame.locator('.managed-app-shell').wait_for(timeout=30000)
            browser.settle()
            expect(browser.frame.locator('dialog.sp-dialog, kickside-agent-onboarding')).to_have_count(0)
            browser.screenshot(f'first-agent-after-reload-{width}x{height}', full_page=False, mask_color='#e5e7eb')
            nav(ctx, 'Agents')
            agents = browser.frame.locator('kickside-agents')
            agents.get_by_role('button', name='New agent', exact=True).first.click()
            reopened = agents.get_by_role('dialog', name='New agent', exact=True)
            expect(reopened.get_by_label('Agent name', exact=True)).to_be_visible(timeout=30000)
            browser.screenshot(f'first-agent-reopened-{width}x{height}', full_page=False)
            reopened.get_by_role('button', name='Cancel', exact=True).click()
    browser.page.set_viewport_size({'width': 1440, 'height': 900})
    browser.login()


def chat_streams(ctx):
    browser, inst = ctx.browser, ctx.inst
    browser.stream('Give ten numbered fictional Northwind office tasks, two sentences each. End with SMOKE-CHAT-READY.', 'SMOKE-CHAT-READY', 'first-reply')
    sessions = inst.request('GET', '/api/v1/sessions?limit=100', label='sessions')['sessions']
    require(len(sessions) == 1, 'First chat creates more than one session: ' + str(sessions))
    ctx.state['session'] = sessions[0]['session_id']
    messages = inst.request('GET', '/api/v1/sessions/messages?session_id=' + ctx.state['session'], label='persisted-messages')['messages']
    require(sum(m['type'] == 'user' for m in messages) == 1, 'First prompt sends more than one user message')
    require(any(m['type'] == 'assistant' and 'SMOKE-CHAT-READY' in str(m.get('data')) for m in messages), 'Streamed reply does not persist')


def pdf_fixture(inst, label='northwind-policy'):
    from reportlab.pdfgen.canvas import Canvas
    path = inst.root / 'temp' / (label + '.pdf')
    canvas = Canvas(str(path))
    canvas.drawString(60, 750, 'Fictional Northwind Policy')
    canvas.drawString(60, 720, 'The document marker is NW-SMOKE-731. The return window is 17 days.')
    canvas.save()
    return path


def upload_attach_send(ctx):
    from playwright.sync_api import expect
    inst, browser = ctx.inst, ctx.browser
    inst.install('chestor/workspace')
    from upload_chat_contract import upload_in_chat
    result = upload_in_chat(ctx, browser.frame.locator('wippy-chat'), pdf_fixture(inst))
    uid = result['uuid']
    def summarized():
        listing = inst.request('GET', '/api/v1/uploads/list?limit=100', label='upload-summary-owner')
        component_id = listing['component_id']
        entry = inst.request('GET', '/api/v1/workspace/entries/' + component_id, label='upload-summary')['entry']
        meta = entry['meta']
        return meta if (meta.get('last_upload_id') == uid and meta.get('status') == 'ready'
                        and int(meta.get('file_count', 0)) >= 1 and int(meta.get('pending_count', -1)) == 0
                        and int(meta.get('total_bytes', 0)) > 0 and int(meta.get('last_event_seq', 0)) > 0) else None
    summary = inst.wait(summarized, 'Committed upload reaches the public summary', seconds=90,
                        pump=browser.page.wait_for_timeout)
    inst.evidence.save('upload-public-summary', summary)
    inst.wait(lambda: any(frame.get('data', {}).get('uuid') == uid and frame.get('data', {}).get('status') == 'completed'
                        for frame in browser.frames if isinstance(frame.get('data'), dict)),
              'Owner upload notification', seconds=90, pump=browser.page.wait_for_timeout)
    expect(browser.frame.locator('wippy-chat')).not_to_contain_text('Processing...', timeout=90000)
    browser.screenshot('attachment-ready')
    browser.stream('Read the attached fictional policy PDF. Quote the exact document marker and give five short bullet points about the policy. End with ATTACHMENT-READY.',
                   'ATTACHMENT-READY', 'attachment-reply')
    messages = inst.request('GET', '/api/v1/sessions/messages?session_id=' + ctx.state['session'], label='attachment-messages')['messages']
    sent = [m for m in messages if m['type'] == 'user' and 'ATTACHMENT-READY' in str(m.get('data'))]
    require(sent and uid in sent[-1].get('file_uuids', []), 'Sent message does not persist attachment UUID')
    require(any(m['type'] == 'assistant' and 'NW-SMOKE-731' in str(m.get('data')) for m in messages), 'Assistant does not quote PDF marker')
    ctx.state['upload'] = uid


def kb_cited_answer(ctx):
    inst = ctx.inst
    created = inst.request('POST', '/api/v1/knowledge/create', {'name': 'Northwind smoke policy', 'kb_type': 'document',
                           'embedding_model': ctx.state['embed'], 'description': 'Fictional smoke policy'}, label='kb-create')
    kid = created['id']
    ctx.state['kb'] = kid
    base = '/api/v1/knowledge/' + kid
    inst.request('PUT', base, {'config': {'analysis_model': ctx.state['model']}}, label='kb-model')
    inst.request('POST', base + '/ingest', {'title': 'Northwind Returns',
                 'content': 'Fictional Northwind return window is 17 days. The stock marker is NW-SMOKE-731.',
                 'external_id': 'northwind-smoke-policy', 'external_version': '1'}, label='kb-ingest')
    inst.wait(lambda: (lambda result: result if 'NW-SMOKE-731' in json.dumps(result) else None)(
        inst.request('POST', base + '/search', {'query': 'Northwind return window NW-SMOKE-731', 'limit': 5}, label='kb-search')),
        'KB indexed marker', seconds=120)
    answer = inst.request('POST', base + '/ask', {'query': 'What is the Northwind return window and stock marker? Cite the source.'}, label='kb-answer', timeout=120)
    require(answer.get('answer') and '17' in str(answer['answer']) and 'NW-SMOKE-731' in str(answer['answer']), 'KB answer lacks grounded policy facts: ' + str(answer))
    require(answer.get('citations') or answer.get('sources') or re.search(r'\[\d+\]|northwind-smoke-policy', str(answer['answer']), re.I),
            'KB answer lacks citations: ' + str(answer))
    ctx.browser.frame.locator('.sidebar').get_by_text('Knowledge', exact=True).click()
    ctx.browser.frame.get_by_text('Northwind smoke policy', exact=True).first.click()


def crm_record(ctx):
    inst = ctx.inst
    inst.install('kickside/crm')
    ctx.browser.reload()
    created = inst.request('POST', '/api/v1/workspace/crm', {'title': 'Northwind smoke CRM', 'idempotency_key': uuid.uuid4().hex}, label='crm-create', expect=(201,))
    cid = created['component_id']
    ctx.state['crm'] = cid
    base = '/api/v1/workspace/crm/' + cid
    current = inst.request('GET', base + '/schema', label='crm-schema')
    declaration = {'version': '5', 'objects': [{'type': 'person', 'label': 'People'}],
                   'attributes': [{'object_type': 'person', 'attr': 'name', 'label': 'Name', 'data_type': 'text'}],
                   'pipelines': [], 'views': [], 'merge_config': {'objects': {}}}
    inst.request('POST', base + '/schema', {'declaration': declaration, 'expected_revision': current['revision']}, label='crm-schema-apply')
    rid = uuid.uuid4().hex
    ctx.state['record'] = rid
    inst.request('POST', base + '/apply', {'commands': [{'type': 'kickside.crm.events:record.created',
                 'body': {'record_id': rid, 'object_type': 'person', 'values': {'name': 'Northwind Demo Contact'}}}]}, label='crm-record-create')
    inst.wait(lambda: (lambda result: result if 'Northwind Demo Contact' in json.dumps(result) else None)(
              inst.request('POST', base + '/query', {'kind': 'record', 'record_id': rid}, label='crm-record-read')),
              'CRM record projection', seconds=60)
    ctx.browser.frame.locator('.sidebar').get_by_text('CRM', exact=True).click()


def restart_continuation(ctx):
    inst, browser = ctx.inst, ctx.browser
    sid = ctx.state['session']
    browser = ctx.restart(hard=True)
    browser.page.goto(inst.base + '/', wait_until='domcontentloaded', timeout=30000)
    browser.frame.locator('.sidebar').wait_for(timeout=30000)
    browser.open_session(sid)
    browser.stream('Give ten short numbered tips about fictional office work. End with SMOKE-RESTART-READY.', 'SMOKE-RESTART-READY', 'restart-reply')
    sessions = inst.request('GET', '/api/v1/sessions?limit=100', label='restart-sessions')['sessions']
    require(len(sessions) == 1 and sessions[0]['session_id'] == sid, 'Restart creates or switches the durable session')
    messages = inst.request('GET', '/api/v1/sessions/messages?session_id=' + sid, label='restart-messages')['messages']
    require(any(m['type'] == 'assistant' and 'SMOKE-CHAT-READY' in str(m.get('data')) for m in messages), 'Restart loses first reply')
    require(any(m['type'] == 'assistant' and 'SMOKE-RESTART-READY' in str(m.get('data')) for m in messages), 'Restart reply does not persist')


def nav(ctx, label):
    ctx.browser.settle()
    toggle = ctx.browser.frame.get_by_role('button', name='Open navigation', exact=True)
    if toggle.is_visible():
        toggle.click()
    navigation = ctx.browser.frame.get_by_role('navigation', name='Main navigation', exact=True)
    target = navigation.get_by_role('button', name=label, exact=True, include_hidden=True).or_(
        navigation.get_by_role('button', include_hidden=True).filter(
            has=ctx.browser.frame.get_by_text(label, exact=True)))
    target.wait_for(state='attached', timeout=30000)
    if not target.is_visible():
        catalog = ctx.inst.request('GET', '/api/v1/nav/list', label='navigation-categories')
        categories = [category for category in catalog['categories']
                      if any(item['label'] == label for item in category['items'])]
        require(len(categories) == 1 and categories[0].get('label'),
                'Hidden navigation surface has no unique category: ' + label)
        navigation.get_by_role('button', name=categories[0]['label'], exact=True).click()
    target.click()


def workflow(ctx, review=False):
    inst = ctx.inst
    inst.install('kickside/workflows')
    wid = inst.request('POST', '/api/v1/workflows', {'title': 'Northwind smoke ' + ('review' if review else 'flow')}, label='workflow-create')['flow_id']
    schema = {'type': 'object', 'additionalProperties': True}
    nodes = [{'type': 'node.add', 'body': {'id': 'start', 'kind': 'start', 'config': {'input_schema': schema}}},
             {'type': 'node.add', 'body': {'id': 'finish', 'kind': 'finish', 'config': {'output_schema': schema}}}]
    def edge(name, source, target, port, target_port='in'):
        return {'type': 'edge.add', 'body': {'id': name, 'source_id': source, 'target_id': target,
                                           'source_port': port, 'target_port': target_port}}
    if review:
        nodes.append({'type': 'node.add', 'body': {'id': 'review', 'kind': 'block', 'ref': 'kickside.inbox:block.review',
                     'config': {'prompt': 'Approve fictional Northwind stock.', 'form': {'type': 'object', 'required': ['reason'],
                                'properties': {'reason': {'type': 'string'}}}}}})
        nodes += [edge('start-review', 'start', 'review', 'out'), edge('approve-finish', 'review', 'finish', 'then'),
                  edge('reject-finish', 'review', 'finish', 'else')]
    else:
        nodes.append(edge('start-finish', 'start', 'finish', 'out'))
    inst.request('POST', f'/api/v1/workflows/{wid}/apply', {'commands': nodes}, label='workflow-apply')
    version = inst.request('POST', f'/api/v1/workflows/{wid}/publish', {}, label='workflow-version')['version']
    return wid, version


def workflow_inbox_restart(ctx):
    inst, browser = ctx.inst, ctx.browser
    wid, version = workflow(ctx, review=True)
    runs = []
    for decision in ('approve', 'reject'):
        run = inst.request('POST', f'/api/v1/workflows/{wid}/run', {'version': version, 'input': {'marker': decision}}, label=decision + '-run')
        def waiting():
            detail = inst.request('GET', f'/api/v1/workflows/{wid}/runs/{run["run_id"]}', label=decision + '-waiting')
            require(detail['run']['status'] not in ('failed', 'completed'), 'Review terminates before response: ' + str(detail))
            return detail if any(node.get('status') == 'waiting' for node in detail['nodes']) else None
        inst.wait(waiting, 'Durable review wait', seconds=60)
        def item():
            matches = [i for i in inst.request('GET', '/api/v1/inbox?status=all&limit=200')['items']
                       if i.get('metadata', {}).get('run_id') == run['run_id']]
            require(len(matches) <= 1, 'Duplicate review items: ' + str(matches))
            return matches[0] if matches else None
        pending = inst.wait(item, 'Review inbox item', seconds=35)
        require(pending['status'] == 'pending', 'Review is not pending')
        runs.append({'decision': decision, 'run': run, 'item': pending})
    nav(ctx, 'Inbox')
    browser.screenshot('reviews-before-restart')
    browser.settle()
    browser.page.goto('about:blank')
    browser = ctx.restart(hard=True)
    browser.frame.locator('.sidebar').wait_for(timeout=30000)
    nav(ctx, 'Inbox')
    browser.screenshot('reviews-after-restart')
    for case in runs:
        rid, item = case['run']['run_id'], case['item']
        matches = [i for i in inst.request('GET', '/api/v1/inbox?status=all&limit=200')['items'] if i['id'] == item['id']]
        require(len(matches) == 1 and matches[0]['status'] == 'pending' and matches[0]['metadata'] == item['metadata'],
                'Review identity/correlation changes after restart')
        path = '/api/v1/inbox/' + quote(item['id'], safe='') + '/' + case['decision']
        body = {'form': {'reason': 'Northwind after restart ' + case['decision']}}
        inst.request('POST', path, body, label=case['decision'] + '-response')
        def terminal():
            result = inst.request('GET', f'/api/v1/workflows/{wid}/runs/{rid}', label=case['decision'] + '-terminal')
            require(result['run']['status'] != 'failed', 'Review workflow fails: ' + str(result))
            return result if result['run']['status'] == 'completed' else None
        completed = inst.wait(terminal, 'Review completion', seconds=60)
        require(case['decision'] in json.dumps(completed['run']['output']), 'Review decision does not reach Finish output')
        repeated = inst.request('POST', path, body, label=case['decision'] + '-repeat', expect=(409,))
        require(repeated.get('success') is False and 'resolved' in str(repeated).lower(), 'Repeated review lacks explicit conflict')
    browser.settle()
    browser.page.goto('about:blank')
    browser = ctx.restart(hard=True)
    browser.frame.locator('.sidebar').wait_for(timeout=30000)
    final = inst.request('GET', f'/api/v1/workflows/{wid}/runs', label='second-restart-runs')['runs']
    require(len(final) == 2 and {r['run_id'] for r in final} == {r['run']['run_id'] for r in runs}
            and all(r['status'] == 'completed' for r in final), 'Completed reviews duplicate or change after second restart')
    nav(ctx, 'Inbox')


def automations(ctx):
    inst = ctx.inst
    created = inst.request('POST', '/api/v1/automations', {'type': 'kickside.automation.recurring:task', 'input': {
        'title': 'Northwind scheduled notice', 'schedule': {'type': 'interval', 'expression': '2s'}, 'enabled': True,
        'target': {'kind': 'action', 'port': 'kickside.inbox.sink:inbox_item', 'config': {},
                   'input': {'content': 'Fictional Northwind scheduled stock check', 'title': 'Northwind scheduled notice'}}}}, label='recurring-create')
    aid = created['id']
    from cleanup_evidence import own_automation
    own_automation(inst, aid)
    ctx.state['automation'] = aid
    def notices():
        items = inst.request('GET', '/api/v1/inbox?status=all&limit=200', label='recurring-inbox')['items']
        found = [item for item in items if item.get('title') == 'Northwind scheduled notice']
        return found if len(found) >= 2 else None
    occurrences = inst.wait(notices, 'Two native scheduled occurrences', seconds=60)
    require(len({item['id'] for item in occurrences}) >= 2, 'Scheduled occurrences share one identity')
    for control in ('pause', 'resume'):
        inst.request('POST', f'/api/v1/automations/{aid}/control', {'control': control, 'request_id': uuid.uuid4().hex}, label='recurring-' + control)
    inst.request('GET', f'/api/v1/automations/{aid}/status', label='recurring-status')
    nav(ctx, 'Automations')


def data_sync(ctx):
    inst = ctx.inst
    cid = ctx.state['crm']
    body = {'type': 'kickside.sync:sync', 'input': {'title': 'Northwind CRM notice sync', 'enabled': True,
        'trigger': {'v': 2, 'source': 'kickside.crm.ports:record_created', 'config': {'thread_id': cid}, 'backfill': 'none'},
        'sink_port': 'kickside.inbox.sink:inbox_item', 'sink_config': {},
        'transform': {'kind': 'expr', 'expr': '{"content":"Fictional Northwind CRM record notice","title":"Northwind sync receipt"}'}}}
    aid = inst.request('POST', '/api/v1/automations', body, label='sync-create')['id']
    rid = uuid.uuid4().hex
    inst.request('POST', f'/api/v1/workspace/crm/{cid}/apply', {'commands': [{'type': 'kickside.crm.events:record.created',
                 'body': {'record_id': rid, 'object_type': 'person', 'values': {'name': 'Northwind Sync Contact'}}}]}, label='sync-trigger')
    def delivered():
        found = [item for item in inst.request('GET', '/api/v1/inbox?status=all&limit=200', label='sync-inbox')['items']
                 if item.get('title') == 'Northwind sync receipt']
        require(len(found) <= 1, 'Data Sync duplicates delivery: ' + str(found))
        return found if found else None
    inst.wait(delivered, 'CRM event to inbox sync', seconds=60)
    inst.request('GET', f'/api/v1/automations/{aid}/status', label='sync-status')
    nav(ctx, 'Automations')


def webhooks(ctx):
    inst = ctx.inst
    inst.install('kickside/webhooks')
    ctx.browser.reload()
    body = {'label': 'Northwind smoke webhook', 'schema': {'type': 'object', 'required': ['product'],
            'properties': {'product': {'type': 'string'}}}}
    webhook = inst.request('POST', '/api/v1/webhooks', body, label='webhook-create')['webhook']
    wid, token = webhook['webhook_id'], webhook['secret']
    inst.redactor.add(token)
    path = '/api/public/hooks/' + wid
    inst.request('POST', path, {'product': 'Northwind Tea'}, auth='', expect=(401,), label='webhook-unauthorized')
    for number in range(2):
        inst.request('POST', path, {'product': 'Northwind Tea'}, auth=token, expect=(202,), label='webhook-delivery-' + str(number))
    inst.request('POST', path, {'product': 42}, auth=token, expect=(400,), label='webhook-invalid-schema')
    inst.request('POST', '/api/v1/webhooks/' + wid + '/revoke', {}, label='webhook-revoke')
    inst.request('POST', path, {'product': 'Northwind Tea'}, auth=token, expect=(401,403,404), label='webhook-revoked')
    inst.request('DELETE', '/api/v1/webhooks/' + wid, label='webhook-delete')
    nav(ctx, 'System')
    ctx.browser.frame.get_by_role('tab', name='Webhooks', exact=True).click()
    ctx.browser.frame.locator('kickside-webhooks').wait_for(timeout=30000)


def assert_scoped_tools(tools):
    controls = {'session_info', 'list_traits', 'find_traits', 'describe_trait', 'use_trait', 'drop_trait',
                'set_traits', 'reset_traits', 'use_profile', 'list_tools', 'find_tools', 'describe_tool', 'call_tool'}
    names = {tool['name'] for tool in tools}
    require('QueryKB' in names, 'Allowed QueryKB tool is absent')
    require(names <= controls | {'QueryKB'}, 'MCP exposes product tools outside the allowed list: ' + str(sorted(names - controls - {'QueryKB'})))


def mcp_scoped_credential(ctx):
    inst = ctx.inst
    inst.install('kickside/mcp')
    kid = ctx.state['kb']
    other = inst.request('POST', '/api/v1/knowledge/create', {'name': 'Northwind outside scope', 'kb_type': 'document',
                         'embedding_model': ctx.state['embed']}, label='outside-kb')['id']
    credential = inst.request('POST', '/api/v1/mcp/credentials', {'label': 'Northwind scoped smoke client',
        'whitelist_component_ids': [kid], 'allowed_tool_ids': ['kickside.knowledge.traits:kb_query_tool']}, label='mcp-mint')['credential']
    token, cid = credential['token'], credential['credential_id']
    inst.redactor.add(token)
    def rpc(method, params, label, expect=(200,)):
        return inst.request('POST', '/api/public/mcp', {'jsonrpc': '2.0', 'id': uuid.uuid4().hex, 'method': method, 'params': params},
                            auth=token, headers={'Accept': 'application/json, text/event-stream'}, label=label, expect=expect, timeout=120)
    try:
        initialized = rpc('initialize', {'protocolVersion': '2025-03-26', 'capabilities': {},
                          'clientInfo': {'name': 'Northwind Smoke', 'version': '1'}}, 'mcp-initialize')
        require('result' in initialized and 'error' not in initialized, 'MCP initialization fails: ' + str(initialized))
        result = rpc('tools/list', {}, 'mcp-tools')
        tools = result.get('result', {}).get('tools', [])
        assert_scoped_tools(tools)
        scoped = rpc('tools/call', {'name': 'QueryKB', 'arguments': {'kb_id': kid, 'query': 'Northwind return window'}}, 'mcp-in-scope')
        require('error' not in scoped and not scoped.get('result', {}).get('isError'), 'In-scope MCP call fails: ' + str(scoped))
        denied = rpc('tools/call', {'name': 'QueryKB', 'arguments': {'kb_id': other, 'query': 'Northwind return window'}}, 'mcp-out-of-scope')
        require(denied.get('error') or denied.get('result', {}).get('isError'), 'MCP permits a component outside the whitelist')
        for name, arguments in [('CRMCreateRecord', {}), ('call_tool', {'id': 'CRMCreateRecord', 'arguments': {}})]:
            denied_tool = rpc('tools/call', {'name': name, 'arguments': arguments}, 'mcp-denied-' + name)
            require(denied_tool.get('error') or denied_tool.get('result', {}).get('isError'), 'MCP permits a product tool outside policy through ' + name)
    finally:
        inst.request('POST', '/api/v1/mcp/credentials/' + str(cid) + '/revoke', {}, label='mcp-revoke')
    rpc('tools/list', {}, 'mcp-after-revoke', expect=(401,403))


def research(ctx):
    inst = ctx.inst
    inst.install('kickside/research')
    ctx.browser.reload()
    created = inst.request('POST', '/api/v1/research', {'prompt': 'Analyze fictional Northwind stock: 47 crates and 12 reserve crates. Use only these facts.',
                 'agent_id': 'kickside.research.agents:researcher'}, label='research-create')
    require(created.get('research_id'), 'Research start returns no research ID: ' + str(created))
    rid = created['research_id']
    def completed():
        detail = inst.request('GET', '/api/v1/research/' + rid + '/events', label='research-events')
        require(detail['run'].get('status') != 'failed', 'Research run fails: ' + str(detail))
        return detail if detail['run'].get('status') == 'complete' else None
    detail = inst.wait(completed, 'Research completion', seconds=180)
    results = [event['body'].get('markdown') for event in detail['events']
               if event['type'] == 'kickside.research.events:result']
    report = detail['run'].get('result_md')
    require(isinstance(report, str) and report in results,
            'Research projected report differs from its durable result event: ' + str(detail))
    require(all(fact in report for fact in ('47', '12')),
            'Research report omits supplied Northwind stock facts: ' + str(detail))
    listed = inst.request('GET', '/api/v1/research/list', label='research-list')
    require(rid in json.dumps(listed), 'Research list does not contain its run')
    nav(ctx, 'Research')
    title = detail['run'].get('title')
    require(isinstance(title, str) and title, 'Research run summary omits its title: ' + str(detail))
    report = ctx.browser.frame.locator('research-app .kx-row')
    from playwright.sync_api import expect
    expect(report).to_have_count(1, timeout=30000)
    name = report.locator('.kx-name')
    expect(name).to_have_attribute('title', title, timeout=30000)
    expect(name).to_have_text(title, timeout=30000)
    report.locator('.kx-row-head').click()
    rendered = report.locator('wc-render-markdown')
    expect(rendered).to_contain_text('47', timeout=30000)
    expect(rendered).to_contain_text('12', timeout=30000)


def all_checks(inst, checks):
    failures = []
    for label, check in checks:
        try:
            check()
        except Exception as error:
            failures.append(label + ': ' + type(error).__name__ + ': ' + str(error))
    inst.evidence.save('checks', {'errors': failures})
    require(not failures, '\n'.join(failures))


def chestor_entries(ctx):
    inst = ctx.inst
    for component in ('chestor/workspace', 'chestor/graph', 'chestor/journal', 'chestor/roadmap', 'chestor/db-schema'):
        inst.install(component)
    ctx.browser.reload()
    folder = inst.request('POST', '/api/v1/workspace/entries', {'type': 'folder', 'name': 'Northwind smoke folder'}, label='folder-create')['entry']['id']
    ctx.state['folder'] = folder
    def create(kind, placement):
        body = {'type': kind, 'name': 'Northwind ' + kind + ' ' + placement}
        if placement == 'folder':
            body['parent_id'] = folder
        entry = inst.request('POST', '/api/v1/workspace/entries', body, label=kind + '-' + placement)['entry']
        require(entry.get('id'), 'Workspace create returns no entry ID')
        ctx.state[kind + '_' + placement] = entry['id']
        inst.request('GET', '/api/v1/workspace/entries/' + entry['id'], label=kind + '-' + placement + '-read')
        crumbs = inst.request('GET', '/api/v1/workspace/entries/' + entry['id'] + '/breadcrumbs', label=kind + '-' + placement + '-breadcrumbs')
        if placement == 'folder':
            require(folder in json.dumps(crumbs), 'Folder entry loses its parent breadcrumb')
    all_checks(inst, [(kind + '-' + placement, lambda k=kind, p=placement: create(k, p))
                     for kind in ('graph', 'journal', 'roadmap', 'database') for placement in ('root', 'folder')])
    nav(ctx, 'Workspace')
    ctx.browser.frame.get_by_text('Northwind smoke folder', exact=True).first.dblclick()
    ctx.browser.frame.get_by_text('Northwind graph folder', exact=True).first.wait_for(timeout=30000)


def dashboard(ctx):
    overview = ctx.inst.request('GET', '/api/v1/dashboard/overview', label='dashboard')
    require(overview.get('count') == len(overview.get('surfaces', [])), 'Dashboard surface count differs')
    nav(ctx, 'Dashboard')
    ctx.browser.frame.locator('kickside-dashboard').wait_for(timeout=30000)


def member_access(ctx):
    inst = ctx.inst
    inst.install('chestor/workspace')
    email = 'member-' + uuid.uuid4().hex[:8] + '@smoke.example'
    created = inst.request('POST', '/api/v1/users', {'email': email, 'password': inst.password,
                 'full_name': 'Northwind Member', 'groups': ['app.security:members']}, expect=(201,), label='member-create')
    token = inst.request('POST', '/api/public/user/token', {'email': email, 'password': inst.password})['token']
    inst.redactor.add(token)
    folder = inst.request('POST', '/api/v1/workspace/entries', {'type': 'folder', 'name': 'Northwind private boundary'}, label='private-folder')['entry']['id']
    checks = []
    for label, method, path, body in [
        ('settings', 'GET', '/api/v1/settings?namespace=kickside.ui&key=completed_splashes', None),
        ('splash-reset', 'DELETE', '/api/v1/ui/splash/seen', None),
        ('private-agent', 'GET', '/api/v1/user-agents/' + ctx.state['agent'], None),
        ('private-kb', 'GET', '/api/v1/knowledge/' + ctx.state['kb'], None),
        ('private-kb-ingest', 'POST', '/api/v1/knowledge/' + ctx.state['kb'] + '/ingest', {'content': 'Unauthorized fictional edit'}),
        ('private-folder', 'GET', '/api/v1/workspace/entries/' + folder, None),
        ('private-folder-write', 'PUT', '/api/v1/workspace/entries/' + folder, {'name': 'Unauthorized Northwind edit'}),
        ('private-crm-write', 'POST', '/api/v1/workspace/crm/' + ctx.state['crm'] + '/apply', {'commands': []})]:
        checks.append((label, lambda m=method, p=path, b=body, l=label: inst.request(m, p, b, auth=token, expect=(403,404), label='member-' + l)))
    all_checks(inst, checks)
    listed = inst.request('GET', '/api/v1/workspace/entries', auth=token, label='member-workspace-filtered')
    require(folder not in json.dumps(listed), 'Member list reveals a private folder')
    browser = ctx.browser
    try:
        browser.login(email, inst.password)
        splash = browser.frame.locator('kickside-agent-onboarding')
        splash.get_by_label('Agent name', exact=True).wait_for(timeout=30000)
        nav_result = inst.request('GET', '/api/v1/nav/list?manage=1', auth=token, label='member-nav')
        require('completed_splashes' not in json.dumps(nav_result), 'Member sees instance completion settings')
        browser.screenshot('member-first-agent')
    finally:
        browser.login()
    browser.frame.locator('.sidebar').wait_for(timeout=30000)


def provider_dialogs(ctx):
    inst, browser = ctx.inst, ctx.browser
    catalog = inst.request('GET', '/api/v1/connections/providers/discover', label='provider-catalog')['providers']
    require(catalog, 'Provider catalog is empty')
    nav(ctx, 'Connections')
    browser.frame.locator('kickside-connections').wait_for(timeout=30000)
    browser.screenshot('provider-picker')
    provider = next(p for p in catalog if p.get('credential_schema', {}).get('fields'))
    browser.frame.locator('.setup-card').filter(has_text=provider['title']).click()
    browser.frame.get_by_role('dialog').wait_for(timeout=30000)
    browser.frame.get_by_role('dialog').get_by_label('Connection name', exact=True).wait_for(timeout=30000)
    browser.screenshot('provider-credential-form')
    browser.page.keyboard.press('Escape')


def converters(ctx):
    inst = ctx.inst
    fixtures = [pdf_fixture(inst, 'northwind-converter')]
    for suffix, text in [('txt', 'Fictional Northwind converter marker NW-SMOKE-731'),
                         ('md', '# Northwind\nFictional marker NW-SMOKE-731'),
                         ('html', '<html><body><p>Fictional Northwind marker NW-SMOKE-731</p></body></html>')]:
        path = inst.root / 'temp' / ('northwind-converter.' + suffix)
        path.write_text(text)
        fixtures.append(path)
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image, ImageDraw
    doc = Document()
    doc.add_paragraph('Fictional Northwind marker NW-SMOKE-731')
    path = inst.root / 'temp' / 'northwind-converter.docx'
    doc.save(path)
    fixtures.append(path)
    book = Workbook()
    book.active.append(['Northwind marker', 'NW-SMOKE-731'])
    path = inst.root / 'temp' / 'northwind-converter.xlsx'
    book.save(path)
    fixtures.append(path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = 'Fictional Northwind NW-SMOKE-731'
    path = inst.root / 'temp' / 'northwind-converter.pptx'
    presentation.save(path)
    fixtures.append(path)
    image = Image.new('RGB', (640, 200), 'white')
    ImageDraw.Draw(image).text((20, 50), 'Fictional Northwind NW-SMOKE-731', fill='black')
    path = inst.root / 'temp' / 'northwind-converter.png'
    image.save(path)
    fixtures.append(path)
    def convert(path):
        uid = inst.upload(path, 'upload-' + path.suffix[1:])['uuid']
        def ready():
            upload = inst.request('GET', '/api/v1/uploads/' + uid, label='conversion-' + path.suffix[1:])['upload']
            require(upload['status'] not in ('error', 'failed'), path.suffix + ' conversion fails: ' + str(upload))
            return upload if upload['status'] == 'completed' else None
        inst.wait(ready, path.suffix + ' conversion', seconds=120)
        artifact = inst.request('GET', '/api/v1/uploads/' + uid + '/artifacts', label='artifacts-' + path.suffix[1:])
        require((artifact.get('primary') or {}).get('raw_url'), path.suffix + ' conversion has no primary artifact')
    all_checks(inst, [(path.suffix, lambda p=path: convert(p)) for path in fixtures])
    nav(ctx, 'Uploads')


# Inventory feature ids (coverage/inventory.json) each scenario exercises.
FEATURES = {
    'automations': (
        'kickside/automation|automation|kickside.automation.recurring:task',
        'kickside/automation|http|GET /api/v1/automations/{id}/status',
        'kickside/automation|http|POST /api/v1/automations',
        'kickside/automation|http|POST /api/v1/automations/{id}/control',
        'kickside/automation|navigation|kickside.automation:nav_item',
        'kickside/automation|workspace_kind|kickside.automation.recurring:task',
        'kickside/inbox|http|GET /api/v1/inbox',
        'kickside/inbox|inbox_kind|kickside.inbox:item_type.notice',
        'kickside/inbox|sink|kickside.inbox.sink:inbox_item',
    ),
    'boot': (
        'kickside/kickside|http|GET /login',
        'kickside/users|http|POST /api/public/user/token',
    ),
    'chat_streams': (
        'kickside/kickside|http|GET /api/v1/ws/join',
    ),
    'chestor_entries': (
        'chestor/db-schema|workspace_kind|chestor.db_schema:database',
        'chestor/graph|workspace_kind|chestor.graph:graph_kind',
        'chestor/journal|workspace_kind|chestor.journal:journal_kind',
        'chestor/roadmap|workspace_kind|chestor.roadmap:roadmap_kind',
        'chestor/workspace|http|GET /api/v1/workspace/entries/{id}',
        'chestor/workspace|http|GET /api/v1/workspace/entries/{id}/breadcrumbs',
        'chestor/workspace|http|POST /api/v1/workspace/entries',
        'chestor/workspace|navigation|chestor.workspace:nav_item',
        'chestor/workspace|workspace_kind|chestor.workspace:folder',
    ),
    'converters': (
        'kickside/doc2md|converter|kickside.convert.doc2md:converter',
        'kickside/uploads|http|GET /api/v1/uploads/{uuid}',
        'kickside/uploads|http|GET /api/v1/uploads/{uuid}/artifacts',
        'kickside/uploads|http|POST /api/v1/uploads',
        'kickside/uploads|navigation|kickside.uploads:nav_item',
        'kickside/uploads|upload_kind|kickside.uploads.documents:types.documents',
        'kickside/uploads|upload_kind|kickside.uploads:types.generic',
        'kickside/uploads|upload_kind|kickside.uploads:types.images',
        'kickside/uploads|upload_kind|kickside.uploads:types.text',
    ),
    'crm_record': (
        'kickside/crm|http|GET /api/v1/workspace/crm/{id}/schema',
        'kickside/crm|http|POST /api/v1/workspace/crm',
        'kickside/crm|http|POST /api/v1/workspace/crm/{id}/apply',
        'kickside/crm|http|POST /api/v1/workspace/crm/{id}/query',
        'kickside/crm|http|POST /api/v1/workspace/crm/{id}/schema',
        'kickside/crm|navigation|kickside.crm.surface:nav_item',
        'kickside/crm|workspace_kind|kickside.crm.binding:crm',
    ),
    'dashboard': (
        'kickside/dashboard|http|GET /api/v1/dashboard/overview',
        'kickside/dashboard|navigation|kickside.dashboard:nav_item',
    ),
    'data_sync': (
        'kickside/automation|http|GET /api/v1/automations/{id}/status',
        'kickside/automation|http|POST /api/v1/automations',
        'kickside/automation|navigation|kickside.automation:nav_item',
        'kickside/crm|http|POST /api/v1/workspace/crm/{id}/apply',
        'kickside/crm|trigger|kickside.crm.ports:record_created',
        'kickside/inbox|http|GET /api/v1/inbox',
        'kickside/inbox|inbox_kind|kickside.inbox:item_type.notice',
        'kickside/inbox|sink|kickside.inbox.sink:inbox_item',
        'kickside/sync|automation|kickside.sync:sync',
        'kickside/sync|workspace_kind|kickside.sync:sync',
        'kickside/transform|transform|kickside.transform:strategy_expr_binding',
    ),
    'first_agent': (
        'kickside/agent-onboarding|navigation|kickside.agent_onboarding:splash.first_agent',
        'kickside/agent-onboarding|navigation|kickside.agent_onboarding:splash.first_agent.step.create',
        'kickside/agents|http|GET /api/v1/agents/traits',
        'kickside/agents|http|GET /api/v1/user-agents/{id}',
        'kickside/agents|http|POST /api/v1/user-agents',
        'kickside/agents|workspace_kind|kickside.agents:user_agent',
    ),
    'kb_cited_answer': (
        'kickside/knowledge|http|POST /api/v1/knowledge/create',
        'kickside/knowledge|http|POST /api/v1/knowledge/{id}/ask',
        'kickside/knowledge|http|POST /api/v1/knowledge/{id}/ingest',
        'kickside/knowledge|http|POST /api/v1/knowledge/{id}/search',
        'kickside/knowledge|http|PUT /api/v1/knowledge/{id}',
        'kickside/knowledge|navigation|kickside.knowledge:nav_item',
    ),
    'mcp_scoped_credential': (
        'kickside/knowledge|http|POST /api/v1/knowledge/create',
        'kickside/knowledge|mcp_tool|kickside.knowledge.traits:kb_query_tool',
        'kickside/mcp|http|POST /api/public/mcp',
        'kickside/mcp|http|POST /api/v1/mcp/credentials',
        'kickside/mcp|http|POST /api/v1/mcp/credentials/{id}/revoke',
    ),
    'member_access': (
        'chestor/workspace|http|GET /api/v1/workspace/entries',
        'chestor/workspace|http|GET /api/v1/workspace/entries/{id}',
        'chestor/workspace|http|POST /api/v1/workspace/entries',
        'chestor/workspace|http|PUT /api/v1/workspace/entries/{id}',
        'kickside/agents|http|GET /api/v1/user-agents/{id}',
        'kickside/crm|http|POST /api/v1/workspace/crm/{id}/apply',
        'kickside/knowledge|http|GET /api/v1/knowledge/{id}',
        'kickside/knowledge|http|POST /api/v1/knowledge/{id}/ingest',
        'kickside/settings|http|GET /api/v1/settings',
        'kickside/ui|http|DELETE /api/v1/ui/splash/seen',
        'kickside/ui|http|GET /api/v1/nav/list',
        'kickside/users|http|POST /api/public/user/token',
        'kickside/users|http|POST /api/v1/users',
    ),
    'platform_onboarding': (
        'kickside/hub|navigation|kickside.hub:splash.welcome.step.modules',
        'kickside/kickside|navigation|app.views:main',
        'kickside/ui|http|GET /api/v1/nav/list',
        'kickside/ui|http|POST /api/v1/ui/splash/{id}/seen',
        'kickside/ui|navigation|kickside.ui:splash.welcome',
        'kickside/ui|navigation|kickside.ui:splash.welcome.step.done',
        'kickside/ui|navigation|kickside.ui:splash.welcome.step.logo',
        'kickside/ui|navigation|kickside.ui:splash.welcome.step.mcp',
        'kickside/ui|navigation|kickside.ui:splash.welcome.step.models',
        'kickside/ui|navigation|kickside.ui:splash.welcome.step.workspace',
        'kickside/users|http|POST /api/public/user/token',
    ),
    'provider_dialogs': (
        'kickside/connection|http|GET /api/v1/connections/providers/discover',
        'kickside/connection|navigation|kickside.connection:nav_item',
    ),
    'research': (
        'kickside/research|http|GET /api/v1/research/list',
        'kickside/research|http|GET /api/v1/research/{id}/events',
        'kickside/research|http|POST /api/v1/research',
        'kickside/research|navigation|kickside.research:nav_item',
        'kickside/research|workspace_kind|kickside.research:research_kind',
    ),
    'restart_continuation': (
        'kickside/kickside|http|GET /api/v1/ws/join',
        'kickside/sessions|navigation|kickside.sessions:nav_item',
    ),
    'upload_attach_send': (
        'chestor/workspace|http|GET /api/v1/workspace/entries/{id}',
        'kickside/uploads|http|GET /api/v1/uploads/list',
        'kickside/doc2md|converter|kickside.convert.doc2md:converter',
        'kickside/uploads|http|GET /api/v1/uploads/{uuid}',
        'kickside/uploads|http|POST /api/v1/uploads',
        'kickside/uploads|upload_kind|kickside.uploads.documents:types.documents',
    ),
    'webhooks': (
        'kickside/webhooks|http|DELETE /api/v1/webhooks/{id}',
        'kickside/webhooks|http|POST /api/public/hooks/{id}',
        'kickside/webhooks|http|POST /api/v1/webhooks',
        'kickside/webhooks|http|POST /api/v1/webhooks/{id}/revoke',
        'kickside/webhooks|navigation|kickside.webhooks:tab.webhooks',
        'kickside/webhooks|workspace_kind|kickside.webhooks:webhook_kind',
    ),
    'workflow_inbox_restart': (
        'kickside/inbox|http|GET /api/v1/inbox',
        'kickside/inbox|http|POST /api/v1/inbox/{id}/{action}',
        'kickside/inbox|inbox_kind|kickside.inbox:item_type.flow_decision',
        'kickside/inbox|navigation|kickside.inbox:nav_item',
        'kickside/inbox|workflow_block|kickside.inbox:block.review',
        'kickside/workflows|http|GET /api/v1/workflows/{id}/runs',
        'kickside/workflows|http|GET /api/v1/workflows/{id}/runs/{run_id}',
        'kickside/workflows|http|POST /api/v1/workflows',
        'kickside/workflows|http|POST /api/v1/workflows/{id}/apply',
        'kickside/workflows|http|POST /api/v1/workflows/{id}/publish',
        'kickside/workflows|http|POST /api/v1/workflows/{id}/run',
        'kickside/workflows|workspace_kind|kickside.workflows:workflow_kind',
    ),
}


REGISTRY = [
    Scenario('boot', 'quick', boot, issues=('R5-RELEASE-FRONTEND-ARTIFACTS', 'SMOKE-SOURCE-FRESH-RESOLUTION'), features=FEATURES['boot']),
    Scenario('platform_onboarding', 'quick', platform_onboarding, ('boot',), issues=('ONBOARDING-DEPENDENCY-INSTALL-CONFLICT', 'SMOKE-ONBOARDING-EXPLICIT-MODEL-LIMITS', 'SMOKE-ONBOARDING-INSTALL-BATCH-READINESS'), features=FEATURES['platform_onboarding'], instance_setup=True),
    Scenario('first_agent', 'quick', first_agent, ('platform_onboarding',), issues=('ONBOARD-FIRST-AGENT', 'R5-RELEASE-FRONTEND-ARTIFACTS', 'FE-SHELL-COMPONENT-LOADER-OWNERSHIP'), features=FEATURES['first_agent'], instance_setup=True, capture_navigation=('session',)),
    Scenario('first_agent_dialog', 'full', first_agent_dialog, ('first_agent',), issues=('ONBOARD-FIRST-AGENT',), features=FEATURES['first_agent'] +
             ('kickside/ui|http|POST /api/v1/ui/splash/{id}/seen', 'kickside/ui|http|GET /api/v1/nav/list')),
    Scenario('chat_streams', 'quick', chat_streams, ('first_agent',), True, issues=('T05', 'ONBOARD-FIRST-AGENT'), features=FEATURES['chat_streams'], navigation=('chat', 'agent', 'session')),
    Scenario('upload_attach_send', 'quick', upload_attach_send, ('chat_streams',), True, features=FEATURES['upload_attach_send'], issues=('EE2-2516', 'R4-P02', 'SMOKE-FRAME-REDACTION-STALL'), navigation=('chat', 'agent', 'session')),
    Scenario('kb_cited_answer', 'quick', kb_cited_answer, ('first_agent',), True, features=FEATURES['kb_cited_answer']),
    Scenario('crm_record', 'quick', crm_record, ('first_agent',), features=FEATURES['crm_record'], instance_setup=True),
    Scenario('restart_continuation', 'quick', restart_continuation, ('chat_streams',), True, features=FEATURES['restart_continuation'], navigation=('chat', 'agent', 'session')),
    Scenario('workflow_inbox_restart', 'full', workflow_inbox_restart, ('first_agent',), features=FEATURES['workflow_inbox_restart'], instance_setup=True),
    Scenario('automations', 'full', automations, ('first_agent',), features=FEATURES['automations']),
    Scenario('data_sync', 'full', data_sync, ('crm_record',), features=FEATURES['data_sync']),
    Scenario('webhooks', 'full', webhooks, ('first_agent',), features=FEATURES['webhooks'], instance_setup=True),
    Scenario('mcp_scoped_credential', 'full', mcp_scoped_credential, ('kb_cited_answer',), True, issues=('SEC-MCP-OWNER-CEILING', 'R4-I10'), features=FEATURES['mcp_scoped_credential'], instance_setup=True),
    Scenario('research', 'full', research, ('first_agent',), True, issues=('RESEARCH-DELIVERABLE-REQUEST-CONSTRAINTS', 'RESEARCH-COMPACT-SUMMARY-TITLE'), features=FEATURES['research'], instance_setup=True),
    Scenario('chestor_entries', 'full', chestor_entries, ('first_agent',), features=FEATURES['chestor_entries'], instance_setup=True),
    Scenario('dashboard', 'full', dashboard, ('first_agent',), features=FEATURES['dashboard']),
    Scenario('member_access', 'full', member_access, ('first_agent', 'kb_cited_answer', 'crm_record'), True, issues=('SEC-MEMBER-MODULE-POLICY-GROUPS', 'SECURITY-SCOPE-PROJECTOR-SURVIVAL', 'SECURITY-SCOPE-PROJECTION-LIVENESS'), features=FEATURES['member_access']),
    Scenario('provider_dialogs', 'full', provider_dialogs, ('first_agent',), features=FEATURES['provider_dialogs']),
    Scenario('converters', 'full', converters, ('first_agent',), True, features=FEATURES['converters']),
]


def member_identity(ctx):
    if 'endpoint_member' not in ctx.state:
        email = 'endpoint-member-' + uuid.uuid4().hex[:8] + '@smoke.example'
        created = ctx.inst.request('POST', '/api/v1/users', {'email': email, 'password': ctx.inst.password,
            'full_name': 'Northwind Endpoint Member', 'groups': ['app.security:members']}, expect=(201,), label='endpoint-member-create')
        token = ctx.inst.request('POST', '/api/public/user/token', {'email': email, 'password': ctx.inst.password}, auth='', label='endpoint-member-login')['token']
        ctx.inst.redactor.add(token)
        ctx.state['endpoint_member'] = {'id': created['user']['user_id'], 'email': email, 'token': token}
    return ctx.state['endpoint_member']


def refresh_member_token(inst, member, path, admitted, label, seconds=60):
    """Mints member tokens through the login issuer until `path` admits one; access granted after a token is minted arrives with the next token."""
    def refreshed():
        token = inst.request('POST', '/api/public/user/token', {'email': member['email'], 'password': inst.password},
                             auth='', label=label + '-login')['token']
        inst.redactor.add(token)
        response = inst.request('GET', path, auth=token, expect=(200, 403), label=label)
        return (token, response) if admitted(response) else None
    member['token'], response = inst.wait(refreshed, 'Member admission through ' + path, seconds=seconds)
    return response
